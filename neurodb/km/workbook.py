"""Reading the KM workbook: the quick check of an upload (in the request), the full check of its rows (the
load's first pass, with no database access), its findings row by row (the load's second pass), and the
Excel template of every column the check expects.

Row numbers are Excel's: the header is row 1, the first data row is row 2. The messages name each column as
the workbook writes it and each sheet as the uploaded file names it; they never quote another cell, only IDs.
"""

from __future__ import annotations

import datetime
import io
import math
import re
import zipfile
from collections import Counter, defaultdict
from collections.abc import Iterator
from dataclasses import dataclass, field

import openpyxl
from openpyxl.styles import Font

from . import columns as cols
from . import words
from .words import counted, quoted, say

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_PARTS = 500  # parts of the zip an .xlsx is (a guard against zip bombs)
MAX_UNPACKED = 300 * 1024 * 1024
OLE = b"\xd0\xcf\x11\xe0"  # an encrypted (password or sensitivity label) Office file is an OLE file
MAX_DOCUMENTS = 5_000
MAX_FINDINGS = 100_000
EXAMPLES = 5  # rows or IDs a grouped problem names
LISTED = 20  # IDs a warning names
MAX_MESSAGES = 100
HEADER_SEARCH = 5  # a header found in rows 2-5 instead of row 1 is named
UNRECOGNISED_SHOWN = 20
UNRECOGNISED_CHARS = 60
LAST_WORDS = {cols.DOCUMENTS: "Documents", cols.FINDINGS: "Findings"}
DATE_FORMATS = (
    "%Y-%m-%d",
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%d.%m.%Y",
    "%d %B %Y",
    "%d %b %Y",
    "%B %d, %Y",
    "%b %d, %Y",
)


class Unreadable(Exception):
    """The stored workbook could not be read again (it was checked when it was uploaded)."""


# ------------------------------------------------------------------------------------------- the layout
@dataclass
class Sheet:
    """A data sheet as found: its title in the file, and where each expected column is."""

    kind: str
    title: str
    index: dict[str, int] = field(default_factory=dict)  # field of the column → its position in row 1
    missing: list[cols.Column] = field(default_factory=list)  # optional columns not found
    unknown: list[str] = field(default_factory=list)  # headers of row 1 that no column matches

    @property
    def columns(self) -> tuple[cols.Column, ...]:
        return cols.COLUMNS[self.kind]


@dataclass
class Checked:
    """What the upload check found: problems (nothing can be loaded), notes (it loads, with these said), the
    data sheets and the Read Me sheet's title ("" when there is none)."""

    problems: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    sheets: dict[str, Sheet] = field(default_factory=dict)
    read_me: str = ""


def _open(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)


def _zip_ok(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            parts = archive.infolist()
    except (zipfile.BadZipFile, OSError, ValueError):
        return False
    return len(parts) <= MAX_PARTS and sum(part.file_size for part in parts) <= MAX_UNPACKED


def _header(row) -> list[str]:
    return [cols.as_text(value) for value in (row or ())]


def _first_row(ws, number: int) -> tuple:
    return next(ws.iter_rows(min_row=number, max_row=number, values_only=True), ())


def _blank(row) -> bool:
    return all(value is None or (isinstance(value, str) and not value.strip()) for value in row)


def check_layout(wb, *, rows: bool = True) -> Checked:
    """The sheets, their headers and (``rows``) that each has a row under its header."""
    checked = Checked()
    titles = list(wb.sheetnames)
    found: dict[str, list[str]] = defaultdict(list)
    for title in titles:
        kind = cols.sheet_kind(title)
        if kind:
            found[kind].append(title)
    for kind in (cols.DOCUMENTS, cols.FINDINGS):
        if not found[kind]:
            checked.problems += [
                say(words.NO_SHEET, last=LAST_WORDS[kind]),
                say(words.SHEETS_ARE, sheets=", ".join(titles)),
            ]
        elif len(found[kind]) > 1:
            checked.problems += [
                say(words.TWO_SHEETS, last=LAST_WORDS[kind]),
                say(words.SHEETS_ARE, sheets=", ".join(titles)),
            ]
    checked.read_me = found[cols.READ_ME][0] if found[cols.READ_ME] else ""
    if checked.problems:
        return checked
    for kind in (cols.DOCUMENTS, cols.FINDINGS):  # both sheets are reported
        ws = wb[found[kind][0]]
        ws.reset_dimensions()  # read every row and column, whatever size the file declares
        sheet, problems = _headers(ws, kind, found[kind][0])
        if not problems and rows and not _has_rows(ws):
            problems.append(say(words.NO_ROWS, sheet=sheet.title))
        checked.problems += problems
        checked.sheets[kind] = sheet
    if checked.problems:
        return checked
    for sheet in checked.sheets.values():
        checked.notes += [
            say(words.MISSING_OPTIONAL, sheet=sheet.title, column=c.header) for c in sheet.missing
        ]
        if sheet.unknown:
            checked.notes.append(say(words.NOT_READ, sheet=sheet.title, headers=_listed(sheet.unknown)))
    if not checked.read_me:
        checked.notes.append(str(words.NO_READ_ME))
    return checked


def _listed(headers: list[str]) -> str:
    return quoted(h[:UNRECOGNISED_CHARS] for h in headers[:UNRECOGNISED_SHOWN]) + (
        ", …" if len(headers) > UNRECOGNISED_SHOWN else ""
    )


def _headers(ws, kind: str, title: str) -> tuple[Sheet, list[str]]:
    columns = cols.COLUMNS[kind]
    by_fold = {column.folded: column for column in columns}
    sheet = Sheet(kind=kind, title=title)
    problems: list[str] = []
    header = _header(_first_row(ws, 1))
    seen: Counter[str] = Counter()
    for position, text in enumerate(header):
        column = by_fold.get(cols.fold(text))
        if column is None:
            if text:
                sheet.unknown.append(text)
            continue
        seen[column.field] += 1
        sheet.index.setdefault(column.field, position)
    if not seen:  # no expected header in row 1: is it further down?
        for number, row in enumerate(
            ws.iter_rows(min_row=2, max_row=HEADER_SEARCH, values_only=True), start=2
        ):
            matched = {by_fold[f].field for f in map(cols.fold, _header(row)) if f in by_fold}
            if len(matched) * 2 >= len(columns):
                return sheet, [say(words.HEADER_ROW, sheet=title, row=number)]
    for column in columns:
        if seen[column.field] > 1:
            problems.append(say(words.REPEATED_COLUMN, column=column.header, sheet=title))
    missing = [c for c in columns if c.field not in sheet.index]
    required = [c.header for c in missing if c.required]
    if required:
        problems.append(say(words.MISSING_COLUMNS, sheet=title, columns=quoted(required)))
        if sheet.unknown:
            problems.append(say(words.UNRECOGNISED, headers=_listed(sheet.unknown)))
    sheet.missing = [c for c in missing if not c.required]
    return sheet, problems


def _has_rows(ws) -> bool:
    return any(not _blank(row) for row in ws.iter_rows(min_row=2, values_only=True))


def check_upload(data: bytes, name: str, size: int | None = None) -> Checked:
    """The quick checks of an upload, in order; a failing group stops the later ones (both data sheets'
    headers are reported together). ``size``: the file's size when ``data`` was not read (too big)."""
    checked = Checked()
    size_problems = []
    size = len(data) if size is None else size
    if size > MAX_UPLOAD_BYTES:
        # rounded up, so a file just over the limit never reads as the limit itself
        size_problems.append(say(words.TOO_BIG, size=f"{math.ceil(size * 10 / (1024 * 1024)) / 10:.1f}"))
    base = re.split(r"[/\\]", name or "")[-1]
    ext = base.rsplit(".", 1)[1] if "." in base.strip(".") else ""
    if not ext:
        size_problems.append(str(words.NO_EXTENSION))
    elif ext.casefold() != "xlsx":
        size_problems.append(say(words.NOT_XLSX, ext=ext.casefold()))
    if size_problems:
        checked.problems = size_problems
        return checked
    if data.startswith(OLE):
        checked.problems = [str(words.PROTECTED)]
        return checked
    if not _zip_ok(data):
        checked.problems = [str(words.NOT_EXCEL)]
        return checked
    try:
        wb = _open(data)
    except Exception:  # noqa: BLE001 - whatever openpyxl raises, the file is not one it can open
        checked.problems = [str(words.NOT_EXCEL)]
        return checked
    try:
        return check_layout(wb)
    finally:
        wb.close()


# ------------------------------------------------------------------------------------------ reading rows
def rows(ws, sheet: Sheet) -> Iterator[tuple[int, cols.Cells]]:
    """Each non-empty row under the header: its Excel row number and its cells read (a missing optional
    column gives its empty values)."""
    present = [(column, sheet.index[column.field]) for column in sheet.columns if column.field in sheet.index]
    empty = cols.empty_values(tuple(sheet.missing))
    for number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if _blank(row):
            continue
        cells = cols.Cells(values=dict(empty))
        width = len(row)
        for column, position in present:
            cols.read(column, row[position] if position < width else None, cells)
        yield number, cols.finish(cells)


@dataclass
class ReadMe:
    prepared_on: datetime.date | None = None
    source_name: str = ""
    stated: dict[str, int] = field(default_factory=dict)  # kind → the rows it says the sheet has


def _date(value) -> datetime.date | None:
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = cols.as_text(value)
    for pattern in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _count(value) -> int | None:
    """A row count the Read Me states ("14,869"), or None (also when it is more than a count can be)."""
    number = cols.whole(value, cols.as_text(value).replace(",", ""))
    return number if number is not None and number <= cols.COUNT_MAX else None


def read_me(ws, sheets: dict[str, Sheet]) -> ReadMe:
    """The Read Me sheet's label and value pairs the KM Database keeps: the prepared-on date, the source
    workbook's name (never its folder) and the row counts it states."""
    found = ReadMe()
    titles = {cols.fold(sheet.title): kind for kind, sheet in sheets.items()}
    prepared, source = cols.fold(cols.PREPARED_ON), cols.fold(cols.SOURCE_WORKBOOK)
    ws.reset_dimensions()
    for row in ws.iter_rows(values_only=True):
        if len(row) < 2:
            continue
        label = cols.fold(cols.as_text(row[0]))
        if not label:
            continue
        if label == prepared and found.prepared_on is None:
            found.prepared_on = _date(row[1])
        elif label == source and not found.source_name:
            found.source_name = re.split(r"[/\\]", cols.as_text(row[1]))[-1][:255]
        else:
            kind = titles.get(label) or {cols.DOCUMENTS: cols.DOCUMENTS, cols.FINDINGS: cols.FINDINGS}.get(
                label.split(" ")[-1]
            )
            number = _count(row[1])
            if kind and number is not None and kind not in found.stated:
                found.stated[kind] = number
    return found


@dataclass
class Report:
    """The load's first pass: the documents read in full, of the findings only what the checks need, and the
    problems (nothing is loaded) and warnings (it loads, with these said)."""

    sheets: dict[str, Sheet] = field(default_factory=dict)
    read_me: ReadMe = field(default_factory=ReadMe)
    documents: list[tuple[int, cols.Cells]] = field(default_factory=list)
    findings: int = 0
    review_findings: int = 0
    review_documents: int = 0
    check_findings: int = 0
    check_documents: int = 0
    problems: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _rows_list(rows_: list[int]) -> str:
    return ", ".join(str(n) for n in rows_[:EXAMPLES])


def _ids(
    sheet: Sheet, entries: list[tuple[int, str]], id_column: cols.Column, report: Report
) -> dict[str, int]:
    """Each ID's first row; the missing, malformed and repeated ones are problems, grouped by kind."""
    first: dict[str, int] = {}
    missing, bad, repeated = [], [], []
    for number, value in entries:
        if not value:
            missing.append(number)
        elif not cols.ID.match(value):
            bad.append(number)
        elif value in first:
            repeated.append((value, first[value], number))
        else:
            first[value] = number
    if missing:
        report.problems.append(
            say(
                words.ROWS_WITHOUT,
                n=len(missing),
                sheet=sheet.title,
                column=id_column.header,
                rows=_rows_list(missing),
            )
        )
    if bad:
        report.problems.append(
            say(words.BAD_IDS, n=len(bad), column=id_column.header, sheet=sheet.title, rows=_rows_list(bad))
        )
    if repeated:
        examples = "; ".join(
            say(words.REPEATED_EXAMPLE, id=value, first=a, second=b) for value, a, b in repeated[:EXAMPLES]
        )
        report.problems.append(
            say(
                words.REPEATED_IDS,
                n=len(repeated),
                column=id_column.header,
                sheet=sheet.title,
                examples=examples,
            )
        )
    return first


def _open_stored(content: bytes):
    try:
        wb = _open(content)
    except Exception as exc:  # noqa: BLE001
        raise Unreadable(str(words.NOT_EXCEL)) from exc
    return wb


def check_rows(content: bytes) -> Report:
    """Read every row once and check them (the load's first pass). Touches no database."""
    report = Report()
    wb = _open_stored(content)
    try:
        layout = check_layout(wb)
        if layout.problems:
            report.problems = layout.problems
            return report
        report.sheets = layout.sheets
        report.warnings += layout.notes
        documents, findings = layout.sheets[cols.DOCUMENTS], layout.sheets[cols.FINDINGS]
        if layout.read_me:
            report.read_me = read_me(wb[layout.read_me], layout.sheets)
        too_many = False
        document_checks, finding_checks = _Tally(), _Tally()
        for number, cells in rows(wb[documents.title], documents):
            if len(report.documents) >= MAX_DOCUMENTS:
                too_many = True
                break
            report.documents.append((number, cells))
            report.review_documents += cells.values["needs_review"]
            if cells.checks:
                report.check_documents += 1
                document_checks.add(cells.values["record_id"] or str(number), cells.checks)
        document_ids = _ids(
            documents, [(n, c.values["record_id"]) for n, c in report.documents], cols.RECORD_ID, report
        )
        finding_entries: list[tuple[int, str]] = []
        record_entries: list[tuple[int, str]] = []
        per_document: Counter[str] = Counter()
        for number, cells in rows(wb[findings.title], findings):
            if len(finding_entries) >= MAX_FINDINGS:
                too_many = True
                break
            values = cells.values
            finding_entries.append((number, values["finding_id"]))
            record_entries.append((number, values["record_id"]))
            per_document[values["record_id"]] += 1
            report.review_findings += values["needs_review"]
            if cells.checks:
                report.check_findings += 1
                finding_checks.add(values["finding_id"] or str(number), cells.checks)
        report.findings = len(finding_entries)
    finally:
        wb.close()
    if too_many:
        report.problems.append(str(words.TOO_MANY))
        return report
    empty = [
        sheet
        for sheet, count in ((documents, len(report.documents)), (findings, report.findings))
        if not count
    ]
    if empty:
        report.problems += [say(words.NO_ROWS, sheet=sheet.title) for sheet in empty]
        return report
    _ids(findings, finding_entries, cols.FINDING_ID, report)
    missing_record = [n for n, value in record_entries if not value]
    if missing_record:
        report.problems.append(
            say(
                words.ROWS_WITHOUT,
                n=len(missing_record),
                sheet=findings.title,
                column=cols.RECORD_ID.header,
                rows=_rows_list(missing_record),
            )
        )
    orphans = [n for n, value in record_entries if value and value not in document_ids]
    if orphans:
        report.problems.append(
            say(
                words.ORPHANS,
                n=len(orphans),
                sheet=findings.title,
                column=cols.RECORD_ID.header,
                documents=documents.title,
                rows=_rows_list(orphans),
            )
        )
    if report.problems:
        return report
    _warnings(report, per_document, document_checks, finding_checks)
    return report


@dataclass
class _Tally:
    """The rows holding a value to check: how many per column, and the first IDs."""

    columns: Counter = field(default_factory=Counter)
    ids: list[str] = field(default_factory=list)

    def add(self, row_id: str, checks: dict) -> None:
        self.columns.update(checks.keys())  # the columns, not their values
        if len(self.ids) < LISTED:
            self.ids.append(row_id)

    def say(self, template, n: int) -> str:
        columns = ", ".join(say(words.CHECK_COLUMN, column=c, n=k) for c, k in self.columns.most_common())
        return say(template, n=n, columns=columns, ids=", ".join(self.ids))


def _warnings(report: Report, per_document: Counter, document_checks: _Tally, finding_checks: _Tally) -> None:
    documents, findings = report.sheets[cols.DOCUMENTS], report.sheets[cols.FINDINGS]
    differs = [
        c.values["record_id"]
        for _n, c in report.documents
        if c.values.get("findings_stated") is not None
        and c.values["findings_stated"] != per_document[c.values["record_id"]]
    ]
    if differs and cols.NUMBER_OF_FINDINGS.field in documents.index:
        report.warnings.append(
            say(
                words.COUNT_DIFFERS,
                n=len(differs),
                column=cols.NUMBER_OF_FINDINGS.header,
                sheet=findings.title,
                ids=", ".join(differs[:LISTED]),
            )
        )
    if report.check_findings:
        report.warnings.append(finding_checks.say(words.FINDINGS_CHECK, report.check_findings))
    if report.check_documents:
        report.warnings.append(document_checks.say(words.DOCUMENTS_CHECK, report.check_documents))
    if report.review_findings or report.review_documents:
        parts = [
            counted(n, noun)
            for n, noun in (
                (report.review_findings, words.FINDINGS_N),
                (report.review_documents, words.DOCUMENTS_N),
            )
            if n
        ]
        report.warnings.append(
            say(
                words.MARKED_FOR_REVIEW,
                n=report.review_findings + report.review_documents,
                what=say(words.BOTH, first=parts[0], second=parts[1]) if len(parts) == 2 else parts[0],
                column=cols.DATA_QUALITY.header,
            )
        )
    lonely = [c.values["record_id"] for _n, c in report.documents if not per_document[c.values["record_id"]]]
    if lonely:
        report.warnings.append(
            say(words.NO_FINDINGS, n=len(lonely), sheet=findings.title, ids=", ".join(lonely[:LISTED]))
        )


def findings(content: bytes, sheets: dict[str, Sheet]) -> Iterator[tuple[int, cols.Cells]]:
    """The findings sheet's rows, read again (the load's second pass), one at a time."""
    wb = _open_stored(content)
    try:
        sheet = sheets[cols.FINDINGS]
        ws = wb[sheet.title]
        ws.reset_dimensions()
        yield from rows(ws, sheet)
    finally:
        wb.close()


# ----------------------------------------------------------------------------------------- the template
TEMPLATE_SHEETS = {cols.DOCUMENTS: "KM Documents", cols.FINDINGS: "KM Findings"}
TEMPLATE_READ_ME = "Read Me"


def template() -> bytes:
    """An empty workbook with every column the check expects (the required ones in bold) and the Read Me
    labels the load reads. No workbook text."""
    wb = openpyxl.Workbook()
    read_me_ws = wb.active
    read_me_ws.title = TEMPLATE_READ_ME
    for label in (cols.PREPARED_ON, cols.SOURCE_WORKBOOK, *TEMPLATE_SHEETS.values()):
        read_me_ws.append([label, None])
    read_me_ws.column_dimensions["A"].width = 24
    bold = Font(bold=True)
    for kind, title in TEMPLATE_SHEETS.items():
        ws = wb.create_sheet(title)
        ws.append([column.header for column in cols.COLUMNS[kind]])
        for cell, column in zip(ws[1], cols.COLUMNS[kind], strict=True):
            if column.required:
                cell.font = bold
        ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
