"""The text of an uploaded document: PDF (its text layer), Word (.docx), PowerPoint (.pptx), Excel
(.xlsx), CSV, Markdown or plain text. Pages (PDF pages, slides, sheets) are separated by a form
feed, so passages can say where they come from."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from xml.etree import ElementTree

PAGE_BREAK = "\f"
MAX_PAGES = 2000
MAX_SHEET_ROWS = 20000
MAX_XML_MB = 60
WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
DRAWING = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


class TextError(Exception):
    """The document cannot be read; the message says why, for the person who added it."""


@dataclass
class Extracted:
    text: str
    pages: int


def decode_text(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _xml(archive: zipfile.ZipFile, name: str) -> ElementTree.Element:
    if archive.getinfo(name).file_size > MAX_XML_MB * 1024 * 1024:
        raise TextError("The document is too large to read.")
    # an upload by a signed-in editor; Python's expat refuses entity expansion attacks
    return ElementTree.fromstring(archive.read(name))  # noqa: S314


def _zip(data: bytes, kind: str) -> zipfile.ZipFile:
    try:
        return zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise TextError(f"The {kind} file could not be opened; save it again or as PDF.") from exc


def docx_text(data: bytes) -> str:
    """The paragraphs of a Word document, one per line (table cells included)."""
    try:
        with _zip(data, "Word") as archive:
            root = _xml(archive, "word/document.xml")
    except (KeyError, ElementTree.ParseError) as exc:
        raise TextError("The Word file could not be opened; save it again as .docx or PDF.") from exc
    lines = []
    for paragraph in root.iter(f"{WORD}p"):
        parts = []
        for node in paragraph.iter():
            if node.tag == f"{WORD}t" and node.text:
                parts.append(node.text)
            elif node.tag == f"{WORD}tab":
                parts.append("\t")
        lines.append("".join(parts))
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def pptx_text(data: bytes) -> tuple[str, int]:
    with _zip(data, "PowerPoint") as archive:
        names = [n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        names.sort(key=lambda n: int(re.search(r"(\d+)\.xml$", n).group(1)))
        slides = []
        for name in names[:MAX_PAGES]:
            try:
                root = _xml(archive, name)
            except ElementTree.ParseError:
                continue
            lines = []
            for paragraph in root.iter(f"{DRAWING}p"):
                line = "".join(t.text or "" for t in paragraph.iter(f"{DRAWING}t")).strip()
                if line:
                    lines.append(line)
            slides.append("\n".join(lines))
    return PAGE_BREAK.join(slides), len(slides)


def xlsx_text(data: bytes) -> tuple[str, int]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any unreadable workbook is reported the same way
        raise TextError("The Excel file could not be opened; save it again as .xlsx.") from exc
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for n, values in enumerate(ws.iter_rows(values_only=True)):
            if n >= MAX_SHEET_ROWS:
                rows.append("[… more rows not read]")
                break
            cells = [str(v).strip() for v in values if v not in (None, "")]
            if cells:
                rows.append(" | ".join(cells))
        sheets.append(f"Sheet: {ws.title}\n" + "\n".join(rows))
    return PAGE_BREAK.join(sheets), len(sheets)


def pdf_text(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise TextError("The PDF is protected by a password; upload an unprotected copy.")
        pages = []
        for page in reader.pages[:MAX_PAGES]:
            try:
                pages.append(page.extract_text() or "")
            except Exception:  # noqa: BLE001 - one unreadable page does not lose the others
                pages.append("")
    except TextError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise TextError("The PDF could not be read; save it again or upload it as Word.") from exc
    if len("".join(pages).strip()) < 20 * max(len(pages), 1) and len("".join(pages).strip()) < 2000:
        raise TextError(
            "The PDF has (almost) no text: it is probably a scan. Upload a version with text (Word, or a "
            "PDF exported from Word), or paste its text."
        )
    return PAGE_BREAK.join(pages), len(pages)


def read(filename: str, data: bytes) -> Extracted:
    name = filename.lower()
    if name.endswith(".pdf"):
        text, pages = pdf_text(data)
    elif name.endswith(".docx"):
        text, pages = docx_text(data), 1
    elif name.endswith(".pptx"):
        text, pages = pptx_text(data)
    elif name.endswith(".xlsx"):
        text, pages = xlsx_text(data)
    elif name.endswith((".txt", ".md", ".csv")):
        text, pages = decode_text(data), 1
    else:
        raise TextError("This kind of file cannot be read: use PDF, Word, PowerPoint, Excel, CSV or text.")
    text = clean(text)
    if not text.replace(PAGE_BREAK, "").strip():
        raise TextError("The document holds no text.")
    return Extracted(text, pages)


def clean(text: str) -> str:
    """NUL characters out (PostgreSQL refuses them), runs of spaces and blank lines shortened."""
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]{2,}", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()
