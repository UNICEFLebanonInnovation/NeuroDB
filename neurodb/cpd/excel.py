"""The results framework as an Excel sheet: the template (pre-filled with what is already entered, so
it is also the export) and its import.

One row per item. Level is Outcome, Output or Indicator; Parent code is the outcome of an output,
and the outcome or output of an indicator. Items are matched by their code within the cycle (an
indicator without a code by its title): an import adds new items and updates existing ones, and
never deletes. The whole sheet is checked first; nothing is saved when a row is wrong.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .models import CountryProgramme, Indicator, Milestone, Origin, Outcome, Output

SHEET = "Results framework"
COLUMNS = [
    ("level", "Level", 12),
    ("code", "Code", 10),
    ("title", "Title", 60),
    ("parent", "Parent code", 12),
    ("unit", "Unit", 12),
    ("direction", "Direction", 12),
    ("accumulate", "Cycle value", 14),
    ("baseline", "Baseline", 11),
    ("baseline_year", "Baseline year", 12),
    ("target", "Target (end of cycle)", 14),
    ("means", "Means of verification", 40),
    ("etools_output", "eTools CP output (outputs only)", 36),
]
LEVELS = ("Outcome", "Output", "Indicator")
UNITS = {"number": "number", "#": "number", "percentage": "percent", "percent": "percent", "%": "percent"}
DIRECTIONS = {"increase": "increase", "up": "increase", "decrease": "decrease", "down": "decrease"}
ACCUMULATE = {"latest": "latest", "level": "latest", "sum": "sum", "cumulative": "sum"}


def milestone_header(year: int) -> str:
    return f"Milestone {year}"


def template(programme: CountryProgramme) -> bytes:
    """The sheet with the cycle's current framework, ready to fill."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    headers = [label for _key, label, _w in COLUMNS] + [milestone_header(y) for y in programme.years]
    ws.append(headers)
    bold = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="2C55A8")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, (_key, _label, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    ws.freeze_panes = "A2"

    def indicator_row(ind: Indicator, parent: str) -> list[Any]:
        milestones = {m.year: m.value for m in ind.milestones.all()}
        return [
            "Indicator",
            ind.code,
            ind.title,
            parent,
            ind.get_unit_display(),
            ind.get_direction_display(),
            "sum" if ind.accumulate == Indicator.Accumulate.SUM else "latest",
            ind.baseline,
            ind.baseline_year,
            ind.target,
            ind.means_of_verification,
            "",
        ] + [milestones.get(y) for y in programme.years]

    for outcome in programme.outcomes.prefetch_related(
        "outputs__indicators__milestones", "indicators__milestones"
    ):
        ws.append(["Outcome", outcome.code, outcome.title, ""])
        for ind in outcome.indicators.all():
            ws.append(indicator_row(ind, outcome.code))
        for output in outcome.outputs.all():
            ws.append(
                [
                    "Output",
                    output.code,
                    output.title,
                    outcome.code,
                    "",
                    "",
                    "",
                    None,
                    None,
                    None,
                    "",
                    output.etools_output,
                ]
            )
            for ind in output.indicators.all():
                ws.append(indicator_row(ind, output.code))
    if ws.max_row == 1:  # an example row, to show the shape
        ws.append(["Outcome", "1", "[Outcome title]", ""])
        ws.append(["Output", "1.1", "[Output title]", "1"])
        ws.append(
            [
                "Indicator",
                "1.1.1",
                "[Indicator title]",
                "1.1",
                "Number",
                "Increase",
                "sum",
                0,
                None,
                None,
                "[source]",
            ]
        )
    last = max(ws.max_row, 500)
    for column, options in (
        ("A", LEVELS),
        ("E", ("Number", "Percentage")),
        ("F", ("Increase", "Decrease")),
        ("G", ("latest", "sum")),
    ):
        rule = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(rule)
        rule.add(f"{column}2:{column}{last}")
    help_ws = wb.create_sheet("How to fill")
    for line in (
        "One row per outcome, output or indicator. Level: Outcome, Output or Indicator.",
        "Parent code: for an output, its outcome's code; for an indicator, its outcome's or output's code.",
        "Unit: Number or Percentage. Direction: Increase or Decrease (e.g. a dropout rate to reduce).",
        "Cycle value: 'latest' when the indicator is a level (a rate, a share); "
        "'sum' when yearly results add up "
        "over the cycle (children reached).",
        "Milestone <year>: optional target for that year. Target: the end-of-cycle target.",
        "eTools CP output: the output's name as eTools writes it on programme documents "
        "(blank: matched by code).",
        "Items are matched by code: importing again updates them; nothing is deleted.",
    ):
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 120
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    errors: list[str] = field(default_factory=list)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _number(value: Any, row: int, column: str, errors: list[str]) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "").replace("%", "").strip())
    except ValueError:
        errors.append(f"Row {row}: {column} is not a number ({value!r}).")
        return None


def _year(value: Any, row: int, column: str, errors: list[str]) -> int | None:
    number = _number(value, row, column, errors)
    if number is None:
        return None
    if not (1990 <= number <= 2100) or not float(number).is_integer():
        errors.append(f"Row {row}: {column} is not a year ({value!r}).")
        return None
    return int(number)


def read_rows(data: bytes) -> tuple[list[dict[str, Any]], dict[int, int], list[str]]:
    """The rows of the sheet (dicts keyed as COLUMNS), the milestone columns and the errors."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - any unreadable file is reported the same way
        return [], {}, ["The file is not an Excel workbook (.xlsx)."]
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], {}, ["The sheet is empty."]
    header = [_text(h).lower() for h in rows[0]]
    positions: dict[str, int] = {}
    for key, label, _w in COLUMNS:
        if label.lower() in header:
            positions[key] = header.index(label.lower())
    missing = [
        label
        for key, label, _w in COLUMNS
        if key in ("level", "code", "title", "parent") and key not in positions
    ]
    if missing:
        return [], {}, [f"Missing column(s): {', '.join(missing)}. Start from the template."]
    milestones: dict[int, int] = {}
    for index, name in enumerate(header):
        if name.startswith("milestone "):
            year = name.removeprefix("milestone ").strip()
            if year.isdigit():
                milestones[int(year)] = index
    out = []
    for number, values in enumerate(rows[1:], start=2):
        if not any(v not in (None, "") for v in values):
            continue
        row = {key: values[i] if i < len(values) else None for key, i in positions.items()}
        row["_row"] = number
        row["_milestones"] = {y: values[i] if i < len(values) else None for y, i in milestones.items()}
        out.append(row)
    return out, milestones, []


def import_framework(programme: CountryProgramme, data: bytes) -> ImportResult:
    result = ImportResult()
    rows, _milestones, errors = read_rows(data)
    result.errors = errors
    if errors:
        return result
    outcomes = {o.code: o for o in programme.outcomes.all()}
    outputs = {o.code: o for o in Output.objects.filter(outcome__programme=programme)}
    plan: list[tuple[str, dict[str, Any]]] = []
    seen_outcomes, seen_outputs = set(outcomes), set(outputs)
    for row in rows:
        n = row["_row"]
        level = _text(row.get("level")).capitalize()
        code, title, parent = _text(row.get("code")), _text(row.get("title")), _text(row.get("parent"))
        if level not in LEVELS:
            errors.append(f"Row {n}: Level must be Outcome, Output or Indicator ({row.get('level')!r}).")
            continue
        if not title:
            errors.append(f"Row {n}: the title is empty.")
            continue
        if level in ("Outcome", "Output") and not code:
            errors.append(f"Row {n}: an {level.lower()} needs a code.")
            continue
        if level == "Outcome":
            seen_outcomes.add(code)
        elif level == "Output":
            if parent not in seen_outcomes:
                errors.append(
                    f"Row {n}: output {code}'s parent outcome {parent!r} is not in the sheet or the cycle."
                )
                continue
            seen_outputs.add(code)
        else:
            if parent not in seen_outcomes and parent not in seen_outputs:
                errors.append(
                    f"Row {n}: indicator {code or title[:40]!r}: "
                    f"parent {parent!r} is not an outcome or output."
                )
                continue
            unit = UNITS.get(_text(row.get("unit")).lower() or "number")
            direction = DIRECTIONS.get(_text(row.get("direction")).lower() or "increase")
            accumulate = ACCUMULATE.get(_text(row.get("accumulate")).lower() or "latest")
            if unit is None or direction is None or accumulate is None:
                errors.append(f"Row {n}: Unit, Direction or Cycle value is not one of the listed choices.")
                continue
            row["_unit"], row["_direction"], row["_accumulate"] = unit, direction, accumulate
            row["_baseline"] = _number(row.get("baseline"), n, "Baseline", errors)
            row["_baseline_year"] = _year(row.get("baseline_year"), n, "Baseline year", errors)
            row["_target"] = _number(row.get("target"), n, "Target", errors)
            row["_milestones"] = {
                y: _number(v, n, f"Milestone {y}", errors)
                for y, v in row["_milestones"].items()
                if v not in (None, "")
            }
        plan.append((level, row))
    if errors:
        return result
    with transaction.atomic():
        for level, row in plan:
            code, title, parent = _text(row.get("code")), _text(row.get("title")), _text(row.get("parent"))
            if level == "Outcome":
                obj, created = Outcome.objects.update_or_create(
                    programme=programme, code=code, defaults={"title": title, "origin": Origin.IMPORTED}
                )
                outcomes[code] = obj
            elif level == "Output":
                obj = outputs.get(code)
                created = obj is None
                obj = obj or Output(code=code)
                obj.outcome, obj.title, obj.origin = outcomes[parent], title, Origin.IMPORTED
                obj.etools_output = _text(row.get("etools_output")) or obj.etools_output
                obj.save()
                outputs[code] = obj
            else:
                existing = Indicator.objects.filter(programme=programme)
                obj = (existing.filter(code=code) if code else existing.filter(title=title, code="")).first()
                created = obj is None
                obj = obj or Indicator(programme=programme)
                obj.outcome = outcomes.get(parent) if parent in outcomes and parent not in outputs else None
                obj.output = outputs.get(parent) if obj.outcome is None else None
                obj.code, obj.title = code, title
                obj.unit, obj.direction, obj.accumulate = row["_unit"], row["_direction"], row["_accumulate"]
                obj.baseline, obj.baseline_year, obj.target = (
                    row["_baseline"],
                    row["_baseline_year"],
                    row["_target"],
                )
                obj.means_of_verification = _text(row.get("means"))
                obj.origin = Origin.IMPORTED
                obj.save()
                for year, value in row["_milestones"].items():
                    if value is not None and year in programme.years:
                        Milestone.objects.update_or_create(
                            indicator=obj, year=year, defaults={"value": value}
                        )
            if created:
                result.created += 1
            else:
                result.updated += 1
    return result
