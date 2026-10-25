"""The Power BI comparison sheet (SPEC_validation §4): one Excel file per page and filter, for the people who
check the figures, with NeuroDB's figures in the PRP Power BI dashboard's order and names, an empty column for
the Power BI value, the difference, and the lists to match row by row (Progress reports: its progress reports
and partners; Partner monitoring: every indicator of Program Performance). NeuroDB never reads Power BI: a
person types Power BI's numbers in.

* **The page's own figures.** Every value comes from the functions the page calls, for its filter
  (``services.partner_reporting`` and ``reporting_table``; ``monitoring.stored_indicators``,
  ``program_performance_rows`` and ``data_quality``), on the day the statuses are read for (Due, Overdue and
  Not yet due depend on it): the page's own day by default, or Power BI's refresh, at most :data:`DAYS_BACK`
  days before. They are read in one read-only snapshot of the data.
* **Power BI's names** are in one place, :data:`POWERBI_PAGE_ORDER`, :data:`POWERBI_VISUALS` (and
  :data:`POWERBI_PAGES` and :data:`POWERBI_PD_PAGE` made from them), :data:`POWERBI_ITEMS` and
  :data:`POWERBI_SLICERS`: the names of the shared report (UNICEF Program Performance), as the user's
  screenshot shows them; correct them there. Every sentence of the sheet that names one reads it there, and so
  does Compare with Power BI (``powerbi_exports``).
* **The filter.** Every field of the page's filter is a Power BI slicer to set (:data:`POWERBI_SLICERS`) or a
  NeuroDB filter to clear before comparing (:data:`NEURODB_ONLY`); a new field in neither stops the sheet.
* **People** never reach the file: partners, programme documents and sections only; the progress reports'
  sheet is the export's row without its focal points. Every text of the lists (a partner, a PD reference, a
  report #, an indicator, its outputs and section) and every value of the filter is cleaned as Raw eTools
  data cleans that column (e-mails, phone numbers and known person names withheld), as the figure trail's
  CSVs clean them; a filter that names a person or a contact (a search holding an e-mail, a phone number or a
  known person's name) has its addresses (the page's and every *Trail* link) not written. Every text from the
  data is written as a text, never a formula; the only formulas are the two difference columns of a sheet,
  written by NeuroDB.

The *Trail* column opens each figure on its figure trail (``datamart.trail``, ``datamart.trail_monitoring``),
for the same filter and day; on the Program Performance rows sheet, each indicator's Reached.
"""

from __future__ import annotations

import dataclasses
import datetime
import io
from collections import Counter
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from typing import Any

from django.conf import settings
from django.db.models import Max
from django.http import QueryDict
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill

from neurodb.partnerships import statuses as pd_statuses
from neurodb.partnerships.models import PCA, PartnerOrganization

from . import models as dm
from . import monitoring, reporting, runfacts
from . import reached as rc

# ------------------------------------------------------------------------------------------- Power BI's names
# The PRP Power BI report, as the user's screenshot of the shared report shows it (25 Oct 2026): its pages in
# their order, and each visual NeuroDB compares, by the key of a block of the sheet or of an export Compare
# with Power BI recognises (``powerbi_exports``). Correct the words here, and only here.
POWERBI_REPORT = "UNICEF Program Performance"
POWERBI_PAGE_ORDER = (
    "Program Performance",
    "Partner Performance",
    "Reporting Progress",
    "Data Quality",
    "PD/SPD Indicator",
    "PD/SPD Report",
    "PD/SPD Report/Location",
    "PD/SPD Report/Disaggregation",
    "Partner Fund Monitoring",
    "Funded Programs",
)
POWERBI_VISUALS: dict[str, tuple[str, str]] = {  # a sheet block or an export -> (Power BI page, visual)
    "pp_indicators": ("Program Performance", "Table"),
    "pp_sections": ("Program Performance", "Section tiles"),
    "partner_cards": (
        "Partner Performance",
        "# of Partners with (Accepted, Submitted, Sent back, Due, and Overdue)",
    ),
    "follow_up": ("Partner Performance", "Partners require follow up (Overdue, Sent back, Submitted & Due)"),
    "status_cards": ("Reporting Progress", "Cards"),
    "by_section": ("Reporting Progress", "Reporting Progress by Sector"),
    "by_month": ("Reporting Progress", "# of Reports by Type over Reporting Date"),
    "dq_tiles": ("Data Quality", "No. of Indicators without Reached"),
    "dq_list": ("Data Quality", "Indicators without Reached or with Zero Reached"),
    "pd_indicators": ("PD/SPD Indicator", "Table"),
    "report_list": ("PD/SPD Report", "Table"),
    "locations": ("PD/SPD Report/Location", "Table"),
    "disaggregations": ("PD/SPD Report/Disaggregation", "Table"),
    "fund_monitoring": ("Partner Fund Monitoring", "Table"),
    "funded_programmes": ("Funded Programs", "Table"),
}
# the blocks of each page's sheet; :data:`POWERBI_PAGES` puts them in Power BI's order, by Power BI page
SHEET_BLOCKS: dict[str, tuple[str, ...]] = {
    "progress-reports": (
        "partner_cards",
        "follow_up",
        "status_cards",
        "by_section",
        "by_month",
        "report_list",
        "locations",
        "disaggregations",
    ),
    "partner-monitoring": ("pp_indicators", "pp_sections", "dq_tiles"),
}


def _powerbi_pages(blocks: tuple[str, ...]) -> tuple[tuple[str, tuple[tuple[str, str], ...]], ...]:
    """The blocks grouped by Power BI page, in Power BI's order, each with its visual's name."""
    found = []
    for name in POWERBI_PAGE_ORDER:
        visuals = tuple((b, POWERBI_VISUALS[b][1]) for b in blocks if POWERBI_VISUALS[b][0] == name)
        if visuals:
            found.append((name, visuals))
    return tuple(found)


# each page's sheet: its Power BI pages in their order, each with its visuals and the block of figures each is
# compared with
POWERBI_PAGES: dict[str, tuple[tuple[str, tuple[tuple[str, str], ...]], ...]] = {
    page: _powerbi_pages(blocks) for page, blocks in SHEET_BLOCKS.items()
}
# Power BI's page with no row in the sheet: NeuroDB shows it on Programmes → Indicators
POWERBI_PD_PAGE = POWERBI_VISUALS["pd_indicators"][0]
# the items Power BI names itself (a status, a section or a month is named by the data)
POWERBI_ITEMS: dict[str, str] = {
    "total": "Total",
    "not_yet_due": "Not Yet Due",
    "sent_back": "Sent back",
    "no_reached": "No. of Indicators without Reached",
    # NeuroDB's own items: no Power BI export names them
    "reports": "Number of reports",
    "lines": "Number of lines",
    "locations": "Number of locations",
    "indicators": "Number of indicators",
    "partners": "Partners",
    "partners_reported": "With at least one report",
    "partners_not_reported": "Not reported",
    "partners_status": "Partners with a report {status}",
    "pds": "Programme documents",
    "running_pds": "Running PDs",
    "active_pds": "Active PDs",
    "ended_pds": "Ended PDs",
    "dq_indicators": "Indicators",
    "with_reached": "With Reached",
    "zero_reached": "Reported zero",
    "without_reached": "Without Reached",
    "with_target": "With a target",
    "without_target": "Without a target",
}
# the page's filter fields Power BI has a slicer for, by the slicer's name: a slicer that means the same. A
# NeuroDB filter whose namesake means something else in Power BI is NeuroDB's own (:data:`NEURODB_ONLY`)
POWERBI_SLICERS: dict[str, dict[str, str]] = {
    "progress-reports": {
        "year": "Reporting Year",
        "pd_statuses": "PD status",
        "sections": "Section",
        "partners": "Partner",
    },
    "partner-monitoring": {
        "year": "Reporting Year",
        "pd_statuses": "PD status",
        "sections": "Section",
        "partners": "Partner",
        "document_types": "Document type",
        "offices": "Office",
    },
}
# Power BI's slicers NeuroDB's page has no filter for: set as said on every page, or cleared. Its Emergency
# indicator is Yes for an SPD (NeuroDB's emergency filter: the indicator's high frequency), its CP Output
# search matches every CP output of the PD (NeuroDB's: the indicator's), and its Indicator Title box the title
# alone (Partner monitoring's search: the title or the output)
POWERBI_ONLY_SLICERS = ("Country name",)  # NeuroDB holds one country
POWERBI_TO_CLEAR: dict[str, tuple[str, ...]] = {
    "progress-reports": ("Emergency indicator", "Document type", "Office", "CP Output", "Indicator Title"),
    "partner-monitoring": ("Emergency indicator", "CP Output", "Indicator Title"),
}
# Power BI's fields (as its "Applied filters" line names them) that a slicer of the report sets, by the
# slicer's name: Compare with Power BI's "What to do now" names the slicer, never the field
POWERBI_FIELD_SLICERS: dict[str, str] = {
    "country_name": POWERBI_ONLY_SLICERS[0],
    "status": POWERBI_SLICERS["partner-monitoring"]["pd_statuses"],
    "reporting_year": POWERBI_SLICERS["partner-monitoring"]["year"],
    "section_name": POWERBI_SLICERS["partner-monitoring"]["sections"],
}
ONE_COUNTRY = "{country}: NeuroDB holds one country"
CLEAR_IN_POWERBI = "Clear these in Power BI: NeuroDB's page has no such filter"
COMPARE_FOR_YOU = (
    "Or let NeuroDB compare for you: upload Power BI's exports in Admin → Data and sync → Compare with Power "
    "BI."
)

# --------------------------------------------------------------------------------------------- NeuroDB's side
# the page's filter fields Power BI has no slicer for, in NeuroDB's words: cleared on NeuroDB before comparing
NEURODB_ONLY: dict[str, dict[str, str]] = {
    "progress-reports": {
        "q": "Search",
        "report_numbers": "Report #",
        "pds": "Programme document (Home's rows)",
        "partner_types": "Partner type",
        "start_from": "Period start from",
        "start_to": "Period start to",
        "end_from": "Period end from",
        "end_to": "Period end to",
        "overdue": "Overdue only",
        "report_types": "Report type",
        "statuses": "Report status",
        "emergency": "Emergency indicator (high frequency)",
        # By disaggregation's own filters
        "indicator": "By disaggregation: indicator",
        "location": "By disaggregation: location",
        "p_code": "By disaggregation: P-code",
        "disaggregation": "By disaggregation: disaggregation",
    },
    "partner-monitoring": {
        "pds": "Programme document",
        "pd_ids": "Programme document (its page)",
        "locations": "Location",
        "report_type": "Report type",
        "scope": "Programme documents",
        "status": "Tracking status",
        "q": "Indicator or output",
        "tags": "Tags",
        "partner_types": "Partner type",
        "emergency": "Emergency (high frequency)",
        "cp_output": "CP output",
    },
}
# the tags of Partner monitoring's filter, as its filter bar names them
TAG_WORDS: dict[str, str] = {
    "gender": "Gender",
    "age_group": "Age group",
    "nationality": "Nationality",
    "disability": "Disability",
}
# the filter dataclasses each page reads its address with
FILTERS: dict[str, tuple[type, ...]] = {
    "progress-reports": (reporting.ReportFilters, reporting.DisaggregationFilters),
    "partner-monitoring": (monitoring.Filters,),
}
PAGES = tuple(POWERBI_PAGES)  # the pages with a sheet (Home has no Power BI page)
TITLES = {"progress-reports": "Progress reports", "partner-monitoring": "Partner monitoring"}
DAYS_BACK = 14  # "statuses read for": the page's day or Power BI's refresh, at most this many days before
NOT_THE_FILTER = ("page", "sort", "export", "tab", "_part", "_eager", "as_of", "day")  # keys of an address
REFUSED_DAY = "Choose a day in the last 14 days: the day of Power BI's refresh, or today."
# written instead of an address whose filter names a person or a contact (:func:`names_someone`)
ADDRESS_WITHHELD = "Not written: the filter names a person or a contact. Copy it from the page."
TRAIL_WITHHELD = "Not written: the filter names a person or a contact. Open the trail from the page."
# a row of a list whose match key lost a name, an e-mail or a phone: Power BI's own export holds them
KEY_WITHHELD = "A name, e-mail or phone is withheld here: match this row by PD and target."
NO_SECTION = "(no section)"
IR, PR = "indicator reports", "progress reports"
OVERDUE_NOTE = (
    "Power BI can count reports owed that the partner reports feed does not carry: see check 11 (Admin → "
    "Data and sync → eTools reporting checks)."
)
PDS_NOTE = (
    "The programme documents of these indicators: a running programme document with no indicator of this "
    "filter is not counted (see check 10: running programme documents with no report row)."
)
# the Power BI pages the reasons name, by the block of figures on each (:func:`powerbi_page`)
WHY_PAGES = {
    "program_performance": "pp_indicators",
    "reporting_progress": "by_section",
    "partner_performance": "partner_cards",
    "pd_report": "report_list",
    "locations": "locations",
    "disaggregations": "disaggregations",
}
WHY_THEY_DIFFER = (  # (why, where to see it); {history} and the Power BI pages of WHY_PAGES are filled in
    (
        "The days differ: Power BI refreshes on its own schedule, NeuroDB each night at 20:30.",
        "This file: Statuses read for, and Power BI refreshed on; Compare with Power BI reads the refresh "
        "from Power BI's files.",
    ),
    (
        "NeuroDB holds partner reports from {history} only: Power BI may show older ones.",
        "This file: Everything NeuroDB holds; Raw eTools data → Partner reports.",
    ),
    (
        "PRP's other feeds list reports still owed that the partner reports feed does not carry.",
        "eTools reporting checks: check 11.",
    ),
    (
        "Draft and cancelled programme documents are never read by NeuroDB.",
        "Raw eTools data → Partner reports, programme document link: draft or cancelled.",
    ),
    (
        "A report's status is its newest word across all its rows.",
        "Raw eTools data → Partner reports: a record's page gives the status every page reads for it.",
    ),
    (
        "Reached: Power BI's {program_performance} counts Accepted reports only (its filter says so); "
        "NeuroDB adds Accepted and Submitted reports, never adds a value repeated on every location, and "
        "takes PRP's cumulative for a partial history.",
        "eTools reporting checks: checks 2, 4 and 5; Compare with Power BI: {program_performance}.",
    ),
    (
        "An indicator is counted once, although eTools lists it per location and disaggregation.",
        "eTools reporting checks: check 7; Raw eTools data → PD indicators.",
    ),
    (
        "{reporting_progress} in Power BI counts progress reports; NeuroDB's tiles and charts count "
        "indicator reports.",
        "This file: Progress reports; Compare with Power BI: {reporting_progress}.",
    ),
    (
        "Power BI's lists call a report not yet due 'Due'.",
        "Compare with Power BI: {partner_performance} and {pd_report}.",
    ),
    (
        "Power BI joins rows that are the same in every column it shows, and repeats a location's row once "
        "per CP output.",
        "Compare with Power BI: {locations}.",
    ),
    (
        "Power BI shows a percentage indicator as a fraction on its {locations} and {disaggregations} pages.",
        "Compare with Power BI: {locations} and {disaggregations}.",
    ),
)
CELL_CHARS = 32_767  # Excel's limit for one cell
YELLOW = PatternFill("solid", fgColor="FFF2CC")  # a cell a person types into
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)


@dataclass(frozen=True)
class Figure:
    """One figure of the page with a Power BI counterpart: ``key`` is the figure's key (the trail's),
    ``item`` Power BI's name for it, ``neurodb`` the page and figure as NeuroDB names them."""

    key: str
    powerbi_page: str
    visual: str
    item: str
    neurodb: str
    value: int | float | None
    unit: str
    note: str = ""


@dataclass
class Read:
    """What the page shows for its filter on the day, read once for the whole file."""

    page: str
    params: QueryDict
    day: datetime.date
    filters: Any
    tab_filters: Any = None
    data: dict[str, Any] = field(default_factory=dict)


# ---------------------------------------------------------------------------------------------------- the day
def page_today(page: str) -> datetime.date:
    """The day the page reads each report's status for, as the page itself does."""
    from . import services

    return services.reporting_today() if page == "progress-reports" else monitoring.today()


def statuses_day(page: str, text: Any) -> datetime.date | None:
    """The day the sheet reads the statuses for: the page's day when none is given, else that day when it is
    at most :data:`DAYS_BACK` days before; None for any other value (refused)."""
    today = page_today(page)
    text = str(text or "").strip()
    if not text:
        return today
    try:
        day = datetime.date.fromisoformat(text[:10])
    except ValueError:
        return None
    return day if today - datetime.timedelta(days=DAYS_BACK) <= day <= today else None


def filter_query(params: Any) -> QueryDict:
    """The page's filter from its address: every key but the page shown, its sort, the export and the day."""
    out = QueryDict(mutable=True)
    if hasattr(params, "lists"):
        lists = params.lists()
    else:  # a plain dict (tests): one value or a list
        lists = ((k, v if isinstance(v, list | tuple) else [v]) for k, v in params.items())
    for key, values in lists:
        if key not in NOT_THE_FILTER:
            out.setlist(key, [str(v) for v in values])
    return out


# ----------------------------------------------------------------------------------------- the page's figures
def read(page: str, params: Any, day: datetime.date) -> Read:
    """The page's figures for ``params`` on ``day``, from the page's own functions."""
    from . import services

    params = filter_query(params)
    if page == "progress-reports":
        found = Read(page, params, day, reporting.ReportFilters.from_params(params))
        found.tab_filters = reporting.DisaggregationFilters.from_params(params)
        found.data = {
            "page": services.partner_reporting(params, day, options=False),
            "left_out": reporting.left_out(found.filters, day),
            "locations": services.reporting_table("locations", params, day)["locations"],
            "disaggregations": services.reporting_table("disaggregations", params, day)["disaggregations"],
        }
        return found
    if page != "partner-monitoring":
        raise ValueError(f"no comparison sheet for the page {page!r}")
    # the grid opens on this year, Program Performance and Data quality on every year (the check line always
    # carries the year the page shows)
    filters = monitoring.Filters.from_params(params, default_year=(params.get("view") or "grid") == "grid")
    every = monitoring.stored_indicators(dataclasses.replace(filters, sections=[]), day)
    found = Read(page, params, day, filters)
    found.data = {
        "program_performance": monitoring.program_performance_rows(filters, day, every=every),
        "section_tiles": monitoring.section_counts(every),
        "quality": monitoring.data_quality(monitoring.stored_indicators(filters, day), filters, day),
    }
    return found


# Each block gives the figures of one Power BI visual: (key, item, NeuroDB page and figure, value, unit, note)
Line = tuple[str, str, str, int | float | None, str, str]


def _status_cards(r: Read) -> Iterator[Line]:
    counts, left_out = r.data["page"]["counts"], r.data["left_out"]
    note = ""
    if left_out["by_pd_status"] or left_out["no_pd"]:
        note = (
            f"NeuroDB leaves out {left_out['by_pd_status']:,} indicator reports of programme documents with "
            f"another status or not in NeuroDB, and {left_out['no_pd']:,} naming no programme document."
        )
    total = counts["indicator_reports"]
    yield (
        "pr.indicator_reports",
        POWERBI_ITEMS["total"],
        "Progress reports → Indicator reports",
        total,
        IR,
        note,
    )
    for s in counts["statuses"]:
        note = OVERDUE_NOTE if s["key"] == rc.OVERDUE else ""
        figure = f"Progress reports → {s['label']}"
        yield (f"pr.status.{s['key']}", s["label"], figure, s["indicator_reports"], IR, note)


def _by_section(r: Read) -> Iterator[Line]:
    table = r.data["page"]["by_section"]
    labels = [rc.STATUS_LABELS[k] for k in reporting.STATUS_KEYS]
    figure, total = "Progress reports → Indicator reports by section", POWERBI_ITEMS["total"]
    for row in table["rows"]:
        section = row["section"]
        name = NO_SECTION if section == reporting.NO_SECTION else section
        for key, label, n in zip(reporting.STATUS_KEYS, labels, row["counts"], strict=True):
            yield (f"pr.section.{section}.{key}", f"{name} · {label}", figure, n, IR, "")
        yield (f"pr.section.{section}.total", f"{name} · {total}", figure, row["total"], IR, "")
    for key, label, n in zip(reporting.STATUS_KEYS, labels, table["totals"], strict=True):
        yield (f"pr.section_total.{key}", f"{total} · {label}", figure, n, IR, "")
    yield ("pr.section_total.total", total, figure, table["total"], IR, "")


def _by_month(r: Read) -> Iterator[Line]:
    chart = r.data["page"]["by_month"]
    keys = {label: key for key, label in rc.STATUS_LABELS.items()}
    figure = "Progress reports → Indicator reports by month of the reporting period's end"
    for i, month in enumerate(chart["labels"]):
        try:
            code = datetime.datetime.strptime(month, "%b %Y").strftime("%Y-%m")
        except ValueError:
            code = "none"  # "No period end"
        for label, values in chart["series"].items():
            yield (f"pr.month.{code}.{keys[label]}", f"{month} · {label}", figure, values[i], IR, "")


def _report_list(r: Read) -> Iterator[Line]:
    n = r.data["page"]["counts"]["progress_reports"]
    note = "The Progress reports sheet lists them, one row each."
    yield (
        "pr.reports_listed",
        POWERBI_ITEMS["reports"],
        "Progress reports → Progress reports table",
        n,
        PR,
        note,
    )


def _partner_cards(r: Read) -> Iterator[Line]:
    p, figure = r.data["page"]["partners"], "Progress reports → Partners"
    note = "The Partners sheet lists them, one row each."
    yield ("pr.partners", POWERBI_ITEMS["partners"], figure, p["partners"], "partners", note)
    note = "Submitted or Accepted."
    yield (
        "pr.partners_reported",
        POWERBI_ITEMS["partners_reported"],
        figure,
        p["reported"],
        "partners",
        note,
    )
    item = POWERBI_ITEMS["partners_not_reported"]
    yield ("pr.partners_not_reported", item, figure, p["not_reported"], "partners", "")
    for s in p["by_status"]:
        item = POWERBI_ITEMS["partners_status"].format(status=s["label"])
        yield (f"pr.partners_status.{s['key']}", item, figure, s["partners"], "partners", "")


def _follow_up(r: Read) -> Iterator[Line]:
    n = len(r.data["page"]["partners"]["follow_up"])
    figure, note = (
        "Progress reports → Partners: reports to follow up",
        "Overdue, Sent back, Due or Submitted.",
    )
    yield ("pr.follow_up", POWERBI_ITEMS["reports"], figure, n, PR, note)


def _too_many_note(found: dict[str, Any]) -> str:
    return (
        f"Not drawn: the filter holds {found['too_many']:,} report rows, more than {found['limit']:,}. "
        "Choose a year, a section or a partner, in NeuroDB and in Power BI."
    )


def _locations(r: Read) -> Iterator[Line]:
    found = r.data["locations"]
    value, note = (None, _too_many_note(found)) if found["too_many"] else (len(found["rows"]), "")
    yield (
        "pr.location_lines",
        POWERBI_ITEMS["lines"],
        "Progress reports → By location",
        value,
        "lines",
        note,
    )


def _disaggregations(r: Read) -> Iterator[Line]:
    found, figure = r.data["disaggregations"], "Progress reports → By disaggregation"
    value, note = (None, _too_many_note(found)) if found["too_many"] else (len(found["blocks"]), "")
    yield ("pr.disaggregation_locations", POWERBI_ITEMS["locations"], figure, value, "locations", note)


def _pp_indicators(r: Read) -> Iterator[Line]:
    n, figure = len(r.data["program_performance"]), "Partner monitoring → Program Performance"
    note = (
        f"Every indicator of the filter: the page lists the first {monitoring.MAX_INDICATORS:,}; the Program "
        "Performance rows sheet holds them all."
    )
    yield ("pp.indicators", POWERBI_ITEMS["indicators"], figure, n, "indicators", note)


def _pp_sections(r: Read) -> Iterator[Line]:
    figure = "Partner monitoring → Program Performance: section tiles"
    note = "Every section the other filters keep, chosen or not."
    for name, n in sorted(r.data["section_tiles"].items()):
        shown = NO_SECTION if name == monitoring.NO_SECTION else name
        yield (f"pp.section_tile.{name}", shown, figure, n, "indicators", note)
        note = ""


def _dq_tiles(r: Read) -> Iterator[Line]:
    q, figure = r.data["quality"], "Partner monitoring → Data quality"
    in_year = f" in {r.filters.year}" if r.filters.year else ""
    pds, items = "programme documents", POWERBI_ITEMS
    yield ("dq.pds", items["pds"], figure, q["programme_documents"], pds, PDS_NOTE)
    yield ("dq.running", items["running_pds"], figure, q["running_pds"], pds, "Active, signed or suspended.")
    yield ("dq.active", items["active_pds"], figure, q["active_pds"], pds, "")
    yield (
        "dq.ended",
        items["ended_pds"],
        figure,
        q["ended_pds"],
        pds,
        "Ended, closed, terminated or expired.",
    )
    yield ("dq.indicators", items["dq_indicators"], figure, q["indicators"], "indicators", "")
    note = "A counted report (Accepted or Submitted) carries a value."
    yield ("dq.with_reached", items["with_reached"] + in_year, figure, q["with_reached"], "indicators", note)
    yield ("dq.zero_reached", items["zero_reached"] + in_year, figure, q["zero_reached"], "indicators", "")
    item = items["without_reached"] + in_year
    yield ("dq.without_reached", item, figure, q["without_reached"], "indicators", OVERDUE_NOTE)
    note = "Without Reached and Reported zero together, as Power BI's card counts them."
    no_reached = q["without_reached"] + q["zero_reached"]
    yield ("dq.no_reached", items["no_reached"] + in_year, figure, no_reached, "indicators", note)
    note = "A target of 0 counts as none."
    yield ("dq.with_target", items["with_target"], figure, q["with_target"], "indicators", note)
    yield ("dq.without_target", items["without_target"], figure, q["without_target"], "indicators", "")


# a sheet row that is no figure of the page, and the figure whose trail it opens
TRAIL_OF = {"dq.no_reached": "dq.without_reached"}
BLOCKS = {
    "status_cards": _status_cards,
    "by_section": _by_section,
    "by_month": _by_month,
    "report_list": _report_list,
    "partner_cards": _partner_cards,
    "follow_up": _follow_up,
    "locations": _locations,
    "disaggregations": _disaggregations,
    "pp_indicators": _pp_indicators,
    "pp_sections": _pp_sections,
    "dq_tiles": _dq_tiles,
}


def visual_name(page: str, block: str) -> tuple[str, str]:
    """The Power BI page and visual a block of figures is compared with (:data:`POWERBI_PAGES`)."""
    return next(
        (name, visual) for name, visuals in POWERBI_PAGES[page] for b, visual in visuals if b == block
    )


def powerbi_page(block: str) -> str:
    """The Power BI page a block of figures is on, as :data:`POWERBI_PAGES` names it."""
    return next(
        name
        for found in POWERBI_PAGES.values()
        for name, visuals in found
        for b, _visual in visuals
        if b == block
    )


def figures(found: Read) -> list[Figure]:
    """The page's figures with a Power BI counterpart, in Power BI's order (:data:`POWERBI_PAGES`)."""
    out = []
    for powerbi_page, visuals in POWERBI_PAGES[found.page]:
        for block, visual in visuals:
            for key, item, neurodb, value, unit, note in BLOCKS[block](found):
                out.append(Figure(key, powerbi_page, visual, item, neurodb, value, unit, note))
    return out


# ------------------------------------------------------------------------------------------------- the filter
def unlisted(page: str) -> list[str]:
    """The page's filter fields that are neither a Power BI slicer nor a NeuroDB filter to clear (none may
    be)."""
    known = set(POWERBI_SLICERS[page]) | set(NEURODB_ONLY[page])
    return [f.name for cls in FILTERS[page] for f in dataclasses.fields(cls) if f.name not in known]


def _default(item: dataclasses.Field) -> Any:
    if item.default_factory is not dataclasses.MISSING:
        return item.default_factory()
    return item.default


def _partner_names(ids: Iterable[str]) -> dict[str, str]:
    wanted = [int(i) for i in ids if str(i).isdigit()]
    if not wanted:
        return {}
    found = PartnerOrganization.objects.filter(pk__in=wanted).values_list("pk", "name")
    return {str(pk): name for pk, name in found}


def _words(page: str, name: str, value: Any, filters: Any, partners: dict[str, str]) -> str:
    """A filter field's value in plain words ("" for none)."""
    if name == "pd_statuses":
        chosen = filters.status_choice if page == "progress-reports" else filters.statuses_read
        if not chosen:
            return "Every status (NeuroDB never reads draft or cancelled programme documents)"
        return ", ".join(pd_statuses.label(s) for s in chosen)
    if name == "year":
        return str(value) if value else "Every year"
    if name == "statuses":
        return ", ".join(rc.STATUS_LABELS.get(s, s) for s in value)
    if name == "partners":
        return ", ".join(partners.get(str(p), f"partner {p}") for p in value)
    if name == "emergency":
        return {"yes": "Yes", "no": "No"}.get(value, "")
    if name == "scope":  # the programme documents kept (monitoring._pd_queryset): a PD status chosen wins
        chosen = bool(filters.pd_statuses)
        if value == "year" and filters.year:
            kept = "of the PD statuses chosen" if chosen else "closed ones too"
            return f"Every programme document running in {filters.year}, {kept}"
        if chosen:
            return "Every programme document of the PD statuses chosen, of every year"
        return "Every programme document, of every status and year"
    if name == "report_type":
        words = monitoring.REPORT_TYPES.get(value, value)
        return words + ": only emergency or cluster indicators" if value == "HR" else words
    if name == "status":
        return monitoring.LABELS.get(value, value)
    if name == "tags":
        return "; ".join(f"{TAG_WORDS[tag]}: {', '.join(values)}" for tag, values in value.items() if values)
    if isinstance(value, bool):
        return "Yes" if value else ""
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, list | tuple):
        return ", ".join(str(v) for v in value)
    return str(value or "")


def settings_rows(found: Read) -> tuple[list[tuple[str, str]], list[tuple[str, str]]]:
    """The two blocks of the filter: the Power BI slicers to set, each with its value (every slicer, also one
    left at every value), and the NeuroDB filters to clear before comparing (those set only). Raises
    ValueError for a filter field in neither list."""
    page = found.page
    missing = unlisted(page)
    if missing:
        raise ValueError(f"filter fields in neither POWERBI_SLICERS nor NEURODB_ONLY: {', '.join(missing)}")
    values: dict[str, tuple[Any, Any, Any]] = {}  # name -> (value, its default, the filter)
    for instance in (found.filters, found.tab_filters):
        for item in dataclasses.fields(instance) if instance is not None else ():
            values[item.name] = (getattr(instance, item.name), _default(item), instance)
    partners = _partner_names(getattr(found.filters, "partners", []))
    slicers = []
    for name, slicer in POWERBI_SLICERS[page].items():
        value, _default_value, instance = values[name]
        slicers.append((slicer, _words(page, name, value, instance, partners) or "Every value: clear it"))
    clear = []
    for name, label in NEURODB_ONLY[page].items():
        value, default, instance = values[name]
        if value == default or (isinstance(value, dict) and not any(value.values())):
            continue
        clear.append((label, _words(page, name, value, instance, partners)))
    return slicers, clear


def _clean(text: str, names: frozenset[str]) -> str:
    """A text of the filter cleaned by the people rule of Raw eTools data (e-mails, phone numbers and known
    person names withheld), as the figure trail's CSVs clean the filter's words."""
    from . import raw_people

    return raw_people.text(text, "partner_reports", "partner_name", names, frozenset())


def names_someone(params: QueryDict, names: frozenset[str]) -> bool:
    """Whether a value of the filter names a person or a contact (a search can name anyone): the sheet then
    writes none of the filter's addresses."""
    return any(_clean(value, names) != value for _key, values in params.lists() for value in values)


# ----------------------------------------------------------------------------------------------- the workbook
class Formula(str):
    """A formula NeuroDB writes (only the difference columns hold one)."""

    __slots__ = ()


class Bold(str):
    """A text written in bold."""

    __slots__ = ()


INPUT = object()  # an empty cell a person types into (yellow)


class Sheet:
    """A write-only sheet, row by row, that knows the number of the row it writes (for the formulas)."""

    def __init__(self, book: Workbook, title: str, widths: Iterable[float], freeze: str = "") -> None:
        self.sheet = book.create_sheet(title=title)
        for i, width in enumerate(widths):
            self.sheet.column_dimensions[chr(ord("A") + i)].width = width
        if freeze:  # a list: its note and header stay in view
            self.sheet.freeze_panes = freeze
        self.row = 0

    def cell(self, value: Any) -> Any:
        if value is INPUT:
            found = WriteOnlyCell(self.sheet, value=None)
            found.fill = YELLOW
        elif isinstance(value, str) and not isinstance(value, Formula):
            # a text from the data is never run as a formula, and loses what Excel refuses (control codes)
            found = WriteOnlyCell(self.sheet, value=ILLEGAL_CHARACTERS_RE.sub("", value)[:CELL_CHARS])
            found.data_type = "s"
            if isinstance(value, Bold):
                found.font = BOLD
        else:
            found = WriteOnlyCell(self.sheet, value=value)
        return found

    def add(self, *values: Any) -> int:
        self.row += 1
        self.sheet.append([self.cell(v) for v in values])
        return self.row

    def title(self, text: str) -> None:
        self.row += 1
        found = self.cell(text)
        found.font = TITLE
        self.sheet.append([found])

    def header(self, *names: str) -> int:
        return self.add(*(Bold(n) for n in names))


def _day(day: datetime.date | datetime.datetime | None) -> str:
    """A day (and a moment's time) in words: 21 Oct 2026, 21 Oct 2026 18:13."""
    if day is None:
        return "—"
    if isinstance(day, datetime.datetime):
        day = timezone.localtime(day) if timezone.is_aware(day) else day
        return f"{day.day} {day:%b %Y %H:%M}"
    return f"{day.day} {day:%b %Y}"


def _runs(page: str) -> list[tuple[str, Any]]:
    """The syncs the page's figures come from: (their dataset in words, the run its records come from)."""
    names = {"partner_reports": "partner reports"}
    if page == "partner-monitoring":
        names["pd_indicators"] = "PD indicators"
    found = runfacts.runs(list(names))
    return [(words, found[name][0]) for name, words in names.items()]


def moment(page: str) -> datetime.datetime | None:
    """The moment of the eTools figures the sheet reads: the end of the latest of its sources' syncs."""
    ends = [run.finished_at for _name, run in _runs(page) if run is not None and run.finished_at]
    return max(ends) if ends else None


def _history(day: datetime.date) -> tuple[datetime.date, str]:
    """The first reporting period NeuroDB holds (``reached.history_from``), with why, in words."""
    held = rc.history_from(day)
    if held.source == rc.FEED:
        why = f"the Datamart has none earlier; NeuroDB asked from {_day(held.asked)}"
    elif held.source == rc.NEXT_SYNC:
        why = "no sync yet: the next one reads from this date"
    else:
        why = "the last sync read from this date"
    return held.date, f"{_day(held.date)} ({why})"


def _compare_sheet(book: Workbook, found: Read, base_url: str) -> None:
    from neurodb.watch import people

    page, day = found.page, found.day
    names = people.known_names()
    withheld = names_someone(found.params, names)
    sheet = Sheet(book, "Compare", (5, 26, 28, 34, 46, 14, 18, 16, 12, 10, 60, 30))
    sheet.title(f"NeuroDB figures for Power BI: {TITLES[page]}")
    path = reverse("reports:partner_reporting" if page == "progress-reports" else "reports:pd_monitoring")
    query = found.params.urlencode()
    address = base_url + path + (f"?{query}" if query else "")
    sheet.add(None, Bold("NeuroDB filter"), ADDRESS_WITHHELD if withheld else address)
    if day == page_today(page):
        reads = (
            "The page reads today, as this file does: for the day of Power BI's refresh, choose it beside "
        )
        reads += "the file's link."
    else:
        reads = f"The page reads today; this file reads {_day(day)}, the day chosen (Power BI's refresh)."
    reads += " Due, Overdue and Not yet due depend on the day."
    sheet.add(None, Bold("Statuses read for"), _day(day), reads)
    for words, run in _runs(page):
        sheet.add(None, Bold("NeuroDB data as of"), f"the last {words} sync: {_day(runfacts.moment(run))}")
    newest = dm.ReportedIndicator.objects.aggregate(newest=Max("last_modify_date"))["newest"]
    words = "the newest change in the partner report rows in NeuroDB (PRP's last change)"
    sheet.add(None, None, f"{words}: {_day(newest)}")
    sheet.add(None, Bold("Power BI refreshed on"), INPUT, "Type the day of Power BI's refresh.")
    sheet.add()
    _everything(sheet, found)
    sheet.add()
    slicers, clear = settings_rows(found)
    sheet.add(None, Bold("Set these in Power BI"))
    sheet.header("", "Power BI slicer", "Value", "Done?")
    for slicer, words in slicers:
        sheet.add(None, slicer, _clean(words, names), INPUT)
    for slicer in POWERBI_ONLY_SLICERS:
        sheet.add(None, slicer, ONE_COUNTRY.format(country=settings.ETOOLS_DATAMART_COUNTRY), INPUT)
    sheet.add()
    if POWERBI_TO_CLEAR[page]:
        sheet.add(None, Bold(CLEAR_IN_POWERBI))
        sheet.header("", "Power BI slicer", "Done?")
        for slicer in POWERBI_TO_CLEAR[page]:
            sheet.add(None, slicer, INPUT)
        sheet.add()
    sheet.add(None, Bold("Clear these in NeuroDB before comparing"))
    if clear:
        sheet.header("", "NeuroDB filter", "Value")
        for label, words in clear:
            sheet.add(None, label, _clean(words, names))
    else:
        sheet.add(None, "Nothing to clear: Power BI has a slicer for every filter set.")
    sheet.add()
    sheet.header(
        "#",
        "Power BI page",
        "Visual",
        "Name in Power BI",
        "NeuroDB page → figure",
        "NeuroDB value",
        "Unit",
        "Power BI value",
        "Difference",
        "Same?",
        "Notes",
        "Trail",
    )
    for n, figure in enumerate(figures(found), start=1):
        row = sheet.row + 1
        numeric = isinstance(figure.value, int | float)
        difference = Formula(f'=IF(ISNUMBER(H{row}),H{row}-F{row},"")') if numeric else None
        same = Formula(f'=IF(ISNUMBER(H{row}),IF(H{row}=F{row},"Same","Differs"),"")') if numeric else None
        # a sheet row that is no figure of the page opens the figure it is made from
        trail = trail_link(found, base_url, TRAIL_OF.get(figure.key, figure.key))
        if trail and withheld:
            trail = TRAIL_WITHHELD
        sheet.add(
            n,
            figure.powerbi_page,
            figure.visual,
            figure.item,
            figure.neurodb,
            figure.value,
            figure.unit,
            INPUT,
            difference,
            same,
            figure.note,
            trail,
        )
    sheet.add()
    sheet.add(None, Bold("Notes"))
    sheet.add(None, "The grid, the Map and Home have no Power BI page: check them against Raw eTools data.")
    sheet.add(
        None,
        f"{POWERBI_PD_PAGE} is on Programmes → Indicators: check it there against the PD indicators (Raw "
        "eTools data → PD indicators, filtered by the programme document).",
    )
    sheet.add(None, "Why NeuroDB and Power BI can differ: the sheet Why they can differ.")
    sheet.add(None, COMPARE_FOR_YOU)


def trail_link(found: Read, base_url: str, key: str) -> str | None:
    """The address of a figure on its figure trail, for the sheet's filter and day (the tab or the view it is
    drawn on, so By location's, By disaggregation's and Program Performance's open); None for a page with no
    trail."""
    from . import trail

    if found.page not in trail.PAGES:
        return None
    query = found.params.copy()
    if found.day != page_today(found.page):
        query["day"] = found.day.isoformat()
    if found.page == "partner-monitoring":
        # the year the sheet read, written out: the view chosen below opens on another year by default
        query["year"] = found.filters.year_param
        view = {"pp.": "table", "dq.": "quality"}.get(key[:3], "")
        if view:
            query["view"] = view
        else:
            query.pop("view", None)
    tabs = {"pr.location": "locations", "pr.disaggregation": "disaggregations"}
    tab = next((t for prefix, t in tabs.items() if key.startswith(prefix)), "")
    if not tab and key.startswith(("pr.partner", "pr.follow_up")):
        tab = "partners"
    if tab:
        query["tab"] = tab
    query["figure"] = key
    path = reverse("reports:check_figures", args=[found.page])
    return f"{base_url}{path}?{query.urlencode()}#{trail.anchor(key)}"


def _everything(sheet: Sheet, found: Read) -> None:
    """Everything NeuroDB holds, whatever the page's filter."""
    from . import services

    every = services.partner_reporting({"pd_status": ""}, found.day, options=False)["counts"]
    first, history = _history(found.day)
    sheet.add(
        None, Bold("Everything NeuroDB holds"), "whatever the page's filter: every PD status and every year"
    )
    sheet.add(None, "Indicator reports", every["indicator_reports"])
    sheet.add(None, "Progress reports", every["progress_reports"])
    if found.page == "partner-monitoring":
        held = PCA.objects.exclude(status__in=pd_statuses.NOT_A_PARTNERSHIP).count()
        sheet.add(None, "Programme documents", held)
    sheet.add(None, "Earliest reporting period", history)
    year = POWERBI_SLICERS[found.page]["year"]
    sheet.add(
        None,
        f"In Power BI, clear every slicer except {year} ≥ {first.year} to compare these. NeuroDB never reads "
        "draft or cancelled programme documents, nor reports naming no programme document.",
    )


def _reports_sheet(book: Workbook, found: Read) -> None:
    """One row per progress report of the filter, as the Progress reports export writes it (without its focal
    points, people cleaned as the trail's CSV of the same rows cleans them), with the key to match the rows of
    Power BI's report list (``report_list``)."""
    from neurodb.reports.views import _report_row
    from neurodb.watch import people

    from .trail_rows import REPORT_TEXTS, cleaned

    names = people.known_names()
    reports = found.data["page"]["reports"]
    vendors = _vendor_numbers({r["partner_id"] for r in reports})
    sheet = Sheet(book, "Progress reports", (18, 36, 14, 10, 10, 12, 12, 12, 14, 12, 26), freeze="A3")
    powerbi_page, visual = visual_name(found.page, "report_list")
    sheet.add(
        f"Every progress report of the filter: {len(reports):,}. Match key: PD reference, a space, report #. "
        f"In Power BI, {powerbi_page} → {visual}: Export data, then one VLOOKUP on the match key."
    )
    sheet.header(
        "PD reference",
        "Partner",
        "Vendor number",
        "Report #",
        "Type",
        "Period end",
        "Due",
        "Submitted",
        "Status",
        "Indicator reports",
        "Match key",
    )
    for report in reports:
        row = cleaned(_report_row(report), REPORT_TEXTS, "partner_reports", names)
        sheet.add(
            row["programme_document"],
            row["partner"],
            vendors.get(report["partner_id"], ""),
            row["report"],
            row["report_type"],
            row["period_end"],
            row["due_date"],
            row["submission_date"],
            row["report_status"],
            row["indicator_reports"],
            f"{row['programme_document']} {row['report']}",
        )


def _vendor_numbers(ids: set[Any]) -> dict[int, str]:
    wanted = [i for i in ids if i]
    if not wanted:
        return {}
    return dict(PartnerOrganization.objects.filter(pk__in=wanted).values_list("pk", "vendor_number"))


def partner_rows(reports: list[dict[str, Any]], names: frozenset[str]) -> tuple[list[dict[str, Any]], int]:
    """One row per partner of the reports, as the Partners tab counts them (``reporting.partner_key``): its
    name (a person named as a partner withheld, as the trail's CSVs withhold it), vendor number, progress
    reports per status, and the match key (the vendor number, else the name); and the number of reports naming
    no partner (in no Partners figure)."""
    found: dict[Any, dict[str, Any]] = {}
    nameless = 0
    for r in reports:
        key = reporting.partner_key(r)
        if key is None:
            nameless += 1
            continue
        if key not in found:
            name = _clean(r["partner_name"] or "", names)
            found[key] = {"partner_id": r["partner_id"], "name": name}
        row = found[key]
        row.setdefault("counts", Counter())[r["status"]] += 1
    vendors = _vendor_numbers({row["partner_id"] for row in found.values()})
    rows = []
    for row in found.values():
        vendor = vendors.get(row["partner_id"]) or ""
        rows.append({**row, "vendor_number": vendor, "match_key": vendor or row["name"]})
    rows.sort(key=lambda row: row["name"].lower())
    return rows, nameless


def _partners_sheet(book: Workbook, found: Read) -> None:
    from neurodb.watch import people

    rows, nameless = partner_rows(found.data["page"]["reports"], people.known_names())
    labels = [rc.STATUS_LABELS[k] for k in reporting.STATUS_KEYS]
    sheet = Sheet(book, "Partners", (36, 14, *([12] * len(labels)), 10, 26), freeze="A3")
    note = (
        f"Every partner with a progress report in the filter: {len(rows):,}. Match key: the vendor number, "
    )
    note += "else the partner's name."
    if nameless:
        note += f" {nameless:,} progress reports name no partner: they are in no Partners figure."
    sheet.add(note)
    sheet.header("Partner", "Vendor number", *labels, "Total", "Match key")
    for row in rows:
        counts = [row["counts"].get(k, 0) for k in reporting.STATUS_KEYS]
        sheet.add(row["name"], row["vendor_number"], *counts, sum(counts), row["match_key"])


def _match_key(row: dict[str, Any]) -> str:
    return f"{row['programme_document']} | {row['indicator']} | {row['pd_output']}"


def _program_performance_sheet(book: Workbook, found: Read, base_url: str) -> None:
    """Every indicator of Program Performance, as its export writes it (people cleaned as the trail's
    indicators CSV cleans them), with the key to match Power BI's rows. Two indicators of one programme
    document can share their title and PD output (check 7): their keys then end in #1, #2..., and a note says
    so; a row whose title or PD output lost a name, an e-mail or a phone says it is matched by hand."""
    from neurodb.watch import people

    from . import trail_monitoring as tm
    from .trail_rows import INDICATOR_TEXTS, cleaned

    names = people.known_names()
    filters, year = found.filters, found.filters.year
    ordered, _sort = monitoring.sorted_rows(found.data["program_performance"], "", year)
    widths = (22, 30, 44, 30, 18, 32, 12, 14, 12, 14, 12, 12, 12, 60, 30, 40)
    sheet = Sheet(book, "Program Performance rows", widths, freeze="A3")
    sheet.add(
        f"Every indicator of the filter: {len(ordered):,} (the page lists the first "
        f"{monitoring.MAX_INDICATORS:,}). Match key: PD number | indicator | PD output."
    )
    sheet.header(
        "Section",
        "CP output",
        "Indicator",
        "PD output",
        "PD",
        "Partner",
        "Target",
        f"NeuroDB Reached in {year}" if year else "NeuroDB Reached",
        "NeuroDB % Reached",
        "Power BI Reached",
        "Power BI %",
        "Difference Reached",
        "Difference %",
        "Match key",
        "Trail",
        "Note",
    )
    written = list(monitoring.table_rows(ordered, filters, base_url))
    rows = [cleaned(row, INDICATOR_TEXTS, tm.SOURCE, names) for row in written]
    shared = Counter(_match_key(row) for row in rows)
    seen: Counter[str] = Counter()
    withheld = names_someone(found.params, names)
    for indicator, row, before in zip(ordered, rows, written, strict=True):
        key, notes = _match_key(row), []
        trail = trail_link(found, base_url, f"pp.reached.{indicator.pd.id}.{indicator.key}")
        if trail and withheld:
            trail = TRAIL_WITHHELD
        if shared[key] > 1:
            seen[key] += 1
            notes.append(
                f"{shared[key]} indicators of this PD share its title and PD output: match them by target."
            )
            key = f"{key} #{seen[key]}"
        if (row["indicator"], row["pd_output"]) != (before["indicator"], before["pd_output"]):
            notes.append(KEY_WITHHELD)
        note = " ".join(notes)
        n = sheet.row + 1
        sheet.add(
            row["section"],
            row["cp_output"],
            row["indicator"],
            row["pd_output"],
            row["programme_document"],
            row["partner"],
            row["target"],
            row["reached_in_year"] if year else row["reached"],
            row["achieved_percent_in_year"] if year else row["achieved_percent"],
            INPUT,
            INPUT,
            Formula(f'=IF(ISNUMBER(J{n}),J{n}-H{n},"")'),
            Formula(f'=IF(ISNUMBER(K{n}),K{n}-I{n},"")'),
            key,
            trail,  # the indicator's Reached, step by step
            note,
        )


def _why_sheet(book: Workbook, found: Read) -> None:
    first, _words = _history(found.day)
    sheet = Sheet(book, "Why they can differ", (4, 100, 70))
    sheet.header("#", "Why NeuroDB and Power BI can differ", "Where to see it")
    pages = {key: powerbi_page(block) for key, block in WHY_PAGES.items()}
    for n, (why, where) in enumerate(WHY_THEY_DIFFER, start=1):
        sheet.add(n, why.format(history=_day(first), **pages), where.format(**pages))


def workbook(page: str, params: Any, day: datetime.date, base_url: str = "") -> bytes:
    """The comparison sheet of ``page`` (one of :data:`PAGES`) for the filter ``params`` on ``day``, as the
    bytes of an Excel file: Compare, then the lists (Progress reports and Partners; Program Performance rows),
    then Why they can differ. Read in one read-only snapshot of the data."""
    from .raw import STREAM_TIMEOUT_MS, snapshot

    if page not in PAGES:
        raise ValueError(f"no comparison sheet for the page {page!r}")
    book = Workbook(write_only=True)
    with snapshot(STREAM_TIMEOUT_MS):
        found = read(page, params, day)
        _compare_sheet(book, found, base_url)
        if page == "progress-reports":
            _reports_sheet(book, found)
            _partners_sheet(book, found)
        else:
            _program_performance_sheet(book, found, base_url)
        _why_sheet(book, found)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def filename(page: str) -> str:
    """``powerbi-comparison-<page>_<YYYY-MM-DD>_<HHMM>.xlsx``, the moment of the eTools figures it reads."""
    from neurodb.reports.exports import export_filename

    return export_filename(f"powerbi-comparison-{page}", "xlsx", moment=moment(page))
