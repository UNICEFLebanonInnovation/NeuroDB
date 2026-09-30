"""Country programme: the cycle and its documents, the results framework (Excel, AI proposal), the
progress rule, the linked sources and the dashboard."""

import datetime
import io
import json
from decimal import Decimal
from types import SimpleNamespace

import pytest
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from neurodb.cpd import excel, extraction, services, sources
from neurodb.cpd.models import (
    CountryProgramme,
    CPDocument,
    FrameworkProposal,
    Indicator,
    Link,
    Milestone,
    Origin,
    Outcome,
    Output,
    Value,
)
from neurodb.datamart import models as dm
from neurodb.datamart import monitoring
from neurodb.partnerships.models import PCA, PartnerOrganization

TODAY = datetime.date(2027, 7, 2)  # half of a 2026-2028 cycle


@pytest.fixture
def media(settings, tmp_path):
    settings.STORAGES = {
        **settings.STORAGES,
        "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    }
    settings.MEDIA_ROOT = tmp_path
    return tmp_path


@pytest.fixture
def cycle(db):
    return CountryProgramme.objects.create(
        name="Lebanon CP 2026-2028", start_year=2026, end_year=2028, current=True, etools_name="LEB CP 2026"
    )


@pytest.fixture
def framework(cycle):
    outcome = Outcome.objects.create(programme=cycle, code="1", title="Children learn")
    output = Output.objects.create(outcome=outcome, code="1.1", title="Access to learning")
    rate = Indicator.objects.create(
        outcome=outcome,
        code="1.a",
        title="Out-of-school rate",
        unit="percent",
        direction="decrease",
        baseline=40,
        baseline_year=2025,
        target=20,
    )
    reached = Indicator.objects.create(
        output=output, code="1.1.1", title="Children enrolled", accumulate="sum", baseline=0, target=30000
    )
    return SimpleNamespace(cycle=cycle, outcome=outcome, output=output, rate=rate, reached=reached)


def _progress(indicator, values=None, today=TODAY):
    indicator = Indicator.objects.prefetch_related("values", "milestones", "links").get(pk=indicator.pk)
    return services.progress(
        indicator, values or {}, today, services.elapsed_share(indicator.programme, today)
    )


# ------------------------------------------------------------------------------------ models
def test_a_cycle_ends_after_it_starts_and_lasts_seven_years_at_most(db):
    with pytest.raises(ValidationError):
        CountryProgramme(name="x", start_year=2028, end_year=2026).full_clean()
    with pytest.raises(ValidationError):
        CountryProgramme(name="x", start_year=2020, end_year=2027).full_clean()
    CountryProgramme(name="x", start_year=2026, end_year=2027).full_clean()  # 2 years


def test_one_cycle_is_current(cycle):
    other = CountryProgramme.objects.create(name="Next", start_year=2029, end_year=2033, current=True)
    cycle.refresh_from_db()
    assert other.current and not cycle.current
    assert other.years == [2029, 2030, 2031, 2032, 2033]


def test_an_indicator_belongs_to_one_outcome_or_one_output(framework):
    both = Indicator(programme=framework.cycle, outcome=framework.outcome, output=framework.output, title="x")
    with pytest.raises(ValidationError):
        both.clean()
    assert framework.reached.programme_id == framework.cycle.pk  # taken from the output


def test_documents_accept_office_files_and_pdfs_only(cycle, media):
    doc = CPDocument(programme=cycle, title="CPD", file=SimpleUploadedFile("cpd.exe", b"x"))
    with pytest.raises(ValidationError):
        doc.full_clean()
    doc = CPDocument(
        programme=cycle, title="CPD", file=SimpleUploadedFile("Lebanon CPD (final).pdf", b"%PDF")
    )
    doc.full_clean()
    doc.save()
    assert doc.file.name.startswith(f"cpd/{cycle.pk}/") and doc.filename.endswith("-Lebanon-CPD--final-.pdf")


# ---------------------------------------------------------------------------------- progress
def test_elapsed_share_runs_from_the_first_january_to_the_last_december(cycle):
    assert services.elapsed_share(cycle, datetime.date(2025, 6, 1)) == 0
    assert services.elapsed_share(cycle, datetime.date(2029, 1, 1)) == 100
    assert 49 < services.elapsed_share(cycle, TODAY) < 51


def test_a_decreasing_rate_on_track_off_track_and_ahead(framework):
    rate = framework.rate  # 40 -> 20, half the cycle elapsed: 30 expected
    Value.objects.create(indicator=rate, year=2027, value=31, source="MEHE")
    p = _progress(rate)
    assert (p.value, p.value_year, p.status) == (31, 2027, services.ON_TRACK)
    assert p.achieved == 45.0
    Value.objects.filter(indicator=rate).update(value=38)
    assert _progress(rate).status == services.OFF_TRACK
    Value.objects.filter(indicator=rate).update(value=19)
    p = _progress(rate)
    assert p.status == services.AHEAD and p.label == "Target reached"


def test_a_milestone_sets_what_is_expected_this_year(framework):
    rate = framework.rate
    Milestone.objects.create(indicator=rate, year=2027, value=36)  # 20 % of the way expected
    Value.objects.create(indicator=rate, year=2027, value=36, source="MEHE")
    p = _progress(rate)
    assert (p.expected, p.status, p.milestone) == (20.0, services.ON_TRACK, 36)


def test_sum_indicators_add_up_the_years_and_levels_take_the_latest(framework):
    reached = framework.reached
    Value.objects.create(indicator=reached, year=2026, value=10000, source="Compiler")
    Value.objects.create(indicator=reached, year=2027, value=5000, source="Compiler")
    p = _progress(reached)
    assert p.value == 15000 and p.status == services.ON_TRACK
    reached.accumulate = Indicator.Accumulate.LATEST
    reached.save()
    assert _progress(reached).value == 5000


def test_a_typed_value_replaces_the_linked_sources_of_its_year(framework):
    reached = framework.reached
    a = Link.objects.create(
        indicator=reached,
        kind="education",
        year=2026,
        education_programme="mscc",
        source_period="2026",
        label="Makani 2026",
    )
    b = Link.objects.create(
        indicator=reached,
        kind="education",
        year=2027,
        education_programme="mscc",
        source_period="2027",
        label="Makani 2027",
    )
    Link.objects.create(
        indicator=reached,
        kind="education",
        year=2027,
        education_programme="bridging",
        source_period="2027",
        label="Dirasa 2027",
        confirmed=False,
    )
    Value.objects.create(indicator=reached, year=2026, value=9000, source="annual report")
    p = _progress(reached, {a.pk: 1000.0, b.pk: 4000.0})
    assert p.years[2026]["value"] == 9000 and p.years[2026]["manual"]
    assert p.years[2027]["value"] == 4000 and len(p.years[2027]["parts"]) == 1  # unconfirmed left out
    assert p.value == 13000


def test_no_target_and_no_data(framework):
    framework.rate.target = None
    framework.rate.save()
    assert _progress(framework.rate).status == services.NO_TARGET
    assert _progress(framework.reached).status == services.NO_DATA


# ---------------------------------------------------------------------------- interventions
def test_an_output_matches_its_etools_name_or_its_code(framework):
    output = framework.output
    assert services.output_matches(output, "1.1 Access to learning for all")
    assert services.output_matches(output, "Output 1.1: Access")
    assert not services.output_matches(output, "1.12 Something else")
    output.etools_output = "Children access quality learning"
    assert services.output_matches(output, "children access quality learning.")
    assert not services.output_matches(output, "1.1 Access to learning for all")


@pytest.fixture
def intervention(framework):
    partner = PartnerOrganization.objects.create(etl_id="1", name="Learning Partner", partner_type="CSO")
    pd = PCA.objects.create(
        etl_id="11",
        partner=partner,
        partner_name=partner.name,
        number="LEB/PCA2027001",
        title="Learning",
        status="active",
        start=datetime.date(2027, 1, 1),
        end=datetime.date(2027, 12, 31),
        country_programme="LEB CP 2026",
        cp_outputs=["1.1 Access to learning"],
    )
    PCA.objects.create(
        etl_id="12",
        partner=partner,
        partner_name=partner.name,
        number="LEB/PCA2019001",
        title="Old",
        status="active",
        start=datetime.date(2019, 1, 1),
        end=datetime.date(2020, 12, 31),
        country_programme="LEB CP 2026",
        cp_outputs=["1.1 Access to learning"],
    )
    dm.PDIndicator.objects.create(
        datamart_id=1,
        source_id=501,
        intervention=pd,
        pd_reference_number=pd.number,
        title="# of children enrolled in non-formal education",
        section_name="Education",
        lower_result_name="1.1 Access",
        target_numerator=Decimal("10000"),
        display_type="number",
        location_name="Akkar",
    )
    dm.ReportedIndicator.objects.create(
        datamart_id=77,
        partner=partner,
        intervention=pd,
        partner_name=partner.name,
        pd_reference_number=pd.number,
        progress_report="PR-1",
        report_number="QPR1",
        report_type="QPR",
        report_status="Accepted",
        period_start=datetime.date(2027, 1, 1),
        period_end=datetime.date(2027, 3, 31),
        indicator="# of children enrolled in non-formal education",
        target="10000",
        location="Akkar",
        achievement_in_period="6000",
        total_cumulative_progress="6000",
        total_cumulative_progress_in_location="6000",
    )
    return pd


def test_interventions_run_in_the_cycle_under_its_country_programme(framework, intervention):
    assert [pd.number for pd in services.interventions(framework.cycle)] == ["LEB/PCA2027001"]
    framework.cycle.etools_name = "Another CP"
    assert services.interventions(framework.cycle) == []


def test_an_etools_link_reads_the_partners_cumulative_progress(framework, intervention):
    row = monitoring.indicators(monitoring.Filters(pd_ids=[intervention.pk], scope="all", year=2027))[0]
    link = sources.apply(Link(indicator=framework.reached, year=2027), f"etools:{intervention.pk}:{row.key}")
    link.full_clean()
    link.save()
    assert link.kind == "etools" and "children enrolled" in link.label
    assert services.link_values([link]) == {link.pk: 6000.0}
    assert sources.value_of(link) == f"etools:{intervention.pk}:{row.key}"


def test_source_choices_list_the_pd_indicators_of_the_output(framework, intervention):
    groups = dict(sources.choices(framework.reached))
    options = groups["eTools PD indicators (partner reporting)"]
    assert options and options[0][0].startswith(f"etools:{intervention.pk}:")


def test_an_activityinfo_link_reads_the_master_indicator(framework, hierarchy):
    master = hierarchy["master"]
    value = f"ai:{master.pk}"
    assert sources.default_year(value) == 2026
    link = sources.apply(Link(indicator=framework.reached, year=2026), value)
    link.save()
    assert services.link_values([link])[link.pk] == 500.0
    assert value in {v for _g, opts in sources.choices(framework.reached) for v, _l in opts}


def test_an_unknown_source_is_refused(framework):
    with pytest.raises(ValueError):
        sources.apply(Link(indicator=framework.reached, year=2026), "etools:999999:1")
    with pytest.raises(ValueError):
        sources.apply(Link(indicator=framework.reached, year=2026), "nope:1")


def test_education_total_reads_the_cube_or_the_blocks():
    cube = {
        "cubes": {
            "enrolment": {
                "dims": ["partner"],
                "measures": ["registrations", "x"],
                "rows": [["A", 10, 1], ["B", 5, 1]],
            }
        }
    }
    assert services.education_total(cube) == 15
    blocks = {
        "blocks": {
            "registrations": {"figures": [{"by": ["partner"], "rows": [["A", 3]]}, {"by": [], "rows": [[7]]}]}
        }
    }
    assert services.education_total(blocks) == 7


# ------------------------------------------------------------------------------------- excel
def test_the_template_round_trips_through_the_import(framework):
    Milestone.objects.create(indicator=framework.rate, year=2027, value=30)
    data = excel.template(framework.cycle)
    ws = load_workbook(io.BytesIO(data))[excel.SHEET]
    header = [c.value for c in ws[1]]
    assert header[-3:] == ["Milestone 2026", "Milestone 2027", "Milestone 2028"]
    assert [row[0] for row in ws.iter_rows(min_row=2, values_only=True)] == [
        "Outcome",
        "Indicator",
        "Output",
        "Indicator",
    ]
    result = excel.import_framework(framework.cycle, data)
    assert (result.created, result.updated, result.errors) == (0, 4, [])
    rate = Indicator.objects.get(pk=framework.rate.pk)
    assert (rate.direction, rate.unit, rate.target, rate.origin) == (
        "decrease",
        "percent",
        20,
        Origin.IMPORTED,
    )
    assert rate.milestones.get(year=2027).value == 30


def _sheet(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = excel.SHEET
    ws.append([label for _k, label, _w in excel.COLUMNS] + ["Milestone 2027"])
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_the_import_adds_a_framework_from_a_filled_sheet(cycle):
    data = _sheet(
        [
            ["Outcome", 2, "Children are protected"],
            [
                "Output",
                2.1,
                "Case management",
                2,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                "2.1 CM services",
            ],
            [
                "Indicator",
                "2.1.1",
                "Children receiving case management",
                2.1,
                "Number",
                "Increase",
                "sum",
                0,
                2025,
                9000,
                "eTools",
                None,
                3000,
            ],
            ["Indicator", None, "Share of cases closed", 2, "%", "Increase", "latest", "50%", 2025, "80%"],
        ]
    )
    result = excel.import_framework(cycle, data)
    assert (result.created, result.errors) == (4, [])
    ind = Indicator.objects.get(code="2.1.1")
    assert ind.output.code == "2.1" and ind.output.etools_output == "2.1 CM services"
    assert ind.milestones.get().value == 3000
    share = Indicator.objects.get(title="Share of cases closed")
    assert share.outcome.code == "2" and (share.baseline, share.target, share.unit) == (50, 80, "percent")


def test_a_wrong_row_saves_nothing(cycle):
    data = _sheet(
        [
            ["Outcome", 1, "Learning"],
            ["Indicator", "1.x", "Rate", 9, "Number"],  # parent 9 does not exist
            ["Indicator", "1.y", "Rate", 1, "Number", "Increase", "latest", "abc"],
            ["Level?", 1, "x"],
        ]
    )
    result = excel.import_framework(cycle, data)
    assert len(result.errors) == 3 and "Row 3" in result.errors[0]
    assert not Outcome.objects.exists()
    assert excel.import_framework(cycle, b"not excel").errors == [
        "The file is not an Excel workbook (.xlsx)."
    ]


# -------------------------------------------------------------------------------- AI proposal
PROPOSED = {
    "outcomes": [
        {
            "code": "1",
            "title": "Children learn",
            "indicators": [
                {
                    "code": "1.a",
                    "title": "Out-of-school rate",
                    "unit": "percent",
                    "direction": "decrease",
                    "baseline": 40,
                    "baseline_year": 2025,
                    "target": 20,
                    "means_of_verification": "MEHE",
                }
            ],
            "outputs": [
                {
                    "code": "1.1",
                    "title": "Access",
                    "indicators": [
                        {
                            "code": "1.1.1",
                            "title": "Children enrolled",
                            "unit": "number",
                            "direction": "increase",
                            "baseline": 0,
                            "baseline_year": None,
                            "target": 30000,
                            "means_of_verification": "",
                        }
                    ],
                },
                {"code": "1.2", "title": "Quality", "indicators": []},
            ],
        }
    ]
}


class FakeResponses:
    def __init__(self, text):
        self.text, self.calls = text, []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(output_text=self.text)


@pytest.fixture
def pdf(cycle, media):
    return CPDocument.objects.create(
        programme=cycle, title="CPD", file=SimpleUploadedFile("cpd.pdf", b"%PDF-1.4")
    )


def test_the_proposal_sends_the_pdf_with_a_strict_schema(pdf, monkeypatch):
    fake = SimpleNamespace(responses=FakeResponses(json.dumps(PROPOSED)))
    monkeypatch.setattr("neurodb.assistant.agent.client", lambda: fake)
    proposal = extraction.run(FrameworkProposal.objects.create(document=pdf))
    assert proposal.status == "ready" and proposal.items == PROPOSED
    call = fake.responses.calls[0]
    assert call["store"] is False and call["text"]["format"]["strict"] is True
    assert call["input"][0]["content"][0]["file_data"].startswith("data:application/pdf;base64,")


def test_a_failed_proposal_says_why(pdf, monkeypatch):
    fake = SimpleNamespace(responses=FakeResponses("not json"))
    monkeypatch.setattr("neurodb.assistant.agent.client", lambda: fake)
    proposal = extraction.run(FrameworkProposal.objects.create(document=pdf))
    assert proposal.status == "failed" and "could not be read" in proposal.error


def test_only_pdfs_are_read(cycle, media):
    doc = CPDocument.objects.create(programme=cycle, title="RRF", file=SimpleUploadedFile("rrf.xlsx", b"x"))
    with pytest.raises(extraction.ExtractionError):
        extraction.extract(doc)


def test_applying_keeps_the_ticked_items_with_their_parents(pdf):
    proposal = FrameworkProposal.objects.create(document=pdf, status="ready", items=PROPOSED)
    keys = [row["key"] for row in extraction.flatten(PROPOSED)]
    assert keys == ["o0", "o0i0", "o0p0", "o0p0i0", "o0p1"]
    made = extraction.apply(proposal, {"o0", "o0p0i0", "o0p1"})  # o0p0 not ticked: its indicator is dropped
    assert made == {"created": 2, "skipped": 0}
    assert list(Output.objects.values_list("code", flat=True)) == ["1.2"]
    assert Outcome.objects.get().origin == Origin.AI
    proposal.refresh_from_db()
    assert proposal.status == "applied" and proposal.applied_at


# --------------------------------------------------------------------------------- admin
def test_admin_downloads_the_template_and_imports_a_sheet(client, admin_user, cycle):
    client.force_login(admin_user)
    response = client.get(reverse("admin:cpd_countryprogramme_template", args=[cycle.pk]))
    assert response.status_code == 200 and "results-framework.xlsx" in response["Content-Disposition"]
    upload = SimpleUploadedFile("f.xlsx", _sheet([["Outcome", 1, "Learning"]]))
    response = client.post(reverse("admin:cpd_countryprogramme_import", args=[cycle.pk]), {"file": upload})
    assert response.status_code == 302 and Outcome.objects.filter(programme=cycle, code="1").exists()
    bad = SimpleUploadedFile("f.xlsx", _sheet([["Output", 1, "x", 7]]))
    response = client.post(reverse("admin:cpd_countryprogramme_import", args=[cycle.pk]), {"file": bad})
    assert response.status_code == 200 and b"Row 2" in response.content


def test_admin_pages_open(client, admin_user, framework, pdf):
    client.force_login(admin_user)
    for name, obj in [
        ("countryprogramme", framework.cycle),
        ("cpdocument", pdf),
        ("outcome", framework.outcome),
        ("output", framework.output),
        ("indicator", framework.reached),
    ]:
        assert client.get(reverse(f"admin:cpd_{name}_change", args=[obj.pk])).status_code == 200
        assert client.get(reverse(f"admin:cpd_{name}_changelist")).status_code == 200


def test_admin_links_an_indicator_to_a_source(client, admin_user, framework, hierarchy):
    client.force_login(admin_user)
    ind = framework.reached
    url = reverse("admin:cpd_indicator_change", args=[ind.pk])
    data = {
        "programme": framework.cycle.pk,
        "output": framework.output.pk,
        "code": ind.code,
        "title": ind.title,
        "unit": "number",
        "direction": "increase",
        "accumulate": "sum",
        "baseline": "0",
        "target": "30000",
        "means_of_verification": "",
        "origin": "manual",
        "milestones-TOTAL_FORMS": "0",
        "milestones-INITIAL_FORMS": "0",
        "values-TOTAL_FORMS": "0",
        "values-INITIAL_FORMS": "0",
        "links-TOTAL_FORMS": "1",
        "links-INITIAL_FORMS": "0",
        "links-0-source": f"ai:{hierarchy['master'].pk}",
        "links-0-year": "",
        "links-0-confirmed": "on",
    }
    response = client.post(url, data)
    assert response.status_code == 302
    link = Link.objects.get(indicator=ind)
    assert (link.kind, link.year, link.master_id) == ("activityinfo", 2026, hierarchy["master"].pk)


def test_admin_reviews_and_applies_a_proposal(client, admin_user, pdf):
    client.force_login(admin_user)
    proposal = FrameworkProposal.objects.create(document=pdf, status="ready", items=PROPOSED)
    url = reverse("admin:cpd_frameworkproposal_review", args=[proposal.pk])
    response = client.get(url)
    assert response.status_code == 200 and b"Out-of-school rate" in response.content
    response = client.post(url, {"keep": ["o0", "o0i0"]})
    assert response.status_code == 302
    assert Indicator.objects.get().origin == Origin.AI


def test_the_admin_action_starts_the_proposal_in_the_background(client, admin_user, pdf, monkeypatch):
    started = []
    monkeypatch.setattr("neurodb.integrations.background.start_command", lambda *a: started.append(a))
    client.force_login(admin_user)
    client.post(
        reverse("admin:cpd_cpdocument_changelist"),
        {"action": "propose_framework", "_selected_action": [pdf.pk]},
    )
    proposal = FrameworkProposal.objects.get()
    assert started == [("propose_cpd_framework", "--proposal", str(proposal.pk))]


# ----------------------------------------------------------------------------------- pages
def test_the_dashboard_shows_where_the_cycle_stands(client_viewer, framework, intervention, pdf):
    Value.objects.create(indicator=framework.rate, year=2026, value=38, source="MEHE")
    response = client_viewer.get(reverse("cpd:dashboard"))
    assert response.status_code == 200
    body = response.content.decode()
    data = response.context["data"]
    assert data["indicator_total"] == 2 and data["pd_total"] == 1
    assert data["outcomes"][0]["outputs"][0]["pds"][0].number == "LEB/PCA2027001"
    assert (
        "Out-of-school rate" in body
        and "LEB/PCA2027001" in body
        and reverse("cpd:document", args=[pdf.pk]) in body
    )


def test_the_dashboard_without_a_cycle_says_what_to_do(client_viewer):
    response = client_viewer.get(reverse("cpd:dashboard"))
    assert response.status_code == 200 and b"No country programme cycle yet" in response.content


def test_the_indicator_page_shows_each_year(client_viewer, framework):
    Value.objects.create(indicator=framework.reached, year=2026, value=12000, source="Compiler")
    response = client_viewer.get(reverse("cpd:indicator", args=[framework.reached.pk]))
    assert response.status_code == 200
    assert [r["year"] for r in response.context["rows"]] == [2026, 2027, 2028]
    assert b"12,000" in response.content


def test_documents_download_for_signed_in_users_only(client, viewer, pdf):
    url = reverse("cpd:document", args=[pdf.pk])
    assert client.get(url).status_code == 302  # to sign in
    client.force_login(viewer)
    response = client.get(url)
    assert response.status_code == 200 and b"".join(response.streaming_content) == b"%PDF-1.4"
