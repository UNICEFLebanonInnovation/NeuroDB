"""Release 2, step 4, stage D1: the exports of Monitoring insights (FMS §13). The Export menu of the page
(D1.1); the Excel workbook with FMS's column names, typed cells and totals that equal the page's (D1.2);
the printable report (D1.3); the Power BI package (D1.4); the Power BI live feed, read with a key only,
throttled, and its keys in the admin (D1.5); and the action points' Excel export and printable report
(D1.6). Since Release 2 step 5 (stage F3b) the main table is the records, one row per record as FMS's
export, the visits stay one row per visit and the rule results are per record. No file, feed or report
holds a person: never the team, the visit lead, an e-mail address or who an action point is assigned
to."""

from __future__ import annotations

import csv
import datetime
import io
import zipfile
from decimal import Decimal

import pytest
from django.db import connection
from django.test import Client, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from neurodb.datamart import models as dm
from neurodb.fmm import exports, metrics, powerbi
from neurodb.fmm.models import PowerBIKey, RecordRuleResult, Visit, VisitEntity
from neurodb.fmm.scope import Scope

from .conftest import CANARIES, LEAD, MEMBER, MEMBER_EMAIL

pytestmark = pytest.mark.django_db
YEAR = {"year": "2026", "section": ""}
ASSIGNEE = "Rania Canary"
ASSIGNEE_EMAIL = "rania.canary@example.org"
PERSON_TEXTS = (*CANARIES, ASSIGNEE_EMAIL)


# ------------------------------------------------------------------------------------------ helpers
@pytest.fixture
def admin_client(client, admin_user):
    client.force_login(admin_user)
    return client


@pytest.fixture
def assigned(built):
    """An FM action point of visit 1722 assigned to a person, its description naming them."""
    dm.ActionPoint.objects.filter(datamart_id=8001).update(
        assigned_to_name=ASSIGNEE,
        description=f"Ask {ASSIGNEE} ({ASSIGNEE_EMAIL}) to send the registers; call 03-123456.",
        reference_number="AP/2026/8001",
    )
    return built


def _scope(**params) -> Scope:
    return Scope.from_params({**YEAR, **params})


def _book(response):
    assert response.status_code == 200, response.content[:300]
    return load_workbook(io.BytesIO(response.content))


def _sheet_rows(book, name: str) -> list[dict]:
    rows = list(book[name].iter_rows(values_only=True))
    header = rows[0]
    return [dict(zip(header, row, strict=True)) for row in rows[1:]]


def _all_text(book) -> str:
    return "\n".join(
        str(value) for sheet in book.worksheets for row in sheet.iter_rows(values_only=True) for value in row
    )


def _csv_rows(text: str) -> list[dict]:
    assert text.startswith("﻿")
    return list(csv.DictReader(io.StringIO(text[1:])))


def _assert_no_person(text: str) -> None:
    for canary in (*PERSON_TEXTS, LEAD, MEMBER, MEMBER_EMAIL):
        assert canary not in text, canary


def _key() -> tuple[PowerBIKey, str]:
    return powerbi.create_key("Office workspace")


# ------------------------------------------------------------------------------------------ D1.1 menu
def test_the_page_header_has_the_export_menu_for_the_filter(built, client_viewer):
    html = client_viewer.get(reverse("fmm:dashboard"), {**YEAR, "governorate": "north"}).content.decode()
    query = _scope(governorate="north").query.replace("&", "&amp;")
    for name in ("fmm:export_xlsx", "fmm:report", "fmm:export_powerbi"):
        assert f'href="{reverse(name)}?{query}" data-current-query=""' in html, name
    assert (
        f"{reverse('fmm:visits')}?{query}&amp;export=csv" in html
        and 'data-current-query="export=csv"' in html
    )
    for label in ("Excel workbook", "PDF report", "Power BI package", "CSV (records list)"):
        assert label in html
    assert "Power BI live connection" not in html  # Administrators only
    # the table toolbar keeps its own CSV and the charts their PNG buttons
    assert "All rows (CSV)" in client_viewer.get(reverse("fmm:dashboard"), {**YEAR, "tab": "visits"}).text


def test_administrators_see_the_live_connection(built, admin_client):
    html = admin_client.get(reverse("fmm:dashboard"), YEAR).content.decode()
    assert "Power BI live connection" in html and reverse("admin:fmm_powerbikey_changelist") in html


def test_the_menu_links_follow_the_address_bar():
    from django.conf import settings

    app = (settings.BASE_DIR / "neurodb" / "web" / "static" / "js" / "app.js").read_text()
    assert "a[data-current-query]" in app and "window.location.search" in app
    assert 'dataset.autoprint === "load"' in app
    charts = (settings.BASE_DIR / "neurodb" / "web" / "static" / "js" / "charts.js").read_text()
    assert "staticPlot: true" in charts and "el.dataset.width" in charts


# ------------------------------------------------------------------------------------------ D1.2 workbook
def test_the_workbook_has_its_sheets_and_fms_columns(built, client_viewer):
    response = client_viewer.get(reverse("fmm:export_xlsx"), YEAR)
    assert response["Content-Disposition"] == (
        f'attachment; filename="monitoring-insights-{timezone.localdate().isoformat()}.xlsx"'
    )
    book = _book(response)
    assert book.sheetnames == [
        "About",
        "Records",
        "Visits",
        "Rule results",
        "Partners",
        "Field offices",
        "Sections",
        "Flags",
        "Action points",
    ]
    header = next(book["Visits"].iter_rows(values_only=True))
    assert list(header) == exports.visit_columns(exports.Context.read().weights)
    for column in (
        "id",
        "country_name",
        "monitoring_activity_id",
        "overall_finding_rating",
        "quality_status",
        "completeness_score",
        "_evidence_score",
        "_alignment_score",
        "_coherence_score",
        "_q3_quality_score",
        "_actionability_score",
        "location_lat",
        "narrative_finding",
        "hact_q3_answer",
        "action_points_assigned",
        "_ai_used",
        "neurodb_url",
    ):
        assert column in header, column
    assert "action_points_assigned_to" not in header
    visits = {row["monitoring_activity_id"]: row for row in _sheet_rows(book, "Visits")}
    assert len(visits) == metrics.kpis(_scope())["visits"] == 8
    v1722 = visits[1722]
    assert v1722["country_name"] == "Lebanon"
    # numbers are numbers, dates are dates
    assert isinstance(v1722["monitoring_activity_end_date"], datetime.datetime)
    assert v1722["monitoring_activity_end_date"].date() == datetime.date(2026, 5, 12)
    stored = Visit.objects.get(activity_id=1722)
    assert v1722["quality_score"] == float(stored.quality_score) and v1722["quality_status"] == "High"
    assert v1722["urgency"] == stored.urgency and isinstance(v1722["urgency"], int)
    assert v1722["entity"] == "Amel Association" and v1722["vendor_number"] == "2500212345"
    assert v1722["neurodb_url"] == reverse("fmm:visit", args=[stored.key])
    assert v1722["_ai_used"] is True
    # texts: the narratives cleaned, the 412 children kept; the HACT answer of the row, or the checklist's
    assert (
        "412 children" in v1722["narrative_finding"]
        and "Classes held as planned" in v1722["narrative_finding"]
    )
    assert "Registers checked." in v1722["hact_q3_answer"]
    assert visits[1727]["hact_q1_answer"] == "Constrained"  # an option's code written as its label
    # the rating as the page counts it: the in-progress visit is not rated yet, the cancelled one blank
    assert visits[1724]["overall_finding_rating"] == "Not rated yet"
    assert visits[1724]["quality_status"] == "Skipped" and visits[1724]["quality_score"] is None
    assert visits[1722]["overall_finding_rating"] == "Off track"
    # a visit's figures from its records: the mean, its lowest record, its records and how many are scored
    assert (v1722["records"], v1722["records_scored"]) == (3, stored.records_scored) == (3, 3)
    assert v1722["lowest_score"] == float(stored.lowest_score)


def test_the_records_sheet_has_one_row_per_record_with_its_own_columns(built, client_viewer):
    """FMS §13.2 per record: one row per record of the filter, ``id`` e<record id> (unique), the visit's
    id and activity id repeated on each of its records, and the record's own entity, type, rating,
    score, flags, urgency, narrative and place; the visit's dates and action points repeated."""
    book = _book(client_viewer.get(reverse("fmm:export_xlsx"), YEAR))
    header = next(book["Records"].iter_rows(values_only=True))
    assert list(header) == exports.record_columns(exports.Context.read().weights)
    assert header[:2] == ("id", "visit_id") and "action_points_assigned_to" not in header
    rows = _sheet_rows(book, "Records")
    records = list(_scope().records().select_related("visit"))
    assert len(rows) == len(records) == metrics.kpis(_scope())["records"] == 12
    assert len({r["id"] for r in rows}) == len(rows)
    assert {r["id"] for r in rows} == {f"e{e.datamart_id}" for e in records}
    by_id = {r["id"]: r for r in rows}
    assert [r["monitoring_activity_id"] for r in rows].count(1722) == 3  # repeated on each record
    visit = Visit.objects.get(activity_id=1722)
    for n in (101, 102, 103):
        row, e = by_id[f"e{n}"], VisitEntity.objects.get(datamart_id=n)
        assert row["visit_id"] == visit.key and row["monitoring_activity_end_date"].date() == visit.end_date
        assert row["quality_score"] == float(e.quality_score) and row["urgency"] == e.urgency
        assert (
            row["quality_flags"] == "; ".join(e.flags) and row["action_points_count"] == visit.action_points
        )
        assert row["neurodb_url"] == f"{reverse('fmm:visit', args=[visit.key])}#record-{n}"
    # each record's own rating and narrative (the visit's is its worst record's, Off track)
    assert (by_id["e101"]["overall_finding_rating"], by_id["e102"]["overall_finding_rating"]) == (
        "On track",
        "Off track",
    )
    assert by_id["e103"]["narrative_finding"] == "The partner keeps its registers up to date."
    assert "Classes held as planned" in by_id["e101"]["narrative_finding"]
    assert "registers" not in by_id["e101"]["narrative_finding"]
    assert (by_id["e101"]["entity_type"], by_id["e102"]["entity_type"], by_id["e103"]["entity_type"]) == (
        "PD/SSFA",
        "CP output",
        "Partner",
    )
    assert by_id["e103"]["entity"] == "Amel Association" and by_id["e103"]["vendor_number"] == "2500212345"
    # the in-progress visit's record is not rated yet and not scored
    in_progress = next(r for r in rows if r["monitoring_activity_id"] == 1724)
    assert (in_progress["overall_finding_rating"], in_progress["quality_status"]) == (
        "Not rated yet",
        "Skipped",
    )
    # the record's own place when eTools gave one, else its visit's
    assert by_id["e171"]["location_name"] == Visit.objects.get(activity_id=1728).place_name


def test_the_workbook_never_holds_a_person(assigned, client_viewer):
    book = _book(client_viewer.get(reverse("fmm:export_xlsx"), YEAR))
    text = _all_text(book)
    _assert_no_person(text)
    points = _sheet_rows(book, "Action points")
    point = next(p for p in points if p["reference"] == "AP/2026/8001")
    assert "Ask" in point["description"] and "send the registers" in point["description"]
    assert "assigned" not in " ".join(book["Action points"][1][i].value or "" for i in range(12)).lower()
    visit = next(r for r in _sheet_rows(book, "Visits") if r["monitoring_activity_id"] == 1722)
    assert "send the registers" in visit["action_points_text"]  # cleaned, the person removed


def test_the_workbook_totals_equal_the_page(built, client_viewer):
    """§0 "Figures agree": the Partners, Field offices, Sections and Flags sheets against the page's own
    blocks, and the visits against the key figures, for two filters."""
    for params in ({}, {"governorate": "bekaa"}):
        scope = _scope(**params)
        book = _book(client_viewer.get(reverse("fmm:export_xlsx"), {**YEAR, **params}))
        assert len(_sheet_rows(book, "Visits")) == metrics.kpis(scope)["visits"]
        records = _sheet_rows(book, "Records")
        assert len(records) == metrics.kpis(scope)["records"]
        # the visits of the Visits sheet are exactly the distinct visits of the records
        assert {r["visit_id"] for r in records} == {r["id"] for r in _sheet_rows(book, "Visits")}
        offices = {r["name"]: r for r in metrics.offices(scope)["rows"]}
        compared = 0
        for row in _sheet_rows(book, "Field offices"):
            if row["field_office"] in offices:
                page = offices[row["field_office"]]
                assert row["visits"] == page["visits"]
                assert row["avg_quality_score"] == (float(page["avg"]) if page["avg"] is not None else None)
                compared += 1
        assert compared and compared == len(offices)  # every office of the page is in the sheet
        sections = {r["name"]: r for r in metrics.sections(scope)}
        for row in _sheet_rows(book, "Sections"):
            name = "none" if row["section"] == "No section" else row["section"]
            assert row["visits"] == sections[name]["visits"]
            assert row["on_track"] + row["constrained"] + row["off_track"] == row["rated"]
        flags = {r["code"]: r for r in metrics.flag_frequency(scope)["rows"]}
        for row in _sheet_rows(book, "Flags"):
            assert row["flagged_records"] == flags[row["rule_id"]]["n"]
            assert row["evaluated_records"] == flags[row["rule_id"]]["evaluated"]
        partners = {r["key"]: r for r in metrics.breakdown(scope, "partner")}
        assert sum(row["records"] for row in _sheet_rows(book, "Partners")) == sum(
            r["records"] for r in partners.values()
        )
        about = {r["name"]: r["value"] for r in _sheet_rows(book, "About")}
        assert about["Visits"] == metrics.kpis(scope)["visits"]
        assert about["Records"] == metrics.kpis(scope)["records"] == len(records)
        assert "No names" in about["Privacy"]


def test_the_breakdown_counts_as_the_page_blocks(built):
    scope = _scope()
    offices = {r["name"]: r for r in metrics.offices(scope)["rows"]}
    for row in metrics.breakdown(scope, "office"):
        if row["key"] in offices:
            assert (row["visits"], row["avg"], row["ratings"]) == (
                offices[row["key"]]["visits"],
                offices[row["key"]]["avg"],
                offices[row["key"]]["ratings"],
            )
    sections = {r["name"]: r for r in metrics.sections(scope)}
    for row in metrics.breakdown(scope, "section"):
        assert (row["visits"], row["avg"]) == (sections[row["key"]]["visits"], sections[row["key"]]["avg"])
    partners = metrics.breakdown(scope, "partner")
    assert {r["name"] for r in partners} == {"AMEL", "MCL"}
    assert sum(r["visits"] for r in partners) == metrics.kpis(scope)["visits"]


def test_rule_results_keep_neurodb_wording_only(built):
    ctx = exports.Context.read()
    rows = list(exports.rule_rows(_scope().records(), ctx))
    assert rows and {r["result"] for r in rows} <= {"passed", "flagged", "not checked", "skipped"}
    # one row per record and rule, with the record and its visit
    assert len(rows) == RecordRuleResult.objects.filter(entity__in=_scope().records()).count()
    assert len({(r["record_id"], r["rule_id"]) for r in rows}) == len(rows)
    r101 = [r for r in rows if r["record_id"] == "e101"]
    assert r101 and {r["visit_id"] for r in r101} == {Visit.objects.get(activity_id=1722).key}
    ai = [r for r in rows if r["rule_id"] in ctx.ai_rules]
    assert ai and all(r["detail"] == "" for r in ai)  # an AI check's explanation may quote a narrative
    assert any(r["ai_used"] == "yes" for r in ai)
    flagged = [r for r in rows if r["result"] == "flagged" and r["rule_id"] not in ctx.ai_rules]
    assert flagged and all(isinstance(r["points_lost"], Decimal) for r in flagged)


def test_a_text_contains_flag_is_not_quoted(built):
    """The flag of a "text contains" check writes the whole text it read (a narrative): the Rule results
    sheet leaves its detail out, and keeps the detail of the check when it passed."""
    from neurodb.fmm.models import RuleSetting

    rule = RuleSetting.objects.filter(type=RuleSetting.Type.REFERENCE).order_by("code").first()
    rule.params = {**(rule.params or {}), "check_type": "string_contains", "contains": "gender"}
    rule.save(update_fields=["params"])
    record = VisitEntity.objects.get(datamart_id=101)
    quote = "The partner said the girls were kept at home during the exams"
    RecordRuleResult.objects.update_or_create(
        entity=record, rule=rule.code, defaults={"status": "fail", "detail": f"Not covered: {quote}"}
    )
    other = VisitEntity.objects.get(datamart_id=111)
    RecordRuleResult.objects.update_or_create(
        entity=other, rule=rule.code, defaults={"status": "pass", "detail": "gender is covered."}
    )
    rows = {(r["record_id"], r["rule_id"]): r for r in exports.rule_rows(_scope().records())}
    assert rows[("e101", rule.code)]["result"] == "flagged"
    assert rows[("e101", rule.code)]["detail"] == ""
    assert rows[("e111", rule.code)]["detail"] == "gender is covered."


def test_query_count_does_not_grow_with_rows(built):
    """A large scope is read in chunks: the queries of 8 visits (12 records) and of 48 (72) are the
    same."""
    ctx = exports.Context.read()

    def queries() -> tuple[int, int, int]:
        with CaptureQueriesContext(connection) as captured:
            rows = list(exports.visit_rows(Visit.objects.all(), ctx))
            records = list(exports.record_rows(VisitEntity.objects.all(), ctx))
            list(exports.rule_rows(VisitEntity.objects.all(), ctx))
            list(exports.action_point_rows(Visit.objects.all(), ctx))
        return len(captured), len(rows), len(records)

    few, n, n_records = queries()
    visits = list(Visit.objects.all())
    entities = list(VisitEntity.objects.all())
    for k in range(5):
        copies = []
        for v in visits:
            copy = Visit.objects.get(pk=v.pk)
            copy.pk, copy.key, copy.activity_id = (
                None,
                f"{v.key}-c{k}",
                (v.activity_id or 0) + 100_000 * (k + 1),
            )
            copies.append(copy)
        Visit.objects.bulk_create(copies)
        moved = {v.pk: c.pk for v, c in zip(visits, copies, strict=True)}
        rows = []
        for e in entities:
            row = VisitEntity.objects.get(pk=e.pk)
            row.pk, row.visit_id, row.datamart_id = None, moved[e.visit_id], e.datamart_id + 100_000 * (k + 1)
            rows.append(row)
        VisitEntity.objects.bulk_create(rows)
    many, m, m_records = queries()
    assert (n, m, n_records, m_records) == (8, 48, 12, 72) and many == few


def test_a_text_that_starts_with_equals_stays_a_text():
    from neurodb.reports.exports import xlsx_response

    response = xlsx_response(
        "x.xlsx",
        [
            (
                "S",
                ["a", "b", "c"],
                [{"a": "=HYPERLINK(1)", "b": datetime.date(2026, 1, 2), "c": Decimal("9.5")}],
            )
        ],
        typed=True,
    )
    book = load_workbook(io.BytesIO(response.content))
    cells = list(book["S"].iter_rows(min_row=2))[0]
    assert (cells[0].value, cells[0].data_type) == ("=HYPERLINK(1)", "s")
    assert cells[1].value == datetime.datetime(2026, 1, 2) and cells[2].value == 9.5
    assert book["S"].freeze_panes == "A2" and book["S"]["A1"].font.b
    # characters Excel refuses are dropped; a long title is cut to Excel's 31 characters
    from neurodb.reports import xlsx

    content = xlsx.workbook([("A title far longer than Excel allows: x/y", ["t"], [{"t": "a\x01b & <c>"}])])
    sheet = load_workbook(io.BytesIO(content)).worksheets[0]
    assert sheet["A2"].value == "ab & <c>" and len(sheet.title) <= 31 and "/" not in sheet.title


# ------------------------------------------------------------------------------------------ D1.3 report
def test_the_report_prints_the_page_figures(assigned, client_viewer):
    response = client_viewer.get(reverse("fmm:report"), YEAR)
    assert response.status_code == 200
    html = response.content.decode()
    assert 'data-autoprint="load"' in html and 'data-static="1"' in html and 'data-width="680"' in html
    for title in (
        "Key figures",
        "Morning briefing",
        "AI monitoring insights",
        "Overall finding rating",
        "Quality bands",
        "Ratings by section",
        "Ratings by field office",
        "Partners",
        "Most frequent flags",
        "The most urgent visits",
        "HACT programmatic visits",
        "Action points",
        "Method",
    ):
        assert title in html, title
    assert "Written by NeuroDB from the figures" in html  # the AI is off: the code-written brief, said so
    assert "Data as of" in html and "fmm-report-data" in html
    assert "Visit 1723" in html  # the most urgent visit
    _assert_no_person(html)


def test_the_report_and_exports_are_read_only_for_a_viewer(built, client_viewer):
    for name in ("fmm:report", "fmm:export_xlsx", "fmm:export_powerbi"):
        assert client_viewer.get(reverse(name), YEAR).status_code == 200
        assert client_viewer.post(reverse(name), YEAR).status_code == 405


def test_the_exports_answer_404_when_switched_off(built, client_viewer):
    with override_settings(FMM_ENABLED=False):
        for name in ("fmm:report", "fmm:export_xlsx", "fmm:export_powerbi"):
            assert client_viewer.get(reverse(name), YEAR).status_code == 404


# ------------------------------------------------------------------------------------------ D1.4 package
def test_the_powerbi_package(assigned, client_viewer):
    response = client_viewer.get(reverse("fmm:export_powerbi"), YEAR)
    assert response.status_code == 200 and response["Content-Type"] == "application/zip"
    assert (
        f"monitoring-insights-powerbi-{timezone.localdate().isoformat()}.zip"
        in response["Content-Disposition"]
    )
    package = zipfile.ZipFile(io.BytesIO(response.content))
    assert sorted(package.namelist()) == [
        "NeuroDB_monitoring.pq",
        "README.txt",
        "data/action_points.csv",
        "data/partners.csv",
        "data/records.csv",
        "data/rule_results.csv",
        "data/visits.csv",
    ]
    texts = {name: package.read(name).decode("utf-8") for name in package.namelist()}
    visits = _csv_rows(texts["data/visits.csv"])
    assert len(visits) == 8
    row = next(r for r in visits if r["monitoring_activity_id"] == "1722")
    assert row["monitoring_activity_end_date"] == "2026-05-12" and "." in row["quality_score"]
    assert row["_ai_used"] == "true"
    book = _book(client_viewer.get(reverse("fmm:export_xlsx"), YEAR))
    sheet = {r["id"]: r for r in _sheet_rows(book, "Visits")}
    assert {r["id"] for r in visits} == set(sheet)
    assert float(row["quality_score"]) == sheet[row["id"]]["quality_score"]
    partners = _csv_rows(texts["data/partners.csv"])
    # each record under its own partner: the records add up, a visit with two partners counts in both
    assert sum(int(p["records"]) for p in partners) == 12
    assert sum(int(p["visits"]) for p in partners) == sum(
        len({e.partner_id for e in v.entity_rows.all() if e.partner_id}) for v in _scope().visits()
    )
    records = _csv_rows(texts["data/records.csv"])
    assert len(records) == 12 and len({r["id"] for r in records}) == 12
    assert {r["visit_id"] for r in records} == {r["id"] for r in visits}
    assert set(_csv_rows(texts["data/rule_results.csv"])[0]) >= {"record_id", "visit_id", "rule_id"}
    script = texts["NeuroDB_monitoring.pq"]
    for part in (
        "RootFolder",
        'Load("records"',
        'Load("visits"',
        'Load("rule_results"',
        'Load("action_points"',
        'Load("partners"',
        '{"quality_score", type number}',
        '{"lowest_score", type number}',
        '{"records", Int64.Type}',
        '{"monitoring_activity_end_date", type date}',
        '{"urgency", Int64.Type}',
        "record_sections",
        "record_offices",
        "record_flags",
        '{{"id", "record_id"}, {column, name}}',
        "Encoding = 65001",
    ):
        assert part in script, part
    assert script.index('records = Load("records"') < script.index('visits = Load("visits"')
    readme = texts["README.txt"]
    assert "Changed in this release: records" in readme and "keeps working" in readme
    assert "data/records.csv" in readme and "records[id] to rule_results[record_id]" in readme
    assert (
        "Blank query" in readme
        and "Advanced editor" in readme
        and "no .pbit" in readme.lower().replace("there is no .pbit", "no .pbit")
    )
    assert "Recommended visuals" in readme
    _assert_no_person("\n".join(texts.values()))


# ------------------------------------------------------------------------------------------ D1.5 feed
FEED = "fmm_powerbi_feed"


def _feed(client, dataset="visits", key=None, header=None, **params):
    extra = {"HTTP_AUTHORIZATION": f"Bearer {header}"} if header else {}
    query = {**params, **({"key": key} if key else {})}
    response = client.get(reverse(FEED, args=[dataset]), query, **extra)
    body = b"".join(response.streaming_content) if response.streaming else response.content
    return response, body.decode("utf-8")


def test_no_key_at_all_answers_not_found(built, client):
    response, body = _feed(client)
    assert response.status_code == 404 and "Amel" not in body


def test_a_missing_wrong_or_revoked_key_gets_nothing(assigned, client):
    row, key = _key()
    for response, body in (_feed(client), _feed(client, key="not-the-key"), _feed(client, header="x" * 43)):
        assert response.status_code == 401 and "Amel" not in body and "1722" not in body
        assert response["Cache-Control"] == "no-store"
    powerbi.revoke(PowerBIKey.objects.filter(pk=row.pk))
    response, body = _feed(client, key=key)
    assert response.status_code == 404 and "1722" not in body  # no key left: the feed is not there
    _key()
    response, body = _feed(client, key=key)
    assert response.status_code == 401 and "1722" not in body
    row.refresh_from_db()
    assert row.uses == 0


ALL_TIME = {"preset": "all_time", "section": ""}


def test_a_good_key_reads_the_feed_by_header_or_query(assigned, client, viewer, caplog):
    row, key = _key()
    by_header, header_body = _feed(client, header=key)
    by_query, query_body = _feed(client, key=key)
    assert by_header.status_code == by_query.status_code == 200
    assert by_header["Cache-Control"] == "no-store" and by_header["Content-Type"].startswith("text/csv")
    assert header_body == query_body
    rows = _csv_rows(header_body)
    # the whole data: every visit, whatever its year, and the same totals as the workbook of all time
    assert (
        len(rows)
        == Visit.objects.count()
        == metrics.kpis(Scope.from_params({"preset": "all_time", "section": ""}))["visits"]
    )
    total = sum(float(r["quality_score"]) for r in rows if r["quality_score"])
    assert total == pytest.approx(
        float(sum(v.quality_score for v in Visit.objects.exclude(quality_score=None)))
    )
    signed_in = Client()  # the workbook of all time, beside the feed read without a session
    signed_in.force_login(viewer)
    book = _book(signed_in.get(reverse("fmm:export_xlsx"), ALL_TIME))
    sheet = {r["id"]: r for r in _sheet_rows(book, "Visits")}
    assert {r["id"] for r in rows} == set(sheet)
    for r in rows:
        assert (float(r["quality_score"]) if r["quality_score"] else None) == sheet[r["id"]]["quality_score"]
    row.refresh_from_db()
    assert row.uses == 2 and row.last_used_at is not None
    for dataset in ("records", "rule_results", "action_points", "partners"):
        response, body = _feed(client, dataset, key=key)
        assert response.status_code == 200 and _csv_rows(body)
        _assert_no_person(body)
    records = _csv_rows(_feed(client, "records", key=key)[1])
    assert len(records) == VisitEntity.objects.count() == len({r["id"] for r in records})
    _assert_no_person(header_body)
    assert key not in caplog.text
    assert _feed(client, "people", key=key)[0].status_code == 404


def test_the_feed_narrows_by_year_and_since(built, client):
    _row, key = _key()
    assert len(_csv_rows(_feed(client, key=key, year="2026")[1])) == 8
    assert _csv_rows(_feed(client, key=key, year="2025")[1]) == []
    # a record follows its visit's date
    assert len(_csv_rows(_feed(client, "records", key=key, year="2026")[1])) == 12
    assert _csv_rows(_feed(client, "records", key=key, year="2025")[1]) == []
    since = _csv_rows(_feed(client, key=key, since="2026-07-01")[1])
    # the visit date: its start date, else its end date (as every period reads it)
    dated = {
        (r["monitoring_activity_start_date"] or r["monitoring_activity_end_date"]) >= "2026-07-01"
        for r in since
    }
    assert dated == {True} and len(since) < 8
    assert _feed(client, key=key, year="twenty")[0].status_code == 400
    assert _feed(client, key=key, since="July")[0].status_code == 400


@override_settings(FMM_POWERBI_REQUESTS_PER_HOUR=2)
def test_the_feed_is_throttled_per_key(built, client):
    _row, key = _key()
    assert _feed(client, key=key)[0].status_code == 200
    assert _feed(client, key=key)[0].status_code == 200
    response, _body = _feed(client, key=key)
    assert response.status_code == 429 and 0 < int(response["Retry-After"]) <= 3600
    _other, other_key = _key()
    assert _feed(client, key=other_key)[0].status_code == 200  # each key its own count


@override_settings(FMM_POWERBI_REQUESTS_PER_HOUR=2)
def test_the_throttle_is_kept_in_the_database(built, client):
    """The count of the hour lives on the key's row, not in one worker's memory: emptying the cache (a
    second worker, another container) does not open the feed again; the next hour does."""
    from django.core.cache import cache

    row, key = _key()
    assert _feed(client, key=key)[0].status_code == 200
    cache.clear()
    assert _feed(client, key=key)[0].status_code == 200
    cache.clear()
    assert _feed(client, key=key)[0].status_code == 429
    row.refresh_from_db()
    assert row.hour_uses == 2 and row.uses == 2
    PowerBIKey.objects.filter(pk=row.pk).update(hour_started=row.hour_started - datetime.timedelta(hours=1))
    assert _feed(client, key=key)[0].status_code == 200
    row.refresh_from_db()
    assert row.hour_uses == 1 and row.uses == 3


def test_the_key_never_reaches_the_access_log_or_a_trace(monkeypatch):
    """Power BI sends the key in the address (?key=): gunicorn's access log writes the request line, and a
    request trace keeps the address. The log hides the key; the feed is not traced."""
    import logging
    import os
    import runpy
    import sys
    import types

    from django.conf import settings

    from config import gunicorn_filters, telemetry

    key = "s3cret-Key_value-0123456789"
    atoms = {
        "h": "10.0.0.1",
        "r": f"GET /powerbi/fmm/visits.csv?year=2026&key={key} HTTP/1.1",
        "q": f"key={key}&year=2026",
        "a": "Microsoft.Data.Mashup",
        "{authorization}i": f"Bearer {key}",
    }
    record = logging.LogRecord(
        "gunicorn.access", logging.INFO, "", 0, '%(h)s "%(r)s" %(q)s %({authorization}i)s', (atoms,), None
    )
    assert gunicorn_filters.HideKeys().filter(record)
    line = record.getMessage()
    assert key not in line and "year=2026" in line and "/powerbi/fmm/visits.csv" in line
    conf = runpy.run_path(str(settings.BASE_DIR / "config" / "gunicorn.conf.py"))["logconfig_dict"]
    assert conf["filters"]["hide_keys"]["()"] == "config.gunicorn_filters.HideKeys"
    assert "hide_keys" in conf["handlers"]["access"]["filters"]

    # Application Insights: the feed's addresses are left out of the traces, beside the operator's own
    calls = []
    stub = types.ModuleType("azure.monitor.opentelemetry")
    stub.configure_azure_monitor = lambda **kwargs: calls.append(kwargs)
    monkeypatch.setitem(sys.modules, "azure.monitor.opentelemetry", stub)
    monkeypatch.setattr(telemetry, "_configured", False)
    monkeypatch.setenv("APPLICATIONINSIGHTS_CONNECTION_STRING", "InstrumentationKey=00000000")
    monkeypatch.setenv("OTEL_PYTHON_EXCLUDED_URLS", "healthz")
    monkeypatch.setenv("OTEL_PYTHON_DJANGO_EXCLUDED_URLS", "set-by-the-test")  # restored afterwards
    monkeypatch.delenv("OTEL_PYTHON_DJANGO_EXCLUDED_URLS")
    monkeypatch.setenv("OTEL_SERVICE_NAME", "neurodb-test")
    monkeypatch.setenv("OTEL_RESOURCE_ATTRIBUTES", "service.version=test")
    telemetry.setup("web")
    assert calls and os.environ["OTEL_PYTHON_DJANGO_EXCLUDED_URLS"] == "healthz,powerbi/fmm/"


def test_a_session_alone_opens_nothing(built, client_viewer, client):
    _key()
    response, body = _feed(client_viewer)
    assert response.status_code == 401 and "Amel" not in body
    # only the feed is left out of the sign-in: the other pages still ask for it
    client.logout()
    for name in ("fmm:dashboard", "fmm:export_xlsx", "fmm:report", "fmm:export_powerbi"):
        assert client.get(reverse(name), YEAR).status_code == 302, name


def test_a_donor_session_does_not_change_the_feed(built, client):
    from neurodb.accounts.models import User
    from neurodb.donors.models import DonorAccount

    user = User.objects.create_user(
        username="donor-x", email="donor-x@example.org", password="donor-pass-12345"
    )
    DonorAccount.objects.create(user=user, name="European Union", donors=["EU"], must_change_password=False)
    client.force_login(user)
    _row, key = _key()
    assert _feed(client)[0].status_code == 401  # no redirect to the donor page, no data
    assert _feed(client, key=key)[0].status_code == 200
    assert client.get(reverse("fmm:export_xlsx"), YEAR).status_code == 302  # the rest stays locked


def test_only_the_hash_of_a_key_is_kept(built):
    row, key = _key()
    assert key not in (row.key_hash, row.prefix) and row.prefix == key[:8] and len(row.key_hash) == 64
    assert powerbi.find_key(key) == row and powerbi.find_key(key[:-1] + "x") is None


# ------------------------------------------------------------------------------------------ admin
def test_an_administrator_creates_a_key_seen_once_and_revokes_it(built, admin_client):
    response = admin_client.post(reverse("admin:fmm_powerbikey_add"), {"name": "Office workspace"})
    assert response.status_code == 200 and "no-store" in response["Cache-Control"]
    html = response.content.decode()
    row = PowerBIKey.objects.get()
    key = html.split('id="powerbi-key" type="text" readonly value="', 1)[1].split('"', 1)[0]
    assert powerbi.find_key(key) == row
    assert "ApiKeyName" in html and "Web.Contents" in html and "powerbi/fmm/" in html
    assert "powerbi/fmm/records.csv" in html and "Load(&quot;records&quot;" in html
    assert "manage.py" not in html
    change = admin_client.get(reverse("admin:fmm_powerbikey_change", args=[row.pk])).content.decode()
    assert key not in change and row.prefix in change and "ApiKeyName" in change
    listing = admin_client.get(reverse("admin:fmm_powerbikey_changelist")).content.decode()
    assert key not in listing and "Office workspace" in listing
    assert admin_client.get(reverse("admin:fmm_powerbikey_revoke", args=[row.pk])).status_code == 200
    admin_client.post(reverse("admin:fmm_powerbikey_revoke", args=[row.pk]))
    row.refresh_from_db()
    assert row.revoked_at is not None and powerbi.find_key(key) == row


def test_the_keys_are_for_administrators_only(built, client_viewer):
    assert client_viewer.get(reverse("admin:fmm_powerbikey_changelist")).status_code in (302, 403)


# ------------------------------------------------------------------------------------------ D1.6 action points
def test_action_points_excel_has_the_csv_columns_without_the_assignee(assigned, client_viewer):
    from neurodb.datamart import services as datamart
    from neurodb.reports import ap_views

    page = reverse("reports:action_points")
    html = client_viewer.get(page, {"module": "fm"}).content.decode()
    assert "Excel of the filter" in html and "Excel of every action point" in html and "PDF report" in html
    for params in ({"module": "fm", "export": "xlsx"}, {"export": "xlsx", "all": "1"}):
        book = _book(client_viewer.get(page, params))
        header = next(book.worksheets[0].iter_rows(values_only=True))
        assert list(header) == list(ap_views.XLSX_COLUMNS)
        assert "assigned_to" not in header and set(header) == set(datamart.AP_EXPORT_COLUMNS) - {
            "assigned_to"
        }
        text = _all_text(book)
        assert ASSIGNEE not in text and ASSIGNEE_EMAIL not in text and "send the registers" in text
    # the CSV keeps its established columns (who an action point is assigned to included)
    response = client_viewer.get(page, {"export": "csv", "all": "1"})
    body = b"".join(response.streaming_content).decode("utf-8")
    assert "assigned_to" in body.splitlines()[0]


def test_the_action_points_report(assigned, client_viewer):
    response = client_viewer.get(
        reverse("reports:action_points_report"), {"module": "fm", "assignee": "Rania"}
    )
    assert response.status_code == 200
    html = response.content.decode()
    assert 'data-autoprint="load"' in html and html.count('data-static="1"') == 6
    assert "Raised from: fm" in html and "Assigned to: filtered by a name" in html
    assert ASSIGNEE not in html and ASSIGNEE_EMAIL not in html
    assert "ap-report-data" in html
