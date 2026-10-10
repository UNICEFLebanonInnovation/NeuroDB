"""Hardening (stage 8b): one privacy canary, end to end. People's names, e-mail addresses, phone numbers
and a link are planted wherever eTools could write them (narratives, the visit lead, the team, checklist
answers short and long, summaries, other keys of the records), the visits are built, and then every path
by which Monitoring insights' data leaves NeuroDB or is kept is taken in turn:

- the AI brief: the request sent, and every field of the brief kept, with a model that writes the
  canaries back on purpose;
- the chat: every request sent over four look-ups, the answer shown, and the question's log, with a
  question and an answer naming people;
- Ask NeuroDB: the four ``fm_*`` look-ups without a chat filter, and the generic ``etools_*`` tools over
  every field monitoring dataset;
- the knowledge hub and What's new, the visits CSV, the NeuroDB Watch check and the log lines written;
- every text, character, list and JSON field of every table of Monitoring insights.

None of the canaries may come out anywhere, while the figures and references written next to them do.
Each path is also checked to have carried something, so that an empty result cannot pass."""

from __future__ import annotations

import datetime
import json
import logging

import pytest
from django.apps import apps
from django.db import models
from django.test import override_settings
from django.urls import reverse

from neurodb.datamart import models as dm
from neurodb.fmm import refresh
from neurodb.fmm import scope as scope_module
from neurodb.fmm.models import Visit
from neurodb.watch import people

from .conftest import CANARIES, LEAD, MEMBER, MEMBER_EMAIL, Q2, _question

TODAY = datetime.date(2026, 10, 5)
SHORT_NAME = "Met Mrs Layla Saab"  # short enough (under 80 characters) to reach Ask redacted, not withheld
SHORT_PHONE = "Call 03-123456"
SHORT_EMAIL = "karim.canary@example.org"
OTHER_LEAD = "Nadia Canary"  # a second lead, written under a key of the record only
EVERYTHING = (*CANARIES, LEAD, MEMBER, MEMBER_EMAIL, OTHER_LEAD, "nadia.canary@example.org")
KEPT = ("2026-05-11", "412 children", "2,800", "PCA2023597")
# Visit.team holds the team's display names on purpose (decision 15: shown to NeuroDB users only, never
# exported nor sent); everything else must be free of them.
TEAM_FIELD = ("fmm", "visit", "team")


def assert_free(where: str, value) -> None:
    blob = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, default=str)
    for canary in EVERYTHING:
        assert canary not in blob, f"{canary!r} reached {where}"


@pytest.fixture
def canary_world(fm_world, monkeypatch):
    """``fm_world`` with more canaries (short answers, other keys of the finding records), built on
    5 October 2026, "this year" being 2026."""
    monkeypatch.setattr(scope_module, "_today", lambda today=None: today or TODAY)
    for row in dm.MonitoringFinding.objects.filter(monitoring_activity_id=1722):
        row.data = {
            **row.data,
            "monitors": [{"name": OTHER_LEAD, "email": "nadia.canary@example.org"}],
            "comments": f"Shared with {OTHER_LEAD} (nadia.canary@example.org)",
            "contact_phone": "+961 3 123 456",
        }
        row.save(update_fields=["data"])
    _question(50201, 1722, "FM-2026-022", Q2, 2, entity="", answer=SHORT_NAME)
    _question(50202, 1723, "FM-2026-023", Q2, 2, answer=SHORT_PHONE, summary=SHORT_EMAIL)
    _question(50203, 1726, "FM-2026-026", (18, "Who was met?"), 6, answer=f"{LEAD}, {MEMBER}")
    run = refresh.run(triggered_by="test", today=TODAY)
    assert run.status == "succeeded", run.error
    people.forget()
    return fm_world


def _year():
    from django.http import QueryDict

    from neurodb.fmm.scope import Scope

    return Scope.from_params(QueryDict("year=2026&section="))


# ------------------------------------------------------------------------------------------ the paths
def _brief(monkeypatch) -> list[dict]:
    """A nightly brief written by a model that names the canaries in every part of its answer."""
    from types import SimpleNamespace

    from neurodb.assistant import agent
    from neurodb.fmm.ai import insights
    from neurodb.fmm.models import Insight

    sent: list[dict] = []
    echo = (
        f"{MEMBER} ({MEMBER_EMAIL}, +961 3 123 456) and Mrs Layla Saab met {LEAD}; see https://evil.example/x"
    )

    def create(**params):
        sent.append(params)
        sentence = {"text": echo, "keys": ["kpi"]}
        answer = {
            "coverage_summary": [sentence, {"text": "Eight visits in the period.", "keys": ["kpi"]}],
            "key_findings": [sentence],
            "challenges": [sentence],
            "recommendations": [sentence],
            "action_points": [
                {
                    "priority": "High",
                    "section": "Education",
                    "action": echo,
                    "owner_role": MEMBER,
                    "timeframe": "within 2 weeks",
                    "keys": ["kpi"],
                },
                {
                    "priority": "High",
                    "section": LEAD,
                    "partner": MEMBER,  # a person is never a partner the facts name: left out
                    "action": "Follow up the visits rated off track",
                    "owner_role": f"{LEAD}, {MEMBER_EMAIL}",
                    "timeframe": "within 2 weeks",
                    "keys": ["kpi"],
                },
            ],
        }
        return SimpleNamespace(output_text=json.dumps(answer), usage=None, status="completed", output=[])

    api = SimpleNamespace(responses=SimpleNamespace(create=create))
    api.with_options = lambda **_: api
    monkeypatch.setattr(agent, "client", lambda: api)
    row = insights.generate(_year(), trigger=Insight.Trigger.NIGHTLY)
    assert row.called and row.status in ("ok", "partial"), (row.status, row.reason)
    return sent


def _chat(monkeypatch, client) -> tuple[list[dict], str]:
    """One chat question naming people, answered after the four look-ups by a model that names them."""
    from neurodb.assistant import agent
    from tests.assistant.test_assistant import FakeClient, call, reply, say

    fake = FakeClient(
        [
            reply(
                call("fm_summary", {"group_by": "partner"}, "call_1"),
                call("fm_visits", {"limit": 30}, "call_2"),
            ),
            reply(
                call("fm_visit", {"visit": "1722"}, "call_3"),
                call("fm_visit", {"visit": "1726"}, "call_4"),
                call("fm_search", {"text": "registers"}, "call_5"),
                call("fm_search", {"text": "Layla"}, "call_6"),
            ),
            reply(
                say(f"{MEMBER} ({MEMBER_EMAIL}) told Mrs Layla Saab: see [Visit 1722](/fmm/visits/1722/).")
            ),
        ]
    )
    monkeypatch.setattr(agent, "client", lambda: fake)
    response = client.post(
        reverse("fmm:chat_stream"),
        {"question": f"What did {MEMBER} ({MEMBER_EMAIL}, 03-123456) find?", "scope": "year=2026&section="},
    )
    body = b"".join(response.streaming_content).decode()
    events = [json.loads(chunk[6:]) for chunk in body.split("\n\n") if chunk.startswith("data: ")]
    done = next(e for e in events if e["type"] == "done")
    outputs = [i for r in fake.requests for i in r["input"] if i.get("type") == "function_call_output"]
    assert len(outputs) >= 6, "the look-ups did not run"
    return fake.requests, json.dumps(done, ensure_ascii=False)


def _ask() -> list:
    """Ask NeuroDB: the four fm_* look-ups without a chat filter, and the generic eTools tools."""
    from neurodb.assistant import tools

    results = [
        tools.run("fm_summary", {"group_by": "partner"}),
        tools.run("fm_visits", {"limit": 30}),
        tools.run("fm_visit", {"visit": "1722"}),
        tools.run("fm_visit", {"visit": "1726"}),
        tools.run("fm_search", {"text": "registers"}),
        tools.run("etools_search", {"text": "Canary"}),
        tools.run("etools_search", {"text": "registers"}),
    ]
    for dataset in ("field_monitoring", "fm_questions", "fm_options", "fm_programme_activities"):
        rows = tools.run("etools_query", {"dataset": dataset, "limit": 50})
        assert rows["rows"], dataset
        results += [rows, tools.run("etools_datasets", {"dataset": dataset})]
        results += [
            tools.run("etools_record", {"dataset": dataset, "record": r["record"]}) for r in rows["rows"]
        ]
    for group in ("answer", "summary", "entity", "narrative_finding"):
        dataset = "field_monitoring" if group in ("entity", "narrative_finding") else "fm_questions"
        results.append(tools.run("etools_query", {"dataset": dataset, "group_by": group}))
    return results


def _hub_and_watch() -> list:
    from neurodb.graph.models import Change, Entity
    from neurodb.watch.models import WatchItem
    from tests.fmm.test_links import _pass
    from tests.graph.conftest import build_hub

    build_hub()
    visits = list(Entity.objects.filter(kind=Entity.Kind.FM_VISIT).values())
    assert visits, "no visit reached the hub"
    _pass(TODAY)
    items = list(WatchItem.objects.filter(detector="fm_follow_up").values())
    assert items, "the Watch check found nothing"
    changes = list(Change.objects.filter(kind=Entity.Kind.FM_VISIT).values())
    return [visits, items, changes]


def _exports(client) -> dict[str, str]:
    """The text of every file of the exports (Release 2 step 5: the records, one row per record with its
    own narrative and HACT answers): the Excel workbook's cells and the Power BI package's files, each
    by where it came from."""
    import io
    import zipfile

    from openpyxl import load_workbook

    params = {"year": "2026", "section": ""}
    out = {}
    response = client.get(reverse("fmm:export_xlsx"), params)
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    for sheet in book.worksheets:
        out[f"the workbook's {sheet.title} sheet"] = "\n".join(
            str(value) for row in sheet.iter_rows(values_only=True) for value in row if value is not None
        )
    response = client.get(reverse("fmm:export_powerbi"), params)
    assert response.status_code == 200
    package = zipfile.ZipFile(io.BytesIO(response.content))
    for name in package.namelist():
        out[f"the Power BI package's {name}"] = package.read(name).decode("utf-8")
    return out


def _explained(viewer) -> list:
    """The Help assistant's explanation of two visits' scores, record by record."""
    from neurodb.assistant import tools as assistant_tools
    from neurodb.help import assistant
    from neurodb.help import tools as help_tools

    with help_tools.bind(viewer):
        return [
            assistant_tools.run("explain_visit_score", {"visit": key}, registry=assistant.REGISTRY)
            for key in ("1722", "1726")
        ]


def _tables() -> dict[str, list]:
    """Every text, character, list and JSON value of every table of Monitoring insights (a table no
    longer there is left out: the AI checks made per visit before records, ``VisitAICheck``, dropped by
    migration 0020 once none is left)."""
    from django.db import connection

    out = {}
    tables = set(connection.introspection.table_names())
    for model in apps.get_app_config("fmm").get_models():
        if model._meta.db_table not in tables:
            continue
        names = [
            f.name
            for f in model._meta.concrete_fields
            if isinstance(f, models.CharField | models.TextField | models.JSONField)
            or f.get_internal_type() == "ArrayField"
        ]
        names = [n for n in names if (model._meta.app_label, model._meta.model_name, n) != TEAM_FIELD]
        if names:
            out[model.__name__] = list(model.objects.values(*names))
    return out


# ------------------------------------------------------------------------------------------ the test
def test_no_canary_leaves_monitoring_insights_by_any_path(canary_world, monkeypatch, client, viewer, caplog):
    with (
        caplog.at_level(logging.INFO),
        override_settings(FMM_AI=True, AI_ASSISTANT_ENABLED=True, OPENAI_API_KEY="x"),
    ):
        brief_requests = _brief(monkeypatch)
        client.force_login(viewer)
        chat_requests, shown = _chat(monkeypatch, client)
        ask = _ask()
        csv = client.get(reverse("fmm:visits"), {"year": "2026", "section": "", "export": "csv"})
        files = _exports(client)
        explained = _explained(viewer)
        hub_watch = _hub_and_watch()

    # the brief: what was sent (the redacted notes keep their figures), and what was kept
    assert len(brief_requests) == 1
    sent = brief_requests[0]["instructions"] + json.dumps(brief_requests[0]["input"], ensure_ascii=False)
    assert_free("the brief's request", sent)
    assert any(figure in sent for figure in KEPT)
    assert "[name withheld]" in sent
    # the chat: every request (instructions, question, history, look-up results) and the answer shown
    for n, request in enumerate(chat_requests):
        assert_free(f"chat request {n}", request["input"])
        assert_free(f"chat request {n}'s instructions", request["instructions"])
    assert_free("the chat's answer", shown)
    assert "/fmm/visits/1722/" in shown
    # Ask NeuroDB
    for n, result in enumerate(ask):
        assert_free(f"Ask result {n}", result)
    asked = json.dumps(ask, ensure_ascii=False)
    assert "Met [name withheld]" in asked and "Call [phone withheld]" in asked  # short values: redacted
    # the CSV of the visits
    assert csv.status_code == 200
    rows = csv.content.decode()
    assert rows.count("\n") > 5 and "Visit 1722" in rows
    assert_free("the visits CSV", rows)
    # the exports: the Excel workbook and the Power BI package, the records' own texts included
    assert "data/records.csv" in " ".join(files) and "Classes held as planned" in json.dumps(files)
    for where, text in files.items():
        assert_free(where, text)
    # the Help assistant's explanation of a score, record by record
    assert all(len(found["records"]) >= 2 for found in explained)
    assert_free("the Help assistant's explanation", explained)
    # the knowledge hub, What's new and NeuroDB Watch
    for where, blob in zip(("the hub's visits", "Watch's records", "What's new"), hub_watch, strict=True):
        assert_free(where, blob)
    # every fmm table, after all of the above wrote to it
    tables = _tables()
    assert tables["Visit"] and tables["KeyProbe"] and tables["Insight"] and tables["ChatQuestion"]
    for model, rows in tables.items():
        assert_free(f"the fmm table {model}", rows)
    # the team is kept as display names only, never an e-mail address
    for team in Visit.objects.values_list("team", flat=True):
        assert not any("@" in name for name in team)
    assert MEMBER in Visit.objects.get(key="1722").team
    # nothing written to the log names anyone either
    assert_free("the log", caplog.text)
