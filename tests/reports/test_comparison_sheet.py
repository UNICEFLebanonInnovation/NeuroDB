"""The Power BI comparison sheet (datamart.comparison, SPEC_validation §4): NeuroDB's figures in Power BI's order,
each the page's own figure for the filter and day; Power BI's cells empty; only the difference columns hold
formulas; every filter field a slicer to set or a NeuroDB filter to clear; the lists to match row by row."""

# ruff: noqa: F401, F811 (the fixtures "data", "more" and "frozen" come from the partner monitoring tests)
from __future__ import annotations

import dataclasses
import datetime
import io
from urllib.parse import unquote_plus

import pytest
from django.contrib.auth.models import Group
from django.urls import reverse
from openpyxl import load_workbook

from neurodb.accounts.models import User
from neurodb.accounts.roles import VALIDATORS, VIEWER
from neurodb.datamart import comparison, monitoring, raw, reporting, services, trail
from neurodb.datamart import models as dm
from neurodb.datamart.tags import TAG_FIELDS
from tests.reports.test_pd_monitoring import TODAY, data, report_rows
from tests.reports.test_powerbi_views import CLASSROOMS, frozen, more

pytestmark = pytest.mark.django_db
PR, PM = "progress-reports", "partner-monitoring"
# the filters each page is compared under (the check line's own queries)
PR_FILTERS = [
    {},
    {"pd_status": ""},
    {"pd_status": "", "section": "Child Protection"},
    {"pd_status": "", "status": "accepted"},
    {"pd_status": "", "status": ["overdue", "not_yet_due"]},
    {"year": "2026"},
    {"report_type": "HR"},
    {"pd_status": "", "overdue": "1"},
]
PM_FILTERS = [
    {"view": "table", "year": "all", "section": ""},
    {"view": "table", "year": "2026", "section": "Education"},
    {"view": "quality", "year": "all", "section": "", "pd_status": ""},
    {"year": "2026", "section": ""},
]


@pytest.fixture
def validator(client, roles):
    user = User.objects.create_user(username="checker", password="pass-123456-x")
    user.groups.set(Group.objects.filter(name__in=(VIEWER, VALIDATORS)))
    client.force_login(user)
    return client


def _figures(page: str, params: dict, day: datetime.date = TODAY) -> dict[str, comparison.Figure]:
    return {f.key: f for f in comparison.figures(comparison.read(page, params, day))}


def _book(client, page: str, params: dict):
    response = client.get(reverse("reports:check_compare", args=[page]), params)
    assert response.status_code == 200, response.content[:200]
    assert response["Content-Disposition"].startswith('attachment; filename="powerbi-comparison-')
    return load_workbook(io.BytesIO(response.content))


def _table(sheet) -> list[tuple]:
    """The figure rows of the Compare sheet (after the header that starts with "#")."""
    rows = list(sheet.iter_rows(values_only=True))
    start = next(i for i, row in enumerate(rows) if row and row[0] == "#")
    out = []
    for row in rows[start + 1 :]:
        if not row or row[0] is None:
            break
        out.append(row)
    return out


def _block(sheet, title: str) -> dict[str, object]:
    """The rows under one of the Compare sheet's bold titles, as {label: value} (until an empty row)."""
    rows = list(sheet.iter_rows(values_only=True))
    start = next(i for i, row in enumerate(rows) if len(row) > 1 and row[1] == title)
    out = {}
    for row in rows[start + 1 :]:
        if not row or all(v is None for v in row):
            break
        out[row[1]] = row[2]
    return out


# ------------------------------------------------------------------------------------------- the page's figures
@pytest.mark.parametrize("params", PR_FILTERS)
def test_progress_reports_figures_are_the_page_s(validator, more, frozen, params):
    found = _figures(PR, params)
    page = validator.get(reverse("reports:partner_reporting"), params).context
    data = page["data"]
    assert found["pr.indicator_reports"].value == data["counts"]["indicator_reports"]
    for s in data["counts"]["statuses"]:
        assert found[f"pr.status.{s['key']}"].value == s["indicator_reports"]
    assert found["pr.reports_listed"].value == page["page_obj"].paginator.count
    for row in data["by_section"]["rows"]:
        for key, n in zip(reporting.STATUS_KEYS, row["counts"], strict=True):
            assert found[f"pr.section.{row['section']}.{key}"].value == n
        assert found[f"pr.section.{row['section']}.total"].value == row["total"]
    assert found["pr.section_total.total"].value == data["by_section"]["total"]
    month_values = [f.value for k, f in found.items() if k.startswith("pr.month.")]
    assert sum(month_values) == sum(sum(v) for v in data["by_month"]["series"].values())
    partners = validator.get(reverse("reports:partner_reporting"), {**params, "tab": "partners"}).context
    perf = partners["data"]["partners"]
    assert found["pr.partners"].value == perf["partners"]
    assert found["pr.partners_reported"].value == perf["reported"]
    assert found["pr.partners_not_reported"].value == perf["not_reported"]
    for s in perf["by_status"]:
        assert found[f"pr.partners_status.{s['key']}"].value == s["partners"]
    assert found["pr.follow_up"].value == partners["page_obj"].paginator.count
    # the tabs' stored tables, as the page reads them
    locations = services.reporting_table("locations", params, TODAY)["locations"]
    assert found["pr.location_lines"].value == len(locations["rows"])
    blocks = services.reporting_table("disaggregations", params, TODAY)["disaggregations"]
    assert found["pr.disaggregation_locations"].value == len(blocks["blocks"])


@pytest.mark.parametrize("params", PM_FILTERS)
def test_partner_monitoring_figures_are_the_page_s(validator, more, frozen, params):
    found = _figures(PM, params)
    table = validator.get(reverse("reports:pd_monitoring"), {**params, "view": "table"}).context
    assert found["pp.indicators"].value == table["page_obj"].paginator.count
    tiles = {k.removeprefix("pp.section_tile."): f.value for k, f in found.items() if "section_tile" in k}
    assert tiles == {t["name"]: t["count"] for t in table["section_tiles"]}
    quality = validator.get(reverse("reports:pd_monitoring"), {**params, "view": "quality"})
    quality = quality.context["quality"]
    pairs = {
        "dq.pds": "programme_documents",
        "dq.running": "running_pds",
        "dq.active": "active_pds",
        "dq.ended": "ended_pds",
        "dq.indicators": "indicators",
        "dq.with_reached": "with_reached",
        "dq.zero_reached": "zero_reached",
        "dq.without_reached": "without_reached",
        "dq.with_target": "with_target",
        "dq.without_target": "without_target",
    }
    for key, name in pairs.items():
        assert found[key].value == quality[name], key


def test_rows_follow_power_bi_s_pages_and_visuals(more, frozen):
    for page in comparison.PAGES:
        order = [(p, v) for p, visuals in comparison.POWERBI_PAGES[page] for _block_, v in visuals]
        seen = []
        for figure in comparison.figures(comparison.read(page, {"pd_status": ""}, TODAY)):
            if not seen or seen[-1] != (figure.powerbi_page, figure.visual):
                seen.append((figure.powerbi_page, figure.visual))
        assert seen == order, page  # every visual, once, in Power BI's order


def test_program_performance_counts_every_indicator_past_the_page_s_list(
    validator, more, frozen, monkeypatch
):
    monkeypatch.setattr(monitoring, "MAX_INDICATORS", 2)
    found = _figures(PM, {"view": "table", "year": "all"})
    assert found["pp.indicators"].value == 3
    assert sum(f.value for k, f in found.items() if k.startswith("pp.section_tile.")) == 3
    book = _book(validator, PM, {"view": "table", "year": "all"})
    rows = list(book["Program Performance rows"].iter_rows(min_row=3, values_only=True))
    assert len(rows) == 3 and {r[2] for r in rows} == {more["girls"], more["boys"], CLASSROOMS}


# --------------------------------------------------------------------------------------------------- the file
def test_the_compare_sheet_leaves_power_bi_s_cells_empty_and_writes_two_formula_columns(
    validator, more, frozen
):
    # a section a spreadsheet would run as a formula: written as a text; a character Excel refuses: left out
    dm.ReportedIndicator.objects.filter(progress_report="PR-2").update(section="=HYPERLINK(1)")
    dm.ReportedIndicator.objects.filter(progress_report="HR-4").update(partner_name="Amel\x0bAssociation")
    book = _book(validator, PR, {"pd_status": ""})
    assert book.sheetnames == ["Compare", "Progress reports", "Partners", "Why they can differ"]
    compare = book["Compare"]
    rows = _table(compare)
    figures = comparison.figures(comparison.read(PR, {"pd_status": ""}, TODAY))
    assert [r[5] for r in rows] == [f.value for f in figures]
    assert all(r[7] is None for r in rows)  # Power BI's value: typed by a person
    # each figure's trail: the figure trail of the same filter, open on that figure
    for row, figure in zip(rows, figures, strict=True):
        assert row[11].startswith("http://testserver/check/progress-reports/?pd_status=&"), row[11]
        assert f"figure={figure.key}" in unquote_plus(row[11])
        assert row[11].endswith("#" + trail.anchor(figure.key))
    assert any(r[3] == "=HYPERLINK(1) · Accepted" for r in rows)
    hr4 = next(r for r in book["Progress reports"].iter_rows(min_row=3, values_only=True) if r[3] == "HR4")
    assert hr4[1] == "AmelAssociation"
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    assert sheet.title == "Compare" and cell.column_letter in ("I", "J"), cell.coordinate
                elif isinstance(cell.value, str) and cell.value.startswith("="):
                    assert cell.data_type == "s"
    first = next(r for r in compare.iter_rows(min_row=1) if r[0].value == 1)
    n = first[0].row
    assert first[8].value == f'=IF(ISNUMBER(H{n}),H{n}-F{n},"")'
    assert first[9].value == f'=IF(ISNUMBER(H{n}),IF(H{n}=F{n},"Same","Differs"),"")'
    assert "eTools reporting checks" in next(r[10] for r in rows if r[3] == "Overdue")


def test_the_program_performance_sheet_writes_two_formula_columns(validator, more, frozen):
    dm.PDIndicator.objects.filter(title=CLASSROOMS).update(section_name="=X")
    book = _book(validator, PM, {"view": "table", "year": "all"})
    assert book.sheetnames == ["Compare", "Program Performance rows", "Why they can differ"]
    sheet = book["Program Performance rows"]
    header = [c.value for c in sheet[2]]
    assert header[7] == "NeuroDB Reached" and header[13] == "Match key"
    for row in sheet.iter_rows(min_row=3):
        assert row[9].value is None and row[10].value is None  # Power BI's: typed by a person
        assert [c.column_letter for c in row if c.data_type == "f"] == ["L", "M"]
    sections = {row[0].value: row[0].data_type for row in sheet.iter_rows(min_row=3)}
    assert sections["=X"] == "s"
    girls = next(r for r in sheet.iter_rows(min_row=3, values_only=True) if r[2] == more["girls"])
    assert girls[7] == 550 and girls[13] == f"LEB/PD1 | {more['girls']} | 1.1 Protection services"
    year = _book(validator, PM, {"view": "table", "year": "2026"})["Program Performance rows"]
    assert year["H2"].value == "NeuroDB Reached in 2026"


def test_same_title_indicators_get_one_match_key_each(validator, more, frozen):
    """Two indicators of one PD with the same title and PD output (check 7): the key is never shared."""
    dm.PDIndicator.objects.create(
        datamart_id=21,
        source_id=600,
        intervention=more["pd"],
        pd_reference_number=more["pd"].number,
        title=more["boys"],
        section_name="Child Protection",
        lower_result_name="1.1 Protection services",
    )
    sheet = _book(validator, PM, {"view": "table", "year": "all"})["Program Performance rows"]
    keys = [r[13] for r in sheet.iter_rows(min_row=3, values_only=True)]
    assert len(keys) == len(set(keys)) == 4
    twins = sorted(k for k in keys if more["boys"] in k)
    assert twins == [f"LEB/PD1 | {more['boys']} | 1.1 Protection services #{n}" for n in (1, 2)]


def test_the_lists_hold_every_report_and_partner_and_no_person(validator, more, frozen):
    params = {"pd_status": ""}
    book = _book(validator, PR, params)
    data = services.partner_reporting(params, TODAY)
    reports = list(book["Progress reports"].iter_rows(min_row=3, values_only=True))
    assert len(reports) == len(data["reports"]) == 5
    assert [c.value for c in book["Progress reports"][2]][-1] == "Match key"
    pr1 = next(r for r in reports if r[0] == "LEB/PD1" and r[3] == "QPR1")
    assert pr1[1:3] == ("Amel Association", "V1") and pr1[9] == 2 and pr1[10] == "LEB/PD1 QPR1"
    partners = list(book["Partners"].iter_rows(min_row=3, values_only=True))
    figures = _figures(PR, params)
    assert len(partners) == figures["pr.partners"].value == 2
    assert sum(r[-2] for r in partners) == figures["pr.reports_listed"].value  # each report has a partner
    assert {r[-1] for r in partners} == {"V1", "V2"}
    text = repr([list(s.iter_rows(values_only=True)) for s in book.worksheets])
    for person in ("Focal Point One", "focal.two@example.org", "Focal Three", "@"):
        assert person not in text


@pytest.mark.parametrize(
    ("q", "cleaned"),
    [
        ("Test Person", "[name withheld]"),  # a known person
        ("test.person@example.org", "[email withheld]"),
        ("+961 3 123 457", "[phone withheld]"),
        ("call 03 123 457 now", "call [phone withheld] now"),
    ],
)
def test_a_search_naming_a_person_or_a_contact_never_reaches_the_sheet(validator, more, frozen, q, cleaned):
    """A search can name anyone: the people rule cleans it in the filters to clear, and the sheet writes no
    address of that filter (the page's, every Trail link)."""
    User.objects.create_user(username="known", first_name="Test", last_name="Person", password="pw-123456-x")
    for page, params in ((PR, {"pd_status": ""}), (PM, {"view": "table", "year": "all"})):
        book = _book(validator, page, {**params, "q": q})
        for sheet in book.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if isinstance(cell, str):
                        for found in (cell, unquote_plus(cell)):
                            assert q.lower() not in found.lower(), (page, sheet.title, cell)
                            assert "123 457" not in found and "123457" not in found, (page, sheet.title, cell)
        compare = book["Compare"]
        clear = _block(compare, "Clear these in NeuroDB before comparing")
        assert clear[comparison.NEURODB_ONLY[page]["q"]] == cleaned
        address = next(r[2] for r in compare.iter_rows(values_only=True) if r[1:2] == ("NeuroDB filter",))
        assert address == comparison.ADDRESS_WITHHELD
        if page == PR:
            assert {r[11] for r in _table(compare)} == {comparison.TRAIL_WITHHELD}
    # a search that names no one keeps its addresses
    compare = _book(validator, PR, {"pd_status": "", "q": "girls"})["Compare"]
    assert all("q=girls" in r[11] for r in _table(compare))


# ------------------------------------------------------------------------------------------- the filter blocks
def test_every_filter_field_is_a_slicer_or_a_neurodb_filter():
    for page in comparison.PAGES:
        assert comparison.unlisted(page) == [], page
        both = set(comparison.POWERBI_SLICERS[page]) & set(comparison.NEURODB_ONLY[page])
        assert both == set(), page


def test_a_new_filter_field_in_neither_list_stops_the_sheet(more, frozen, monkeypatch):
    planted = dataclasses.make_dataclass("Planted", [("new_choice", str, dataclasses.field(default=""))])
    monkeypatch.setitem(comparison.FILTERS, PR, (*comparison.FILTERS[PR], planted))
    assert comparison.unlisted(PR) == ["new_choice"]
    with pytest.raises(ValueError, match="new_choice"):
        comparison.settings_rows(comparison.read(PR, {}, TODAY))


def test_the_slicers_to_set_and_the_neurodb_filters_to_clear(validator, more, frozen):
    params = {
        "year": "2026",
        "section": "Child Protection",
        "partner": str(more["partner"].pk),
        "q": "girls",
        "report_number": "QPR1",
        "overdue": "1",
        "indicator": "girls",
    }
    slicers, clear = comparison.settings_rows(comparison.read(PR, params, TODAY))
    slicers = dict(slicers)
    assert slicers["Reporting year"] == "2026" and slicers["Section"] == "Child Protection"
    assert slicers["Partner"] == "Amel Association"
    assert slicers["PD status"] == "Active, Signed, Suspended"  # the page's default: the running PDs
    assert slicers["Report status"] == "Every value: clear it"
    assert dict(clear) == {
        "Search": "girls",
        "Report #": "QPR1",
        "Overdue only": "Yes",
        "By disaggregation: indicator": "girls",
    }
    compare = _book(validator, PR, params)["Compare"]
    assert _block(compare, "Clear these in NeuroDB before comparing")["Search"] == "girls"
    _slicers, clear = comparison.settings_rows(comparison.read(PR, {}, TODAY))
    assert clear == []
    pm = {"year": "2026", "scope": "year", "report_type": "HR", "status": "on_track", "cp_output": "x"}
    pm |= {"gender": "Female", "age_group": ["12-17", "Under 5"]}
    slicers, clear = comparison.settings_rows(comparison.read(PM, pm, TODAY))
    assert dict(slicers)["Reporting year"] == "2026" and dict(slicers)["PD status"].startswith("Every status")
    assert dict(clear) == {
        "Programme documents": "Every programme document running in 2026, closed ones too",
        "Report type": "Humanitarian reports (monthly): only emergency or cluster indicators",
        "Tracking status": "On track",
        "Tags": "Gender: Female; Age group: 12-17, Under 5",  # the filter bar's words
        "CP output": "x",
    }
    assert set(comparison.TAG_WORDS) == set(TAG_FIELDS)
    slicers, _clear = comparison.settings_rows(comparison.read(PM, {"year": "all"}, TODAY))
    assert dict(slicers)["Reporting year"] == "Every year"
    # the programme documents the scope keeps, in words that say what the figures keep: every one with "all",
    # the closed programme document that ran in 2024 too; only those running in the year with "year"
    scopes = {
        "Every programme document, of every status and year": ({"scope": "all", "year": "2026"}, 2),
        "Every programme document running in 2026, closed ones too": ({"scope": "year", "year": "2026"}, 1),
        "Every programme document of the PD statuses chosen, of every year": (
            {"scope": "all", "year": "2026", "pd_status": ["active", "closed"]},
            2,
        ),
        "Every programme document running in 2026, of the PD statuses chosen": (
            {"scope": "year", "year": "2026", "pd_status": ["active", "closed"]},
            1,
        ),
    }
    for words, (params, pds) in scopes.items():
        params = {**params, "view": "table"}
        _slicers, clear = comparison.settings_rows(comparison.read(PM, params, TODAY))
        assert dict(clear)["Programme documents"] == words, params
        assert _figures(PM, params)["dq.pds"].value == pds, params
    _slicers, clear = comparison.settings_rows(comparison.read(PM, {"scope": "all", "year": "all"}, TODAY))
    assert dict(clear)["Programme documents"] == "Every programme document, of every status and year"
    compare = _book(validator, PM, {**pm, "scope": "all"})["Compare"]
    clear = _block(compare, "Clear these in NeuroDB before comparing")
    assert clear["Programme documents"] == "Every programme document, of every status and year"
    assert clear["Tags"] == "Gender: Female; Age group: 12-17, Under 5"


def test_power_bi_s_names_come_from_one_place(validator, more, frozen, monkeypatch):
    """Every Power BI page, visual and slicer either workbook names is read from POWERBI_PAGES,
    POWERBI_PD_PAGE and POWERBI_SLICERS: renamed there, no old name is left in a cell."""
    renamed: dict[str, str] = {}

    def new(name: str) -> str:
        return renamed.setdefault(name, f"Renamed {len(renamed) + 1:02}")

    pages = {
        page: tuple((new(name), tuple((block, new(v)) for block, v in visuals)) for name, visuals in found)
        for page, found in comparison.POWERBI_PAGES.items()
    }
    slicers = {
        page: {k: new(v) for k, v in found.items()} for page, found in comparison.POWERBI_SLICERS.items()
    }
    monkeypatch.setattr(comparison, "POWERBI_PAGES", pages)
    monkeypatch.setattr(comparison, "POWERBI_SLICERS", slicers)
    monkeypatch.setattr(comparison, "POWERBI_PD_PAGE", new(comparison.POWERBI_PD_PAGE))
    # NeuroDB's own words too (columns of the lists, a filter, a view of Partner monitoring): not searched
    own = {"Partner", "Section", "PD status", "Program Performance"}
    text = ""
    for page in comparison.PAGES:
        book = _book(validator, page, {"pd_status": ""})
        text += repr(
            [cell for sheet in book.worksheets for row in sheet.iter_rows(values_only=True) for cell in row]
        )
    left = [old for old in renamed if old not in own and old in text]
    assert left == []
    assert [name for name in renamed.values() if name not in text] == []  # every name is shown


# ------------------------------------------------------------------------------------------------- the day
def test_statuses_read_for_a_past_day_read_as_the_page_reads_them(validator, more, frozen):
    # due on 25 June, not submitted: Due on 22 June, Overdue on 1 July (the page's day)
    report_rows(
        more["pd"],
        more["partner"],
        more["boys"],
        "PR-5",
        "QPR5",
        "QPR",
        datetime.date(2026, 6, 10),
        {"Akkar": 0},
        20,
        status="Due",
    )
    dm.ReportedIndicator.objects.filter(progress_report="PR-5").update(due_date=datetime.date(2026, 6, 25))
    past = datetime.date(2026, 6, 22)
    then, now = _figures(PR, {}, past), _figures(PR, {})
    for day, found in ((past, then), (TODAY, now)):
        counts = services.partner_reporting({}, day)["counts"]
        assert {k: found[f"pr.status.{k}"].value for k in reporting.STATUS_KEYS} == {
            k: counts[k] for k in reporting.STATUS_KEYS
        }
    assert then["pr.status.overdue"].value < now["pr.status.overdue"].value
    compare = _book(validator, PR, {"day": "2026-06-22"})["Compare"]
    shown = {r[3]: r[5] for r in _table(compare)}
    assert shown["Overdue"] == then["pr.status.overdue"].value
    assert compare["B3"].value == "Statuses read for" and compare["C3"].value == "22 Jun 2026"
    url = reverse("reports:check_compare", args=[PR])
    for refused in ("2026-06-16", "2026-07-02", "yesterday"):  # 15 days before, tomorrow, not a day
        response = validator.get(url, {"day": refused})
        assert response.status_code == 400 and response.content.decode() == comparison.REFUSED_DAY
    assert validator.get(url, {"day": "2026-06-17"}).status_code == 200  # 14 days before


def test_everything_neurodb_holds_ignores_the_filter(validator, more, frozen):
    every = services.partner_reporting({"pd_status": ""}, TODAY)["counts"]
    for params in ({}, {"section": "Education", "year": "2025"}):
        held = _block(_book(validator, PR, params)["Compare"], "Everything NeuroDB holds")
        assert held["Indicator reports"] == every["indicator_reports"] == 6
        assert held["Progress reports"] == every["progress_reports"] == 5
    held = _block(_book(validator, PM, {"year": "all"})["Compare"], "Everything NeuroDB holds")
    assert held["Programme documents"] == 2 and held["Indicator reports"] == 6


# ---------------------------------------------------------------------------------------------- the address
def test_home_and_other_pages_have_no_sheet(validator, more):
    for page in ("home", "field-monitoring", "nothing"):
        assert validator.get(reverse("reports:check_compare", args=[page])).status_code == 404


def test_a_third_large_download_waits(validator, more, frozen, monkeypatch):
    monkeypatch.setattr(raw, "take_slot", lambda: None)
    response = validator.get(reverse("reports:check_compare", args=[PR]))
    assert response.status_code == 503 and response.content.decode() == raw.BUSY


@pytest.mark.parametrize("params", PR_FILTERS)
def test_each_progress_reports_value_is_its_trail_s_figure(validator, more, frozen, params):
    """Stage 3: every NeuroDB value of the Progress reports sheet is the figure its trail shows on the page and
    rebuilds from the rows; its Trail link opens that trail."""
    found = trail.progress_trail(PR, params, TODAY)
    tables = {tab: trail.table_trail(tab, params, TODAY) for tab in trail.TABLES}
    figures = {f.key: f for f in found.figures}
    for t in tables.values():
        figures.update({f.key: f for f in t.block.figures})
    book = _book(validator, PR, params)
    for row in _table(book["Compare"]):
        key = unquote_plus(row[11].split("figure=", 1)[1].split("#", 1)[0].split("&", 1)[0])
        figure = figures[key]
        assert row[5] == figure.page == figure.rebuilt, key
        if key.startswith(("pr.location", "pr.disaggregation")):
            assert "tab=" in row[11]
    link = _table(book["Compare"])[0][11].removeprefix("http://testserver")
    assert validator.get(link).status_code == 200


@pytest.mark.parametrize("params", PM_FILTERS)
def test_each_partner_monitoring_value_is_its_trail_s_figure(validator, more, frozen, params):
    """Stage 4: every NeuroDB value of the Partner monitoring sheet is the figure its trail shows on the page and
    rebuilds from the indicators kept (Program Performance's from its own part); each Trail link opens that
    trail, on Program Performance or Data quality and the year the sheet read; each indicator of the Program
    Performance rows sheet links to its Reached."""
    from neurodb.datamart import trail_monitoring as tm

    found = tm.monitoring_trail(params, TODAY)
    figures = {f.key: f for f in found.figures}
    figures.update({f.key: f for f in tm.performance_trail(params, TODAY).block.figures})
    book = _book(validator, PM, params)
    for row in _table(book["Compare"]):
        key = unquote_plus(row[11].split("figure=", 1)[1].split("#", 1)[0].split("&", 1)[0])
        figure = figures[key]
        assert row[5] == figure.page == figure.rebuilt, key
        assert ("view=table" if key.startswith("pp.") else "view=quality") in row[11], key
        assert "year=" in row[11]
        link = row[11].removeprefix("http://testserver")
        assert validator.get(link).status_code == 200
    sheet = book["Program Performance rows"]
    rows = list(sheet.iter_rows(min_row=3, values_only=True))
    assert rows and all("figure=pp.reached." in r[14] for r in rows)
    for r in rows:
        key = unquote_plus(r[14].split("figure=", 1)[1].split("#", 1)[0].split("&", 1)[0])
        assert key in figures, key
