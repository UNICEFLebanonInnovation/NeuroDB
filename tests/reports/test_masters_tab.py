"""The Master indicators tab of Partner monitoring (``view=masters``): its two parts, the figures they show (those of
neurodb.datamart.masters.figures over partner monitoring's stored rows), its own filter (the status, the search and
the CP output choose master rows, never their values; a partner's contribution with a share of sums only), the
subs read from another section, the detail modal, the export, who may change what, and that a change shows at once.

The world: three running programme documents of Partners A, B and C (PD/2026/001 to 003) with Education and Child
Protection PD indicators reported quarterly in Akkar and Bekaa; a sum master and a percentage master in Education;
a suggestion, a proposal and an indicator in no master. Generic names only."""

from __future__ import annotations

import csv
import io
import re
from collections import Counter
from urllib.parse import parse_qs, urlsplit

import openpyxl
import pytest
from django.contrib.auth.models import Group
from django.contrib.messages import get_messages
from django.test import Client
from django.urls import reverse

from neurodb.accounts.models import Section, User
from neurodb.accounts.roles import SECTION_EDITOR
from neurodb.datamart import curation, master_links, masters, masters_tab, monitoring
from neurodb.datamart import models as dm
from neurodb.donors.models import DonorAccount
from neurodb.watch.models import SectionMatch
from tests.datamart import masters_world as w
from tests.datamart.masters_world import day
from tests.web.parts import HX, get_part, placeholders, raw_get, text

pytestmark = pytest.mark.django_db

PAGE = "/partner-monitoring/"
ENROLLED = "# of children enrolled in NFE"
BOTH = "# of girls and boys enrolled in NFE"
PASSING = "% of children passing the end-of-year exam"
KITS = "# of children receiving recreational kits"
TEACHERS = "# of teachers trained on inclusive education"
CAREGIVERS = "# of caregivers trained on positive parenting"
SUM_NAME, PERCENT_NAME, MOVED_NAME = (
    "Children enrolled in NFE",
    "Children passing the exam",
    "Children with kits",
)
PERSON = "staff.member@example.org"  # who updated or confirmed: never on the page or in the export


def quarterly(pd, source_id, title, values, places=("Akkar",)):
    """Quarterly reports of one indicator: quarter -> value at each place (Q1 ends 31 March)."""
    for quarter, value in values.items():
        end = {1: day(3, 31), 2: day(6, 30), 3: day(9, 30)}[quarter]
        w.report(pd, source_id, title, f"QPR{quarter}", end, dict.fromkeys(places, value))


@pytest.fixture
def world(db, roles):
    a, b, c = (w.programme(letter, n) for n, letter in enumerate("ABC", start=1))
    # a sum: A 500 (Akkar and Bekaa), B 200 (Akkar): up to 700, at least 500
    w.indicator(a, 101, ENROLLED, target=(1000, 1), places=("Akkar", "Bekaa"))
    w.report(a, 101, ENROLLED, "QPR1", day(3, 31), {"Akkar": 100, "Bekaa": 50})
    quarterly(a, 101, ENROLLED, {2: 200, 3: 150})
    w.indicator(b, 201, BOTH, target=(600, 1))
    quarterly(b, 201, BOTH, {1: 80, 2: 120})
    w.indicator(c, 301, "# of children enrolled in NFE classes", target=(300, 1))
    quarterly(c, 301, "# of children enrolled in NFE classes", {1: 70})
    enrolled = w.master(SUM_NAME, target=1000, code="ED-1.1", cp_output="1.1 Children access learning")
    w.link(enrolled, a, 101)
    w.link(enrolled, b, 201)
    suggestion = w.link(enrolled, c, 301, status="suggested", unit_class="")
    suggestion.score, suggestion.reason = 0.95, "Same CP output, type and title as 2 confirmed indicators"
    suggestion.save()
    # a percentage: 40/100 and 10/50 give 33.3 %; C sent the percentage only: left out, "counted 2 of 3 PDs"
    for pd, source_id, value in ((a, 131, w.pct(40, 100)), (b, 231, w.pct(10, 50)), (c, 331, {"c": 45})):
        w.indicator(pd, source_id, PASSING, display_type="percentage", target=(60, 100))
        quarterly(pd, source_id, PASSING, {1: value})
    passing = w.master(PERCENT_NAME, "percent", target=50, cp_output="1.2 Children learn")
    for pd, source_id in ((a, 131), (b, 231), (c, 331)):
        w.link(passing, pd, source_id, unit_class="percentage")
    # an Education master whose PD indicator is in Child Protection in eTools now
    w.indicator(a, 105, KITS, section=w.CP, target=(400, 1))
    quarterly(a, 105, KITS, {1: 90})
    moved = w.master(MOVED_NAME, target=500)
    w.link(moved, a, 105)
    # in no master; and a proposal of two programme documents
    w.indicator(b, 202, TEACHERS)
    w.indicator(a, 141, CAREGIVERS)
    w.indicator(b, 241, CAREGIVERS)
    proposal = w.master(
        "Caregivers trained on positive parenting",
        status="proposed",
        origin="proposed",
        proposal_key="k" * 40,
    )
    for pd, source_id in ((a, 141), (b, 241)):
        w.link(proposal, pd, source_id, status="suggested", unit_class="")
    # who did it, which nothing but the administration shows
    dm.EtoolsMaster.objects.update(updated_by=PERSON)
    dm.EtoolsMasterLink.objects.update(confirmed_by=PERSON)
    # a Child Protection master, never on the Education tab
    w.indicator(c, 311, "# of children reached with PSS", section=w.CP)
    quarterly(c, 311, "# of children reached with PSS", {1: 30})
    protection = w.master("Children reached with PSS", section=w.CP)
    w.link(protection, c, 311)
    return {
        "pds": (a, b, c),
        "enrolled": enrolled,
        "passing": passing,
        "moved": moved,
        "proposal": proposal,
        "protection": protection,
        "suggestion": suggestion,
    }


def _signed_in(user) -> Client:
    client = Client()
    client.force_login(user)
    return client


def _editor(section_name: str, code: str, etools_name: str) -> User:
    """A Section editor of the NeuroDB section an eTools section name is confirmed to mean."""
    section = Section.objects.create(name=section_name, code=code)
    SectionMatch.objects.create(etools_name=etools_name, section=section, how="exact", confirmed=True)
    user = User.objects.create_user(username=f"editor-{code}", password="editor-pass-123456", section=section)
    user.groups.add(Group.objects.get(name=SECTION_EDITOR))
    return user


def _tab(client, **params):
    query = {"view": "masters", "section": w.EDUCATION, **params}
    return client.get(PAGE, query)


def _lines(response) -> dict[str, masters_tab.Line]:
    return {line.master.name: line for line in response.context["tab"].lines}


def _expected() -> dict[str, masters.Figure]:
    """masters.figures of the shown masters over the year's rows of every section (the tab's section rows and
    its read by programme document hold the same rows)."""
    shown = dm.EtoolsMaster.objects.filter(status="active", section=w.EDUCATION).prefetch_related(
        "links", "years"
    )
    return {f.master.name: f for f in masters.figures(shown, w.rows(), w.YEAR, w.TODAY)}


# ------------------------------------------------------------------------------------------------- the frame
@pytest.mark.shell
def test_the_tab_is_a_view_of_partner_monitoring_with_its_two_parts_on_one_year(client_viewer, world):
    html = text(raw_get(client_viewer, PAGE, {"view": "masters", "section": w.EDUCATION}))
    found = placeholders(html)
    assert list(found) == ["masters-summary", "masters-table"]
    for url in found.values():
        query = parse_qs(urlsplit(url).query, keep_blank_values=True)
        assert (
            query["view"] == ["masters"] and query["year"] == ["2026"] and query["section"] == [w.EDUCATION]
        )
    assert 'data-current-query="view=masters"' in html and "Master indicators" in html
    # one year: no "All years" in its filter, and "All years" from another tab is this year
    assert '<option value="all"' not in html
    every = raw_get(client_viewer, PAGE, {"view": "masters", "year": "all", "section": w.EDUCATION})
    assert every.status_code == 302 and "year=2026" in every["Location"]
    table = raw_get(client_viewer, PAGE, {"view": "table", "year": "all"})
    assert parse_qs(urlsplit(table.context["tab_urls"]["masters"]).query)["year"] == ["2026"]
    assert table.context["tab_urls"]["grid"].endswith("year=all")  # the other tabs keep every year
    # the filters the tab does not apply are not in its filter bar
    for field in ('id="f-pd-status"', 'name="partner_type"', 'name="gender"', 'id="f-emergency"'):
        assert field not in html, field
    assert 'name="partner"' in html and 'name="location"' in html


def test_the_tab_opens_on_the_users_own_section(world):
    editor = _editor("Education", "ED", w.EDUCATION)
    html = text(raw_get(_signed_in(editor), PAGE + "?view=masters"))
    assert all(f"section={w.EDUCATION}" in url for url in placeholders(html).values())


# ----------------------------------------------------------------------------------------------- the figures
def test_the_parts_show_the_figures_of_masters_figures(client_viewer, world):
    response = _tab(client_viewer)
    assert response.status_code == 200
    lines, expected = _lines(response), _expected()
    assert set(lines) == {SUM_NAME, PERCENT_NAME, MOVED_NAME}  # not the proposal, not Child Protection's
    for name, line in lines.items():
        want = expected[name]
        assert (line.figure.value, line.figure.at_least, line.figure.quarters) == (
            want.value,
            want.at_least,
            want.quarters,
        )
        assert (line.figure.status, line.figure.achieved, line.figure.coverage) == (
            want.status,
            want.achieved,
            want.coverage,
        )
        assert [(s.sub.link.pk, s.sub.counted, s.sub.reason) for s in line.subs] == [
            (s.link.pk, s.counted, s.reason) for s in want.subs
        ]
    total, share = lines[SUM_NAME].figure, lines[PERCENT_NAME].figure
    assert (total.value, total.at_least, total.quarters[1]) == (700.0, 500.0, 230.0)
    assert share.value == pytest.approx(50 / 150 * 100) and share.coverage == "counted 2 of 3 PDs"
    html = text(response)
    for shown in (
        "at least 500",
        "700",
        "33.3%",
        "Reporting 2 of 2",
        "counted 2 of 3 PDs",
        "Reported without numerator and denominator",
        "Suggested, not confirmed",
        "Now in section Child Protection",
        "40 / 100",
        "ED-1.1",
        "1.1 Children access learning",
    ):
        assert shown in html, shown
    # the tiles: three masters, the suggestion waiting, the proposal and the indicators in a master
    tab = response.context["tab"]
    assert (tab.status_counts["no_target"], tab.suggested, len(tab.proposals), tab.proposed_links) == (
        0,
        1,
        1,
        2,
    )
    assert (tab.in_master, tab.indicators) == (5, 9)  # Education's PD indicators of the year
    assert "5 of 9" in html and 'id="pm-masters-data"' in html


def test_every_pd_indicator_in_no_master_is_listed_with_what_neurodb_says_of_it(client_viewer, world):
    tab = _tab(client_viewer).context["tab"]
    notes = {u.row.title: u.note for u in tab.unattached}
    assert notes[TEACHERS] == "No suggestion yet"
    assert notes[CAREGIVERS] == "In the proposed master indicator Caregivers trained on positive parenting"
    assert notes["# of children enrolled in NFE classes"].startswith(
        "Suggested for Children enrolled in NFE (0.95)"
    )
    assert tab.unattached_count == len(tab.unattached) == 4


def test_two_masters_that_fit_as_well_are_said_ambiguous(client_viewer, world):
    w.master("Teachers trained on inclusive education")
    w.master("Teachers trained in inclusive education")
    notes = {u.row.title: u.note for u in _tab(client_viewer).context["tab"].unattached}
    assert notes[TEACHERS].startswith("Ambiguous: Teachers trained")


def test_status_search_and_cp_output_choose_master_rows_never_their_values(client_viewer, world):
    every = {name: line.figure.value for name, line in _lines(_tab(client_viewer)).items()}
    for params, names in (
        ({"q": "passing"}, {PERCENT_NAME}),
        ({"q": "ED-1.1"}, {SUM_NAME}),
        ({"cp_output": "1.2 children"}, {PERCENT_NAME}),
        ({"status": "on_track"}, {SUM_NAME, MOVED_NAME}),  # 70 % and 18 % of the year's share of the target
        ({"status": "off_track"}, {PERCENT_NAME}),  # a level: 33.3 % of a target of 50 %
        ({"gender": "Girls", "pd_status": "active", "office": "Beirut"}, set(every)),
    ):
        response = _tab(client_viewer, **params)
        lines = _lines(response)
        assert set(lines) == names, params
        assert {name: line.figure.value for name, line in lines.items()} == {n: every[n] for n in names}
    note = text(_tab(client_viewer, gender="Girls", office="Beirut"))
    assert "Not applied here: office, gender." in note


def test_a_partner_filter_shows_its_contribution_with_a_share_of_sums_only(client_viewer, world):
    a = world["pds"][0]
    lines = _lines(_tab(client_viewer, partner=str(a.partner_id)))
    total, share = lines[SUM_NAME], lines[PERCENT_NAME]
    assert (total.figure.value, total.whole.value) == (500.0, 700.0)
    assert total.share == pytest.approx(500 / 700 * 100)
    assert [s.partner for s in total.subs] == ["Partner A"]
    assert share.figure.value == pytest.approx(40.0) and share.share is None
    assert share.whole.status == _expected()[PERCENT_NAME].status  # the master's own status
    html = text(_tab(client_viewer, partner=str(a.partner_id)))
    assert "71% of 700" in html


def test_a_location_filter_reads_reacheds_own_and_never_calls_an_indicator_gone(client_viewer, world):
    lines = _lines(_tab(client_viewer, location="Bekaa"))
    total = lines[SUM_NAME]
    assert total.figure.value == 50.0  # A's Bekaa report; B plans Akkar only
    assert [s.pd_number for s in total.subs] == ["PD/2026/001"]
    assert all(s.sub.reason != masters.GONE for line in lines.values() for s in line.subs)


def test_the_humanitarian_reports_never_call_an_indicator_without_them_gone(client_viewer, world):
    """Reports = humanitarian reads high-frequency and cluster PD indicators only. A confirmed PD indicator that
    is neither, in a programme document that has one, is still in eTools: like a location filter, it is not
    listed (on the tab, in the detail and in the export); one that has left eTools still reads so."""
    a, b, _c = world["pds"]
    enrolled = world["enrolled"]
    outreach = "# of children reached with emergency learning kits"
    w.indicator(a, 161, outreach, is_high_frequency=True)
    w.report(a, 161, outreach, "HR1", day(1, 31), {"Akkar": 40}, kind="HR")
    w.link(enrolled, a, 161)
    clustered = "# of children enrolled in cluster NFE classes"
    w.indicator(b, 261, clustered, cluster_name="Education")
    w.report(b, 261, clustered, "HR1", day(1, 31), {"Akkar": 25}, kind="HR")
    w.link(enrolled, b, 261)
    w.link(enrolled, a, 199)  # its indicator is no longer in eTools
    gone = "No longer in eTools"

    lines = _lines(_tab(client_viewer, report_type="HR"))
    total = lines[SUM_NAME]
    assert {s.link.ident: s.sub.reason for s in total.subs} == {"s161": "", "s261": "", "s199": masters.GONE}
    assert total.figure.value == 65.0
    # PD/2026/001's quarterly percentage and its indicator moved to Child Protection are not humanitarian
    assert not lines[PERCENT_NAME].subs and not lines[MOVED_NAME].subs
    reasons = [s.sub.reason for line in lines.values() for s in line.subs]
    assert reasons.count(masters.GONE) == 1

    url = reverse("reports:etools_master", args=[enrolled.pk])
    modal = text(client_viewer.get(url, {"year": "2026", "report_type": "HR"}, **HX))
    assert modal.count(gone) == 1 and "PD/2026/002" in modal

    response = client_viewer.get(
        PAGE, {"view": "masters", "section": w.EDUCATION, "report_type": "HR", "export": "csv"}
    )
    rows = list(csv.DictReader(io.StringIO(b"".join(response.streaming_content).decode("utf-8-sig"))))
    subs = [r for r in rows if r["level"] == "sub"]
    assert [r["counted"] for r in subs].count(gone) == 1
    assert {r["indicator"] for r in subs} >= {outreach, clustered}
    assert not {ENROLLED, PASSING, KITS} & {r["indicator"] for r in subs}

    # the quarterly reports read the same links: the others count again
    quarterly_lines = _lines(_tab(client_viewer, report_type="QPR"))
    assert {"s101", "s201"} <= {s.link.ident for s in quarterly_lines[SUM_NAME].subs}
    assert quarterly_lines[PERCENT_NAME].subs and quarterly_lines[MOVED_NAME].subs


def test_subs_in_another_section_are_read_with_one_more_read_by_programme_document(
    client_viewer, world, monkeypatch
):
    calls: list[monitoring.Filters] = []
    real = monitoring.stored_indicators

    def counted(filters, *args, **kwargs):
        calls.append(filters)
        return real(filters, *args, **kwargs)

    monkeypatch.setattr(monitoring, "stored_indicators", counted)
    response = get_part(client_viewer, PAGE + f"?view=masters&section={w.EDUCATION}", "masters-table")
    assert response.status_code == 200 and response["X-NeuroDB-Part"] == "masters-table"
    assert len(calls) == 2  # the section's rows, then the moved indicator's programme document
    assert calls[0].sections == [w.EDUCATION] and calls[0].scope == "year" and not calls[0].pd_ids
    assert calls[1].pd_ids == [world["pds"][0].pk] and not calls[1].sections
    moved = _lines(_tab(client_viewer))[MOVED_NAME]
    assert moved.figure.value == 90.0 and moved.subs[0].sub.note == "Now in section Child Protection"


def test_each_part_loads_on_its_own(client_viewer, world):
    for name in ("masters-summary", "masters-table"):
        response = get_part(client_viewer, PAGE + f"?view=masters&section={w.EDUCATION}", name)
        assert response.status_code == 200 and response["X-NeuroDB-Part"] == name


# -------------------------------------------------------------------------------------------------- the detail
def test_the_detail_shows_how_it_is_worked_out_its_quarters_and_subs(client_viewer, world):
    m = world["enrolled"]
    url = reverse("reports:etools_master", args=[m.pk]) + "?year=2026"
    modal = client_viewer.get(url, **HX)
    html = text(modal)
    assert modal.status_code == 200 and 'class="modal-header"' in html
    assert "Up to 700: the sum over the programme documents" in html and "At least 500" in html
    assert 'id="master-quarters"' in html and "230" in html
    assert "PD/2026/001 and PD/2026/002 both plan Akkar" in html
    assert "for reference only" in html and "1,600" in html  # the PDs' own targets
    assert "In the total" in html and "Suggested" in html
    page = client_viewer.get(url)
    assert page.status_code == 200 and "<title>Children enrolled in NFE" in text(page)
    dm.EtoolsMaster.objects.filter(pk=m.pk).update(status="dismissed")
    assert client_viewer.get(url).status_code == 404


# -------------------------------------------------------------------------------------------------- the export
@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_the_export_has_one_row_per_master_and_per_sub_and_no_person(client_viewer, world, fmt):
    response = client_viewer.get(PAGE, {"view": "masters", "section": w.EDUCATION, "export": fmt})
    assert response.status_code == 200
    if fmt == "csv":
        body = b"".join(response.streaming_content).decode("utf-8-sig")
        rows = list(csv.DictReader(io.StringIO(body)))
        header = body.splitlines()[0].split(",")
    else:
        sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
        values = list(sheet.values)
        header = list(values[0])
        rows = [
            {k: ("" if v is None else str(v)) for k, v in zip(header, row, strict=True)} for row in values[1:]
        ]
        body = "\n".join(",".join("" if v is None else str(v) for v in row) for row in values)
    assert tuple(header) == masters_tab.EXPORT_COLUMNS
    assert Counter(r["level"] for r in rows) == {"master": 3, "sub": 7}
    master = next(r for r in rows if r["level"] == "master" and r["master"] == SUM_NAME)
    assert (master["code"], master["method"], master["coverage"]) == ("ED-1.1", "Sum", "counted 2 of 2 PDs")
    assert float(master["value"]) == 700.0 and float(master["at_least"]) == 500.0
    assert float(master["to_q1"]) == 230.0 and master["status"] == "On track"
    left = next(
        r
        for r in rows
        if r["level"] == "sub" and r["programme_document"] == "PD/2026/003" and r["master"] == PERCENT_NAME
    )
    assert left["counted"] == "Reported without numerator and denominator" and left["partner"] == "Partner C"
    suggested = next(
        r
        for r in rows
        if r["level"] == "sub" and r["counts_as"] and r["counted"] == "Suggested, not confirmed"
    )
    assert suggested["programme_document"] == "PD/2026/003"
    for word in (PERSON, "updated_by", "confirmed_by", "@"):
        assert word not in body, word


def test_the_export_of_a_partner_adds_its_share(client_viewer, world):
    a = world["pds"][0]
    response = client_viewer.get(
        PAGE, {"view": "masters", "section": w.EDUCATION, "partner": str(a.partner_id), "export": "csv"}
    )
    rows = list(csv.DictReader(io.StringIO(b"".join(response.streaming_content).decode("utf-8-sig"))))
    total = next(r for r in rows if r["level"] == "master" and r["master"] == SUM_NAME)
    assert float(total["partner_share_percent"]) == pytest.approx(71.4)


# ---------------------------------------------------------------------------------------------- who may change
def _confirm(client, world, **extra):
    m, link = world["enrolled"], world["suggestion"]
    return client.post(
        reverse("reports:etools_master_edit", args=[m.pk]), {"action": "confirm", "link": link.pk, **extra}
    )


def test_viewers_editors_of_another_section_and_donors_are_refused(client_viewer, world):
    link = world["suggestion"]
    assert _confirm(client_viewer, world).status_code == 403
    other = _signed_in(_editor("Child protection", "CP", w.CP))
    assert _confirm(other, world).status_code == 403
    donor = User.objects.create_user(username="eu-donor", password="donor-pass-123456")
    DonorAccount.objects.create(user=donor, name="Donor A", donors=["Donor A"], must_change_password=False)
    refused = _confirm(_signed_in(donor), world)
    assert refused.status_code in (302, 403) and reverse("reports:pd_monitoring") not in refused.get(
        "Location", ""
    )
    for client in (client_viewer, other):
        assert (
            client.post(reverse("reports:etools_master_attach"), {"master": world["enrolled"].pk}).status_code
            == 403
        )
        assert (
            client.post(reverse("reports:etools_masters_confirm"), {"section": w.EDUCATION}).status_code
            == 403
        )
    link.refresh_from_db()
    assert link.status == "suggested"
    # the buttons are not drawn for them
    assert reverse("reports:etools_master_edit", args=[world["enrolled"].pk]) not in text(_tab(client_viewer))


@pytest.mark.parametrize("who", ["editor", "admin"])
def test_a_section_editor_of_the_section_and_an_administrator_confirm_and_it_shows_at_once(
    world, admin_user, settings, who
):
    settings.DEBUG = False  # the PD indicators stored for every worker, as in production
    user = _editor("Education", "ED", w.EDUCATION) if who == "editor" else admin_user
    client = _signed_in(user)
    before = _lines(_tab(client))[SUM_NAME]
    assert before.figure.value == 700.0
    assert reverse("reports:etools_master_edit", args=[world["enrolled"].pk]) in text(_tab(client))
    back = PAGE + f"?view=masters&section={w.EDUCATION}"
    response = _confirm(client, world, next=back)
    assert response.status_code == 302 and response["Location"] == back
    link = world["suggestion"]
    link.refresh_from_db()
    assert link.status == "confirmed" and link.unit_class == "number"
    after = _lines(_tab(client))[SUM_NAME]
    assert (
        after.figure.value == 770.0 and after.figure.counted_pds == 3
    )  # the stored rows unchanged, the link read anew


def test_a_section_editor_curates_on_the_page_without_the_administration(world):
    client = _signed_in(_editor("Education", "ED", w.EDUCATION))
    enrolled, proposal = world["enrolled"], world["proposal"]
    edit = reverse("reports:etools_master_edit", args=[enrolled.pk])
    # the year's code and target, and the name
    client.post(
        edit,
        {
            "action": "settings",
            "year": "2026",
            "name": "Children enrolled in non-formal education",
            "code": "ED-9",
            "target": "2,000",
        },
    )
    enrolled.refresh_from_db()
    settings_row = enrolled.years.get(year=2026)
    assert (enrolled.name, settings_row.code, settings_row.target) == (
        "Children enrolled in non-formal education",
        "ED-9",
        2000.0,
    )
    client.post(
        edit, {"action": "settings", "year": "2026", "name": enrolled.name, "code": "ED-9", "target": ""}
    )
    assert _lines(_tab(client))[enrolled.name].whole.status == "no_target"
    # a role, then taking a sub out
    b_link = enrolled.links.get(ident="s201")
    client.post(edit, {"action": "role", "link": b_link.pk, "role": "shown"})
    b_link.refresh_from_db()
    assert b_link.role == "shown"
    assert _lines(_tab(client))[enrolled.name].figure.value == 500.0
    client.post(edit, {"action": "remove", "link": b_link.pk})
    assert not enrolled.links.filter(ident="s201").exists()
    # attach the indicator in no master, as a total, confirmed at once
    b = world["pds"][1]
    teachers = w.master("Teachers trained")
    client.post(
        reverse("reports:etools_master_attach"),
        {"master": teachers.pk, "role": "total", "key": [f"{b.pk}|s202"]},
    )
    assert teachers.links.get().status == "confirmed"
    # accept the proposal, as its row posts it (with the PD indicators it shows): active, they are confirmed
    accept_url = reverse("reports:etools_master_edit", args=[proposal.pk])
    (accept,) = _posted(text(_tab(client)), accept_url)
    client.post(accept_url, {**accept, "action": "accept"})
    proposal.refresh_from_db()
    assert proposal.status == "active" and set(proposal.links.values_list("status", flat=True)) == {
        "confirmed"
    }
    # who did it is never on the page
    assert PERSON not in text(_tab(client))


def test_dismiss_keeps_the_proposal_dismissed_and_a_percentage_refuses_a_number_total(world, admin_user):
    client = _signed_in(admin_user)
    proposal, passing = world["proposal"], world["passing"]
    client.post(reverse("reports:etools_master_edit", args=[proposal.pk]), {"action": "dismiss"})
    proposal.refresh_from_db()
    assert proposal.status == "dismissed" and not proposal.links.exists()
    assert not _tab(client).context["tab"].proposals
    # a number as a percentage master's total: left out and said why
    b = world["pds"][1]
    response = client.post(
        reverse("reports:etools_master_attach"),
        {"master": passing.pk, "role": "total", "key": [f"{b.pk}|s202"], "next": PAGE},
        follow=True,
    )
    assert not passing.links.filter(ident="s202").exists()
    assert "left out" in text(response)


def _posted(html: str, action: str) -> list[dict[str, list[str]]]:
    """The hidden fields each form posting to ``action`` sends, as a browser sends them."""
    out = []
    for form in re.findall(r'<form method="post" action="' + re.escape(action) + r'".*?</form>', html, re.S):
        fields: dict[str, list[str]] = {}
        for name, value in re.findall(r'<input type="hidden" name="([a-z_]+)" value="([^"]*)"', form):
            fields.setdefault(name, []).append(value)
        out.append(fields)
    return out


def test_confirm_the_strong_matches_of_a_section(world, admin_user):
    client = _signed_in(admin_user)
    html = text(_tab(client))
    assert "strong match" in html
    (form,) = _posted(html, reverse("reports:etools_masters_confirm"))
    assert form["link"] == [str(world["suggestion"].pk)]
    client.post(reverse("reports:etools_masters_confirm"), form)
    link = world["suggestion"]
    link.refresh_from_db()
    assert link.status == "confirmed"


def _strong(master, pd, source_id, title):
    w.indicator(pd, source_id, title)
    link = w.link(master, pd, source_id, title=title, status="suggested", unit_class="")
    dm.EtoolsMasterLink.objects.filter(pk=link.pk).update(score=0.95)
    return link


def test_the_confirm_buttons_confirm_the_suggestions_the_tab_shows_and_nothing_else(world, admin_user):
    """With a partner filter the buttons count that partner's suggestions: they confirm those, never
    another partner's the page does not show; a search hides a master, and its suggestions with it."""
    a, b, _c = world["pds"]
    enrolled, unseen = world["enrolled"], world["suggestion"]  # Partner C's
    mine = _strong(enrolled, a, 102, "# of girls enrolled in NFE")
    teachers = w.master("Teachers trained", target=40)
    hidden = _strong(teachers, b, 203, "# of teachers trained on child safeguarding")
    client = _signed_in(admin_user)
    edit = reverse("reports:etools_master_edit", args=[enrolled.pk])
    response = _tab(client, partner=str(a.partner_id))
    line = _lines(response)[SUM_NAME]
    assert (line.suggested, line.suggested_links) == (1, [mine.pk])
    html = text(response)
    assert "Confirm the suggestion" in html and "Confirm the 1 strong match (0.9 or more)" in html
    (confirm_all,) = [f for f in _posted(html, edit) if f.get("action") == ["confirm_all"]]
    assert confirm_all["link"] == [str(mine.pk)]
    done = client.post(edit, confirm_all, follow=True)
    assert "1 PD indicator(s) confirmed: they count from now on." in text(done)
    statuses = dict(
        dm.EtoolsMasterLink.objects.filter(pk__in=[mine.pk, unseen.pk]).values_list("pk", "status")
    )
    assert statuses == {mine.pk: "confirmed", unseen.pk: "suggested"}

    html = text(_tab(client, q="enrolled"))  # Teachers trained is not shown
    (section,) = _posted(html, reverse("reports:etools_masters_confirm"))
    assert section["link"] == [str(unseen.pk)] and "Confirm the 1 strong match (0.9 or more)" in html
    done = client.post(reverse("reports:etools_masters_confirm"), section, follow=True)
    assert "1 PD indicator(s) confirmed: they count from now on." in text(done)
    unseen.refresh_from_db()
    hidden.refresh_from_db()
    assert (unseen.status, hidden.status) == ("confirmed", "suggested")
    # nothing shown, nothing confirmed: a post without the links confirms none
    nothing = client.post(reverse("reports:etools_masters_confirm"), {"section": w.EDUCATION}, follow=True)
    assert "None of the strong matches shown is waiting any more." in text(nothing)
    hidden.refresh_from_db()
    assert hidden.status == "suggested"


def test_accept_confirms_the_pd_indicators_the_tab_showed_and_none_proposed_since(world, admin_user):
    """The night (or Suggest now) adds a PD indicator to a proposal while the page is open: Accept confirms
    the ones the page listed, as its count said, and leaves the new one as a suggestion of the master, said
    so. A post that names none of them (a page from before) accepts nothing."""
    proposal = world["proposal"]
    _a, _b, c = world["pds"]
    client = _signed_in(admin_user)
    edit = reverse("reports:etools_master_edit", args=[proposal.pk])
    html = text(_tab(client))
    (form,) = _posted(html, edit)
    shown = sorted(proposal.links.values_list("pk", flat=True))
    assert sorted(map(int, form["link"])) == shown and "2 PD indicators in 2 programme documents" in html
    w.indicator(c, 341, CAREGIVERS)
    late = w.link(proposal, c, 341, status="suggested", unit_class="")

    stale = client.post(edit, {"action": "accept", "next": PAGE}, follow=True)
    assert "None of the PD indicators shown is in the proposal any more" in text(stale)
    proposal.refresh_from_db()
    assert proposal.status == "proposed"

    done = client.post(edit, {**form, "action": "accept"}, follow=True)
    said = text(done)
    assert "is active: 2 PD indicator(s) confirmed, they count from now on." in said
    assert "1 PD indicator(s) added to the proposal after the page was shown are left as suggestions" in said
    proposal.refresh_from_db()
    assert proposal.status == "active"
    statuses = dict(proposal.links.values_list("pk", "status"))
    assert statuses == {shown[0]: "confirmed", shown[1]: "confirmed", late.pk: "suggested"}


def test_a_pd_indicator_without_an_etools_id_is_attached_from_the_tab_and_counts(world, admin_user):
    b = world["pds"][1]
    title = "# of children enrolled in catch-up classes"
    w.indicator(b, None, title)
    w.report(b, None, title, "QPR1", day(3, 31), {"Akkar": 40})
    catch_up = w.master("Children in catch-up classes", target=100)
    client = _signed_in(admin_user)
    (row,) = [u for u in _tab(client).context["tab"].unattached if u.row.title == title]
    assert f'name="key" value="{row.key}"' in text(_tab(client))
    response = client.post(
        reverse("reports:etools_master_attach"),
        {"master": catch_up.pk, "role": "total", "key": [row.key], "next": PAGE},
        follow=True,
    )
    assert "1 PD indicator(s) attached to “Children in catch-up classes”" in text(response)
    assert catch_up.links.get().status == "confirmed"
    assert _lines(_tab(client))["Children in catch-up classes"].figure.value == 40.0


def test_an_indicator_rejected_for_a_master_is_attached_to_it_again(world, admin_user):
    enrolled, link = world["enrolled"], world["suggestion"]
    c = world["pds"][2]
    client = _signed_in(admin_user)
    edit = reverse("reports:etools_master_edit", args=[enrolled.pk])
    client.post(edit, {"action": "reject", "link": link.pk})
    link.refresh_from_db()
    assert link.status == "rejected"
    notes = {u.key: u.note for u in _tab(client).context["tab"].unattached}
    assert notes[f"{c.pk}|s301"] == "No suggestion yet"  # listed in no master, never suggested here again
    response = client.post(
        reverse("reports:etools_master_attach"),
        {"master": enrolled.pk, "role": "total", "key": [f"{c.pk}|s301"], "next": PAGE},
        follow=True,
    )
    assert "1 PD indicator(s) attached" in text(response)
    link.refresh_from_db()
    assert (link.status, link.role) == ("confirmed", "total")  # the same link, a person's choice now
    assert _lines(_tab(client))[SUM_NAME].figure.value == 770.0


def test_a_digit_that_is_no_number_gets_the_pages_own_message_never_a_server_error(world, admin_user):
    """The character "²" is a digit to str.isdigit() and no number to int(): a hand-edited form that sends it
    in a number's field gets the message any wrong value gets, and the tab's partner filter reads it as no
    partner."""
    client = _signed_in(admin_user)
    enrolled, proposal = world["enrolled"], world["proposal"]
    b = world["pds"][1]
    edit = reverse("reports:etools_master_edit", args=[enrolled.pk])
    attach = reverse("reports:etools_master_attach")
    sent = [
        (edit, {"action": "confirm", "link": "²"}, "Choose one of this master indicator's PD indicators."),
        (edit, {"action": "settings", "year": "²", "name": enrolled.name}, "Choose the year these settings"),
        (edit, {"action": "confirm_all", "link": "²"}, "None of the suggestions shown is waiting any more."),
        (
            reverse("reports:etools_master_edit", args=[proposal.pk]),
            {"action": "accept", "link": "²"},
            "None of the PD indicators shown is in the proposal any more",
        ),
        (attach, {"master": "²", "role": "total", "key": [f"{b.pk}|s202"]}, "Choose the master indicator"),
        (attach, {"master": enrolled.pk, "role": "total", "key": ["²|s202"]}, "Tick the PD indicators"),
        (
            reverse("reports:etools_masters_confirm"),
            {"section": w.EDUCATION, "link": "²"},
            "None of the strong matches shown is waiting any more.",
        ),
    ]
    for url, data, message in sent:
        response = client.post(url, data)
        assert response.status_code == 302, data
        assert any(message in str(m) for m in get_messages(response.wsgi_request)), data
    proposal.refresh_from_db()
    assert proposal.status == "proposed" and enrolled.links.filter(status="suggested").count() == 1
    assert curation.parse_key("²|s202") is None
    tab = _tab(client, partner="²")
    assert tab.status_code == 200 and not tab.context["tab"].partner


@pytest.mark.parametrize(
    "typed, shown", [("1234567", "1234567"), ("123456.7", "123456.7"), ("2,000", "2000")]
)
def test_the_target_box_shows_every_digit_so_saving_it_unchanged_keeps_the_target(world, typed, shown):
    client = _signed_in(_editor("Education", "ED", w.EDUCATION))
    enrolled = world["enrolled"]
    edit = reverse("reports:etools_master_edit", args=[enrolled.pk])
    client.post(
        edit, {"action": "settings", "year": "2026", "name": enrolled.name, "code": "ED-1.1", "target": typed}
    )
    stored = enrolled.years.get(year=2026).target
    html = text(_tab(client))
    boxes = re.findall(r'name="target" value="([^"]*)"', html)
    assert shown in boxes and not any("e+" in box for box in boxes)
    assert _lines(_tab(client))[SUM_NAME].target_input == shown
    client.post(
        edit,
        {"action": "settings", "year": "2026", "name": "Children in NFE", "code": "ED-1.1", "target": shown},
    )
    assert enrolled.years.get(year=2026).target == stored == float(typed.replace(",", ""))


def test_a_filter_that_hides_every_master_says_so_and_the_administration_only_to_who_can_open_it(
    world, admin_user
):
    admin = text(_tab(_signed_in(admin_user), q="zzz"))
    assert "No master indicator matches these filters" in admin
    assert "3 master indicators are listed in 2026: clear the search" in admin
    assert "No master indicator in this year" not in admin
    editor = _signed_in(_editor("Education", "ED", w.EDUCATION))
    assert "No master indicator matches these filters" in text(_tab(editor, partner="999999"))
    empty_year = text(_tab(editor, year="2027"))
    assert "No master indicator in this year" in empty_year
    assert "ask an Administrator to make one" in empty_year and "Reporting setup" not in empty_year
    assert "make one in the administration (Reporting setup)" in text(
        _tab(_signed_in(admin_user), year="2027")
    )


def test_a_proposals_links_are_confirmed_by_accepting_it_and_one_stuck_in_its_total_says_so(
    world, admin_user
):
    proposal = world["proposal"]
    client = _signed_in(admin_user)
    edit = reverse("reports:etools_master_edit", args=[proposal.pk])
    link = proposal.links.order_by("id").first()
    response = client.post(edit, {"action": "confirm", "link": link.pk, "next": PAGE}, follow=True)
    assert "Accept this proposed master indicator first" in text(response)
    link.refresh_from_db()
    assert link.status == "suggested"
    # one confirmed there before the rule: it counts nowhere, and the tab says why
    dm.EtoolsMasterLink.objects.filter(pk=link.pk).update(status="confirmed")
    notes = {(u.row.pd.id, u.row.title): u.note for u in _tab(client).context["tab"].unattached}
    assert notes[(link.pd_id, CAREGIVERS)] == (
        "In the total of Caregivers trained on positive parenting, which is proposed: it counts nowhere until "
        "that master indicator is active, or the PD indicator is taken out of it."
    )


def test_the_detail_words_the_pd_targets_by_method_and_three_pds_all_plan_a_place(client_viewer, world):
    url = reverse("reports:etools_master", args=[world["passing"].pk]) + "?year=2026"
    html = text(client_viewer.get(url, **HX))
    assert (
        "The programme documents' own targets come to 60% (their target numerators over their target" in html
    )
    assert "add up to 60" not in html
    master_links.confirm([world["suggestion"]], "staff")  # Partner C plans Akkar too
    url = reverse("reports:etools_master", args=[world["enrolled"].pk]) + "?year=2026"
    html = text(client_viewer.get(url, **HX))
    assert "PD/2026/001, PD/2026/002 and PD/2026/003 all plan Akkar" in html and "both plan" not in html


@pytest.mark.parametrize(
    "field, method, reasons",
    [
        ("calculation_across_periods", "avg", masters.PERIODS_COMBINED),
        ("calculation_across_locations", "avg", masters.LOCATIONS_COMBINED),
        ("calculation_across_locations", "max", masters.LOCATIONS_COMBINED),
    ],
    ids=["averages its periods", "averages its locations", "takes its largest location"],
)
def test_a_percentage_whose_pd_averages_its_periods_says_why_on_the_tab_and_in_the_export(
    client_viewer, world, field, method, reasons
):
    """Its reports carry their numerator and denominator at each location, but its PD averages its periods,
    or its locations (or takes the largest): the row says so ("not combined", never "not reported"), and so
    does the export."""
    a, b, _c = world["pds"]
    title = "% of children attending regularly"
    for pd, source_id in ((a, 171), (b, 271)):
        w.indicator(pd, source_id, title, display_type="percentage", places=("Akkar", "Bekaa"))
        for number, end, value in (
            ("QPR1", day(3, 31), w.pct(30, 100)),
            ("QPR2", day(6, 30), w.pct(50, 100)),
        ):
            places = {"Akkar": value, "Bekaa": w.pct(10, 100)}
            w.report(pd, source_id, title, number, end, places, **{field: method})
    regular = w.master("Children attending regularly", "percent", target=60)
    for pd, source_id in ((a, 171), (b, 271)):
        w.link(regular, pd, source_id, unit_class="percentage")
    reason = reasons[method]

    response = _tab(client_viewer)
    line = _lines(response)["Children attending regularly"]
    assert line.figure.value is None and line.figure.coverage == "counted 0 of 2 PDs"
    assert [s.sub.reason for s in line.subs] == [reason, reason] and all(
        s.periods_combined for s in line.subs
    )
    html = text(response)
    table = html[html.index(f'id="master-{regular.pk}-subs-table"') :]
    table = table[: table.index("</table>")]
    assert table.count("not combined") == 2 and "not reported" not in table and reason in table
    assert "not reported" in html  # PD/2026/003 sent its percentage alone
    response = client_viewer.get(PAGE, {"view": "masters", "section": w.EDUCATION, "export": "csv"})
    rows = list(csv.DictReader(io.StringIO(b"".join(response.streaming_content).decode("utf-8-sig"))))
    mine = [r for r in rows if r["level"] == "sub" and r["master"] == "Children attending regularly"]
    assert [r["counted"] for r in mine] == [reason, reason]


def _row(html: str, master) -> str:
    """The tab's row of one master, from its cells to the end of the row."""
    found = html[html.index(f'<tr id="master-{master.pk}">') :]
    return found[: found.index("</tr>")]


def test_an_average_of_percentages_and_a_percentage_masters_target_read_with_a_percent_sign(
    client_viewer, world
):
    a, b, _c = world["pds"]
    title = "% of children completing catch-up classes"
    for pd, source_id, value in ((a, 181, w.pct(40, 100)), (b, 281, w.pct(10, 50))):
        w.indicator(pd, source_id, title, display_type="percentage")
        w.report(pd, source_id, title, "QPR1", day(3, 31), {"Akkar": value})
    average = w.master("Children completing catch-up classes", "average", target=50)
    for pd, source_id in ((a, 181), (b, 281)):
        w.link(average, pd, source_id, unit_class="percentage")

    response = _tab(client_viewer)
    lines = _lines(response)
    assert lines["Children completing catch-up classes"].figure.value == 30.0  # (40 + 20) / 2
    assert lines["Children completing catch-up classes"].percentage and lines[PERCENT_NAME].percentage
    assert not lines[SUM_NAME].percentage
    row = _row(text(response), average)
    assert (
        '<td class="num">50%</td>' in row
        and '<td class="num">30%</td>' in row
        and "<strong>30%</strong>" in row
    )
    assert '<td class="num">1,000</td>' in _row(text(response), world["enrolled"])
    for master, value in ((average, "30.0%"), (world["passing"], "33.3%")):
        url = reverse("reports:etools_master", args=[master.pk]) + "?year=2026"
        modal = text(client_viewer.get(url, **HX))
        assert f'<div class="kpi__value">{value}' in modal
        assert '<div class="kpi__value">50%' in modal  # the year's target, as the tab's row reads it
    modal = text(client_viewer.get(reverse("reports:etools_master", args=[average.pk]) + "?year=2026", **HX))
    assert '<td class="num">30%</td>' in modal  # its quarters


@pytest.mark.parametrize("who", ["admin", "editor", "viewer"])
def test_with_no_master_and_no_proposal_whoever_can_make_one_is_told_how(db, roles, admin_user, viewer, who):
    a = w.programme("A", 1)
    w.indicator(a, 101, ENROLLED)
    w.report(a, 101, ENROLLED, "QPR1", day(3, 31), {"Akkar": 10})
    user = {"admin": admin_user, "viewer": viewer}.get(who) or _editor("Education", "ED", w.EDUCATION)
    html = text(_tab(_signed_in(user)))
    assert "No master indicator in this year" in html and "Accept a proposed" not in html
    told = {
        "admin": "Make one in the administration (Reporting setup).",
        "editor": "Ask an Administrator to make one.",
        "viewer": "Change the filters, or ask your section",
    }
    for person, words in told.items():
        assert (words in html) is (person == who), (person, who)


def test_a_master_switched_off_for_the_year_lists_its_pd_indicators_in_no_master(client_viewer, world):
    """Unticking "shown this year": the master is not listed, so its PD indicators count nowhere that year.
    They are listed under Not in a master, saying why, and the tile counts them out."""
    a, b, c = world["pds"]
    dm.EtoolsMasterYear.objects.filter(master=world["enrolled"], year=2026).update(is_active=False)
    response = _tab(client_viewer)
    tab = response.context["tab"]
    assert SUM_NAME not in _lines(response) and tab.listed == 2
    notes = {(u.row.pd.id, u.row.title): u.note for u in tab.unattached}
    stuck = (
        "In the total of Children enrolled in NFE, which is not shown in 2026: it counts nowhere in 2026 until "
        "that master indicator is shown this year again, or the PD indicator is taken out of it."
    )
    assert notes[(a.pk, ENROLLED)] == notes[(b.pk, BOTH)] == stuck
    assert notes[(c.pk, "# of children enrolled in NFE classes")] == (
        "Suggested for Children enrolled in NFE (0.95), which is not shown in 2026"
    )
    assert (tab.in_master, tab.indicators, tab.unattached_count) == (3, 9, 6)
    assert world["enrolled"].pk not in {m.pk for m in tab.attach}
    html = text(response)
    assert "3 of 9" in html and stuck in html
    # every PD indicator of the year is in a listed master or listed in no master, never both
    year = {masters.key_of(row) for row in w.rows(sections=[w.EDUCATION])}
    linked = {(s.pd_id, s.link.ident) for line in tab.lines for s in line.subs if s.sub.status == "confirmed"}
    listed = {masters.key_of(u.row) for u in tab.unattached}
    assert len(linked & year) == tab.in_master and (linked & year) | listed == year and not linked & listed
    dm.EtoolsMasterYear.objects.filter(master=world["enrolled"], year=2026).update(is_active=True)
    assert _tab(client_viewer).context["tab"].in_master == 5
