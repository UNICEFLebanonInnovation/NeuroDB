"""The eTools Power BI views on NeuroDB's pages (SPEC6 R3): Program Performance and Data quality on Partner
monitoring, Reporting Progress, Partner Performance, the report table and Report/Location on Progress
reports, and the flat PD/SPD Indicators list on Programmes. The same filter gives the same figures on every
view, in their exports and in Ask NeuroDB; nothing is added across indicators; no e-mail address is shown;
donors never reach them."""

# ruff: noqa: F401, F811 (the partner monitoring fixture "data" comes from test_pd_monitoring)
from __future__ import annotations

import csv
import datetime
import io
from decimal import Decimal

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from neurodb.assistant import tools
from neurodb.datamart import models as dm
from neurodb.datamart import monitoring, reporting, services
from neurodb.partnerships import services as partnerships
from neurodb.partnerships import statuses
from neurodb.partnerships.models import PCA, PartnerOrganization
from tests.reports.test_pd_monitoring import TODAY, _FrozenDatetime, data, report_rows

pytestmark = pytest.mark.django_db
HX = {"HX-Request": "true"}
CLASSROOMS = "# of classrooms rehabilitated"


@pytest.fixture
def frozen(monkeypatch):
    """Today is 1 July 2026 for partner monitoring and the Progress reports page."""
    monkeypatch.setattr(monitoring, "datetime", _FrozenDatetime)
    monkeypatch.setattr(services, "datetime", _FrozenDatetime)


@pytest.fixture
def more(data):
    """The fixture's PDs and reports, plus: an Education indicator without a target or a report in the
    active PD, its PD's attributes (office, CP output, focal points), the closed PD's attributes and an
    overdue report of the closed PD."""
    pd, closed = data["pd"], data["closed"]
    PCA.objects.filter(pk=pd.pk).update(
        offices_set=["Beirut"],
        cp_outputs=["3.1 Children at risk reached with protection services"],
        section_names=["Child Protection"],
        unicef_focal_points=["Focal Point One", "focal.two@example.org", "Focal Three (three@example.org)"],
    )
    PCA.objects.filter(pk=closed.pk).update(document_type="SSFA", offices_set=["Zahle"])
    PartnerOrganization.objects.filter(pk=closed.partner_id).update(cso_type="National")
    dm.PDIndicator.objects.create(
        datamart_id=20,
        source_id=500,
        intervention=pd,
        pd_reference_number=pd.number,
        title=CLASSROOMS,
        section_name="Education",
        lower_result_name="2.1 Learning spaces",
        target_numerator=None,
        location_name="Akkar",
    )
    dm.PDIndicator.objects.filter(source_id=100).update(cp_output="1.2 Psychosocial support for children")
    report_rows(
        closed,
        closed.partner,
        "# of schools",
        "PR-9",
        "QPR1",
        "QPR",
        datetime.date(2024, 3, 31),
        {"Beirut": 0},
        0,
        status="Due",
    )
    data["pd"].refresh_from_db()
    data["closed"].refresh_from_db()
    return data


def _rows(**params):
    return monitoring.indicators(monitoring.Filters.from_params(params, today=TODAY), today=TODAY)


# ---------------------------------------------------------------------------------- PD statuses
def test_one_definition_of_active_and_running():
    assert statuses.RUNNING == monitoring.ACTIVE_PD_STATUSES == partnerships.RUNNING_STATUSES
    assert statuses.ACTIVE == partnerships.ACTIVE_STATUSES == ("active",)
    assert set(statuses.ACTIVE) < set(statuses.RUNNING)


def test_pd_status_filter_defaults_to_the_running_pds(more, frozen):
    signed = PCA.objects.create(etl_id="13", number="LEB/PD3", title="Signed", status="signed")
    dm.PDIndicator.objects.create(datamart_id=30, source_id=700, intervention=signed, title="# of kits")
    titles = {r.title for r in _rows()}
    assert "# of kits" in titles and "# of schools" not in titles  # running: active, signed, suspended
    # "pd_status=" with nothing chosen is every status; a chosen status keeps only its PDs
    assert "# of schools" in {r.title for r in _rows(pd_status="")}
    assert {r.title for r in _rows(pd_status="closed")} == {"# of schools"}
    assert {r.title for r in _rows(pd_status="signed")} == {"# of kits"}
    # the filter shows the default as chosen statuses, and an old "scope=all" link as every status
    assert monitoring.Filters.from_params({}).status_choice == list(statuses.RUNNING)
    assert monitoring.Filters.from_params({"scope": "all"}).status_choice == []
    # the Programmes page's "Active" is the eTools status alone
    assert list(partnerships.programme_documents(scope="active")) == [more["pd"]]


# --------------------------------------------------------------------------- Program Performance
def test_program_performance_filters(more, frozen):
    girls, boys = more["girls"], more["boys"]
    # the Emergency indicator is the PD being an SPD, as Power BI's (the user's decision): the girls'
    # indicator is high frequency, which is not it
    assert {r.title for r in _rows(emergency="yes")} == set()
    assert {r.title for r in _rows(emergency="no")} == {girls, boys, CLASSROOMS}
    assert {r.title for r in _rows(pd_status="", document_type="SSFA")} == {"# of schools"}
    assert {r.title for r in _rows(pd_status="", office="Zahle")} == {"# of schools"}
    assert {r.title for r in _rows(pd_status="", partner_type="National")} == {"# of schools"}
    assert {r.title for r in _rows(cp_output="psychosocial")} == {girls}
    # an indicator not reported under a CP output takes its PD's, when the PD has only one
    by_title = {r.title: r for r in _rows()}
    assert by_title[boys].cp_output == "3.1 Children at risk reached with protection services"
    assert by_title[girls].cp_output == "1.2 Psychosocial support for children"
    assert {r.title for r in _rows(cp_output="children at risk")} == {boys, CLASSROOMS}
    # an SPD's indicators are the emergency ones (an "shpd" written in any case too), in the column and export
    PCA.objects.filter(pk=more["closed"].pk).update(document_type=" shpd ")
    assert {r.title for r in _rows(pd_status="", emergency="yes")} == {"# of schools"}
    assert {r.title for r in _rows(pd_status="", emergency="no")} == {girls, boys, CLASSROOMS}
    every = _rows(pd_status="")
    assert {r.title for r in every if r.emergency} == {"# of schools"}
    exported = {r["indicator"]: r["emergency"] for r in monitoring.table_rows(every, monitoring.Filters())}
    assert exported == {girls: "No", boys: "No", CLASSROOMS: "No", "# of schools": "Yes"}


def test_program_performance_table_sorts_and_counts_sections(client_viewer, more, frozen):
    page = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "table"})
    assert page.status_code == 200 and page.context["view"] == "table"
    # Reporting year "All" by default, as Power BI opens Program Performance; the grid opens on the year
    assert page.context["filters"].year is None
    assert client_viewer.get(reverse("reports:pd_monitoring")).context["filters"].year == TODAY.year
    rows = list(page.context["page_obj"])
    # % Reached, highest first; an indicator without a value last
    assert [r.title for r in rows] == [more["girls"], more["boys"], CLASSROOMS]
    assert page.context["sort"] == "-achieved"
    for heading in (
        "Section",
        "CP output",
        "Emergency",
        "Indicator",
        "Target",
        "Reached",
        "% Reached",
        "Partner",
    ):
        assert f">{heading}</a>" in page.text
    ascending = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "table", "sort": "achieved"})
    assert [r.title for r in ascending.context["page_obj"]] == [more["boys"], more["girls"], CLASSROOMS]
    by_title = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "table", "sort": "indicator"})
    assert [r.title for r in by_title.context["page_obj"]] == sorted(
        [more["girls"], more["boys"], CLASSROOMS], key=str.lower
    )
    # section tiles count every section the other filters keep, the chosen one or not
    tiles = {t["name"]: t["count"] for t in page.context["section_tiles"]}
    assert tiles == {"Child Protection": 2, "Education": 1}
    education = client_viewer.get(
        reverse("reports:pd_monitoring"), {"view": "table", "section": "Education"}, headers=HX
    )
    assert [r.title for r in education.context["page_obj"]] == [CLASSROOMS]
    assert {t["name"]: t["count"] for t in education.context["section_tiles"]} == tiles
    assert "<html" not in education.text and 'aria-current="true"' in education.text


def test_program_performance_export_holds_every_row_with_the_grid_figures(
    client_viewer, more, frozen, monkeypatch
):
    monkeypatch.setattr(monitoring, "PAGE_SIZE", 1)
    page = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "table"})
    assert len(page.context["page_obj"].object_list) == 1  # one row on the page...
    response = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "table", "export": "csv"})
    assert response.status_code == 200 and response["Content-Type"].startswith("text/csv")
    lines = list(csv.DictReader(io.StringIO(b"".join(response.streaming_content).decode("utf-8-sig"))))
    assert len(lines) == 3  # ... every indicator of the filter in the file
    grid = {r.title: r for r in _rows(year="all")}
    for line in lines:
        row = grid[line["indicator"]]
        assert line["reached"] == ("" if row.reached is None else str(row.reached))
        assert line["programme_document"] == "LEB/PD1" and line["pd_status"] == "Active"
    assert "reached_in_year" not in lines[0]  # every year: no "Reached in" column
    workbook = client_viewer.get(
        reverse("reports:pd_monitoring"), {"view": "table", "export": "xlsx", "year": "2026"}
    )
    sheet = load_workbook(io.BytesIO(workbook.content))["Program Performance"]
    header = [c.value for c in next(sheet.iter_rows(max_row=1))]
    assert "reached_in_year" in header and sheet.max_row == 4
    values = {r[header.index("indicator")]: r for r in sheet.iter_rows(min_row=2, values_only=True)}
    assert values[more["girls"]][header.index("reached")] == 550
    assert values[more["girls"]][header.index("reached_in_year")] == 550


# ---------------------------------------------------------------------------------- Data quality
def test_data_quality_tiles_and_lists(client_viewer, more, frozen):
    page = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "quality"})
    assert page.status_code == 200
    q = page.context["quality"]
    assert (q["programme_documents"], q["active_pds"], q["ended_pds"]) == (1, 1, 0)
    assert (q["indicators"], q["with_reached"], q["without_reached"]) == (3, 2, 1)
    assert (q["with_target"], q["without_target"]) == (2, 1)
    assert [r.title for r in q["missing_target"]] == [CLASSROOMS]
    assert [r.title for r in q["no_reached"]] == [CLASSROOMS]
    # the same filter gives the same figures as the grid's summary
    grid = client_viewer.get(reverse("reports:pd_monitoring"), {"year": "all"}).context["data"]
    assert (grid["indicators"], grid["with_reached"], grid["with_target"]) == (3, 2, 2)
    # every status: the closed PD is ended, its indicator without Reached
    everything = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "quality", "pd_status": ""})
    assert (
        everything.context["quality"]["programme_documents"],
        everything.context["quality"]["ended_pds"],
    ) == (
        2,
        1,
    )
    for label in ("Active PDs", "Ended PDs", "With Reached", "Without Reached", "Without a target"):
        assert f">{label}<" in page.text


def test_data_quality_neurodb_checks(client_viewer, more, frozen, monkeypatch):
    # PRP's cumulative on the latest counted report says 700, the reports add up to 550
    dm.ReportedIndicator.objects.filter(progress_report="PR-2", indicator=more["girls"]).update(
        total_cumulative_progress="700"
    )
    dm.PDIndicator.objects.create(
        datamart_id=21,
        source_id=600,
        intervention=more["pd"],
        title=more["boys"],
        section_name="Child Protection",
    )
    q = monitoring.data_quality(_rows(), monitoring.Filters(), TODAY)
    assert [r.title for r in q["mismatch"]] == [more["girls"]] and q["mismatch"][0].prp_cumulative == 700
    (twice,) = q["same_title"]
    assert (twice["title"], twice["count"], twice["pd"]) == (more["boys"], 2, more["pd"])
    assert q["partial"] == []
    # reports held only from March 2026: the PD (started January) has a partial history
    march = datetime.date(2026, 3, 1)
    monkeypatch.setattr(
        monitoring.rc, "history_from", lambda today=None: monitoring.rc.HistoryFrom(march, "read_from")
    )
    q = monitoring.data_quality(_rows(), monitoring.Filters(), TODAY)
    assert q["partial"] == [more["pd"]] and (q["read_from"], q["read_from_source"]) == (march, "read_from")
    page = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "quality"})
    assert "Partial history" in page.text and "Same title twice" in page.text
    assert (
        "These programme documents started before 1 Mar 2026, the first reporting period NeuroDB holds: their"
        in (" ".join(page.text.split()))
    )
    # the Datamart holds no earlier report than March 2026, although the sync asked from 1 Jan 2025: said so
    asked = datetime.date(2025, 1, 1)
    monkeypatch.setattr(
        monitoring.rc, "history_from", lambda today=None: monitoring.rc.HistoryFrom(march, "feed", asked)
    )
    q = monitoring.data_quality(_rows(), monitoring.Filters(), TODAY)
    assert (q["read_from"], q["read_from_source"], q["read_from_asked"]) == (march, "feed", asked)
    page = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "quality"})
    assert (
        "These programme documents started before 1 Mar 2026, the first reporting period NeuroDB holds: the "
        "Datamart holds no earlier partner report (NeuroDB asked from 1 Jan 2025), so their Reached is PRP's own "
        "cumulative when PRP gives one." in " ".join(page.text.split()).replace("&#x27;", "'")
    )


def test_data_quality_in_a_year_reads_reached_in_that_year(more, frozen):
    q = monitoring.data_quality(_rows(year="2025"), monitoring.Filters(year=2025), TODAY)
    assert (q["with_reached"], q["without_reached"]) == (0, 3)  # nothing reported in 2025
    q = monitoring.data_quality(_rows(year="2026"), monitoring.Filters(year=2026), TODAY)
    assert (q["with_reached"], q["without_reached"]) == (2, 1)


# ---------------------------------------------------------------------------- Reporting Progress
def _reporting(**params):
    """The Progress reports page's data for ``params``, every PD status unless one is chosen (the page
    opens on the running PDs: :func:`test_progress_reports_open_on_the_running_pds`, so the closed PD's
    PR-9 is only read with every status)."""
    params.setdefault("pd_status", "")
    return services.partner_reporting(params, today=TODAY)


def test_tiles_count_progress_reports_with_indicator_reports_under(more):
    """The tiles, the table by section and the chart by month count progress reports, as Power BI's
    Reporting Progress does (the user's decision of 26 Oct 2026); each tile gives its indicator reports."""
    data = _reporting()
    counts = {s["key"]: (s["indicator_reports"], s["progress_reports"]) for s in data["counts"]["statuses"]}
    # PR-1 (girls and boys), PR-2 and HR-4 accepted; PR-3 still running; PR-9 (closed PD) overdue
    assert counts == {
        "accepted": (4, 3),
        "submitted": (0, 0),
        "sent_back": (0, 0),
        "due": (0, 0),
        "overdue": (1, 1),
        "not_yet_due": (1, 1),
    }
    assert (data["counts"]["indicator_reports"], data["counts"]["progress_reports"]) == (6, 5)
    assert data["summary"]["reports"] == 5  # progress reports per status, as before
    # the tiles' values are progress reports
    assert {k: data["counts"][k] for k in reporting.STATUS_KEYS} == {
        "accepted": 3,
        "submitted": 0,
        "sent_back": 0,
        "due": 0,
        "overdue": 1,
        "not_yet_due": 1,
    }
    # one status per report: the section table and the months add up to the tiles
    assert data["by_section"]["total"] == 5
    assert data["by_section"]["totals"] == [3, 0, 0, 0, 1, 1]
    assert sum(sum(v) for v in data["by_month"]["series"].values()) == 5
    assert "Not yet due" in data["by_month"]["series"]
    # each indicator report is counted once, however many locations it has
    pr1 = next(r for r in data["reports"] if r["progress_report"] == "PR-1")
    assert pr1["indicators"] == 2
    # the page opens on the running PDs: the closed PD's overdue PR-9 is left out
    running = services.partner_reporting({}, today=TODAY)["counts"]
    assert (running["indicator_reports"], running["progress_reports"], running["overdue"]) == (5, 4, 0)


def test_section_table_and_section_filter(more):
    dm.ReportedIndicator.objects.filter(progress_report="PR-2").update(section="Education")
    data = _reporting()
    rows = {row["section"]: row for row in data["by_section"]["rows"]}
    # each progress report once, under its Section (a report without a section of its own under its PD's
    # sections; the closed PD has none): nothing is counted twice
    assert {s: row["total"] for s, row in rows.items()} == {"Education": 1, "Child Protection": 3, "—": 1}
    assert sum(row["total"] for row in data["by_section"]["rows"]) == data["by_section"]["total"] == 5
    pr1 = next(r for r in data["reports"] if r["progress_report"] == "PR-1")
    assert pr1["section"] == "Child Protection"  # from the PD, for the table
    assert {r["progress_report"] for r in _reporting(section="Education")["reports"]} == {"PR-2"}
    # a report without a section of its own is in its PD's sections, by the rule of the table
    assert {r["progress_report"] for r in _reporting(section="Child Protection")["reports"]} == {
        "PR-1",
        "HR-4",
        "PR-3",
    }
    # a section in a list of sections ("C4D, Education") matches, a section containing the word does not
    dm.ReportedIndicator.objects.filter(progress_report="PR-2").update(section="C4D, Education")
    assert {r["progress_report"] for r in _reporting(section="Education")["reports"]} == {"PR-2"}
    assert _reporting(section="Educ")["reports"] == []


def test_filters_on_status_number_partner_and_dates(more):
    assert [r["progress_report"] for r in _reporting(status="not_yet_due")["reports"]] == ["PR-3"]
    # PRP's word for a status is understood too
    assert {r["progress_report"] for r in _reporting(status="Accepted")["reports"]} == {
        "PR-1",
        "PR-2",
        "HR-4",
    }
    assert {r["progress_report"] for r in _reporting(report_number="QPR1")["reports"]} == {"PR-1", "PR-9"}
    amel = str(more["partner"].id)
    assert "PR-9" not in {r["progress_report"] for r in _reporting(partner=amel)["reports"]}
    assert [r["progress_report"] for r in _reporting(partner_type="National")["reports"]] == ["PR-9"]
    assert [r["progress_report"] for r in _reporting(pd_status="closed")["reports"]] == ["PR-9"]
    later = {r["progress_report"] for r in _reporting(end_from="2026-04-01")["reports"]}
    assert later == {"PR-2", "HR-4", "PR-3"}
    assert {r["progress_report"] for r in _reporting(start_to="2026-01-01")["reports"]} == {"PR-1", "PR-9"}
    assert [r["progress_report"] for r in _reporting(report_type="HR")["reports"]] == ["HR-4"]
    # the options: six statuses, report numbers in order
    options = _reporting()["options"]
    assert [s["value"] for s in options["statuses"]] == list(reporting.rc.STATUS_LABELS)
    assert options["report_numbers"] == ["HR4", "QPR1", "QPR2", "QPR3"]


def test_report_table_columns_and_focal_point_names_only(client_viewer, more, frozen):
    page = client_viewer.get(reverse("reports:partner_reporting"))
    assert page.status_code == 200
    pr1 = next(r for r in page.context["data"]["reports"] if r["progress_report"] == "PR-1")
    assert (pr1["pd_status_label"], pr1["document_type"], pr1["focal_points"]) == (
        "Active",
        "PD",
        ["Focal Point One"],
    )
    for heading in (
        "Programme document",
        "Section",
        "Document type",
        "Period start",
        "Period end",
        "UNICEF focal points",
    ):
        assert f">{heading}<" in page.text
    assert (
        "Focal Point One" in page.text
        and "focal.two@example.org" not in page.text
        and "three@example.org" not in page.text
        and "Focal Three" not in page.text
    )
    for label in ("Progress reports", "Accepted", "Submitted", "Sent back", "Due", "Overdue", "Not yet due"):
        assert f">{label}<" in page.text
    assert "4 indicator reports" in page.text  # the hint under Accepted (3 progress reports)
    assert "Progress reports by section" in page.text and "Indicator reports by section" not in page.text


def test_report_exports_hold_every_report_and_the_tile_counts(client_viewer, more, frozen, monkeypatch):
    monkeypatch.setattr("neurodb.reports.views.PAGE_SIZE", 2)
    every = {"pd_status": ""}  # every PD status: the closed PD's PR-9 too
    response = client_viewer.get(reverse("reports:partner_reporting"), {"export": "csv", **every})
    text = b"".join(response.streaming_content).decode("utf-8-sig")
    lines = list(csv.DictReader(io.StringIO(text)))
    assert len(lines) == 5 and "@" not in text
    pr1 = next(line for line in lines if line["report"] == "QPR1" and line["programme_document"] == "LEB/PD1")
    # the focal points as a count: an entry holding an e-mail address is no person NeuroDB names
    assert pr1["indicator_reports"] == "2" and pr1["unicef_focal_points_count"] == "1"
    assert "unicef_focal_points" not in pr1 and "Focal Point One" not in text
    workbook = load_workbook(
        io.BytesIO(
            client_viewer.get(reverse("reports:partner_reporting"), {"export": "xlsx", **every}).content
        )
    )
    assert workbook.sheetnames == ["Progress reports", "By status", "By section"]
    header = next(workbook["By status"].iter_rows(max_row=1, values_only=True))
    assert header[:3] == ("status", "progress_reports", "indicator_reports")
    by_status = {r[0]: (r[1], r[2]) for r in workbook["By status"].iter_rows(min_row=2, values_only=True)}
    assert by_status["Accepted"] == (3, 4) and by_status["Overdue"] == (1, 1)
    # By section counts progress reports, as the page's table
    by_section = {r[0]: r[-1] for r in workbook["By section"].iter_rows(min_row=2, values_only=True)}
    assert sum(by_section.values()) == 5
    # Ask NeuroDB gives the tiles' figures
    out = tools.run("partner_reporting", {})
    assert out["tiles"]["Accepted"] == {"progress_reports": 3, "indicator_reports": 4}


# --------------------------------------------------------------------------- Partner Performance
def test_partners_tab(client_viewer, more, frozen):
    perf = _reporting()["partners"]
    assert (perf["partners"], perf["reported"], perf["not_reported"]) == (2, 1, 1)
    assert {s["key"]: s["partners"] for s in perf["by_status"]} == {
        "accepted": 1,
        "submitted": 0,
        "sent_back": 0,
        "due": 0,
        "overdue": 1,
        "not_yet_due": 1,
    }
    assert [p["name"] for p in perf["not_reported_list"]] == ["Himaya"]
    # Overdue, Sent back, Due and Submitted reports, the most pressing first; a report not yet due (PR-3) is
    # one of the Due ones, listed as Due, as Power BI's follow-up list (the user's decision)
    assert [r["progress_report"] for r in perf["follow_up"]] == ["PR-9", "PR-3"]
    assert [r["status_label"] for r in perf["follow_up"]] == ["Overdue", "Due"]
    dm.ReportedIndicator.objects.filter(progress_report="PR-2").update(report_status="Sent back")
    assert [r["progress_report"] for r in _reporting()["partners"]["follow_up"]] == ["PR-9", "PR-2", "PR-3"]
    # the reports' own list keeps Not yet due
    pr3 = next(r for r in _reporting()["reports"] if r["progress_report"] == "PR-3")
    assert pr3["status"] == "not_yet_due"
    page = client_viewer.get(reverse("reports:partner_reporting"), {"tab": "partners"})
    assert page.status_code == 200 and page.context["tab"] == "partners"
    for label in ("Partners", "With at least one report", "Not reported"):
        assert f">{label}<" in page.text
    workbook = load_workbook(
        io.BytesIO(
            client_viewer.get(
                reverse("reports:partner_reporting"), {"tab": "partners", "export": "xlsx", "pd_status": ""}
            ).content
        )
    )
    assert workbook.sheetnames[0] == "Follow up" and "Partners not reported" in workbook.sheetnames
    statuses = [r[3] for r in workbook["Follow up"].iter_rows(min_row=2, values_only=True)]
    assert statuses == ["Overdue", "Sent back", "Due"]


def test_the_due_filter_keeps_the_reports_the_follow_up_list_calls_due(client_viewer, more, frozen):
    """Power BI's Report status slicer set to Due keeps, on its follow-up list, the reports not yet due the list
    calls Due (the user's decision of 26 Oct 2026). So do the Partners tab's reports to follow up and their
    export: PR-3, not yet due, stays under Report status Due, listed as Due, as under its own status. The tiles
    and the other tabs keep the six statuses: no report is Due there."""
    due = _reporting(status="due")
    assert [(r["progress_report"], r["status_label"]) for r in due["partners"]["follow_up"]] == [
        ("PR-3", "Due")
    ]
    assert due["reports"] == [] and due["counts"]["progress_reports"] == 0
    not_yet = _reporting(status="not_yet_due")["partners"]["follow_up"]
    assert [(r["progress_report"], r["status_label"]) for r in not_yet] == [("PR-3", "Due")]
    both = _reporting(status=["due", "overdue"])["partners"]["follow_up"]
    assert [r["progress_report"] for r in both] == ["PR-9", "PR-3"]
    assert [r["progress_report"] for r in _reporting(status="overdue")["partners"]["follow_up"]] == ["PR-9"]
    url = reverse("reports:partner_reporting")
    page = client_viewer.get(url, {"tab": "partners", "status": "due", "pd_status": ""})
    assert page.context["page_obj"].paginator.count == 1
    assert "every report in the filter is Accepted" not in page.text
    workbook = load_workbook(
        io.BytesIO(
            client_viewer.get(
                url, {"tab": "partners", "status": "due", "export": "xlsx", "pd_status": ""}
            ).content
        )
    )
    assert [r[3] for r in workbook["Follow up"].iter_rows(min_row=2, values_only=True)] == ["Due"]
    # a filter that keeps no report at all says so; "every report is Accepted" is for Accepted reports only
    nothing = client_viewer.get(url, {"tab": "partners", "status": "sent_back", "pd_status": ""})
    assert "every report in the filter is Accepted" not in nothing.text
    assert "No progress report in the filter." in nothing.text
    accepted = client_viewer.get(url, {"tab": "partners", "status": "accepted", "pd_status": ""})
    assert "Nothing to follow up: every report in the filter is Accepted." in accepted.text


# ------------------------------------------------------------------------------- Report/Location
def test_report_by_location_reads_each_location_once(client_viewer, more, frozen):
    girls = more["girls"]
    # a second row for Akkar in QPR1 (another disaggregation group): never added, the largest is used
    copy = dm.ReportedIndicator.objects.get(progress_report="PR-1", indicator=girls, location="Akkar")
    copy.pk, copy.datamart_id = None, 999_001
    copy.disaggregation = {"()": {"v": 50, "d": 1, "c": 50}}
    copy.save()
    # HR4: the report total repeated on both locations, no value of their own
    dm.ReportedIndicator.objects.filter(progress_report="HR-4").update(
        disaggregation={},
        achievement_in_period="100",
        total_cumulative_progress_in_location="",
        previous_location_progress="",
    )
    data = _reporting()
    found = reporting.location_rows(
        data["rows"], {r["progress_report"]: r["status"] for r in data["reports"]}, TODAY
    )
    rows = {(r["progress_report"], r["indicator"], r["location"]): r for r in found["rows"]}
    akkar = rows[("PR-1", girls, "Akkar")]
    assert (akkar["value_in_period"], akkar["rows"]) == (200.0, 2)
    assert rows[("PR-1", girls, "Bekaa")]["value_in_period"] == 100.0
    hr = rows[("HR-4", girls, "Akkar")]
    assert hr["repeated"] and hr["value_in_period"] is None
    assert rows[("PR-9", "# of schools", "Beirut")]["report_status"] == "Overdue"
    assert not rows[("PR-9", "# of schools", "Beirut")]["counted"]
    # one row per indicator report and location: PR-1 girls has two, not three
    assert sum(1 for k in rows if k[:2] == ("PR-1", girls)) == 2
    page = client_viewer.get(reverse("reports:partner_reporting"), {"tab": "locations", "status": "accepted"})
    assert page.status_code == 200
    assert {r["report_status"] for r in page.context["page_obj"]} == {"Accepted"}
    assert "report total" in page.text and "is-uncounted" not in page.text
    response = client_viewer.get(
        reverse("reports:partner_reporting"), {"tab": "locations", "export": "csv", "pd_status": ""}
    )
    lines = list(csv.DictReader(io.StringIO(b"".join(response.streaming_content).decode("utf-8-sig"))))
    assert len(lines) == len(found["rows"])
    noted = next(
        line
        for line in lines
        if line["report"] == "QPR1" and line["location"] == "Akkar" and line["indicator"] == girls
    )
    assert noted["value_in_period"] == "200.0" and "never their sum" in noted["note"]


def test_report_by_location_asks_to_narrow_a_large_filter(client_viewer, more, frozen, monkeypatch):
    monkeypatch.setattr(reporting, "LOCATION_LIMIT", 2)
    page = client_viewer.get(reverse("reports:partner_reporting"), {"tab": "locations"})
    assert page.status_code == 200 and "choose a reporting year" in page.text
    response = client_viewer.get(reverse("reports:partner_reporting"), {"tab": "locations", "export": "xlsx"})
    assert response.status_code == 400


# --------------------------------------------------------------------------- PD/SPD Indicators
def test_programme_indicators_tab(client_viewer, more):
    pds = partnerships.programme_documents(scope="all")
    rows = reporting.pd_indicators(pds, {})
    # one row per indicator, though eTools lists the girls indicator once per location
    assert sorted(r["indicator"] for r in rows) == sorted(
        [more["girls"], more["boys"], CLASSROOMS, "# of schools"]
    )
    girls = next(r for r in rows if r["indicator"] == more["girls"])
    assert (
        girls["programme_document"],
        girls["pd_status"],
        girls["office"],
        girls["target"],
        girls["emergency"],
    ) == (
        "LEB/PD1",
        "Active",
        "Beirut",
        1000.0,
        "No",  # high frequency, but not an SPD's: the Emergency indicator is the PD being an SPD
    )
    assert reporting.pd_indicators(pds, {"emergency": "yes"}) == []
    PCA.objects.filter(pk=more["pd"].pk).update(document_type="SPD")
    pds = partnerships.programme_documents(scope="all")
    assert {r["indicator"] for r in reporting.pd_indicators(pds, {"emergency": "yes"})} == {
        more["girls"],
        more["boys"],
        CLASSROOMS,
    }
    assert {r["emergency"] for r in reporting.pd_indicators(pds, {"emergency": "no"})} == {"No"}
    assert [r["indicator"] for r in reporting.pd_indicators(pds, {"section": "Education"})] == [
        CLASSROOMS,
        "# of schools",
    ]
    assert [r["indicator"] for r in reporting.pd_indicators(pds, {"pd_q": "PD2"})] == ["# of schools"]
    assert [r["indicator"] for r in reporting.pd_indicators(pds, {"indicator": "syrian"})] == [more["boys"]]
    page = client_viewer.get(reverse("reports:programmes"), {"tab": "indicators"})
    assert page.status_code == 200 and page.context["tab"] == "indicators" and more["girls"] in page.text
    partial = client_viewer.get(
        reverse("reports:programmes"), {"tab": "indicators", "emergency": "no"}, headers=HX
    )
    assert (
        more["girls"] not in partial.text and "# of schools" in partial.text and "<html" not in partial.text
    )
    response = client_viewer.get(reverse("reports:programmes"), {"tab": "indicators", "export": "csv"})
    lines = list(csv.DictReader(io.StringIO(b"".join(response.streaming_content).decode("utf-8-sig"))))
    assert len(lines) == 4 and set(reporting.INDICATOR_COLUMNS) == set(lines[0])
    # the PD list keeps working, with the PD reference search
    documents = client_viewer.get(reverse("reports:programmes"), {"pd_q": "PD2"})
    assert [r["number"] for r in documents.context["rows"]] == ["LEB/PD2"]


# --------------------------------------------------------------------------------- who sees them
NEW_VIEWS = [
    ("reports:pd_monitoring", {"view": "table"}),
    ("reports:pd_monitoring", {"view": "quality"}),
    ("reports:pd_monitoring", {"view": "table", "export": "xlsx"}),
    ("reports:partner_reporting", {"tab": "partners"}),
    ("reports:partner_reporting", {"tab": "locations"}),
    ("reports:partner_reporting", {"export": "csv"}),
    ("reports:programmes", {"tab": "indicators"}),
    ("reports:programmes", {"tab": "indicators", "export": "xlsx"}),
]


@pytest.mark.parametrize(("name", "params"), NEW_VIEWS)
def test_every_new_view_for_a_viewer(client_viewer, more, frozen, name, params):
    assert client_viewer.get(reverse(name), params).status_code == 200


@pytest.mark.parametrize(("name", "params"), NEW_VIEWS)
def test_donors_never_reach_the_new_views(client, more, name, params):
    from tests.donors.test_donor_access import make_donor

    client.force_login(make_donor().user)
    response = client.get(reverse(name), params)
    assert response.status_code == 302 and response["Location"] == reverse("donors:page")
    assert client.get(reverse(name), params, headers=HX).status_code == 403


@pytest.mark.parametrize(("name", "params"), NEW_VIEWS[:2] + NEW_VIEWS[3:5] + NEW_VIEWS[6:7])
def test_the_new_views_render_without_data(client_viewer, name, params):
    assert client_viewer.get(reverse(name), params).status_code == 200


def test_same_filter_same_numbers_across_the_new_views(client_viewer, more, frozen):
    """Grid, Program Performance, its export, Data quality and Ask NeuroDB agree for one filter."""
    params = {"year": "2026", "section": "Child Protection"}
    grid = client_viewer.get(reverse("reports:pd_monitoring"), params).context
    table = client_viewer.get(reverse("reports:pd_monitoring"), {**params, "view": "table"}).context
    quality = client_viewer.get(reverse("reports:pd_monitoring"), {**params, "view": "quality"}).context
    grid_reached = {r.title: r.reached for r in grid["rows"]}
    assert (
        {r.title: r.reached for r in table["rows"]}
        == grid_reached
        == {more["girls"]: 550.0, more["boys"]: 20.0}
    )
    assert {r.title: r.reached_in_year for r in table["rows"]} == {
        r.title: r.reached_in_year for r in grid["rows"]
    }
    assert quality["quality"]["indicators"] == grid["data"]["indicators"] == 2
    assert quality["quality"]["with_reached"] == grid["data"]["with_reached"] == 2
    response = client_viewer.get(
        reverse("reports:pd_monitoring"), {**params, "view": "table", "export": "csv"}
    )
    lines = list(csv.DictReader(io.StringIO(b"".join(response.streaming_content).decode("utf-8-sig"))))
    assert {line["indicator"]: float(line["reached"]) for line in lines} == grid_reached
    ask = tools.run("pd_indicator_progress", {"partner": "amel", "year": 2026})
    assert {i["indicator"]: i["reached"] for i in ask["indicators_detail"]} == {
        **grid_reached,
        CLASSROOMS: None,
    }


# ------------------------------------------------------------------------------ the stage check
def test_indicators_without_a_section_have_a_tile_that_lists_them(client_viewer, more, frozen):
    dm.PDIndicator.objects.filter(source_id=500).update(section_name="")
    page = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "table"})
    tiles = {t["name"]: t for t in page.context["section_tiles"]}
    assert {name: t["count"] for name, t in tiles.items()} == {
        "Child Protection": 2,
        monitoring.NO_SECTION: 1,
    }
    # the tile of the indicators that name no section lists them (it used to list nothing)
    listed = client_viewer.get(
        reverse("reports:pd_monitoring") + tiles[monitoring.NO_SECTION]["url"], headers=HX
    )
    assert [r.title for r in listed.context["page_obj"]] == [CLASSROOMS]
    assert {t["name"]: t["count"] for t in listed.context["section_tiles"]} == {
        name: t["count"] for name, t in tiles.items()
    }
    # the grid, Data quality and the export read the same filter alike
    assert [r.title for r in _rows(section=monitoring.NO_SECTION)] == [CLASSROOMS]
    quality = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "quality", "section": "—"})
    assert quality.context["quality"]["indicators"] == 1


def test_program_performance_export_is_not_cut_at_the_page_limit(client_viewer, more, frozen, monkeypatch):
    monkeypatch.setattr(monitoring, "MAX_INDICATORS", 2)
    page = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "table"})
    assert len(page.context["rows"]) == 2 and page.context["truncated"]  # the page lists two...
    assert "The Excel and CSV files hold every indicator" in page.text
    response = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "table", "export": "csv"})
    lines = list(csv.DictReader(io.StringIO(b"".join(response.streaming_content).decode("utf-8-sig"))))
    assert sorted(line["indicator"] for line in lines) == sorted(  # ... the file holds all three
        [more["girls"], more["boys"], CLASSROOMS]
    )
    workbook = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "table", "export": "xlsx"})
    assert load_workbook(io.BytesIO(workbook.content))["Program Performance"].max_row == 4
    # the page lists the first two in its sort (% Reached, highest first), not two taken before sorting
    assert [r.title for r in page.context["rows"]] == [more["girls"], more["boys"]]
    # every count reads the whole filter, whatever a page lists: grid tiles, section tiles, Data quality
    # and Ask NeuroDB agree on three indicators
    assert sum(t["count"] for t in page.context["section_tiles"]) == 3
    grid = client_viewer.get(reverse("reports:pd_monitoring"), {"year": "all"})
    assert grid.context["data"]["indicators"] == 3 and len(grid.context["rows"]) == 2
    assert "The figures above count every indicator" in grid.text
    quality = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "quality"})
    assert quality.context["quality"]["indicators"] == 3 and not quality.context["truncated"]
    assert tools.run("pd_indicator_progress", {"partner": "amel"})["indicators"] == 3


def test_ask_counts_progress_reports_by_status_like_its_totals(more):
    out = tools.run("partner_reporting", {})
    # "by_status" counts progress reports, as "reports", the status keys and the page's tiles do ("tiles",
    # with the indicator reports in them). Like the page, it reads the running PDs (the closed PD's PR-9 is
    # left out)
    assert sum(out["by_status"].values()) == out["reports"] == 4
    assert out["pd_statuses"] == ["active", "signed", "suspended"]
    assert out["by_status"]["Accepted"] == out["accepted"] == 3
    assert out["tiles"]["Accepted"] == {"progress_reports": 3, "indicator_reports": 4}
    # the reports to follow up, as the Partners tab lists them: PR-3, not yet due, is one of the Due ones
    assert out["to_follow_up"] == {"Due": 1}


def test_ask_reads_the_emergency_indicator_as_the_pages_do(more, frozen):
    """Ask NeuroDB's Emergency indicator is the pages' (the user's decision of 26 Oct 2026): the programme
    document is an SPD. Its indicator and progress report tools filter by it as Program Performance and Progress
    reports do, and each indicator says it, as Program Performance's column does; eTools' high frequency (the
    girls' indicator) is not it."""
    girls, boys = more["girls"], more["boys"]
    PCA.objects.filter(pk=more["pd"].pk).update(document_type="SPD")
    detail = tools.run("pd_indicator_progress", {"include_closed": True})["indicators_detail"]
    assert {i["indicator"]: i["emergency"] for i in detail} == {
        girls: True,
        boys: True,
        CLASSROOMS: True,
        "# of schools": False,  # the closed SSFA's
    }
    yes = tools.run("pd_indicator_progress", {"include_closed": True, "emergency": "yes"})
    assert {i["indicator"] for i in yes["indicators_detail"]} == {girls, boys, CLASSROOMS}
    assert yes["indicators"] == len(_rows(pd_status="", emergency="yes")) == 3
    no = tools.run("pd_indicator_progress", {"include_closed": True, "emergency": "no"})
    assert [i["indicator"] for i in no["indicators_detail"]] == ["# of schools"]
    # Progress reports' filter: the reports of SPDs, the same count as the page's
    reports = tools.run("partner_reporting", {"emergency": "yes"})
    page = services.partner_reporting({"emergency": "yes"}, today=TODAY)
    assert reports["reports"] == page["summary"]["reports"] == 4
    assert tools.run("partner_reporting", {"emergency": "no"})["reports"] == 0
    with pytest.raises(tools.ToolInputError):
        tools.run("partner_reporting", {"emergency": "high frequency"})


def test_report_by_location_keeps_each_report_by_its_own_status(client_viewer, more, frozen):
    # one row of PR-1 names no status and no submission date: alone, it would read as overdue, but the
    # progress report is Accepted (its other rows say so), as the tiles and the report table count it
    dm.ReportedIndicator.objects.filter(progress_report="PR-1", indicator=more["boys"]).update(
        report_status="", submission_date=None
    )
    data = _reporting()
    pr1 = next(r for r in data["reports"] if r["progress_report"] == "PR-1")
    assert pr1["status"] == "accepted"
    assert (
        data["statuses"][pr1["key"]] == "accepted"
    )  # each report's status by its key (PD, type, number, period)
    page = client_viewer.get(
        reverse("reports:partner_reporting"), {"tab": "locations", "status": "overdue", "pd_status": ""}
    )
    assert {(r["progress_report"], r["report_status"]) for r in page.context["page_obj"]} == {
        ("PR-9", "Overdue")
    }
    accepted = client_viewer.get(
        reverse("reports:partner_reporting"), {"tab": "locations", "status": "accepted"}
    )
    boys = [r for r in accepted.context["page_obj"] if r["indicator"] == more["boys"]]
    assert [(r["progress_report"], r["report_status"]) for r in boys] == [("PR-1", "Accepted")]


def test_partner_monitoring_view_links_follow_the_filter_in_the_address_bar(client_viewer, more, frozen):
    """A filter changed through HTMX changes the address, not the links drawn with the page: the view
    links take the address bar's filter when clicked (data-current-query), so switching from Program
    Performance to Data quality keeps the filter just chosen."""
    page = client_viewer.get(reverse("reports:pd_monitoring"), {"view": "table", "year": "2026"})
    for view in ("grid", "table", "quality"):
        assert f'data-current-query="view={view}"' in page.text
    assert page.context["tab_urls"]["quality"].endswith("?view=quality&year=2026")


def test_programme_indicators_tab_offers_the_sections_indicators_name(client_viewer, more):
    documents = client_viewer.get(reverse("reports:programmes"))
    assert documents.context["options"]["sections"] == ["Child Protection"]  # the PDs' sections
    indicators = client_viewer.get(reverse("reports:programmes"), {"tab": "indicators"})
    # the Indicators tab filters on each indicator's own section: those are offered too
    assert indicators.context["options"]["sections"] == ["Child Protection", "Education"]
    chosen = client_viewer.get(reverse("reports:programmes"), {"tab": "indicators", "section": "Education"})
    assert sorted(r["indicator"] for r in chosen.context["page_obj"]) == sorted([CLASSROOMS, "# of schools"])


def test_country_figures_read_every_indicator_whatever_a_page_lists(more, frozen, monkeypatch):
    """The 2,000 rows a page lists never cut a count or a total: the country overview, the home page,
    the map and Ask NeuroDB read every indicator of their filter (a country with more indicators than a
    page lists would otherwise lose the rest from its figures)."""
    from neurodb.web import services as web_services

    monkeypatch.setattr(monitoring, "MAX_INDICATORS", 1)
    assert len(_rows()) == 3
    assert monitoring.map_points(monitoring.Filters(year=None), TODAY)["totals"]["indicators"] == 3
    assert web_services._etools_counts(None)["pd_indicators"] == 3  # PD1 runs in this year
    assert tools.run("pd_indicator_progress", {"partner": "amel"})["indicators"] == 3
