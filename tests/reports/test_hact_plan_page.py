"""The "HACT visit plan" tab of the Field monitoring page (``/field-monitoring/?tab=plan``): who reaches it,
its figures against the service's (``neurodb.datamart.hact_plan``), the filters and quick views, the list's
links, the by-partner view, the section comments written from the list, the Excel and CSV files (the
e-mail's columns, and never a staff name), the strip on the default tab, the cross-links from the other
HACT pages and the number of queries.

Synthetic data only: "Partner A", references like TEST/PCA2026001/PD2026001, focal point "Focal Person
Test"."""

import csv
import datetime
import io
import itertools
import json
import re
import shutil
import subprocess
from pathlib import Path

import pytest
from django.conf import settings as django_settings
from django.contrib.auth.models import Group
from django.http import QueryDict
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from neurodb.accounts.models import Section, User
from neurodb.accounts.roles import ADMIN, SECTION_EDITOR, VIEWER
from neurodb.datamart import fm, hact_plan
from neurodb.datamart import models as dm
from neurodb.donors.models import DonorAccount
from neurodb.partnerships.models import PCA, PartnerOrganization
from neurodb.watch.models import SectionMatch

pytestmark = pytest.mark.django_db

PAGE = reverse("reports:monitoring")
# the suite's day (tests/conftest.py SITE_NOW), not the real clock: this is read at import, before it travels
TODAY = datetime.date(2026, 10, 15)
YEAR = TODAY.year
D = datetime.date
FOCAL = "Focal Person Test"
FOCAL_EMAIL = "focal.person@example.org"
# the files' columns: the team's workbook's, then NeuroDB's (SPEC_HACT §6, with the partner's field
# monitoring visits that name none of its PDs among the "not counted" figures)
FILE_COLUMNS = [
    "partner", "vendor_number", "reference_number", "status", "sections", "document_type", "start_date",
    "end_date", "total_budget", "budget_currency", "risk_rating", "required_pv", "completed_pv",
    "completed_pv_out_of_minimum", "still_needed_by_31_dec", "q1_planned", "q2_planned", "q3_planned",
    "q4_planned", "due_by_end_of_quarter", "remark", "not_visited", "completed_fm", "completed_tpm",
    "awaiting_approval", "fm_completed_not_programmatic", "fm_partner_visits_naming_no_pd",
    "staff_trip_activities_not_counted", "tpm", "unicef_focal_points_count", "section_comment",
    "section_comment_updated",
]  # fmt: skip
_ids = itertools.count(1)


# ---------------------------------------------------------------------------------- the world
def partner(name: str, rating: str = "Low") -> PartnerOrganization:
    n = next(_ids)
    return PartnerOrganization.objects.create(
        etl_id=f"p{n}",
        name=name,
        partner_type="Civil Society Organization",
        vendor_number=f"V{n:05d}",
        rating=rating,
    )


def programme(owner, *, status="active", sections=("Education",), document_type="PD", end=None) -> PCA:
    n = next(_ids)
    return PCA.objects.create(
        etl_id=f"pd{n}",
        partner=owner,
        partner_name=owner.name,
        number=f"TEST/PCA{YEAR}{n:03d}/PD{YEAR}{n:03d}",
        title=f"Test programme {n}",
        status=status,
        document_type=document_type,
        start=D(YEAR, 1, 1),
        end=end or D(YEAR, 12, 31),
        section_names=list(sections),
        total_budget="120,000.50",
        budget_currency="USD",
        unicef_focal_points=[FOCAL, FOCAL_EMAIL],
    )


def plan(pd: PCA, quarters=(1, 1, 1, 1), year: int = YEAR) -> None:
    q1, q2, q3, q4 = quarters
    dm.PlannedVisits.objects.create(
        datamart_id=next(_ids),
        intervention=pd,
        partner_id=pd.partner_id,
        pd_reference_number=pd.number,
        year=year,
        q1=q1,
        q2=q2,
        q3=q3,
        q4=q4,
    )


def fm_visit(pd: PCA, *, status="completed", programmatic=True) -> None:
    """A one-row field monitoring visit about ``pd``, ending in the year (before today)."""
    activity = next(_ids)
    end = D(YEAR, 1, 15)
    dm.MonitoringFinding.objects.create(
        datamart_id=next(_ids),
        partner=pd.partner,
        vendor_number=pd.partner.vendor_number,
        entity=pd.number,
        entity_type="PD/SSFA",
        intervention=pd,
        monitoring_activity=f"FM-TEST-{activity}",
        monitoring_activity_id=activity,
        status=status,
        is_programmatic_visit=programmatic,
        start_date=end - datetime.timedelta(days=1),
        end_date=end,
        visit_lead="Lead Person Test",
    )


def tpm(pd: PCA, *, status="unicef_approved") -> None:
    dm.TPMActivity.objects.create(
        datamart_id=next(_ids),
        partner=pd.partner,
        intervention=pd,
        visit_reference_number=f"TPM/{YEAR}/{next(_ids)}",
        visit_status=status,
        is_programmatic_visit=True,
        pd_reference_number=pd.number,
        date=D(YEAR, 1, 20),
    )


def hact_year(owner, *, required, completed, risk="Medium", **spot_checks) -> None:
    dm.PartnerHACTYear.objects.create(
        datamart_id=next(_ids),
        partner=owner,
        partner_name=owner.name,
        vendor_number=owner.vendor_number,
        year=YEAR,
        risk_rating=risk,
        pv_required=required,
        pv_completed=completed,
        **spot_checks,  # sc_required, sc_completed
    )


@pytest.fixture
def world(db):
    """Three partners and five programme documents of this year:

    * ``a1`` (Partner A, Education): 4 planned, 1 field monitoring visit and 1 report awaiting approval:
      3 still needed. Partner A's eTools minimum (2 of 2) is met: "minimum met, PD visits needed".
    * ``a2`` (Partner A, Health, SPD): 1 planned in Q1, 1 approved TPM visit: completed.
    * ``b`` (Partner B, Education and Health, ended): 2 planned, none visited: not visited, ended.
    * ``c1`` (Partner C, rated High, Health): no plan, 1 visit: "No visit planned".
    * ``c2`` (Partner C, Child Protection): 2 planned in Q1, none visited: not visited; Partner C's
      eTools minimum is 5, above its PDs' plan (2).
    """
    a, b, c = partner("Partner A"), partner("Partner B"), partner("Partner C", rating="High")
    pds = {
        "a1": programme(a),
        "a2": programme(a, sections=("Health",), document_type="SPD"),
        "b": programme(b, status="ended", sections=("Education", "Health"), end=D(YEAR, 3, 31)),
        "c1": programme(c, sections=("Health",)),
        "c2": programme(c, sections=("Child Protection",)),
    }
    plan(pds["a1"], (1, 1, 1, 1))
    fm_visit(pds["a1"])
    fm_visit(pds["a1"], status="submitted")
    plan(pds["a2"], (1, 0, 0, 0))
    tpm(pds["a2"])
    plan(pds["b"], (1, 1, 0, 0))
    fm_visit(pds["c1"])
    plan(pds["c2"], (2, 0, 0, 0))
    hact_year(a, required=2, completed=2, risk="Low")
    hact_year(c, required=5, completed=1, risk="High")
    return {"partners": {"a": a, "b": b, "c": c}, **pds}


def _person(username, role, section=None) -> User:
    user = User.objects.create_user(
        username=username, email=f"{username}@example.org", password="pass-123456-abc"
    )
    if role:
        user.groups.add(Group.objects.get(name=role))
    if section is not None:
        user.section = section
        user.save(update_fields=["section"])
    return User.objects.get(pk=user.pk)


@pytest.fixture
def people(roles):
    education = Section.objects.create(name="Education", code="EDU")
    health = Section.objects.create(name="Health", code="HLT")
    SectionMatch.objects.create(etools_name="Education", section=education, how="exact", confirmed=True)
    SectionMatch.objects.create(etools_name="Health", section=health, how="exact", confirmed=True)
    donor = _person("donor-test", None)
    DonorAccount.objects.create(user=donor, name="Donor Test", donors=["EU"], must_change_password=False)
    return {
        "admin": _person("admin-test", ADMIN),
        "editor": _person("edu-editor", SECTION_EDITOR, education),
        "other_editor": _person("hlt-editor", SECTION_EDITOR, health),
        "viewer": _person("edu-viewer", VIEWER, education),
        "donor": User.objects.get(pk=donor.pk),
    }


def _client(user) -> Client:
    client = Client()
    client.force_login(user)
    return client


def _ids_of(response) -> set[int]:
    return {row["pd_id"] for row in response.context["rows"]}


def _plan(client, **params):
    return client.get(PAGE, {"tab": "plan", **params})


# ---------------------------------------------------------------------------------- access
@pytest.mark.parametrize("who", ["admin", "editor", "viewer"])
def test_staff_open_the_plan(world, people, who):
    response = _plan(_client(people[who]))
    assert response.status_code == 200
    html = response.content.decode()
    assert f"HACT visit plan {YEAR}" in html and world["a1"].number in html
    assert 'aria-current="page">HACT visit plan</a>' in html
    # the subtitle names the plan; the default tab stays the first link
    assert "and the HACT programmatic visit plan" in html


def test_donors_and_anonymous_users_are_kept_out(world, people):
    donor = _client(people["donor"])
    for params in ({"tab": "plan"}, {"tab": "plan", "export": "xlsx"}, {"tab": "plan", "export": "csv"}):
        response = donor.get(PAGE, params)
        assert response.status_code == 302 and response["Location"] == reverse("donors:page")
        assert donor.get(PAGE, params, HTTP_HX_REQUEST="true").status_code == 403
    anonymous = Client().get(PAGE, {"tab": "plan", "export": "csv"})
    assert anonymous.status_code == 302 and anonymous["Location"].startswith(reverse("account_login"))


# ---------------------------------------------------------------------------------- figures
def test_tiles_quick_views_and_by_section_agree_with_the_service(world, people):
    client = _client(people["viewer"])
    rows = hact_plan.build(YEAR, TODAY)
    figures = hact_plan.headline(rows, YEAR)
    response = _plan(client)
    assert response.context["figures"] == figures
    assert figures["not_visited"]["pds"] == 2 and figures["not_visited"]["partners"] == 2
    assert figures["still_needed"] == {
        "visits": 7,
        "pds": 3,
        "partners": 3,
        "visits_ended": 2,
        "pds_ended": 1,
    }
    assert figures["etools_behind"] == {"partners": 1, "with_row": 2}
    by_section = [{k: v for k, v in s.items() if k != "url"} for s in response.context["by_section"]]
    assert by_section == hact_plan.by_section(rows)
    assert [(q["count"], q["active"]) for q in response.context["quick"]] == [
        (2, False),
        (3, False),
        (5, True),
    ]
    html = response.content.decode()
    assert "Not visited (2)" in html and "End-of-year missing visits (3)" in html and "All (5)" in html

    # the quick views narrow the list, never the tiles nor the by-section panel
    for params, expected in (
        ({"zero": "1"}, {world["b"].pk, world["c2"].pk}),
        ({"remark": ["needed", "ended"]}, {world["a1"].pk, world["b"].pk, world["c2"].pk}),
    ):
        narrowed = _plan(client, **params)
        assert _ids_of(narrowed) == expected
        assert narrowed.context["figures"] == figures
        assert [
            {k: v for k, v in s.items() if k != "url"} for s in narrowed.context["by_section"]
        ] == by_section
        assert [q["active"] for q in narrowed.context["quick"]].count(True) == 1
    # a quick view keeps the rest of the filter
    health = _plan(client, section="Health")
    eoy = health.context["quick"][1]["url"]
    assert "section=Health" in eoy and "remark=needed" in eoy and "remark=ended" in eoy and "tab=plan" in eoy


def test_each_filter_narrows_the_list(world, people):
    client = _client(people["viewer"])
    every = {pd.pk for key, pd in world.items() if key != "partners"}
    assert _ids_of(_plan(client)) == every
    cases = [
        ({"section": "health"}, {world["a2"].pk, world["b"].pk, world["c1"].pk}),  # any case
        ({"partner": str(world["partners"]["c"].pk)}, {world["c1"].pk, world["c2"].pk}),
        ({"risk": "high"}, {world["c1"].pk, world["c2"].pk}),
        ({"remark": "completed"}, {world["a2"].pk}),
        ({"remark": "no_plan"}, {world["c1"].pk}),
        ({"status": "ended"}, {world["b"].pk}),
        ({"type": "spd"}, {world["a2"].pk}),
        ({"q": world["a1"].number.lower()}, {world["a1"].pk}),
        ({"q": world["partners"]["b"].vendor_number}, {world["b"].pk}),
        ({"q": "partner a"}, {world["a1"].pk, world["a2"].pk}),
        ({"zero": "1"}, {world["b"].pk, world["c2"].pk}),
        ({"year": str(YEAR + 1)}, set()),  # no programme document runs next year
        ({"partner": "not-a-number", "remark": "nonsense", "year": "²"}, every),  # bad values are dropped
    ]
    for params, expected in cases:
        response = _plan(client, **params)
        assert response.status_code == 200, params
        assert _ids_of(response) == expected, params
    assert "No programme document runs in" in _plan(client, year=str(YEAR + 1)).content.decode()
    assert "No programme document matches" in _plan(client, q="nothing like it").content.decode()
    # the filter bar keeps the tab, offers the year's values and resets to the plan
    html = _plan(client).content.decode()
    assert '<input type="hidden" name="tab" value="plan">' in html
    assert (
        ">Child Protection</option>" in html
        and ">Partner C</option>" in html
        and ">PD ended, visits missing</option>" in html
    )
    assert f'href="{PAGE}?tab=plan">Reset</a>' in html


def test_htmx_answers_the_results_and_the_tabs_out_of_band(world, people):
    client = _client(people["viewer"])
    response = client.get(PAGE, {"tab": "plan", "section": "Education"}, HTTP_HX_REQUEST="true")
    html = response.content.decode()
    assert response.status_code == 200 and "<html" not in html
    assert [t.name for t in response.templates][0] == "reports/partials/monitoring_plan.html"
    assert 'id="monitoring-tabs" aria-label="Field monitoring views" hx-swap-oob="true"' in html
    assert f'href="{PAGE}?year={YEAR}"' in html  # "Findings and visits" keeps the year
    assert world["a1"].number in html and world["c1"].number not in html


def test_too_many_programme_documents_ask_for_a_narrower_filter(world, people, monkeypatch):
    monkeypatch.setattr(hact_plan, "MAX_ROWS", 3)
    client = _client(people["viewer"])
    response = _plan(client)
    html = response.content.decode()
    assert "5 programme documents match: choose a section or a partner first." in html
    assert response.context["rows"] == [] and "kpi-grid" in html
    assert _ids_of(_plan(client, section="Education")) == {world["a1"].pk, world["b"].pk}
    lines = _csv(_plan(client, export="csv"))
    assert len(lines) == 5  # the file still holds every row


def test_a_year_without_a_plan_says_so(people):
    pd = programme(partner("Partner A"))
    html = _plan(_client(people["viewer"])).content.decode()
    assert f"eTools has no planned visits for {YEAR} yet: Required is 0 for every PD." in html
    assert pd.number in html and reverse("reports:data_health") in html


def _unlinked_plan(year: int, quarters) -> None:
    """A planned-visit row eTools sent whose PD reference matches no programme document (not linked)."""
    q1, q2, q3, q4 = quarters
    dm.PlannedVisits.objects.create(
        datamart_id=next(_ids),
        intervention=None,
        pd_reference_number=f"TEST/PCA{year}999/PD{year}999-1",
        year=year,
        q1=q1,
        q2=q2,
        q3=q3,
        q4=q4,
    )


def test_planned_visits_matching_no_pd_are_named_on_the_tab_and_in_the_file(world, people):
    """eTools' planned-visit rows that the sync could not link to a programme document are in no figure:
    the tab says how many rows of the year there are and how many visits they plan, with the Data health
    link, and so does the Excel file's About sheet. When no plan row of the year is linked, that warning
    replaces "eTools has no planned visits", which would be wrong."""
    client = _client(people["viewer"])
    response = _plan(client)
    before = response.context["figures"]
    assert 'id="plan-unlinked"' not in response.content.decode()
    assert "Planned visits matching no programme document" not in _about(_plan(client, export="xlsx"))
    _unlinked_plan(YEAR, (2, 1, 0, 0))
    _unlinked_plan(YEAR - 1, (5, 0, 0, 0))  # another year's: not this year's
    response = _plan(client)
    html = response.content.decode()
    assert response.context["figures"] == before  # left out of every figure
    alert = re.search(r'<div [^>]*id="plan-unlinked">(.*?)</div>', html, re.S).group(1)
    assert " ".join(re.sub(r"<[^>]+>", " ", alert).split()) == (
        f"eTools sent 1 planned-visit row for {YEAR} (3 visits) that matches no programme document: it is "
        "left out of every figure. Data health"
    )
    assert f'href="{reverse("reports:data_health")}"' in alert
    assert _about(_plan(client, export="xlsx"))["Planned visits matching no programme document"] == (
        f"eTools sent 1 planned-visit row for {YEAR} (3 visits) matching no programme document: left out "
        "of every figure (Data health shows the rows not linked)."
    )
    _unlinked_plan(YEAR, (1, 0, 0, 0))
    assert "eTools sent 2 planned-visit rows for" in _plan(client).content.decode()
    # no plan row of the year linked: the rows exist, so eTools did send a plan
    dm.PlannedVisits.objects.exclude(intervention=None).delete()
    html = _plan(client).content.decode()
    assert 'id="plan-unlinked"' in html and "eTools has no planned visits for" not in html
    dm.PlannedVisits.objects.filter(intervention=None).delete()
    assert "eTools has no planned visits for" in _plan(client).content.decode()


def test_the_partners_behind_tile_says_how_it_differs_from_the_assurance_page_it_links_to(people):
    """The tile counts the partners of the plan's programme documents behind on programmatic visits; the
    Assurance page it links to counts every partner with a HACT year of the year, behind on programmatic
    visits or spot checks. Both counts stand (SPEC_HACT §2.2), and the words say how they differ."""
    a, b, c = partner("Partner A"), partner("Partner B"), partner("Partner C")
    plan(programme(a))
    plan(programme(b))
    hact_year(a, required=2, completed=1)  # behind on programmatic visits
    hact_year(b, required=1, completed=1, sc_required=1, sc_completed=0)  # behind on spot checks only
    hact_year(c, required=3, completed=0)  # behind, but no programme document this year
    client = _client(people["viewer"])
    response = _plan(client)
    assert response.context["figures"]["etools_behind"] == {"partners": 1, "with_row": 2}
    html = response.content.decode()
    label = re.escape('<div class="kpi__label">Partners behind their HACT minimum (eTools)</div>')
    tile = re.search(label + r".*?</a>\s*</div>", html, re.S).group(0)
    assert '<div class="kpi__value">1</div>' in tile
    assert (
        "programmatic visits only, of 2 partners of these programme documents with an eTools HACT year"
        in tile
    )
    assert f'href="{reverse("reports:assurance")}?hact_year={YEAR}">Assurance →</a>' in tile
    assurance = client.get(reverse("reports:assurance"), {"hact_year": str(YEAR)}).content.decode()
    assert "3 partners behind their minimum requirements" in assurance
    # the help explains the difference; it no longer calls them the same figure
    guide = client.get(reverse("help:page", args=["field-monitoring"])).content.decode()
    assert "the partner-level figure of the" not in guide
    assert "Assurance counts spot checks too, and every partner with a HACT" in " ".join(guide.split())


def test_no_visit_on_any_pd_counts_the_partners_visits_outside_the_filter(world, people):
    client = _client(people["viewer"])
    # Partner C: its Child Protection PD is not visited, its Health PD was visited once
    response = _plan(client, section="Child Protection")
    assert response.context["figures"]["not_visited"] == {
        "pds": 1,
        "partners": 1,
        "partners_no_visit": 0,
        "ended": 0,
    }
    assert "0 of them with no visit on any programme document" in response.content.decode()
    # Partner B has no visit anywhere
    assert _plan(client).context["figures"]["not_visited"]["partners_no_visit"] == 1


def test_by_section_under_a_section_filter_keeps_the_chosen_sections_only(world, people):
    client = _client(people["viewer"])
    every = {s["section"]: s for s in hact_plan.by_section(hact_plan.build(YEAR, TODAY))}
    response = _plan(client, section="Education")  # b is of Education and Health
    lines = [s for s in response.context["by_section"] if not s["total"]]
    assert [s["section"] for s in lines] == ["Education"]
    assert {k: v for k, v in lines[0].items() if k != "url"} == every["Education"]
    html = response.content.decode()
    panel = re.search(r'<table class="table table-sm" id="plan-sections">.*?</table>', html, re.S).group(0)
    assert ">Education</a>" in panel and ">Health</a>" not in panel
    book = _workbook(_plan(client, section="Education", export="xlsx"))
    assert [c.value for c in book["By section"]["A"]] == [
        "section",
        "Education",
        "Total (each programme document once)",
    ]
    assert [c.value for c in book["By section"][2]] == ["Education", 2, 2, 1, 5]
    # with no section chosen, every section keeps its line
    assert [s["section"] for s in _plan(client).context["by_section"] if not s["total"]] == [
        "Child Protection",
        "Education",
        "Health",
    ]


def test_a_change_of_year_redraws_the_filter_bar_with_that_years_lists(world, people):
    d = partner("Partner D")  # a programme document of last year only
    last = programme(d, sections=("Nutrition",))
    PCA.objects.filter(pk=last.pk).update(start=D(YEAR - 1, 1, 1), end=D(YEAR - 1, 12, 31))
    plan(last, year=YEAR - 1)
    client = _client(people["viewer"])
    full = _plan(client)
    assert full.context["breadcrumbs"][-1] == {"label": "HACT visit plan", "url": None}  # no year in it
    assert (
        '<div id="plan-filters">\n<div class="filters-collapse" data-filters-collapse>'
        in full.content.decode()
    )

    def hx(year, came_from=None, **params):
        headers = {"HTTP_HX_REQUEST": "true"}
        if came_from is not None:
            headers["HTTP_HX_CURRENT_URL"] = f"http://testserver{PAGE}?tab=plan&year={came_from}"
        return client.get(PAGE, {"tab": "plan", "year": str(year), **params}, **headers).content.decode()

    # the year select changed: the filter bar comes again, out of band, with last year's lists
    html = hx(YEAR - 1, came_from=YEAR)
    # the box around the folding wrapper is swapped, so the page's script enhances the wrapper inside it
    bar = re.search(
        r'<div id="plan-filters" hx-swap-oob="true">\s*<div class="filters-collapse".*?</form>', html, re.S
    )
    assert bar
    assert ">Partner D</option>" in bar.group(0) and ">Partner A</option>" not in bar.group(0)
    assert ">Nutrition</option>" in bar.group(0) and ">Child Protection</option>" not in bar.group(0)
    assert f'<option value="{YEAR - 1}" selected>' in bar.group(0)
    assert last.number in html and world["a1"].number not in html
    # back to this year: this year's lists again
    html = hx(YEAR, came_from=YEAR - 1)
    assert 'id="plan-filters"' in html and ">Partner A</option>" in html and ">Partner D</option>" not in html
    # any other change of the filter (the search box is typed in): the bar stays as it is
    assert 'id="plan-filters"' not in hx(YEAR, came_from=YEAR, q="partner")
    assert 'id="plan-filters"' not in hx(YEAR - 1, came_from=YEAR - 1, section="Nutrition")
    assert 'id="plan-filters"' in hx(YEAR)  # where the request came from is not known: drawn again


def _chosen(html: str, field: str) -> list[str]:
    """The options the filter bar marks in its list ``field`` (``f-partner``...), by their text."""
    select = re.search(rf'<select [^>]*id="{field}".*?</select>', html, re.S).group(0)
    return re.findall(r'<option value="[^"]*" selected>([^<]*)</option>', select)


def test_a_value_the_new_year_lacks_stays_chosen_in_the_filter_bar(world, people):
    """The partner page links to this year's plan on its partner; the year is then changed to one the
    partner has no programme document in. The list is still filtered on the partner (none matches), so
    the bar must still show it chosen, where it can be taken off: never a bar with nothing chosen over
    an empty list, which the next change of the bar would quietly drop."""
    d = partner("Partner D")  # a programme document of last year only
    last = programme(d, sections=("Nutrition",))
    PCA.objects.filter(pk=last.pk).update(start=D(YEAR - 1, 1, 1), end=D(YEAR - 1, 12, 31))
    plan(last, year=YEAR - 1)
    a = world["partners"]["a"]
    client = _client(people["viewer"])
    assert _chosen(_plan(client, partner=str(a.pk)).content.decode(), "f-partner") == ["Partner A"]

    # the year select changed: the bar comes again with last year's lists, Partner A still chosen
    response = client.get(
        PAGE,
        {"tab": "plan", "year": str(YEAR - 1), "partner": str(a.pk)},
        HTTP_HX_REQUEST="true",
        HTTP_HX_CURRENT_URL=f"http://testserver{PAGE}?tab=plan&year={YEAR}&partner={a.pk}",
    )
    html = response.content.decode()
    assert response.context["f"].partners == (a.pk,) and response.context["count"] == 0
    assert "No programme document matches" in html
    bar = re.search(r'<div id="plan-filters" hx-swap-oob="true">.*?</form>', html, re.S).group(0)
    assert _chosen(bar, "f-partner") == ["Partner A"] and ">Partner D</option>" in bar
    # the same page loaded whole (a reload of that address) agrees too
    whole = _plan(client, year=str(YEAR - 1), partner=str(a.pk)).content.decode()
    assert _chosen(whole, "f-partner") == ["Partner A"]
    # the Excel file's About sheet names it as the bar does, never by its id; so too for a partner of
    # the year that a section filter leaves without a programme document
    about = _about(_plan(client, year=str(YEAR - 1), partner=str(a.pk), export="xlsx"))
    assert about["Filter"] == "Partners: Partner A"
    about = _about(_plan(client, partner=str(a.pk), section="Child Protection", export="xlsx"))
    assert about["Filter"] == "Sections: Child Protection; Partners: Partner A"

    # a section, risk rating, status and type the year lacks; a section typed in another case
    html = _plan(
        client, year=str(YEAR - 1), section="education", risk="High", status="signed", type="SPD"
    ).content.decode()
    assert _chosen(html, "f-section") == ["education"] and _chosen(html, "f-risk") == ["High"]
    assert _chosen(html, "f-status") == ["Signed"] and _chosen(html, "f-type") == ["SPD"]
    assert "No programme document matches" in html
    html = _plan(client, section="health").content.decode()  # select() reads any case: "Health" is it
    assert _chosen(html, "f-section") == ["Health"] and html.count(">health</option>") == 0
    assert _ids_of(_plan(client, section="health")) == {world["a2"].pk, world["b"].pk, world["c1"].pk}


APP_JS = Path(django_settings.BASE_DIR) / "neurodb" / "web" / "static" / "js" / "app.js"
PARTS_DOM = Path(django_settings.BASE_DIR) / "tests" / "web" / "parts_dom.mjs"
NODE = shutil.which("node")
# app.js's filter bar folding (initFiltersCollapse) and keepFiltersOpen, run in node on the stand-in page
# of tests/web/parts_dom.mjs; the bar is swapped out of band as htmx 2.0.10 does it: htmx:oobBeforeSwap on
# the element it replaces, with the new content as a fragment, then htmx:load on the new element.
FOLD_SCRIPT = """
const dom = await import(process.argv[1]);
// what the stand-in lacks and initFiltersCollapse uses: Element.before() and classList.toggle()
dom.Element.prototype.before = function (node) { this.parentNode.insertBefore(node, this); };
const { doc } = dom.installPage();
const create = doc.createElement.bind(doc);
doc.createElement = (tag) => {
  const el = create(tag);
  el.classList.toggle = (name) => {
    const on = !el.classList.contains(name);
    if (on) el.classList.add(name);
    else el.classList.remove(name);
    return on;
  };
  return el;
};
%(code)s
const bar = () => doc.h("div", { id: "plan-filters" },
  doc.h("div", { class: "filters-collapse", "data-filters-collapse": true },
    doc.h("form", { class: "filter-bar", "data-filter-bar": true }, doc.h("select", { id: "f-year" }))));
const state = () => {
  const wrap = doc.querySelector("[data-filters-collapse]");
  const button = wrap.querySelector("button");
  return { open: wrap.classList.contains("is-open"), expanded: button.getAttribute("aria-expanded") };
};
const out = {};
for (const opened of [true, false]) {
  doc.querySelector("#plan-filters")?.remove();
  doc.body.append(bar());
  initFiltersCollapse(doc.body);
  if (opened) doc.querySelector("[data-filters-collapse]").querySelector("button").click();
  const before = state();
  const old = doc.querySelector("#plan-filters");
  const incoming = bar();
  const fragment = doc.h("div", {}, incoming);  // the fragment holds the new element, as htmx's does
  keepFiltersOpen({ detail: { shouldSwap: true, target: old, fragment } });
  old.parentNode.insertBefore(incoming, old);
  old.remove();
  initFiltersCollapse(incoming);  // htmx:load on the new element: enhance() reaches what is inside it
  out[opened ? "opened" : "folded"] = { before, after: state() };
}
console.log(JSON.stringify(out));
"""


def _app_function(app: str, name: str) -> str:
    """The source of app.js's top-level function ``name``."""
    body = app.split(f"\nfunction {name}(", 1)[1].split("\n}\n", 1)[0]
    return f"function {name}({body}\n}}\n"


@pytest.mark.skipif(NODE is None, reason="node is not installed")
def test_a_filter_bar_opened_on_a_phone_stays_open_when_the_year_redraws_it():
    """On a phone the plan's filter bar folds behind its "Filters" button. A change of year draws the bar
    again (out of band, in a new wrapper): the bar the user opened stays open, its button pressed, and a
    folded one stays folded."""
    app = APP_JS.read_text()
    assert 'document.body.addEventListener("htmx:oobBeforeSwap", keepFiltersOpen);' in app
    helpers = "\n".join(line for line in app.splitlines() if line.startswith(("const $ =", "const $$ =")))
    code = helpers + "\n" + _app_function(app, "initFiltersCollapse") + _app_function(app, "keepFiltersOpen")
    result = subprocess.run(
        [NODE, "--input-type=module", "-e", FOLD_SCRIPT % {"code": code}, PARTS_DOM.as_uri()],
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert result.returncode == 0, result.stderr
    out = json.loads(result.stdout)
    assert out["opened"] == {
        "before": {"open": True, "expanded": "true"},
        "after": {"open": True, "expanded": "true"},
    }
    assert out["folded"] == {
        "before": {"open": False, "expanded": "false"},
        "after": {"open": False, "expanded": "false"},
    }


# ---------------------------------------------------------------------------------- the list
def _row_html(html: str, pd) -> str:
    match = re.search(rf'<tr id="pd-{pd.pk}">.*?</tr>', html, re.S)
    assert match, pd.number
    return match.group(0)


def test_the_list_links_the_partner_the_pd_and_its_visits(world, people, settings):
    client = _client(people["viewer"])
    settings.FMM_ENABLED = True
    html = _plan(client).content.decode()
    row = _row_html(html, world["a1"])
    pd_url = reverse("reports:programme_detail", args=[world["a1"].pk])
    assert f'href="{reverse("reports:partner_profile", args=[world["partners"]["a"].pk])}"' in row
    assert f'href="{pd_url}" hx-get="{pd_url}" hx-target="#modal-content"' in row
    completed = re.search(r'<a href="(/fmm/\?[^"]+)"', row).group(1).replace("&amp;", "&")
    assert QueryDict(completed.split("?", 1)[1]).dict() == {
        "programmatic": "1",
        "status": "reported",
        "year": str(YEAR),
        "section": "",
        "pd": str(world["a1"].pk),
    }
    # the row's figures: required, completed (FM 1, TPM 0, one awaiting approval), counted, still needed
    assert "FM 1 · TPM 0" in row and "1 awaiting approval" in row
    assert '<span class="pill pill--sm pill--warning">3</span>' in row
    assert "HACT FMM needed" in row and "eTools HACT" in row and "120,000.5" in row
    assert "not visited yet" in _row_html(html, world["b"]) and "PD ended, visits missing" in _row_html(
        html, world["b"]
    )
    assert "No visit planned" in _row_html(html, world["c1"]) and "1 completed visit" in _row_html(
        html, world["c1"]
    )
    assert "partner rating" not in _row_html(html, world["c1"])  # Partner C has a HACT year
    # this quarter's header is marked; the due column names it
    quarter = (TODAY.month - 1) // 3 + 1
    assert f'class="num plan-q--now" title="This quarter">Q{quarter}</th>' in html
    assert f"Due by end of Q{quarter}" in html

    settings.FMM_ENABLED = False
    row = _row_html(_plan(client).content.decode(), world["a1"])
    # its findings and TPM activities, searched by the PCA/PD numbers every spelling of the reference holds
    assert f'href="{PAGE}?q={_pair(world["a1"])}&amp;year={YEAR}"' in row
    assert "/fmm/" not in row


def _pair(pd) -> str:
    """The PCA/PD numbers of a test reference ("TEST/PCA…/PD…"): what the findings tab is searched by."""
    return pd.number.removeprefix("TEST/")


def _figure_href(row: str) -> str:
    """Where the Completed figure of a row's HTML links to."""
    cell = re.search(r'<td class="num" data-value="\d+">(<a [^>]*>)\d+</a>', row).group(1)
    return re.search(r'href="([^"]+)"', cell).group(1).replace("&amp;", "&")


def test_completed_of_a_pd_visited_by_tpm_points_to_its_tpm_activities(world, people, settings):
    """Monitoring insights builds its visits from the field monitoring findings only, never from the TPM
    activities. With it on, "TPM m" under the figure opens the PD's TPM activities on the first tab, and
    so does the Completed figure itself of a PD visited by TPM alone (a2: FM 0, TPM 1): never an empty
    list in Monitoring insights. A PD with a field monitoring report there to list (d: TPM 1, an FM report
    awaiting approval) keeps the figure on Monitoring insights. The note under the list says so."""
    d = programme(partner("Partner D"))
    plan(d, (1, 1, 0, 0))
    tpm(d)
    fm_visit(d, status="submitted")
    client = _client(people["viewer"])
    settings.FMM_ENABLED = True
    html = _plan(client).content.decode()
    row = _row_html(html, world["a2"])
    link = re.search(r'FM 0 · <a href="([^"]+)" title="[^"]+">TPM 1</a>', row).group(1).replace("&amp;", "&")
    assert link == f"{PAGE}?q={_pair(world['a2'])}&year={YEAR}#tpm-activities"
    # the figure: the same TPM activities, never Monitoring insights' list, which holds no visit of a2
    assert _figure_href(row) == link and "/fmm/" not in row
    tpm_table = re.search(
        r'<table [^>]*id="tpm-activities".*?</table>', client.get(link.split("#")[0]).content.decode(), re.S
    ).group(0)
    assert world["a2"].number in tpm_table and "No TPM activities." not in tpm_table
    assert _figure_href(_row_html(html, world["a1"])).startswith("/fmm/?")  # FM 1: Monitoring insights
    assert "#tpm-activities" not in _row_html(html, world["a1"])  # no TPM visit: no link
    assert "FM 1 · TPM 0" in _row_html(html, world["a1"])
    row = _row_html(html, d)  # FM 0 and TPM 1, but its FM report awaiting approval is listed there
    assert "FM 0 · <a" in row and "1 awaiting approval" in row
    assert _figure_href(row).startswith("/fmm/?") and _figure_href(row).endswith(f"&pd={d.pk}")
    assert "TPM visits are not in Monitoring insights" in html
    assert "and so does the figure when TPM visits are all there is to list" in html
    # with Monitoring insights off the figure itself opens the same list: no second link
    settings.FMM_ENABLED = False
    html = _plan(client).content.decode()
    assert "#tpm-activities" not in html and "TPM visits are not in Monitoring insights" not in html
    assert "FM 0 · TPM 1" in _row_html(html, world["a2"])


def _finding(pd, owner, activity: int, *, entity, entity_type="PD/SSFA", start, end) -> None:
    """One finding row of a completed programmatic field monitoring visit."""
    dm.MonitoringFinding.objects.create(
        datamart_id=next(_ids),
        partner=owner,
        vendor_number=owner.vendor_number,
        entity=entity,
        entity_type=entity_type,
        intervention=pd,
        monitoring_activity=f"FM-TEST-{activity}",
        monitoring_activity_id=activity,
        status="completed",
        is_programmatic_visit=True,
        start_date=start,
        end_date=end,
    )


def test_completed_with_monitoring_insights_off_finds_every_spelling_of_the_reference(people, settings):
    """With Monitoring insights off, Completed opens the Findings and visits tab searched by the PD's
    PCA/PD numbers, the pair its visits are linked by, not by its whole reference: a finding that names
    an earlier amendment under another prefix, and a TPM activity whose reference has no amendment,
    are both listed. That tab dates a field monitoring visit by its start date, so a visit that started
    the year before and is counted this year (by its end date) is listed under that year: the note
    under the list says so."""
    settings.FMM_ENABLED = False
    a = partner("Partner A")
    pd = programme(a)
    pair = _pair(pd)
    PCA.objects.filter(pk=pd.pk).update(number=f"TEST/{pair}-2")
    pd.refresh_from_db()
    plan(pd)
    first, second = next(_ids), next(_ids)
    # its finding names the first amendment, under another prefix: linked to the PD all the same
    _finding(pd, a, first, entity=f"TESTA/{pair}-1", start=D(YEAR, 3, 1), end=D(YEAR, 3, 2))
    # started on 30 December of the year before, ended on 2 January: counted this year
    _finding(pd, a, second, entity=pd.number, start=D(YEAR - 1, 12, 30), end=D(YEAR, 1, 2))
    tpm(pd)
    dm.TPMActivity.objects.filter(intervention=pd).update(pd_reference_number=f"TESTA/{pair}")
    client = _client(people["viewer"])
    html = _plan(client).content.decode()
    row = _row_html(html, pd)
    assert "FM 2 · TPM 1" in row
    href = _figure_href(row)
    assert href == f"{PAGE}?q={pair}&year={YEAR}"
    found = client.get(href).context
    assert {f.monitoring_activity_id for f in found["page_obj"].object_list} == {first}
    assert [t.intervention_id for t in found["data"]["tpm_activities"]] == [pd.pk]
    found = client.get(f"{PAGE}?q={pair}&year={YEAR - 1}").context  # where that tab dates the second
    assert {f.monitoring_activity_id for f in found["page_obj"].object_list} == {second}
    assert "That tab dates a field monitoring visit by its start date" in html


def test_a_visit_to_the_partner_naming_no_pd_shows_beside_its_pds_not_counted(people):
    """A completed programmatic field monitoring visit whose findings name only the partner and a CP
    output counts for the partner on the Assurance page and for none of its programme documents, which
    stay "not visited yet": each of the partner's rows says so under Completed (not counted), and so do
    the By partner view and the files. A visit that also names a programme document is that PD's."""
    a, b = partner("Partner A"), partner("Partner B")
    a1, a2, b1 = programme(a), programme(a, sections=("Health",)), programme(b)
    for pd in (a1, a2, b1):
        plan(pd, (1, 1, 0, 0))
    partner_level, named = next(_ids), next(_ids)
    when = {"start": D(YEAR, 2, 1), "end": D(YEAR, 2, 2)}
    _finding(None, a, partner_level, entity=a.name, entity_type="Partner", **when)
    _finding(None, a, partner_level, entity="1.1 Test output", entity_type="CP Output", **when)
    _finding(None, b, named, entity=b.name, entity_type="Partner", **when)
    _finding(b1, b, named, entity=b1.number, **when)
    assert fm.programmatic_visits_by_partner(YEAR) == {a.pk: 1, b.pk: 1}  # the Assurance page's

    client = _client(people["viewer"])
    response = _plan(client)
    html = response.content.decode()
    rows = {r["pd_id"]: r for r in response.context["rows"]}
    for pd in (a1, a2):
        row = rows[pd.pk]
        assert (row["completed"], row["zero"], row["fm_partner_level"]) == (0, True, 1)
        row = _row_html(html, pd)
        assert "1 FM visit to the partner naming no PD (not counted)" in row and "not visited yet" in row
    assert (rows[b1.pk]["completed_fm"], rows[b1.pk]["fm_partner_level"]) == (1, 0)
    assert "naming no PD" not in _row_html(html, b1)
    assert "visits to the partner whose findings name none of its programme documents" in html  # the note

    partners = {p["partner"]: p for p in _plan(client, view="partners").context["partners"]}
    assert (partners["Partner A"]["fm_partner_level"], partners["Partner B"]["fm_partner_level"]) == (1, 0)
    assert "FM visits naming no PD" in _plan(client, view="partners").content.decode()
    lines = {line["reference_number"]: line for line in _csv(_plan(client, export="csv"))}
    assert [lines[pd.number]["fm_partner_visits_naming_no_pd"] for pd in (a1, a2, b1)] == ["1", "1", "0"]
    lines = _csv(_plan(client, view="partners", export="csv"))
    assert [(line["partner"], line["fm_partner_visits_naming_no_pd"]) for line in lines] == [
        ("Partner A", "1"),
        ("Partner B", "0"),
    ]


def test_the_focal_points_show_on_screen_only(world, people):
    html = _plan(_client(people["viewer"])).content.decode()
    header = re.search(r'<th scope="col"[^>]*>UNICEF focal points</th>', html).group(0)
    assert 'class="d-none d-md-table-cell"' in header and 'data-export="no"' in header
    cell = re.search(
        r'<td class="small d-none d-md-table-cell" data-export="no">(.*?)</td>', _row_html(html, world["a1"])
    )
    assert cell and cell.group(1) == FOCAL  # the e-mail address is never shown
    assert FOCAL_EMAIL not in html


def test_by_partner_shows_the_etools_minimum_and_the_two_flags(world, people):
    client = _client(people["viewer"])
    response = _plan(client, view="partners")
    html = response.content.decode()
    partners = {p["partner"]: p for p in response.context["partners"]}
    assert set(partners) == {"Partner A", "Partner B", "Partner C"}
    a, c = partners["Partner A"], partners["Partner C"]
    assert (a["pds"], a["required"], a["completed"], a["still_needed"]) == (2, 5, 2, 3)
    assert (a["etools_required"], a["etools_completed"], a["etools_gap"]) == (2, 2, 0)
    assert a["minimum_met_pd_needed"] and not a["plans_below_minimum"]
    assert c["plans_below_minimum"] and partners["Partner B"]["etools_gap"] is None
    for label in ("eTools HACT required", "eTools HACT completed", "eTools gap"):
        assert label in html
    assert "Partner minimum met in eTools, PD visits still needed" in html
    assert "PD plans below the partner minimum" in html
    assert f'href="{reverse("reports:assurance")}?hact_year={YEAR}"' in html
    # a partner's PDs open the list by programme document on that partner
    pds_url = a["pds_url"]
    assert "view=" not in pds_url and f"partner={world['partners']['a'].pk}" in pds_url
    assert _ids_of(client.get(PAGE + pds_url)) == {world["a1"].pk, world["a2"].pk}
    assert 'name="view" value="partners"' in html  # the filter bar keeps the view
    assert 'id="plan-pds"' not in html and response.context["rows"] == []


def test_the_partner_flags_read_every_pd_of_the_partner_whatever_the_filter(world, people):
    client = _client(people["viewer"])
    # Partner A: 4 + 1 visits planned against an eTools minimum of 2; Partner C: 0 + 2 against 5
    for params in (
        {},
        {"section": "Health"},
        {"remark": "completed"},
        {"zero": "1"},
        {"remark": ["needed", "ended"]},
        {"section": "Health", "remark": "completed"},
    ):
        response = _plan(client, view="partners", **params)
        partners = {p["partner"]: p for p in response.context["partners"]}
        assert partners, params
        assert not partners.get("Partner A", {}).get("plans_below_minimum"), params
        if "Partner C" in partners:
            assert partners["Partner C"]["plans_below_minimum"], params
        for row in re.findall(r"<tr>.*?</tr>", response.content.decode(), re.S):
            if ">Partner A</a>" in row:
                assert "PD plans below the partner minimum" not in row, params
    # the files say the same: by partner under a section filter, Excel and CSV
    book = _workbook(_plan(client, section="Health", export="xlsx"))
    header = [c.value for c in book["By partner"][1]]
    lines = {
        line["partner"]: line
        for line in (
            dict(zip(header, (c.value for c in row), strict=True))
            for row in book["By partner"].iter_rows(min_row=2)
        )
    }
    assert lines["Partner A"]["required_pv"] == 1  # the filter's PDs
    assert lines["Partner A"]["pd_plans_below_etools_minimum"] == "no"
    assert lines["Partner C"]["pd_plans_below_etools_minimum"] == "yes"
    csv_lines = {
        line["partner"]: line
        for line in _csv(_plan(client, view="partners", remark="completed", export="csv"))
    }
    assert csv_lines["Partner A"]["pd_plans_below_etools_minimum"] == "no"


def test_the_partner_flags_are_readable_in_both_themes(world, people):
    html = _plan(_client(people["viewer"]), view="partners").content.decode()
    flags = re.findall(r'<span class="cell-sub ([^"]*)"[^>]*>(?:Partner minimum met|PD plans below)', html)
    # Bootstrap's text-warning is yellow on white (1.6:1); its emphasis colour follows the theme
    assert flags == ["text-warning-emphasis", "text-warning-emphasis"]


def test_copy_and_csv_of_the_list_keep_both_dates_and_the_sections_apart(world, people):
    html = _plan(_client(people["viewer"])).content.decode()
    header = re.search(r'<table class="table table-sm table-hover" id="plan-pds".*?</thead>', html, re.S)
    headings = re.findall(r"<th scope=\"col\"[^>]*>(.*?)</th>", header.group(0), re.S)
    assert "Start" in headings and "End" in headings
    row = _row_html(html, world["b"])
    # one cell per heading, so Copy and CSV keep the columns in line
    assert len(re.findall(r"<td\b", row)) == len(headings)
    assert f'<td class="small text-nowrap" data-value="{YEAR}-01-01">1 Jan {YEAR}</td>' in row
    assert f'<td class="small text-nowrap" data-value="{YEAR}-03-31">31 Mar {YEAR}</td>' in row
    assert '<td data-value="Education, Health"><span class="chips">' in row


def test_a_long_comment_wraps_at_a_readable_width():
    from pathlib import Path

    css = (Path(__file__).resolve().parents[2] / "neurodb/web/static/css/app.css").read_text()
    rule = re.search(r"\.plan-table \.plan-comment \{([^}]*)\}", css).group(1)
    assert "width: 20rem" in rule and "max-width" not in rule


# ---------------------------------------------------------------------------------- section comments
def test_the_comment_button_shows_only_to_those_who_may_comment(world, people):
    button = 'hx-target="#modal-content" hx-push-url="false">Comment</a>'
    url = reverse("reports:visit_plan_comment", args=[world["a1"].pk])
    admin = _plan(_client(people["admin"])).content.decode()
    assert all(button in _row_html(admin, world[k]) for k in ("a1", "a2", "b", "c1", "c2"))
    editor = _plan(_client(people["editor"])).content.decode()  # Education
    assert f'href="{url}?year={YEAR}"' in _row_html(editor, world["a1"])
    assert button in _row_html(editor, world["b"])  # Education and Health
    assert button not in _row_html(editor, world["a2"]) and button not in _row_html(editor, world["c2"])
    viewer = _plan(_client(people["viewer"])).content.decode()
    assert button not in viewer and "/comment/" not in viewer


def test_an_editor_comments_from_the_list_and_the_row_swaps_in_place(world, people):
    client = _client(people["editor"])
    pd = world["a1"]
    url = reverse("reports:visit_plan_comment", args=[pd.pk])
    form = client.get(url, {"year": YEAR}, HTTP_HX_REQUEST="true").content.decode()
    assert f'hx-post="{url}" hx-target="#pd-{pd.pk}" hx-swap="outerHTML"' in form

    swapped = client.post(url, {"text": "Visit planned for November.", "year": YEAR}, HTTP_HX_REQUEST="true")
    html = swapped.content.decode()
    assert swapped.status_code == 200 and html.strip().startswith(f'<tr id="pd-{pd.pk}">')
    assert html.count("<tr") == 1 and "Visit planned for November." in html
    assert f"updated {timezone.localdate():%-d %b %Y}" in html and ">Edit</a>" in html
    assert FOCAL in html  # the row as the list draws it
    assert "edu-editor" not in html  # never who wrote it

    page = _plan(client).content.decode()
    assert "Visit planned for November." in _row_html(page, pd)
    # without JavaScript: back to the PD's row
    response = client.post(url, {"text": "x" * 600, "year": YEAR})
    assert response.status_code == 302 and response["Location"] == f"{PAGE}?tab=plan&year={YEAR}#pd-{pd.pk}"
    assert dm.VisitPlanComment.objects.get().text == "x" * 500  # capped
    removed = client.post(url, {"remove": "1", "year": YEAR}, HTTP_HX_REQUEST="true").content.decode()
    assert ">Comment</a>" in removed and not dm.VisitPlanComment.objects.exists()


@pytest.mark.parametrize("who", ["other_editor", "viewer"])
def test_others_get_403_on_the_comment(world, people, who):
    client = _client(people[who])
    url = reverse("reports:visit_plan_comment", args=[world["a1"].pk])
    assert client.get(url, {"year": YEAR}, HTTP_HX_REQUEST="true").status_code == 403
    assert client.post(url, {"text": "Not allowed", "year": YEAR}, HTTP_HX_REQUEST="true").status_code == 403
    assert not dm.VisitPlanComment.objects.exists()


MARKUP = '</textarea><script>x</script>" onmouseover="y'
ESCAPED = "&lt;/textarea&gt;&lt;script&gt;x&lt;/script&gt;&quot; onmouseover=&quot;y"


def test_markup_in_a_comment_is_shown_as_text_everywhere(world, people):
    client = _client(people["editor"])
    pd = world["a1"]
    url = reverse("reports:visit_plan_comment", args=[pd.pk])
    swapped = client.post(url, {"text": MARKUP, "year": YEAR}, HTTP_HX_REQUEST="true").content.decode()
    assert dm.VisitPlanComment.objects.get().text == MARKUP  # kept as written: it is shown as text
    pages = {
        "the row swapped in": swapped,
        "the list": _row_html(_plan(client).content.decode(), pd),
        "the form in the window": client.get(url, {"year": YEAR}, HTTP_HX_REQUEST="true").content.decode(),
        "the form page": client.get(url, {"year": YEAR}).content.decode(),
    }
    for where, html in pages.items():
        assert ESCAPED in html, where
        for raw in ("<script>x", "</textarea><", '" onmouseover="y'):
            assert raw not in html, (where, raw)
    # in the list, both the cell's text and its data-value (what Copy and CSV write)
    assert pages["the list"].count(ESCAPED) == 2 and pages["the row swapped in"].count(ESCAPED) == 2


def test_a_deactivated_user_cannot_comment(world, people):
    url = reverse("reports:visit_plan_comment", args=[world["a1"].pk])
    for who in ("editor", "admin"):
        client = _client(people[who])
        assert client.get(url, {"year": YEAR}).status_code == 200  # signed in and allowed
        User.objects.filter(pk=people[who].pk).update(is_active=False)  # deactivated: its session stays
        for response in (
            client.get(url, {"year": YEAR}),
            client.post(url, {"text": "Visit planned.", "year": YEAR}),
        ):
            assert response.status_code == 302 and response["Location"].startswith(reverse("account_login"))
        response = client.post(url, {"text": "Visit planned.", "year": YEAR}, HTTP_HX_REQUEST="true")
        assert "Visit planned." not in response.content.decode()
    assert not dm.VisitPlanComment.objects.exists()


# ---------------------------------------------------------------------------------- the files
def _csv(response) -> list[dict]:
    assert response.status_code == 200 and response["Content-Type"].startswith("text/csv")
    text = b"".join(response.streaming_content).decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))


def _about(response) -> dict:
    """The About sheet of an Excel file of the plan, {name: value}."""
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    return {row[0].value: row[1].value for row in book["About"].iter_rows(min_row=2)}


def _workbook(response):
    assert response.status_code == 200
    assert (
        response["Content-Disposition"]
        == f'attachment; filename="hact-visit-plan-{YEAR}_{datetime.date.today():%Y-%m-%d}.xlsx"'
    )
    return load_workbook(io.BytesIO(response.content))


def test_the_files_hold_the_emails_columns_and_never_a_staff_name(world, people):
    client = _client(people["editor"])
    hact_plan.set_comment(world["a1"], YEAR, "Visit planned for November.", people["editor"])
    book = _workbook(_plan(client, export="xlsx"))
    assert book.sheetnames == ["HACT visit plan", "By partner", "By section", "About"]
    sheet = book["HACT visit plan"]
    header = [cell.value for cell in sheet[1]]
    assert header == FILE_COLUMNS
    lines = [
        dict(zip(header, (cell.value for cell in row), strict=True)) for row in sheet.iter_rows(min_row=2)
    ]
    assert len(lines) == 5
    a1 = next(line for line in lines if line["reference_number"] == world["a1"].number)
    assert (
        a1["required_pv"],
        a1["completed_pv"],
        a1["completed_pv_out_of_minimum"],
        a1["still_needed_by_31_dec"],
    ) == (4, 1, 1, 3)
    assert (a1["remark"], a1["not_visited"], a1["awaiting_approval"], a1["unicef_focal_points_count"]) == (
        "HACT FMM needed",
        "no",
        1,
        1,
    )
    assert a1["section_comment"] == "Visit planned for November."
    assert a1["section_comment_updated"].date() == timezone.localdate()
    assert a1["start_date"].date() == D(YEAR, 1, 1) and float(a1["total_budget"]) == 120000.5  # typed
    b = next(line for line in lines if line["reference_number"] == world["b"].number)
    assert (b["remark"], b["not_visited"], b["status"]) == ("PD ended, visits missing", "yes", "ended")
    assert [c.value for c in book["By section"]["A"]][-1] == "Total (each programme document once)"
    assert [c.value for c in book["By partner"][1]][9:13] == [
        "fm_partner_visits_naming_no_pd",
        "etools_hact_required",
        "etools_hact_completed",
        "etools_hact_gap",
    ]

    # every cell of every sheet: no focal point, no e-mail, no visit lead, no username
    texts = [
        str(cell.value)
        for ws in book.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    for secret in (FOCAL, FOCAL_EMAIL, "Lead Person Test", "edu-editor", "@"):
        assert not any(secret in text for text in texts), secret

    lines = _csv(_plan(client, export="csv"))
    assert list(lines[0]) == FILE_COLUMNS and len(lines) == 5
    flat = " ".join(" ".join(line.values()) for line in lines)
    for secret in (FOCAL, FOCAL_EMAIL, "Lead Person Test", "edu-editor", "@"):
        assert secret not in flat, secret
    assert "Visit planned for November." in flat


def test_the_files_follow_the_filter_and_the_view(world, people):
    client = _client(people["viewer"])
    lines = _csv(_plan(client, zero="1", export="csv"))
    assert {line["reference_number"] for line in lines} == {world["b"].number, world["c2"].number}
    book = _workbook(_plan(client, section="Health", export="xlsx"))
    assert book["HACT visit plan"].max_row == 1 + 3
    about = {row[0].value: row[1].value for row in book["About"].iter_rows(min_row=2)}
    assert about["Year"] == YEAR and about["Filter"] == "Sections: Health"
    assert "required_pv" in about and "completed_pv" in about and "Not counted" in about
    partners = _csv(_plan(client, view="partners", export="csv"))
    assert [line["partner"] for line in partners] == ["Partner A", "Partner B", "Partner C"]
    assert partners[0]["etools_minimum_met_pd_visits_still_needed"] == "yes"
    assert partners[2]["pd_plans_below_etools_minimum"] == "yes"


# ---------------------------------------------------------------------------------- the default tab
def test_the_default_tab_has_the_strip_and_the_tabs(world, people):
    client = _client(people["viewer"])
    response = client.get(PAGE)
    html = response.content.decode()
    strip = re.search(r'<p class="panel plan-strip".*?</p>', html, re.S).group(0)
    text = " ".join(re.sub(r"<[^>]+>", " ", strip).split())
    assert f"HACT visit plan {YEAR}: 2 programme documents (2 partners) not visited yet" in text
    assert "7 visits still needed (3 partners). Open the plan →" in text
    assert f'href="{PAGE}?tab=plan&amp;year={YEAR}"' in strip
    assert 'aria-current="page">Findings and visits</a>' in html
    assert f'href="{PAGE}?tab=plan"' in html  # no year chosen: the plan opens on this year
    # the strip follows the year filter, and a year without a plan has none
    other = client.get(PAGE, {"year": str(YEAR - 1)}, HTTP_HX_REQUEST="true").content.decode()
    assert "plan-strip" not in other and 'hx-swap-oob="true"' in other
    assert f'href="{PAGE}?tab=plan&amp;year={YEAR - 1}"' in other
    dm.PlannedVisits.objects.all().delete()
    assert "plan-strip" not in client.get(PAGE).content.decode()


# ---------------------------------------------------------------------------------- cross-links
def test_the_other_hact_pages_link_the_plan(world, people, settings):
    client = _client(people["viewer"])
    assurance = client.get(reverse("reports:assurance")).content.decode()
    assert f'href="{PAGE}?tab=plan&amp;year={YEAR}">Visit plan per programme document →</a>' in assurance
    a = world["partners"]["a"]
    profile = client.get(reverse("reports:partner_profile", args=[a.pk])).content.decode()
    assert (
        f'href="{PAGE}?tab=plan&amp;partner={a.pk}&amp;year={TODAY.year}">Visit plan per programme document →</a>'
        in profile
    )
    settings.FMM_ENABLED = True
    pd = world["a1"]
    detail = client.get(reverse("reports:programme_detail", args=[pd.pk])).content.decode()
    assert f'href="{PAGE}?tab=plan&amp;q={pd.number}&amp;year={YEAR}">HACT visit plan →</a>' in detail
    # Monitoring insights' HACT block: tests/fmm/test_hact_plan_link.py (it needs the FMM test world)


# ---------------------------------------------------------------------------------- queries
def _many(n: int, author: User) -> None:
    owners = [partner(f"Partner {k}") for k in range(3)]
    for k in range(n):
        pd = programme(owners[k % 3], sections=("Education", "Health")[: 1 + k % 2])
        plan(pd)
        fm_visit(pd)
        tpm(pd)
        hact_plan.set_comment(pd, YEAR, f"Comment {k}", author)
    for owner in owners:
        hact_year(owner, required=2, completed=1)


@pytest.mark.parametrize("who", ["admin", "editor"])
def test_the_plan_takes_a_fixed_number_of_queries(people, who, django_assert_max_num_queries):
    client = _client(people[who])
    _many(3, people[who])
    _plan(client)  # the session's first request
    with django_assert_max_num_queries(PLAN_QUERIES[who]):
        assert _plan(client).status_code == 200
    _many(30, people[who])
    with django_assert_max_num_queries(PLAN_QUERIES[who]):
        response = _plan(client)
    assert len(response.context["rows"]) == 33


# measured (18, 20) + 2 at first; now reached: the partners' field monitoring visits naming no PD (in
# build) and the planned-visit rows matching no PD (the tab) take one query each
PLAN_QUERIES = {"admin": 20, "editor": 22}
