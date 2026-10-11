"""eTools reporting, stage R4 (SPEC6): the disaggregation labels read by the Datamart sync, the Report/
Disaggregation tab (a subtotal per disaggregation and the location total, never a grand sum), Reached by
disaggregation on the indicator page, Partner Fund Monitoring and Funded Programs on the Funds page (the
programme document's own amounts counted once), and the two privacy leaks closed (no e-mail address as a
focal point's name, no contact details in the typed tables' records)."""

from __future__ import annotations

import csv
import datetime
import importlib
import io
from decimal import Decimal

import pytest
import responses
from django.db import connection
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from neurodb.core.models import SyncRun
from neurodb.datamart import funds as fund_views
from neurodb.datamart import models as dm
from neurodb.datamart import reporting
from neurodb.integrations.etools import datamart_sync as sync
from neurodb.integrations.etools.datamart import DatamartClient
from neurodb.integrations.http import make_session
from neurodb.partnerships.models import PCA, PartnerOrganization

pytestmark = pytest.mark.django_db
HX = {"HX-Request": "true"}
BASE = "https://datamart.test"
INDICATOR = "5.2.1 # of adolescents and youth reached with outreach info sessions"
SEX, NATIONALITY = 7, 8
FEMALE, MALE, LEBANESE = 11, 12, 21


def _csv(response) -> list[dict[str, str]]:
    text = b"".join(response.streaming_content).decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))


# ------------------------------------------------------------------------------------------ fixtures
@pytest.fixture
def labelled(db):
    """Two reports of one indicator at Akkar (P-code LB71) and Bekaa, as PRP keys them: QPR1 Accepted
    (Akkar: female 5, Male 15, Lebanese 20, total 20), QPR2 Sent back; the labels of the values read."""
    partner = PartnerOrganization.objects.create(
        etl_id="1", name="Youth Partner", partner_type="Civil Society Organization", vendor_number="V1"
    )
    pd = PCA.objects.create(
        etl_id="11",
        partner=partner,
        partner_name=partner.name,
        number="LEB/PCA2026/PD1",
        title="Youth outreach",
        status="active",
        section_names=["Adolescent & Youth"],
        start=datetime.date(2026, 1, 1),
        end=datetime.date(2027, 12, 31),
    )
    dm.PDIndicator.objects.create(
        datamart_id=1, source_id=500, intervention=pd, title=INDICATOR, display_type="number"
    )
    for kind, pk, label, type_id in (
        ("value", FEMALE, "female", SEX),
        ("value", MALE, "Male", SEX),
        ("value", LEBANESE, "Lebanese", NATIONALITY),
        ("type", SEX, "Sex", None),
        ("type", NATIONALITY, "Nationality", None),
    ):
        dm.DisaggregationLabel.objects.create(kind=kind, prp_id=pk, label=label, disaggregation_id=type_id)

    def row(n, report, number, status, location, p_code, values, end):
        return dm.ReportedIndicator.objects.create(
            datamart_id=n,
            partner=partner,
            intervention=pd,
            partner_name=partner.name,
            pd_reference_number=pd.number,
            progress_report=report,
            indicator_report=f"9{n}",
            report_number=number,
            report_type="QPR",
            report_status=status,
            period_start=end - datetime.timedelta(days=89),
            period_end=end,
            due_date=end + datetime.timedelta(days=30),
            submission_date=end + datetime.timedelta(days=20) if status == "Accepted" else None,
            section="Adolescent & Youth",
            indicator=INDICATOR,
            etools_indicator_id="500",
            location=location,
            p_code=p_code,
            calculation_across_locations="sum",
            calculation_across_periods="sum",
            disaggregation=values,
            last_modify_date=timezone.make_aware(datetime.datetime(2026, 10, 8, 9)),
        )

    q1 = datetime.date(2026, 3, 31)
    row(1, "PR-1", "QPR1", "Accepted", "Akkar", "LB71", {"()": {"v": 20, "d": 1, "c": 20}, "(11,)": {"v": 5, "d": 1, "c": 5}, "(12,)": {"v": 15, "d": 1, "c": 15}, "(21,)": {"v": 20, "d": 1, "c": 20}, "(99,)": {"v": 3, "d": 1, "c": 3}}, q1)  # fmt: skip
    row(2, "PR-1", "QPR1", "Accepted", "Bekaa", "LB81", {"()": {"v": 30, "d": 1, "c": 30}, "(11,)": {"v": 10, "d": 1, "c": 10}, "(12,)": {"v": 20, "d": 1, "c": 20}}, q1)  # fmt: skip
    row(3, "PR-2", "QPR2", "Sent back", "Akkar", "LB71", {"()": {"v": 7, "d": 1, "c": 7}, "(11,)": {"v": 7, "d": 1, "c": 7}}, datetime.date(2026, 6, 30))  # fmt: skip
    return {"partner": partner, "pd": pd}


@pytest.fixture
def funded(db):
    """Two programme documents with their budgets, FR lines on two grants, FR headers, and PRP's
    financial contribution repeated on every indicator of each progress report."""
    shift = PartnerOrganization.objects.create(
        etl_id="21",
        name="Alpha Association",
        partner_type="Civil Society Organization",
        cso_type="National",
        vendor_number="V21",
    )
    gov = PartnerOrganization.objects.create(
        etl_id="22", name="Ministry", partner_type="Government", vendor_number="V22"
    )
    nawaya = PCA.objects.create(
        etl_id="31",
        partner=shift,
        partner_name=shift.name,
        number="LEBA/PCA2023583/PD20261704-3",
        title="Youth network",
        status="active",
        start=datetime.date(2026, 3, 6),
        end=datetime.date(2027, 12, 31),
        total_unicef_budget="582888.90",
        total_budget="644,148.90",
        cso_contribution="61260",
        budget_currency="USD",
    )
    ended = PCA.objects.create(
        etl_id="32",
        partner=gov,
        partner_name=gov.name,
        number="LEBA/PCA2022100/PD2022100",
        title="Ended",
        status="ended",
        start=datetime.date(2022, 1, 1),
        end=datetime.date(2024, 12, 31),
        total_unicef_budget="1000",
        total_budget="1000",
        cso_contribution="0",
        budget_currency="LBP",
    )
    PCA.objects.create(etl_id="33", number="LEBA/DRAFT", title="Draft", status="draft", total_budget="5")
    for n, (pd, grant, donor, amount) in enumerate(
        [
            (nawaya, "SC260014", "United States Fund for UNICEF", "50418.50"),
            (nawaya, "SC260017", "Netherlands", "600000"),
            (nawaya, "SC260017", "Netherlands", "93754.80"),
            (ended, "SM220001", "EU", "1000"),
        ],
        start=1,
    ):
        dm.FundsReservation.objects.create(
            datamart_id=n,
            intervention=pd,
            fr_number=f"04000{n}",
            line_item=n,
            donor=donor,
            grant_number=grant,
            overall_amount=Decimal(amount),
            total_amt=Decimal("999999"),  # the FR total repeated on every line: never added
            start_date=datetime.date(2026, 3, 6),
        )
    dm.FundsReservationHeader.objects.create(
        datamart_id=1, intervention=nawaya, fr_number="040001", actual_amt=Decimal("300000")
    )
    dm.FundsReservationHeader.objects.create(datamart_id=2, intervention=nawaya, fr_number="040002", actual_amt=Decimal("308936.90"))  # fmt: skip

    def prp(n, number, status, end, value, indicator):
        dm.DatamartDocument.objects.create(
            dataset="prp_indicator_reports",
            record_key=str(n),
            source_id=n,
            intervention=nawaya,
            data={
                "pd_reference_number": nawaya.number,
                "report_type": "QPR",
                "report_number": number,
                "report_status": status,
                "time_period_end": end,
                "due_date": end,
                "financial_contribution_to_date": value,
                "financial_contribution_currency": "USD",
                "performance_indicator": indicator,
            },
        )

    prp(1, 1, "Acc", "2026-06-30", "100000.50", "A")
    prp(2, 1, "Acc", "2026-06-30", "100000.50", "B")  # the same report, another indicator: read once
    prp(3, 2, "Sub", "2026-09-30", "224108.86", "A")
    prp(4, 2, "Sub", "2026-09-30", "224108.86", "B")
    prp(5, 3, "Due", "2026-12-31", "999999", "A")  # not submitted: never read
    return {"nawaya": nawaya, "ended": ended, "shift": shift, "gov": gov}


# ------------------------------------------------------------------------------- the label sync
def _client() -> DatamartClient:
    return DatamartClient(
        BASE, "svc-user", "not-a-real-secret", country="Lebanon", session=make_session("", backoff=0)
    )


SOURCES = f"{BASE}/api/latest/sources/prp"


@responses.activate
def test_the_partner_reports_sync_reads_the_labels_of_the_ids_it_sees(labelled, settings):
    """The ids of the report rows' keys are read one by one, then their disaggregations; an id PRP does
    not have is marked; labels already read are not read again before a month."""
    responses.get(f"{BASE}/api/latest/prp/datareport/", json={"results": [], "next": None})
    responses.get(f"{SOURCES}/indicatordisaggregationvalue/99/", json={"id": 99, "value": "Syrian", "disaggregation": NATIONALITY, "active": True})  # fmt: skip
    dm.DisaggregationLabel.objects.filter(kind="type", prp_id=NATIONALITY).delete()
    responses.get(
        f"{SOURCES}/indicatordisaggregation/{NATIONALITY}/", json={"id": NATIONALITY, "name": "Nationality"}
    )
    (run,) = sync.sync_all(only=["partner_reports"], triggered_by="test", client=_client())
    assert run.status == SyncRun.Status.SUCCEEDED, run.details
    labels = run.details["disaggregation_labels"]
    assert labels == {"values_seen": 4, "types_seen": 2, "read": 2, "not_found": 0, "failed": 0, "left": 0}
    assert dm.DisaggregationLabel.objects.get(kind="value", prp_id=99).label == "Syrian"
    assert dm.DisaggregationLabel.objects.get(kind="type", prp_id=NATIONALITY).label == "Nationality"
    # a value PRP does not have: marked, the rest still read
    dm.ReportedIndicator.objects.filter(datamart_id=2).update(disaggregation={"(77,)": 1})
    responses.get(f"{SOURCES}/indicatordisaggregationvalue/77/", status=404)
    assert sync.sync_disaggregation_labels(_client())["not_found"] == 1
    assert dm.DisaggregationLabel.objects.get(kind="value", prp_id=77).not_found
    # read more than a month ago: read again
    dm.DisaggregationLabel.objects.filter(prp_id=FEMALE).update(
        synced_at=timezone.now() - datetime.timedelta(days=40)
    )
    responses.get(f"{SOURCES}/indicatordisaggregationvalue/{FEMALE}/", json={"id": FEMALE, "value": "Female", "disaggregation": SEX})  # fmt: skip
    assert sync.sync_disaggregation_labels(_client())["read"] == 1
    assert dm.DisaggregationLabel.objects.get(kind="value", prp_id=FEMALE).label == "Female"


@responses.activate
def test_labels_that_cannot_be_read_never_fail_the_partner_reports(labelled, monkeypatch):
    responses.get(f"{BASE}/api/latest/prp/datareport/", json={"results": [], "next": None})
    dm.DisaggregationLabel.objects.all().delete()
    for pk in (FEMALE, MALE, LEBANESE, 99):
        responses.get(f"{SOURCES}/indicatordisaggregationvalue/{pk}/", status=403)
    monkeypatch.setattr(sync, "LABEL_FAILURES_IN_A_ROW", 3)
    (run,) = sync.sync_all(only=["partner_reports"], triggered_by="test", client=_client())
    assert run.status == SyncRun.Status.PARTIAL
    assert "stopped after 3 failures in a row" in run.details["disaggregation_labels"]["error"]
    assert dm.ReportedIndicator.objects.count() == 3  # the reports are kept
    # a budget per run, of requests and of time: the rest wait for the next sync
    monkeypatch.setattr(sync, "LABEL_REQUESTS_PER_RUN", 1)
    monkeypatch.setattr(sync, "LABEL_FAILURES_IN_A_ROW", 10)
    assert sync.sync_disaggregation_labels(_client())["left"] == 3
    monkeypatch.setattr(sync, "LABEL_REQUESTS_PER_RUN", 2000)
    monkeypatch.setattr(sync, "LABEL_SECONDS_PER_RUN", -1)
    assert sync.sync_disaggregation_labels(_client())["left"] == 4


@responses.activate
def test_a_value_may_name_its_disaggregation_as_a_nested_record(labelled):
    dm.DisaggregationLabel.objects.all().delete()
    dm.ReportedIndicator.objects.update(disaggregation={"()": 3, "(31,)": 1, "(32,)": 2})
    responses.get(f"{SOURCES}/indicatordisaggregationvalue/31/", json={"id": 31, "value": "under 5", "disaggregation": {"id": 30, "name": "Age"}})  # fmt: skip
    responses.get(f"{SOURCES}/indicatordisaggregationvalue/32/", json={"id": 32, "value": "5 to 17", "disaggregation_id": 30})  # fmt: skip
    found = sync.sync_disaggregation_labels(_client())
    assert found["read"] == 2 and found["types_seen"] == 1  # the name came with the value: not read again
    labels = {
        (x.kind, x.prp_id): (x.label, x.disaggregation_id) for x in dm.DisaggregationLabel.objects.all()
    }
    assert labels == {
        ("value", 31): ("under 5", 30),
        ("value", 32): ("5 to 17", 30),
        ("type", 30): ("Age", None),
    }


# ------------------------------------------------------------------------- Report/Disaggregation
def test_the_disaggregation_tab_shows_a_subtotal_per_disaggregation_and_the_location_total(
    client_viewer, labelled
):
    url = reverse("reports:partner_reporting")
    page = client_viewer.get(url, {"tab": "disaggregations"})
    assert page.status_code == 200
    html = page.text
    assert "By disaggregation" in html and "female" in html and "Lebanese" in html
    assert "Sex:</span> subtotal" in html and "Nationality:</span> subtotal" in html
    assert "Location total" in html and "1 value without its label yet, not listed" in html
    data = page.context["disaggregations"]
    akkar = next(b for b in data["blocks"] if b["location"] == "Akkar" and b["report"] == "QPR1")
    assert akkar["total"] == 20 and akkar["total_is_prp"] and akkar["unlabelled"] == 1
    subtotals = {g.name: g.subtotal for g in akkar["groups"]}
    assert subtotals == {"Sex": 20, "Nationality": 20}  # never 40
    assert [ln["line"] for ln in akkar["table_lines"]] == [
        "value",
        "subtotal",
        "value",
        "value",
        "subtotal",
        "total",
    ]
    sent_back = next(b for b in data["blocks"] if b["report"] == "QPR2")
    assert not sent_back["counted"] and sent_back["report_status"] == "Sent back"
    assert "is-uncounted" in html
    # the tab's own filters: location, P-code, disaggregation
    only = client_viewer.get(url, {"tab": "disaggregations", "location": "bekaa"}).context["disaggregations"]
    assert {b["location"] for b in only["blocks"]} == {"Bekaa"}
    by_code = client_viewer.get(url, {"tab": "disaggregations", "p_code": "lb71"}).context["disaggregations"]
    assert {b["p_code"] for b in by_code["blocks"]} == {"LB71"}
    lebanese = client_viewer.get(url, {"tab": "disaggregations", "disaggregation": "lebanese"}).context[
        "disaggregations"
    ]
    (block,) = lebanese["blocks"]
    assert [(g.name, [ln.label for ln in g.lines], g.subtotal) for g in block["groups"]] == [
        ("Nationality", ["Lebanese"], 20)
    ]
    # a filter that leaves values out keeps the subtotal of every value, and says so
    female = client_viewer.get(url, {"tab": "disaggregations", "disaggregation": "female"})
    akkar = next(
        b
        for b in female.context["disaggregations"]["blocks"]
        if b["location"] == "Akkar" and b["report"] == "QPR1"
    )
    (sex,) = akkar["groups"]
    assert [ln.label for ln in sex.lines] == ["female"] and sex.subtotal == 20 and sex.hidden == 1
    assert "every value of the disaggregation: 1 value is left out by the filter" in female.text
    none = client_viewer.get(url, {"tab": "disaggregations", "indicator": "classrooms"}).context[
        "disaggregations"
    ]
    assert none["blocks"] == []
    accepted = client_viewer.get(url, {"tab": "disaggregations", "status": "accepted"}).context[
        "disaggregations"
    ]
    assert {b["report"] for b in accepted["blocks"]} == {"QPR1"}
    partial = client_viewer.get(url, {"tab": "disaggregations"}, headers=HX)
    assert "<html" not in partial.text and "Location total" in partial.text


def test_the_disaggregation_export_keeps_values_subtotals_and_totals_apart(client_viewer, labelled):
    url = reverse("reports:partner_reporting")
    lines = _csv(client_viewer.get(url, {"tab": "disaggregations", "export": "csv"}))
    akkar = [ln for ln in lines if ln["location"] == "Akkar" and ln["report"] == "QPR1"]
    assert [ln["line"] for ln in akkar] == [
        "Value",
        "Subtotal",
        "Value",
        "Value",
        "Subtotal",
        "Location total",
    ]
    assert [ln["value_in_period"] for ln in akkar if ln["line"] == "Value"] == ["20.0", "5.0", "15.0"]
    assert [ln["subtotal_of_the_disaggregation"] for ln in akkar if ln["line"] == "Subtotal"] == [
        "20.0",
        "20.0",
    ]
    total = akkar[-1]
    assert (
        total["location_total"] == "20.0"
        and total["value_in_period"] == ""
        and total["last_modified"] == "2026-10-08"
    )
    assert total["note"] == "1 value without its label yet is not listed"
    assert {ln["emergency"] for ln in lines} == {"No"} and "cp_output" in lines[0]
    female = _csv(
        client_viewer.get(url, {"tab": "disaggregations", "disaggregation": "female", "export": "csv"})
    )
    subtotal = next(
        ln for ln in female if (ln["location"], ln["report"], ln["line"]) == ("Akkar", "QPR1", "Subtotal")
    )
    assert subtotal["note"] == "The subtotal of every value of the disaggregation; 1 left out by the filter"
    book = load_workbook(
        io.BytesIO(client_viewer.get(url, {"tab": "disaggregations", "export": "xlsx"}).content)
    )
    sheet = book["Report by disaggregation"]
    header = [c.value for c in sheet[1]]
    assert header[:3] == ["programme_document", "pd_status", "partner"] and "location_total" in header
    assert sheet.max_row == len(lines) + 1


def test_reached_by_disaggregation_on_the_indicator_page(client_viewer, labelled):
    pd = labelled["pd"]
    html = client_viewer.get(reverse("reports:pd_indicator", args=[pd.pk, "500"]), {"year": "all"}).text
    assert "Reached by disaggregation" in html and "female" in html and "Nationality" in html
    detail = client_viewer.get(reverse("reports:pd_indicator", args=[pd.pk, "500"]), {"year": "all"}).context
    akkar = next(p for p in detail["by_disaggregation"] if p["location"] == "Akkar")
    sex = next(g for g in akkar["groups"] if g["name"] == "Sex")
    assert akkar["reached"] == 20 and sex["subtotal"] == 20 and akkar["unlabelled"] == 1  # Sent back left out


# -------------------------------------------------------------------------- Partner Fund Monitoring
TODAY = datetime.date(2026, 10, 9)


def test_partner_fund_monitoring_reads_the_budget_and_prp_once_per_report(funded):
    found = fund_views.partner_contributions(fund_views.FundFilters(), TODAY)
    by_pd = {r["programme_document"]: r for r in found["rows"]}
    assert set(by_pd) == {"LEBA/PCA2023583/PD20261704-3", "LEBA/PCA2022100/PD2022100"}  # no draft
    nawaya = by_pd["LEBA/PCA2023583/PD20261704-3"]
    assert (nawaya["unicef_planned"], nawaya["total"], nawaya["partner_contribution"]) == (
        Decimal("582888.90"),
        Decimal("644148.90"),
        Decimal("61260"),
    )
    assert round(nawaya["partner_share"], 1) == Decimal("9.5")
    assert (nawaya["partner_type"], nawaya["cso_type"]) == ("Civil Society Organization", "National")
    reported = nawaya["reported"]
    # the latest counted report (Submitted, awaiting acceptance), once: never QPR1 + QPR2, never the Due QPR3
    assert reported["value"] == Decimal("224108.86") and reported["report"] == "QPR2"
    assert (
        reported["status"] == "submitted"
        and reported["as_of"] == datetime.date(2026, 9, 30)
        and not reported["varies"]
    )
    assert by_pd["LEBA/PCA2022100/PD2022100"]["reported"] is None
    # amounts in two currencies are never added together
    assert {t["currency"]: t["total"] for t in found["totals"]} == {
        "USD": Decimal("644148.90"),
        "LBP": Decimal("1000"),
    }
    assert found["reported_totals"] == [{"currency": "USD", "value": Decimal("224108.86"), "pds": 1}]


def test_the_funds_filter_chooses_the_programme_documents(funded):
    def numbers(**params):
        filters = fund_views.FundFilters.from_params(params)
        return {r["programme_document"] for r in fund_views.partner_contributions(filters, TODAY)["rows"]}

    assert numbers(pd_status=["ended"]) == {"LEBA/PCA2022100/PD2022100"}
    assert numbers(partner_type=["National"]) == {"LEBA/PCA2023583/PD20261704-3"}
    assert numbers(partner=[str(funded["gov"].pk)]) == {"LEBA/PCA2022100/PD2022100"}
    assert numbers(donor=["Netherlands"]) == {"LEBA/PCA2023583/PD20261704-3"}
    assert numbers(grant=["SM220001"]) == {"LEBA/PCA2022100/PD2022100"}
    assert numbers(fr="040004") == {"LEBA/PCA2022100/PD2022100"}
    assert numbers(q="youth network") == {"LEBA/PCA2023583/PD20261704-3"}
    assert numbers(pd_status=["draft"]) == {"LEBA/DRAFT"}  # chosen on purpose


def test_funded_programmes_one_row_per_grant_and_the_pd_amounts_once(funded):
    found = fund_views.funded_programmes(fund_views.FundFilters(), TODAY)
    rows = [(r["programme_document"], r["grant"], r["reserved"], r["pd_amounts"]) for r in found["rows"]]
    assert rows == [
        ("LEBA/PCA2022100/PD2022100", "SM220001", Decimal("1000"), True),
        ("LEBA/PCA2023583/PD20261704-3", "SC260014", Decimal("50418.50"), True),
        ("LEBA/PCA2023583/PD20261704-3", "SC260017", Decimal("693754.80"), False),
    ]
    nawaya = next(g for g in found["groups"] if g["programme_document"].endswith("1704-3"))
    assert nawaya["days_to_end"] == 448  # 31 Dec 2027 from 9 Oct 2026
    assert nawaya["disbursed"] == Decimal("608936.90") and nawaya["tone"] == "on_track"
    assert [r["donor"] for r in nawaya["rows"]] == ["United States Fund for UNICEF", "Netherlands"]
    ended = next(g for g in found["groups"] if g["programme_document"].endswith("2022100"))
    assert ended["days_to_end"] == -1 and ended["tone"] == "stale"
    # the disbursement counted once per programme document, the reservations per grant
    assert found["disbursed"] == Decimal("608936.90") and found["reserved"] == Decimal("745173.30")
    assert found["pds"] == 2 and found["grants"] == 3
    netherlands = fund_views.funded_programmes(fund_views.FundFilters(donors=["Netherlands"]), TODAY)
    assert [(r["grant"], r["pd_amounts"]) for r in netherlands["rows"]] == [("SC260017", True)]


def test_the_funds_tabs_render_and_export_the_whole_filter(client_viewer, funded):
    url = reverse("reports:funds")
    for tab, expected in (
        ("reservations", "Funds reservations"),
        ("contributions", "Financial contribution during reporting periods"),
        ("programmes", "Days to end"),
    ):
        page = client_viewer.get(url, {"tab": tab})
        assert page.status_code == 200 and expected in page.text, tab
        partial = client_viewer.get(url, {"tab": tab}, headers=HX)
        assert "<html" not in partial.text and expected in partial.text
    html = client_viewer.get(url, {"tab": "contributions"}).text
    assert "9.5%" in html and "awaiting acceptance" in html and "224,109" in html
    lines = _csv(client_viewer.get(url, {"tab": "contributions", "export": "csv"}))
    nawaya = next(ln for ln in lines if ln["programme_document"].endswith("1704-3"))
    assert nawaya["partner_contribution_share_percent"] == "9.5"
    assert nawaya["financial_contribution_during_reporting_periods"] == "224108.86"
    assert nawaya["financial_contribution_report"] == "QPR2"
    rows = _csv(client_viewer.get(url, {"tab": "programmes", "export": "csv"}))
    nawaya_rows = [r for r in rows if r["programme_document"].endswith("1704-3")]
    assert [r["actual_disbursement_usd"] for r in nawaya_rows] == ["608936.90", ""]
    assert [r["total_budget"] for r in nawaya_rows] == ["644148.90", ""]
    assert nawaya_rows[1]["note"] == "Programme document amounts are on its first row"
    book = load_workbook(io.BytesIO(client_viewer.get(url, {"tab": "programmes", "export": "xlsx"}).content))
    assert book["Funded programmes"].max_row == 4
    # the FR tab keeps its filter and its figures, and takes the new ones
    assert "040001" in client_viewer.get(url, {"pd_status": "active"}).text
    assert "040001" not in client_viewer.get(url, {"pd_status": "ended"}).text
    assert "040001" not in client_viewer.get(url, {"partner_type": "Government"}).text
    assert "040001" in client_viewer.get(url, {"fr": "0400"}).text
    assert "040001" not in client_viewer.get(url, {"fr": "040002"}).text


@pytest.mark.parametrize(
    ("name", "params"),
    [
        ("reports:partner_reporting", {"tab": "disaggregations"}),
        ("reports:funds", {"tab": "contributions"}),
        ("reports:funds", {"tab": "programmes"}),
    ],
)
def test_donors_never_reach_the_r4_views(client, funded, name, params):
    from tests.donors.test_donor_access import make_donor

    client.force_login(make_donor().user)
    response = client.get(reverse(name), params)
    assert response.status_code == 302 and response["Location"] == reverse("donors:page")
    assert client.get(reverse(name), params, headers=HX).status_code == 403


@pytest.mark.parametrize(
    ("name", "params"),
    [
        ("reports:partner_reporting", {"tab": "disaggregations"}),
        ("reports:funds", {"tab": "contributions"}),
        ("reports:funds", {"tab": "programmes"}),
    ],
)
def test_the_r4_views_render_without_data(client_viewer, name, params):
    assert client_viewer.get(reverse(name), params).status_code == 200


# ------------------------------------------------------------------------------------------- privacy
def test_a_focal_point_is_never_stored_under_an_email_address():
    assert sync.person_names(
        [
            {"first_name": "Nadia", "last_name": "Madeup", "email": "r@example.org"},
            {"email": "only.an.address@example.org"},
            {"name": "Someone (someone@example.org)"},
            "plain@example.org",
            "Karim  Invented",
            {"name": "Karim Invented"},
        ]
    ) == ["Nadia Madeup", "Karim Invented"]


@responses.activate
def test_the_partner_reports_keep_who_submitted_by_name_only(labelled, client_viewer):
    """A ``submitted_by`` holding an e-mail address is not kept; and the progress report pop-up never
    shows the partner's staff (only UNICEF focal points are named, on the report table)."""
    common = {
        "partner_vendor_number": "V1",
        "intervention_reference_number": "LEB/PCA2026/PD1",
        "report_type": "QPR",
        "report_status": "Accepted",
        "reporting_period_start_date": "2026-01-01",
        "reporting_period_end_date": "2026-03-31",
        "performance_indicator": INDICATOR,
    }
    responses.get(
        f"{BASE}/api/latest/prp/datareport/",
        json={
            "results": [
                {
                    **common,
                    "id": 101,
                    "progress_report": "PR-9",
                    "report_number": "QPR9",
                    "submitted_by": "staff@example.org",
                },  # fmt: skip
                {
                    **common,
                    "id": 102,
                    "progress_report": "PR-8",
                    "report_number": "QPR8",
                    "submitted_by": "Partner Staff",
                },  # fmt: skip
            ],
            "next": None,
        },
    )
    sync.sync_all(only=["partner_reports"], triggered_by="test", client=_client())
    by_id = dict(dm.ReportedIndicator.objects.filter(datamart_id__in=[101, 102]).values_list("datamart_id", "submitted_by"))  # fmt: skip
    assert by_id == {101: "", 102: "Partner Staff"}
    # nor in the record kept beside it (which Ask NeuroDB reads)
    records = dict(dm.ReportedIndicator.objects.filter(datamart_id__in=[101, 102]).values_list("datamart_id", "data"))  # fmt: skip
    assert (records[101]["submitted_by"], records[102]["submitted_by"]) == ("", "Partner Staff")
    pop_up = client_viewer.get(reverse("reports:progress_report"), {"report": "PR-8"}, headers=HX)
    assert pop_up.status_code == 200 and "Partner Staff" not in pop_up.text


@responses.activate
def test_the_typed_tables_keep_no_contact_details(db):
    responses.get(
        f"{BASE}/api/latest/datamart/funds/grants/",
        json={
            "results": [
                {
                    "id": 1,
                    "name": "SC1",
                    "donor": "EU",
                    "focal_point_email": "x@example.org",
                    "contacts": [{"name": "A", "phone_number": "+961 1", "Mobile": "+961 2"}],
                }
            ],
            "next": None,
        },
    )
    (run,) = sync.sync_all(only=["grants"], triggered_by="test", client=_client())
    assert run.status == SyncRun.Status.SUCCEEDED
    assert dm.Grant.objects.get().data == {"id": 1, "name": "SC1", "donor": "EU", "contacts": [{"name": "A"}]}


def test_the_migration_removes_the_contact_details_already_stored(funded):
    migration = importlib.import_module("neurodb.datamart.migrations.0013_scrub_contact_details")
    assert migration.Migration.atomic is False
    grant = dm.Grant.objects.create(datamart_id=9, name="SC9", data={"name": "SC9", "Email": "x@example.org", "list": [{"phone": 1, "n": 2}]})  # fmt: skip
    clean = dm.Grant.objects.create(datamart_id=10, name="SC10", data={"name": "SC10"})
    engagement = dm.AuditEngagement.objects.create(datamart_id=1, details={"audits": {"auditor_email": "a@example.org", "x": 1}})  # fmt: skip
    PCA.objects.filter(pk=funded["nawaya"].pk).update(
        unicef_focal_points=["Nadia Madeup", "r@example.org", "Karim Invented"]
    )
    named = dm.ReportedIndicator.objects.create(datamart_id=7, submitted_by="Partner Staff")
    addressed = dm.ReportedIndicator.objects.create(datamart_id=8, submitted_by="staff@example.org")

    class Editor:
        pass

    Editor.connection = connection
    from django.apps import apps

    for _ in range(2):  # running it twice changes nothing more
        migration.scrub_records(apps, Editor())
        migration.drop_focal_point_addresses(apps, Editor())
    grant.refresh_from_db()
    engagement.refresh_from_db()
    assert grant.data == {"name": "SC9", "list": [{"n": 2}]}
    assert dm.Grant.objects.get(pk=clean.pk).data == {"name": "SC10"}
    assert engagement.details == {"audits": {"x": 1}}
    assert PCA.objects.get(pk=funded["nawaya"].pk).unicef_focal_points == ["Nadia Madeup", "Karim Invented"]
    named.refresh_from_db()
    addressed.refresh_from_db()
    assert (named.submitted_by, addressed.submitted_by) == ("Partner Staff", "")


def test_no_email_address_on_the_r4_views(client_viewer, labelled, funded):
    PCA.objects.filter(pk=funded["nawaya"].pk).update(
        unicef_focal_points=["Nadia Madeup", "focal.point@example.org"]
    )
    for name, params in (
        ("reports:partner_reporting", {"tab": "disaggregations"}),
        ("reports:funds", {"tab": "contributions"}),
        ("reports:funds", {"tab": "programmes"}),
    ):
        assert "focal.point@example.org" not in client_viewer.get(reverse(name), params).text


def test_reporting_module_reads_the_same_location_values_on_both_tabs(labelled):
    """Report/Location and Report/Disaggregation read one location's value the same way."""
    rows = dm.ReportedIndicator.objects.all()
    today = datetime.date(2026, 10, 9)
    locations = reporting.location_rows(rows, {}, today)["rows"]
    blocks = reporting.disaggregation_rows(rows, {}, today)["blocks"]
    by_location = {(r["report"], r["location"]): r["value_in_period"] for r in locations}
    by_block = {(b["report"], b["location"]): b["total"] for b in blocks}
    assert by_location == by_block == {("QPR1", "Akkar"): 20, ("QPR1", "Bekaa"): 30, ("QPR2", "Akkar"): 7}


# ---------------------------------------------------------------------------------- help and PD page
@pytest.mark.parametrize(
    ("question", "expected"),
    [
        ("How is Reached calculated from the progress reports?", "reached-one-rule-for-every-page"),
        ("Why does the disaggregation subtotal differ from the location total?", "progress-reports"),
        ("What is the financial contribution during reporting periods?", "funds"),
    ],
)
def test_the_help_guide_explains_reached_and_the_r4_views(question, expected):
    from neurodb.help import guide

    found = [(s.page, s.slug) for s in guide.search(question)]
    assert ("partner-progress", expected) in found, found


def test_the_programme_page_shows_the_partner_contribution(client_viewer, funded):
    html = client_viewer.get(reverse("reports:programme_detail", args=[funded["nawaya"].pk])).text
    assert "Partner contribution" in html and "61,260" in html
    assert "tab=contributions&amp;q=LEBA/PCA2023583/PD20261704-3" in html


def test_the_reporting_checks_say_which_disaggregation_labels_are_read(labelled):
    from neurodb.datamart import checks

    found = next(c for c in checks.run_all() if c["key"] == "3")
    assert "error" not in found, found
    figures = dict(found["figures"])
    assert figures["Disaggregation values named by the reports"] == 4  # female, Male, Lebanese, 99
    assert figures["… whose label is read"] == 3
    assert "The labels of 3 of the 4 disaggregation values named are read." in found["verdict"]
    rows = {key: label for key, label, _ in found["table"]["rows"]}
    assert rows["()"] == "Location total" and rows["(11,)"] == "female" and rows["(99,)"] is None
