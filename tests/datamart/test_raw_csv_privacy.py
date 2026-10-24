"""The privacy sweep of every download and every page Raw eTools data writes, of the two exports the standing
rule changed and of the Power BI comparison sheet: people planted in a record of every source, and in every text
column of NeuroDB's own tables, never reach a cell, a column's name, a source's list or a record's page.
Synthetic people only."""

from __future__ import annotations

import csv
import datetime
import io

import pytest
from django.contrib.auth.models import Group
from django.db import models
from django.urls import reverse

from neurodb.accounts.models import User
from neurodb.accounts.roles import VALIDATORS, VIEWER
from neurodb.datamart import catalogue, comparison, raw, raw_people
from neurodb.datamart import models as dm
from neurodb.integrations.etools.datamart_sync import ENTITY_SYNCS
from neurodb.partnerships.models import PCA, PartnerOrganization

pytestmark = pytest.mark.django_db

NAME, EMAIL = "Person One", "person.one@example.org"
PHONES = ("+961 3 123 457", "03 123 457")
UNKNOWN = "Zorba Quintel"
STRANGER = "Test Stranger"  # a name the record holds nowhere else: only the field name rule finds it
# a record naming people every way eTools does; the unknown name only where a person or a field's name is, since
# no rule can know it inside a text
PLANTED = {
    "id": 4242,
    "source_id": 4242,
    "submitted_by": NAME,
    "visit_lead": NAME,
    "team_members": [NAME, UNKNOWN],
    "assigned_to": {"name": NAME, "email": EMAIL},
    "unicef_focal_points_data": [{"first_name": "Person", "last_name": "One", "email": EMAIL}],
    "supervisor_name": UNKNOWN,
    "participants": [UNKNOWN],
    "interviewees": [{"full": UNKNOWN}],
    "pme": UNKNOWN,
    "beneficiary": UNKNOWN,
    "attendees": UNKNOWN,
    "focalPerson": UNKNOWN,
    "approver_full_name": UNKNOWN,
    UNKNOWN: "approved",
    EMAIL: "approved",
    "contact_email": EMAIL,
    "email": EMAIL,
    "phone": PHONES[0],
    "mobile": PHONES[1],
    "comments": f"Called {UNKNOWN} at {PHONES[1]}",
    "narrative": f"Visit with {NAME} ({EMAIL}), call {PHONES[0]} or {PHONES[1]}; Mrs Person One agreed.",
    "location": {"name": "Tripoli", "focal_point": UNKNOWN},
    "scores": {STRANGER: 4},
    "participants_count": 12,
    "status": "Accepted",
}
FORBIDDEN = (NAME, EMAIL, UNKNOWN, STRANGER, *PHONES, "Quintel", "123 457")
# what every text column of NeuroDB's own tables holds (cut to the column's length)
TYPED_TEXT = f"Note by {NAME} ({EMAIL}) call {PHONES[0]} or {PHONES[1]}"
CONTACT = ("email", "phone", "mobile")


@pytest.fixture
def client_validator(client, roles):
    user = User.objects.create_user(username="validator", password="pass-123456-x", is_staff=True)
    user.groups.set(Group.objects.filter(name__in=(VIEWER, VALIDATORS)))
    User.objects.create_user(username="known", first_name="Person", last_name="One", password="pass-123456-x")
    client.force_login(user)
    return client


def _plant(name: str) -> int:
    """A record of ``name`` naming people every way, in its data and in each of its text columns: its pk."""
    if raw.rest(name):  # as the eTools REST copy writes it: its title built from raw values
        return dm.RestRecord.objects.create(
            entity=name,
            record_key="4242",
            source_id="4242",
            reference=f"REF/1 {NAME}",
            title=f"Report · {NAME}",
            data=catalogue.scrub(PLANTED),
        ).pk
    if raw.typed(name):
        model = raw.model_of(name)
        fields = {"datamart_id": 4242, "data": PLANTED}
        if model is dm.HACTAggregate:
            fields["year"] = 2026
        for f in model._meta.concrete_fields:
            if isinstance(f, models.CharField | models.TextField) and not f.is_relation:
                fields[f.name] = TYPED_TEXT[: f.max_length] if f.max_length else TYPED_TEXT
        return model.objects.create(**fields).pk
    return dm.DatamartDocument.objects.create(
        dataset=name, record_key="4242", title=f"Report · {NAME}", data=PLANTED
    ).pk


def _download(client, name: str) -> tuple[list[str], list[list[str]]]:
    response = client.get(reverse("admin:datamart_raw_csv", args=[name]))
    assert response.status_code == 200, name
    rows = list(csv.reader(io.StringIO(b"".join(response.streaming_content).decode("utf-8-sig"))))
    return rows[0], rows[1:]


SOURCES = [*ENTITY_SYNCS, *raw.REST_ENTITIES]  # every Datamart dataset and every eTools REST entity


@pytest.mark.parametrize("name", SOURCES)
def test_no_person_reaches_a_raw_csv(client_validator, name):
    _plant(name)
    header, rows = _download(client_validator, name)
    assert len(rows) == 1, name
    text = repr(rows)
    for planted in FORBIDDEN:
        assert planted not in text, (name, planted)
    for column in header:
        assert "@" not in column and UNKNOWN not in column, (name, column)
        bare = column.removesuffix(" (people)").removesuffix(" (text withheld)")
        assert bare not in CONTACT and not bare.startswith("contact"), (name, column)
        if (
            bare == column
            and column not in raw.NEURODB_COLUMNS + raw.PARTNER_REPORT_COLUMNS + raw.TAIL_COLUMNS
        ):
            assert not raw_people.names_people(column), (
                name,
                column,
            )  # a person field is a count, or withheld
    if name in catalogue.SAFEGUARDING:
        assert set(header) - set(raw.NEURODB_COLUMNS) <= catalogue.SAFEGUARDING_KEYS | {
            raw_people.WITHHELD_KEYS
        }
    else:
        row = dict(zip(header, rows[0], strict=True))
        assert row["participants_count"] == "12"
        assert row["submitted_by (people)"] == "1" and row["team_members (people)"] == "2"
        assert row["people_inside"] in ("0", "1")  # the location's focal point, when the location is shown


@pytest.mark.parametrize("name", SOURCES)
def test_no_person_reaches_a_source_s_list_or_a_record_s_page(client_validator, name):
    pk = _plant(name)
    for url in (
        reverse("admin:datamart_raw_source", args=[name]),
        reverse("admin:datamart_raw_record", args=[name, pk]),
    ):
        response = client_validator.get(url)
        assert response.status_code == 200, url
        html = response.content.decode()
        for planted in FORBIDDEN:
            assert planted not in html, (url, planted)


def test_an_engagement_s_records_are_read_by_their_own_source_s_rule(client_validator):
    """An engagement's ``details`` holds its audit, spot check... records whole, as eTools sent them: on the
    engagement's page and in its CSV each is read as its own source's page reads it, so a field nobody reviewed
    (an interlocutor, a consultant, enumerators, a rating in words) reads "not reviewed", a reviewed one keeps
    its value, and a superuser's review of a field shows it on both pages alike."""
    audit = {
        **PLANTED,
        "audited_expenditure": 1200,
        "interlocutor": UNKNOWN,
        "consultant": f"{UNKNOWN} (lead)",
        "enumerators": [UNKNOWN, STRANGER],
        "visit_rating": "Good",
    }
    pk = _plant("engagements")
    dm.AuditEngagement.objects.filter(pk=pk).update(
        details={"audits": catalogue.scrub(audit), "not_a_source": {"note": UNKNOWN}}
    )
    own = dm.DatamartDocument.objects.create(dataset="audits", record_key="4243", data=audit).pk
    page = reverse("admin:datamart_raw_record", args=["engagements", pk])
    audits_page = reverse("admin:datamart_raw_record", args=["audits", own])

    def details(html: str) -> str:
        return html.split("visit_rating", 1)[1][:200] if "visit_rating" in html else ""

    for reviewed in (False, True):
        if reviewed:
            dm.RawFieldReview.objects.create(source="audits", key="visit_rating")
        html = client_validator.get(page).content.decode()
        _header, rows = _download(client_validator, "engagements")
        for planted in FORBIDDEN:
            assert planted not in html and planted not in repr(rows), (reviewed, planted)
        assert "1200" in html and "not reviewed" in html  # a field the page reads, and those nobody reviewed
        assert ("Good" in details(html)) is reviewed, reviewed
        assert ("Good" in details(client_validator.get(audits_page).content.decode())) is reviewed


def test_the_progress_reports_export_counts_the_focal_points(client_viewer):
    partner = PartnerOrganization.objects.create(etl_id="1", name="Amel Association", vendor_number="V1")
    pd = PCA.objects.create(
        etl_id="11",
        partner=partner,
        number="LEB/PD1",
        title="A",
        status="active",
        unicef_focal_points=[NAME, EMAIL, "Second Person"],
    )
    dm.ReportedIndicator.objects.create(
        datamart_id=1,
        partner=partner,
        intervention=pd,
        pd_reference_number=pd.number,
        progress_report="PR-1",
        report_number="QPR1",
        report_type="QPR",
        report_status="Accepted",
        period_end=datetime.date(2026, 3, 31),
        indicator="# reached",
    )
    response = client_viewer.get(reverse("reports:partner_reporting"), {"export": "csv"})
    text = b"".join(response.streaming_content).decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(text)))
    assert (
        rows[0]["unicef_focal_points_count"] == "2"
    )  # an entry holding an e-mail address is no person NeuroDB names
    assert "unicef_focal_points" not in rows[0] and NAME not in text and "Second Person" not in text


def test_the_action_points_csv_says_whether_one_is_assigned_never_to_whom(client_viewer):
    dm.ActionPoint.objects.create(
        datamart_id=1, reference_number="AP/1", description=f"Call {NAME} back", assigned_to_name=NAME
    )
    dm.ActionPoint.objects.create(datamart_id=2, reference_number="AP/2", description="Follow up")
    response = client_viewer.get(reverse("reports:action_points"), {"export": "csv", "all": "1"})
    text = b"".join(response.streaming_content).decode("utf-8-sig")
    rows = {r["reference"]: r for r in csv.DictReader(io.StringIO(text))}
    assert rows["AP/1"]["assigned_to_count"] == "1" and rows["AP/2"]["assigned_to_count"] == "0"
    assert "assigned_to" not in rows["AP/1"] and NAME not in text


@pytest.mark.parametrize(
    ("page", "params"),
    [("progress-reports", {"pd_status": ""}), ("partner-monitoring", {"view": "table", "year": "all"})],
)
def test_no_person_reaches_the_power_bi_comparison_sheet(client_validator, page, params):
    """Every sheet of the comparison sheet, its lists' match keys among them: no planted name, e-mail or phone,
    in a record, a programme document's focal points, an indicator's title, PD output or CP output, an unlinked
    programme document's partner or a person named as a report's partner."""
    from openpyxl import load_workbook

    partner = PartnerOrganization.objects.create(etl_id="1", name="Amel Association", vendor_number="V1")
    pd = PCA.objects.create(
        etl_id="11",
        partner=partner,
        partner_name=partner.name,
        number="LEB/PD1",
        title="A",
        status="active",
        start=datetime.date(2026, 1, 1),
        end=datetime.date(2026, 12, 31),
        unicef_focal_points=[NAME, EMAIL, f"{UNKNOWN} ({EMAIL})"],
    )
    dm.PDIndicator.objects.create(
        datamart_id=1,
        source_id=100,
        intervention=pd,
        pd_reference_number=pd.number,
        title="# reached",
        section_name="Education",
        data=PLANTED,
    )
    # an indicator whose own texts name a person, an e-mail and phones, and a programme document with no linked
    # partner, named after a person (as the Partner monitoring trail's privacy test plants them)
    dm.PDIndicator.objects.create(
        datamart_id=3,
        source_id=300,
        intervention=pd,
        pd_reference_number=pd.number,
        title=f"# of sessions (contact {EMAIL}, {PHONES[1]})",
        lower_result_name=f"Output led by {NAME} {PHONES[0]}",
        cp_output=f"CP output of {NAME}",
        section_name="Education",
        target_numerator=5,
    )
    unlinked = PCA.objects.create(
        etl_id="12",
        partner_name=f"Group of {NAME}",
        number="LEB/PD2",
        title="Water",
        status="active",
        start=datetime.date(2026, 1, 1),
        end=datetime.date(2026, 12, 31),
    )
    dm.PDIndicator.objects.create(datamart_id=4, source_id=400, intervention=unlinked, title="# of wells")
    # two report rows: the partner's, and one naming a person as its partner (no eTools partner)
    for n, (held, name) in enumerate(((partner, partner.name), (None, NAME)), start=1):
        dm.ReportedIndicator.objects.create(
            datamart_id=n,
            partner=held,
            intervention=pd,
            partner_name=name,
            pd_reference_number=pd.number,
            progress_report=f"PR-{n}",
            report_number=f"QPR{n}",
            report_type="QPR",
            report_status="Accepted",
            period_end=datetime.date(2026, 3, 31),
            indicator="# reached",
            section="Education",
            submitted_by=NAME,
            narrative=TYPED_TEXT,
            narrative_assessment=TYPED_TEXT,
            data=PLANTED,
        )
    response = client_validator.get(reverse("reports:check_compare", args=[page]), params)
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    text = repr([list(sheet.iter_rows(values_only=True)) for sheet in book.worksheets])
    assert "Amel Association" in text and "LEB/PD1" in text  # partners and programme documents stay
    for planted in FORBIDDEN:
        assert planted not in text, (page, planted)
    assert "@" not in text
    if page == "progress-reports":  # the report of the person named as a partner is still listed
        reports = list(book["Progress reports"].iter_rows(values_only=True))
        assert any(r[0] == "LEB/PD1" and r[3] == "QPR2" for r in reports)
        assert (
            len(list(book["Partners"].iter_rows(values_only=True))) == 4
        )  # a note, the header, two partners
    else:  # the indicator whose title named a contact: its texts cleaned, its row says to match it by hand
        rows = list(book["Program Performance rows"].iter_rows(values_only=True))
        (row,) = [r for r in rows if str(r[2]).startswith("# of sessions (contact")]
        assert row[4] == "LEB/PD1" and row[15] == comparison.KEY_WITHHELD
        assert any(r[4] == "LEB/PD2" and r[2] == "# of wells" for r in rows)


def test_no_person_reaches_a_figure_trail_or_its_csvs(client_validator):
    """Every CSV of the Progress reports and Home trails (the kept rows, each step's rows, each figure's rows,
    the rows that set a status, the progress reports, the lines and location blocks) and the trail pages: no
    planted name, e-mail or phone."""
    from neurodb.datamart import trail

    partner = PartnerOrganization.objects.create(etl_id="1", name="Amel Association", vendor_number="V1")
    pd = PCA.objects.create(
        etl_id="11",
        partner=partner,
        partner_name=partner.name,
        number="LEB/PD1",
        title="A",
        status="active",
        unicef_focal_points=[NAME, EMAIL, f"{UNKNOWN} ({EMAIL})"],
    )
    for n, (ref, held) in enumerate((("LEB/PD1", pd), ("LEB/GONE", None), ("", None)), start=1):
        dm.ReportedIndicator.objects.create(
            datamart_id=n,
            partner=partner,
            intervention=held,
            partner_name=partner.name,
            pd_reference_number=ref,
            progress_report=f"PR-{n}",
            report_number="QPR1",
            report_type="QPR",
            report_status="Accepted",
            period_start=datetime.date(2026, 1, 1),
            period_end=datetime.date(2026, 3, 31),
            indicator="# reached",
            section="Education",
            location="Akkar",
            submitted_by=NAME,
            narrative=TYPED_TEXT,
            narrative_assessment=TYPED_TEXT,
            disaggregation={"()": {"v": 3, "d": 1, "c": 3}},
            data=PLANTED,
        )
    # a row whose own texts name a person, an e-mail and phones: its indicator, CP output and location
    dm.ReportedIndicator.objects.create(
        datamart_id=4,
        partner=partner,
        intervention=pd,
        partner_name=partner.name,
        pd_reference_number="LEB/PD1",
        progress_report="PR-4",
        report_number="QPR4",
        report_type="QPR",
        report_status="Submitted",
        period_start=datetime.date(2026, 1, 1),
        period_end=datetime.date(2026, 3, 31),
        indicator=f"# of sessions with {NAME} ({EMAIL}, {PHONES[1]})",
        cp_output=f"CP output of {NAME}",
        section="Education",
        location=f"Akkar, near {NAME}'s house {PHONES[0]}",
        disaggregation={"()": {"v": 2, "d": 1, "c": 2}},
    )
    params = {"pd_status": "", "section": "Education"}
    found = trail.progress_trail("progress-reports", params, datetime.date.today())
    whats = ["kept", "status", "reports", "follow_up", "lines", "blocks", "step:start", "reason:no_pd"]
    whats += [s.what for s in found.steps if s.what] + [f.what for f in found.figures if f.what]
    url = reverse("reports:check_rows", args=["progress-reports"])
    for what in dict.fromkeys(whats):
        response = client_validator.get(url, {**params, "what": what})
        assert response.status_code == 200, what
        text = b"".join(response.streaming_content).decode("utf-8-sig")
        header = text.splitlines()[0]
        if what == "kept":  # the records as Raw eTools data writes them: people as counts
            row = next(csv.DictReader(io.StringIO(text)))
            assert row["participants_count"] == "12" and row["submitted_by (people)"] == "1"
        for planted in FORBIDDEN:
            assert planted not in text, (what, planted)
        for column in header.split(","):
            bare = column.removesuffix(" (people)").removesuffix(" (text withheld)")
            assert "@" not in column and bare not in CONTACT, (what, column)
    home = reverse("reports:check_rows", args=["home"])
    response = client_validator.get(home, {"block": "reporting", "what": "kept"})
    assert response.status_code == 200
    assert not any(p in b"".join(response.streaming_content).decode() for p in FORBIDDEN)
    for page, query in (("progress-reports", params), ("home", {"block": "reporting"})):
        html = client_validator.get(reverse("reports:check_figures", args=[page]), query).content.decode()
        for planted in FORBIDDEN:
            assert planted not in html, (page, planted)


def test_a_person_named_as_a_partner_or_searched_never_reaches_a_trail_csv(client_validator):
    """A report row naming a person as its partner (no eTools partner), and a search for a person: the partner's
    figures, the search's step and the kept rows name them in no cell and no file name."""
    from neurodb.datamart import trail

    pd = PCA.objects.create(etl_id="11", number="LEB/PD1", title="A", status="active")
    for n, (partner, status) in enumerate(((NAME, ""), ("Org One", "Accepted")), start=1):
        dm.ReportedIndicator.objects.create(
            datamart_id=n,
            intervention=pd,
            partner_name=partner,
            pd_reference_number="LEB/PD1",
            progress_report=f"PR-{n}",
            report_number=f"QPR{n}",
            report_type="QPR",
            report_status=status,
            period_start=datetime.date(2026, 1, 1),
            period_end=datetime.date(2026, 3, 31),
            due_date=datetime.date(2026, 4, 15),
            indicator="# reached",
            section="Education",
            location="Akkar",
            disaggregation={"()": {"v": 3, "d": 1, "c": 3}},
        )
    url = reverse("reports:check_rows", args=["progress-reports"])
    for params in ({"pd_status": ""}, {"pd_status": "", "q": NAME}):
        found = trail.progress_trail("progress-reports", params, datetime.date.today())
        whats = ["kept", "status", "reports", "follow_up", "lines", "blocks"]
        whats += [*(s.what for s in found.steps if s.what), *(f.what for f in found.figures)]
        if "q" not in params:  # the person's figures are there, by the partner's name
            assert any(w.startswith("figure:pr.partner.") for w in whats)
        for what in dict.fromkeys(w for w in whats if w):
            response = client_validator.get(url, {**params, "what": what})
            assert response.status_code == 200, what
            assert "person" not in response["Content-Disposition"].lower(), what
            text = b"".join(response.streaming_content).decode("utf-8-sig")
            assert NAME.lower() not in text.lower(), what
            if what in ("reports", "follow_up", "lines", "blocks") and "q" not in params:
                assert "LEB/PD1" in text, what  # the rows are there, cleaned
                assert ("Org One" in text) == (what != "follow_up"), what  # Org One's report is accepted


def test_no_person_reaches_a_partner_monitoring_trail_or_its_csvs(client_validator):
    """Every CSV of the Partner monitoring trail, Program Performance's, the map's and Home's Indicator progress
    block's (kept indicators, each step's programme documents, PD indicator rows and indicators, each figure's,
    the Reached and Overdue reports folds' partner report rows, the map's places) and the trail pages: no
    planted name, e-mail or phone, in a record or in an indicator's title, PD output, CP output or unlinked
    partner; the programme documents' focal points as a count."""
    from neurodb.datamart import trail_monitoring as tm

    partner = PartnerOrganization.objects.create(etl_id="1", name="Amel Association", vendor_number="V1")
    pd = PCA.objects.create(
        etl_id="11",
        partner=partner,
        partner_name=partner.name,
        number="LEB/PD1",
        title=f"Shelter with {NAME} ({EMAIL})",
        status="active",
        start=datetime.date(2026, 1, 1),
        end=datetime.date(2026, 12, 31),
        unicef_focal_points=[NAME, EMAIL, f"{UNKNOWN} ({EMAIL})"],
    )
    dm.PDIndicator.objects.create(
        datamart_id=1,
        source_id=100,
        intervention=pd,
        pd_reference_number=pd.number,
        title="# reached",
        section_name="Education",
        location_name="Akkar",
        target_numerator=10,
        data=PLANTED,
    )
    dm.PDIndicator.objects.create(datamart_id=2, source_id=200, title="# elsewhere", data=PLANTED)
    # an indicator of the programme document kept whose own texts name a person, an e-mail and phones, and a
    # programme document with no linked partner, named after a person
    dm.PDIndicator.objects.create(
        datamart_id=3,
        source_id=300,
        intervention=pd,
        pd_reference_number=pd.number,
        title=f"# of sessions (contact {EMAIL}, {PHONES[1]})",
        lower_result_name=f"Output led by {NAME} {PHONES[0]}",
        cp_output=f"CP output of {NAME}",
        section_name="Education",
        location_name="Akkar",
        target_numerator=5,
    )
    unlinked = PCA.objects.create(
        etl_id="12",
        partner_name=f"Group of {NAME}",
        number="LEB/PD2",
        title="Water",
        status="active",
        start=datetime.date(2026, 1, 1),
        end=datetime.date(2026, 12, 31),
    )
    dm.PDIndicator.objects.create(
        datamart_id=4, source_id=400, intervention=unlinked, title="# of wells", location_name="Akkar"
    )
    # two indicators of the programme document kept sharing a title that holds an e-mail and a phone (Data
    # quality's "Same title twice in one PD"; the trail's pages name a figure by the title as the page does)
    shared = f"# of visits (contact {EMAIL}, {PHONES[1]})"
    for n, output in ((5, "1.1 Visits"), (6, "1.2 Visits")):
        dm.PDIndicator.objects.create(
            datamart_id=n,
            source_id=n * 100,
            intervention=pd,
            pd_reference_number=pd.number,
            title=shared,
            lower_result_name=output,
            section_name="Education",
            location_name="Akkar",
            target_numerator=5,
        )
    for n, (status, due, title) in enumerate(
        (
            ("Accepted", datetime.date(2026, 4, 15), "# reached"),
            ("", datetime.date(2026, 1, 15), "# unknown"),
        ),
        start=1,
    ):
        dm.ReportedIndicator.objects.create(
            datamart_id=n,
            partner=partner,
            intervention=pd,
            partner_name=partner.name,
            pd_reference_number=pd.number,
            progress_report=f"PR-{n}",
            report_number=f"QPR{n}",
            report_type="QPR",
            report_status=status,
            period_start=datetime.date(2026, 1, 1),
            period_end=datetime.date(2026, 3, 31),
            due_date=due,
            submission_date=datetime.date(2026, 4, 10) if status else None,
            indicator=title,
            section="Education",
            location="Akkar",
            submitted_by=NAME,
            narrative=TYPED_TEXT,
            narrative_assessment=TYPED_TEXT,
            disaggregation={"()": {"v": 3, "d": 1, "c": 3}},
            data=PLANTED,
        )
    params = {"year": "2026", "section": ""}
    found = tm.monitoring_trail(params, datetime.date.today())
    table = tm.performance_trail({**params, "view": "table"}, datetime.date.today())
    on_map = tm.map_trail(params, datetime.date.today())
    whats: list[tuple[str, dict]] = [("kept", {}), ("no_partner", {})]
    for step in (*found.steps, *found.reached, *found.overdue):
        whats += [(step.what, {})] if step.what else []
        whats += [(line.what, {}) for line in step.lines if line.what]
    whats += [(f.what, {}) for f in found.figures if f.what]
    whats += [(s.what, {"of": "performance"}) for s in table.steps if s.what]
    whats += [(f.what, {"of": "performance"}) for f in table.block.figures if f.what]
    whats += [(f.what, {"of": "map"}) for f in on_map.block.figures if f.what]
    assert len(whats) > 20
    url = reverse("reports:check_rows", args=["partner-monitoring"])
    for what, extra in dict.fromkeys((w, tuple(e.items())) for w, e in whats):
        response = client_validator.get(url, {**params, "what": what, **dict(extra)})
        assert response.status_code == 200, what
        text = b"".join(response.streaming_content).decode("utf-8-sig")
        header = text.splitlines()[0]
        for planted in FORBIDDEN:
            assert planted not in text, (what, planted)
        for column in header.split(","):
            bare = column.removesuffix(" (people)").removesuffix(" (text withheld)")
            assert "@" not in column and bare not in CONTACT, (what, column)
        if what == "step:M1":
            row = next(csv.DictReader(io.StringIO(text)))
            assert row["unicef_focal_points (people)"] == "3" and "Shelter with" in row["title"]
        if what == "kept":  # the indicators' own texts, cleaned: what is left of them is still there
            rows = {r["indicator_key"]: r for r in csv.DictReader(io.StringIO(text))}
            assert rows[f"{pd.pk}|300"]["indicator"].startswith("# of sessions (contact")
            assert rows[f"{pd.pk}|300"]["pd_output"].startswith("Output led by")
            assert rows[f"{unlinked.pk}|400"]["indicator"] == "# of wells"
    # Data quality's titles shared in one programme document: a unit per title, never the title's words
    quality = {**params, "view": "quality"}
    figure = tm.monitoring_trail(quality, datetime.date.today()).figure("dq.list.same_title")
    assert figure.page == 1
    response = client_validator.get(url, {**quality, "what": figure.what})
    text = b"".join(response.streaming_content).decode("utf-8-sig")
    for planted in FORBIDDEN:
        assert planted not in text, ("same_title", planted)
    rows = list(csv.DictReader(io.StringIO(text)))
    assert len(rows) == 2 and len({r["unit_key"] for r in rows}) == figure.page
    assert sum(int(r["counts_once"]) for r in rows) == figure.page
    home = reverse("reports:check_rows", args=["home"])
    response = client_validator.get(home, {**params, "scope": "year", "block": "indicators", "what": "kept"})
    assert response.status_code == 200
    assert not any(p in b"".join(response.streaming_content).decode() for p in FORBIDDEN)
    pages = (
        ("partner-monitoring", params),
        ("partner-monitoring", {**params, "view": "table"}),
        ("partner-monitoring", {**params, "view": "map"}),
        ("home", {**params, "scope": "year", "block": "indicators"}),
    )
    for page, query in pages:
        html = client_validator.get(reverse("reports:check_figures", args=[page]), query).content.decode()
        for planted in FORBIDDEN:
            assert planted not in html, (page, planted)
