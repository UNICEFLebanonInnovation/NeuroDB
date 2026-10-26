"""Compare with Power BI's pages (Admin → Data and sync; datamart.powerbi_admin): who opens them, the sidebar and
the admin home's link, the upload page (what to do now, one line per Power BI page in Power BI's order, the files
not read, the hidden files, the columns dropped, the reload while busy), an export's page, how to export, and
that nothing of these tables reaches the AI or a developer table. Generated fake exports only."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from django.conf import settings
from django.contrib import admin
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse

from neurodb.datamart import comparison, powerbi_filters
from neurodb.datamart import powerbi_compare as pc
from neurodb.datamart import powerbi_exports as px
from neurodb.datamart import powerbi_words as words
from neurodb.datamart.models import PowerBIDifference, PowerBIExport, PowerBIUpload
from neurodb.integrations import background
from neurodb.web import admin_site
from tests.datamart import powerbi_factory as pf

pytestmark = pytest.mark.django_db
INDEX = reverse("admin:datamart_powerbi")
ROOT = Path(settings.BASE_DIR) / "neurodb"


@pytest.fixture
def started(monkeypatch):
    found = []
    monkeypatch.setattr(background, "start_command", lambda *args: found.append(args) or 4242)
    return found


@pytest.fixture
def checker(client, roles):
    client.force_login(pf.user("checker"))
    return client


def _done(client, *files) -> PowerBIUpload:
    response = pf.post(client, *files)
    assert response.status_code == 302, response.content[:300]
    upload = PowerBIUpload.objects.order_by("-pk").first()
    pc.run(upload.pk)
    upload.refresh_from_db()
    return upload


# ------------------------------------------------------------------------------------------------- access
def test_only_validators_with_staff_status_open_compare_with_power_bi(client, roles, admin_user, started):
    owner = pf.user("owner")
    client.force_login(owner)
    upload = _done(client, ("cards.xlsx", pf.export("partner_cards")))
    export = PowerBIExport.objects.get(upload=upload)
    other = _done(client, ("sector.xlsx", pf.export("by_sector")))
    client.logout()
    pages = [
        ("get", INDEX),
        ("post", INDEX),
        ("get", reverse("admin:datamart_powerbi_upload", args=[upload.pk])),
        ("get", reverse("admin:datamart_powerbi_export", args=[upload.pk, export.pk])),
        ("get", reverse("admin:datamart_powerbi_xlsx", args=[upload.pk])),
        ("post", reverse("admin:datamart_powerbi_again", args=[upload.pk])),
        ("get", reverse("admin:datamart_powerbi_again", args=[upload.pk])),  # whatever the method
        ("get", reverse("admin:datamart_powerbi_again", args=[upload.pk + 1000])),  # an upload that is not
    ]

    def statuses():
        return [getattr(client, method)(url).status_code for method, url in pages]

    assert statuses() == [302] * 8  # anonymous: the admin's sign-in
    client.force_login(pf.user("plain-validator", staff=False))
    assert statuses() == [302] * 8  # the admin asks for staff status
    for user in (pf.user("staff-viewer", validator=False), admin_user):  # an Administrator not ticked: no
        client.force_login(user)
        assert statuses() == [404] * 8, user
    for n, user in enumerate((pf.user("staff-validator"), pf.user("root", superuser=True))):
        client.force_login(user)
        found = statuses()
        assert found[0] == 200 and found[2] == 200 and found[3] == 200, (user, found)
        assert found[1] == 200  # a POST without files: "Choose the files to upload."
        # the Excel file of every difference; once Compare again queued the upload, none until it is done
        assert found[4] == (200 if n == 0 else 404)
        assert found[5] in (200, 302) and found[6] == found[7] == 405  # Compare again is a POST
    # an export of another upload is not found
    assert client.get(reverse("admin:datamart_powerbi_export", args=[other.pk, export.pk])).status_code == 404


def test_the_sidebar_and_the_admin_home_link_it_for_validators(client, roles, admin_user, rf):
    from neurodb.web.admin_site import DATA_AND_SYNC, sidebar_navigation

    def pages(user) -> list[str]:
        request = rf.get("/")
        request.user = user
        groups = [g for g in sidebar_navigation(request) if str(g["title"]) == str(DATA_AND_SYNC)]
        return [str(i["title"]) for g in groups for i in g["items"]][:3]

    checker = pf.user("checker")
    assert pages(checker) == ["eTools reporting checks", "Raw eTools data", "Compare with Power BI"]
    assert "Compare with Power BI" not in pages(admin_user)
    client.force_login(checker)
    html = client.get(reverse("admin:index")).content.decode()
    assert f'href="{INDEX}"' in html
    raw = client.get(reverse("admin:datamart_raw")).content.decode()
    assert f'href="{INDEX}"' in raw and words.RAW_LINK in raw
    client.force_login(admin_user)
    assert f'href="{INDEX}"' not in client.get(reverse("admin:index")).content.decode()


# --------------------------------------------------------------------------------------------- the pages
def test_the_first_page_says_how_to_export_every_visual_in_power_bi_s_order(checker):
    html = checker.get(INDEX).content.decode()
    visuals = re.findall(r"<li>(.*?)</li>", html.split("data-pbi-visuals", 1)[1].split("</ol>", 1)[0])
    assert len(visuals) == 13
    order = [comparison.POWERBI_PAGE_ORDER.index(v.split(" → ")[0]) for v in visuals]
    assert order == sorted(order)
    assert words.UPLOAD_BUTTON in html and 'enctype="multipart/form-data"' in html


def test_the_upload_page_lists_every_power_bi_page_in_order(checker, started):
    upload = _done(
        checker,
        (
            "exports.zip",
            pf.archive(
                ("cards.xlsx", pf.export("partner_cards")),
                ("list.xlsx", pf.export("report_list", people=True)),
            ),
        ),
    )
    html = checker.get(reverse("admin:datamart_powerbi_upload", args=[upload.pk])).content.decode()
    table = html.split("data-pbi-pages", 1)[1].split("</table>", 1)[0]
    pages = re.findall(r'<th scope="row" rowspan="\d+">(.*?)</th>', table)
    assert pages == [p.replace("&", "&amp;") for p in comparison.POWERBI_PAGE_ORDER]
    assert table.count("<td>—</td><td>—</td><td>—</td>") == 11  # 13 visuals, 2 uploaded
    verdicts = re.findall(r'data-pbi-verdict="(\w+)"', table)
    assert len(verdicts) == 10 and verdicts.count("not_uploaded") == 8  # 2 of the 10 pages have an export
    assert "cards.xlsx" in table and "list.xlsx" in table
    todo = _todo(html)
    assert 1 <= len(todo) <= 4
    assert todo[0].startswith("Needs a look: ")  # the fake rows against an empty NeuroDB
    assert "1 file was not read: see why below, then export it again." in todo
    assert "11 visuals are not uploaded: export them too (see How to export)." in todo
    assert words.NOT_EXCEL.format(name="notes.txt") in html
    assert "Skipped 3 hidden files a Mac adds to a zip." in html
    assert (
        "Dropped as the file was read: UNICEF Focal Points (people&#x27;s names). It was never stored."
        in html
    )
    assert 'http-equiv="refresh"' not in html and words.AGAIN_BUTTON in html


def test_the_upload_page_reloads_itself_while_the_job_runs(checker, started):
    pf.post(checker, ("cards.xlsx", pf.export("partner_cards")))
    upload = PowerBIUpload.objects.get()
    html = checker.get(reverse("admin:datamart_powerbi_upload", args=[upload.pk])).content.decode()
    assert '<meta http-equiv="refresh" content="5">' in html and words.STARTED in html
    assert words.AGAIN_BUTTON not in html


def _todo(html: str) -> list[str]:
    if "data-pbi-todo" not in html:
        return []
    return re.findall(r"<li>(.*?)</li>", html.split("data-pbi-todo", 1)[1].split("</ul>", 1)[0])


def test_a_filter_neurodb_cannot_apply_is_said_on_the_upload_and_export_pages(checker, started):
    """What to do now names the slicer as the report shows it (Section, PD status); Power BI's field names stay on
    the export's page, in its filter table."""
    filters = (
        "status is not active",
        "section_name is Fake Section Nobody Has",
        "actual_value is between 1 and 2",
        "name is Zq Unknownperson",
    )
    upload = _done(checker, ("sector.xlsx", pf.export("by_sector", filters=filters)))
    html = checker.get(reverse("admin:datamart_powerbi_upload", args=[upload.pk])).content.decode()
    todo = _todo(html)
    assert todo[0].startswith("Needs a look: ") and todo[1:] == [
        words.plural(12, words.TODO_MISSING_ONE, words.TODO_MISSING),
        words.TODO_CANNOT.format(
            clauses="the PD status slicer; the Section slicer; the actual_value filter in the Filters pane; "
            "the name filter in the Filters pane"
        ).replace("'", "&#x27;"),
    ]
    assert "section_name" not in html and "status is not" not in html
    export = PowerBIExport.objects.get(upload=upload)
    page = checker.get(reverse("admin:datamart_powerbi_export", args=[upload.pk, export.pk])).content.decode()
    assert words.CLAUSE_STATES["cannot"] in page and words.CLAUSE_STATES["unknown_field"] in page
    table = page.split("data-pbi-clauses", 1)[1].split("</table>", 1)[0]
    assert "status is not …" in table and "section_name is …" in table and "name is …" in table
    assert "Columns kept: Section, Accepted, Submitted, Due, Overdue" in page
    assert "Unknownperson" not in page and "Unknownperson" not in html
    assert "Fake Section Nobody Has" not in page and "Fake Section Nobody Has" not in html


def test_every_slicer_what_to_do_now_names_is_on_the_report():
    """Power BI's fields with a slicer, by the slicer's name in comparison.py (the one place for Power BI's
    names)."""
    slicers = {*comparison.POWERBI_ONLY_SLICERS, *comparison.POWERBI_SLICERS["partner-monitoring"].values()}
    assert set(comparison.POWERBI_FIELD_SLICERS.values()) <= slicers
    assert set(comparison.POWERBI_FIELD_SLICERS) <= set(powerbi_filters.FIELDS)


@pytest.mark.parametrize("how", ["stale", "start_failed", "locked", "other_job_later", "stopped"])
def test_a_failed_upload_says_why_first_and_never_that_there_is_nothing_to_do(checker, monkeypatch, how):
    """The error's own instruction comes first; an export a failed upload never read says "Not read" with what to
    do, never "Waiting to be read"; every Power BI page of its files says the comparison stopped before it, never
    "Not uploaded"."""
    if how == "start_failed":

        def broken(*args):
            raise OSError("no process")

        monkeypatch.setattr(background, "start_command", broken)
    else:
        monkeypatch.setattr(background, "start_command", lambda *args: 4242)
    pf.post(checker, *pf.every_export())
    upload = PowerBIUpload.objects.get()
    if how == "stale":
        PowerBIUpload.objects.filter(pk=upload.pk).update(status=PowerBIUpload.Status.READING)
        assert pc.close_stale() == 1
    elif how == "locked":
        monkeypatch.setattr(pc, "_lock", lambda: False)
        pc.run(upload.pk)
    elif how == "other_job_later":
        monkeypatch.setattr(pc, "heavy_job_running", lambda: True)
        pc.run(upload.pk)
    elif how == "stopped":

        def too_much(upload):
            raise MemoryError

        monkeypatch.setattr(pc, "_read_staged", too_much)
        pc.run(upload.pk)
    upload.refresh_from_db()
    assert upload.status == PowerBIUpload.Status.FAILED
    html = checker.get(reverse("admin:datamart_powerbi_upload", args=[upload.pk])).content.decode()
    todo = _todo(html)
    assert todo and todo[0] == upload.error.replace("'", "&#x27;") and words.TODO_NOTHING not in todo
    assert words.STATE_WORDS["staged"] not in html
    table = html.split("data-pbi-pages", 1)[1].split("</table>", 1)[0]
    assert words.NOT_UPLOADED not in table and table.count('data-pbi-verdict="stopped"') == 10  # every page
    if how == "other_job_later":  # read already, never compared: Compare again (each export and each page)
        assert upload.error == words.OTHER_JOB_LATER and words.AGAIN_BUTTON in html
        assert table.count(words.STOPPED_BEFORE) == 13 + 10
        assert table.count(words.STOPPED_BEFORE_ACTION) == 13 + 10
    else:  # the copies were emptied: upload the files again
        assert "upload the files again" in upload.error and words.AGAIN_BUTTON not in html
        assert table.count(words.NOT_READ_STATE) == 13 and table.count(words.NOT_READ_ACTION) == 13 + 10
        assert words.COVERS_EVERY_YEAR not in table  # nothing of their filters was read


def test_the_history_counts_the_exports_read_and_the_files_not_read(checker, started):
    pf.post(checker, ("cards.xlsx", pf.export("partner_cards")), ("sector.xlsx", pf.export("by_sector")))
    upload = PowerBIUpload.objects.get()
    PowerBIExport.objects.filter(upload=upload, kind="by_sector").update(content=b"PK broken")
    pc.run(upload.pk)
    upload.refresh_from_db()
    assert upload.status == PowerBIUpload.Status.DONE
    html = checker.get(INDEX).content.decode()
    table = html.split("data-pbi-history", 1)[1].split("</table>", 1)[0]
    headers = re.findall(r'<th scope="col">(.*?)</th>', table)
    cells = re.findall(r"<td>(.*?)</td>", table)
    found = dict(zip(headers[1:], cells, strict=False))
    assert found["Exports read"] == "1" and found["Files not read"] == "1"


def test_an_export_whose_comparison_stopped_is_never_called_a_file_not_read(checker, started, monkeypatch):
    """Its file was read: What to do now says to press Compare again, the history counts it among the exports
    read, and only a file NeuroDB could not read is listed (with why) under the files not read."""
    from neurodb.datamart import powerbi_kinds

    def broken(job):
        raise TypeError("fake")

    monkeypatch.setitem(powerbi_kinds.COMPARERS, "by_sector", broken)
    pf.post(checker, ("cards.xlsx", pf.export("partner_cards")), ("sector.xlsx", pf.export("by_sector")))
    upload = PowerBIUpload.objects.get()
    pc.run(upload.pk)
    html = checker.get(reverse("admin:datamart_powerbi_upload", args=[upload.pk])).content.decode()
    todo = _todo(html)
    assert words.TODO_NOT_COMPARED_ONE in todo and not any("not read" in line for line in todo)
    assert "data-pbi-not-read" not in html
    assert words.VERDICT_ACTIONS["failed"] in html and "send the Excel file" not in html
    table = checker.get(INDEX).content.decode().split("data-pbi-history", 1)[1].split("</table>", 1)[0]
    headers = re.findall(r'<th scope="col">(.*?)</th>', table)
    found = dict(zip(headers[1:], re.findall(r"<td>(.*?)</td>", table), strict=False))
    assert found["Exports read"] == "2" and found["Files not read"] == "0"
    # a file the job could not read: listed, with why
    PowerBIExport.objects.filter(upload=upload, kind="by_sector").update(
        error=words.COULD_NOT_READ.format(name="sector.xlsx", cls="ValueError")
    )
    html = checker.get(reverse("admin:datamart_powerbi_upload", args=[upload.pk])).content.decode()
    assert words.TODO_NOT_READ_ONE.replace("'", "&#x27;") in _todo(html)
    listed = html.split("data-pbi-not-read", 1)[1].split("</ul>", 1)[0]
    assert "NeuroDB could not read sector.xlsx (ValueError)." in listed


def test_an_upload_with_nothing_recognised_keeps_nothing(checker, started):
    response = pf.post(checker, ("notes.txt", b"fake notes"))
    assert response.status_code == 200
    html = response.content.decode()
    assert words.NOT_EXCEL.format(name="notes.txt") in html and words.NOTHING_KEPT in html
    assert PowerBIUpload.objects.count() == 0 and started == []


def test_the_pages_never_read_the_stored_rows_or_copies(checker, started):
    upload = _done(
        checker,
        ("cards.xlsx", pf.export("partner_cards")),
        ("list.xlsx", pf.export("report_list", people=True)),
    )
    export = PowerBIExport.objects.filter(upload=upload).first()
    for url in (
        INDEX,
        reverse("admin:datamart_powerbi_upload", args=[upload.pk]),
        reverse("admin:datamart_powerbi_export", args=[upload.pk, export.pk]),
    ):
        with CaptureQueriesContext(connection) as queries:
            assert checker.get(url).status_code == 200
        selects = [  # the columns read (a filter on whether a copy is still there reads no bytes)
            q["sql"].split(" FROM ", 1)[0]
            for q in queries.captured_queries
            if "datamart_powerbiexport" in q["sql"] and q["sql"].startswith("SELECT")
        ]
        assert all('"content"' not in sql and '"rows"' not in sql for sql in selects), url
    assert {"content", "rows"} <= pc.exports_of(upload).query.deferred_loading[0]


# ---------------------------------------------------------------------------------------- kept apart from the AI
def test_nothing_of_these_tables_reaches_the_ai_or_a_developer_table():
    names = ("PowerBIUpload", "PowerBIExport", "PowerBIDifference", "powerbi_compare", "powerbi_read")
    for folder in ("assistant", "watch", "fmm", "knowledge", "graph", "insights"):
        for path in (ROOT / folder).rglob("*.py"):
            text = path.read_text(encoding="utf-8")
            assert not any(name in text for name in names), path
    catalogue = (ROOT / "datamart" / "catalogue.py").read_text(encoding="utf-8")
    assert not any(name in catalogue for name in names)
    for model in (PowerBIUpload, PowerBIExport, PowerBIDifference):
        assert model not in admin.site._registry
        assert f"datamart.{model.__name__}" not in admin_site.DEVELOPER_ONLY_KEEP


def test_every_kind_has_a_power_bi_visual_and_every_page_is_in_power_bi_s_order():
    for kind in px.KINDS.values():
        page, visual = kind.visual
        assert page in comparison.POWERBI_PAGE_ORDER and visual
    assert {page for page, _visual in comparison.POWERBI_VISUALS.values()} <= set(
        comparison.POWERBI_PAGE_ORDER
    )
    assert [k.order for k in px.KINDS.values()] == list(range(13))
    pages = [comparison.POWERBI_PAGE_ORDER.index(k.visual[0]) for k in px.KINDS.values()]
    assert pages == sorted(pages)


def test_the_help_says_what_works_now_and_keeps_the_sheet_until_neurodb_compares():
    """The help's section opens with what Compare with Power BI does today; it says the manual check is no longer
    needed only once NeuroDB compares the exports with its pages (powerbi_compare.COMPARERS)."""
    text = (ROOT / "help" / "guide" / "checking-figures.md").read_text(encoding="utf-8")
    section = text.split("## Compare with Power BI\n", 1)[1].split("\n## ", 1)[0]
    opening = section.strip().split("\n\n", 1)[0]
    if not pc.COMPARERS:
        assert "no longer need" not in section
        assert "comparison sheet" in opening and "people's names" in opening.lower()


# ----------------------------------------------------------------------------------- the results (stage 2)
@pytest.fixture
def fake_world(db, monkeypatch):
    import datetime

    from neurodb.datamart import reached

    monkeypatch.setattr(
        reached,
        "history_from",
        lambda today=None: reached.HistoryFrom(datetime.date(2026, 1, 1), "read_from"),
    )
    return pf.world()


REPORT_DROP = ("pd_status", "partner", "section", "document_type", "due_date", "period_start")


def _report(**values):
    import datetime

    base = {
        "pd": pf.WORLD_PD,
        "report_status": "Accepted",
        "period_end": datetime.date(2026, 3, 31),
        "report": "QPR1",
    }
    return pf.row("report_list", **{**base, **values})


def test_an_exports_page_lists_its_lines_needs_a_look_first_with_two_values_and_what_to_do(
    checker, started, fake_world
):
    import datetime

    rows = [
        _report(),
        _report(report="HR1", report_status="Overdue", period_end=datetime.date(2026, 8, 31)),  # owed
        _report(report="HR2", report_status="Accepted", period_end=datetime.date(2026, 8, 31)),  # not in the read
        *[_report(report=f"HR{n}", report_status="Overdue", period_end=datetime.date(2026, 8, 31)) for n in range(3, 120)],
        _report(pd="=SUM(1,2)", report="+1-2"),  # a formula-like text stays a text
    ]  # fmt: skip
    pf.sync(
        datetime.datetime(2026, 10, 15, 8, 0)
    )  # after Power BI's refresh: a row it did not bring needs a look
    upload = pf.compared(
        checker,
        ("list.xlsx", pf.export("report_list", rows, drop=REPORT_DROP)),
        refreshed=datetime.datetime(2026, 10, 14, 6, 0),
    )
    export = PowerBIExport.objects.get(upload=upload)
    url = reverse("admin:datamart_powerbi_export", args=[upload.pk, export.pk])
    html = checker.get(url).content.decode()
    reasons = html.split("data-pbi-reasons", 1)[1].split("</ul>", 1)[0]
    assert "Needs a look: Not in NeuroDB: not in the Datamart read (2)" in reasons
    expected = html.split("data-pbi-expected", 1)[1].split("</p>", 1)[0]
    assert "A report still owed (118)" in expected and "?reason=owed&amp;group=1" in expected
    table = html.split("data-pbi-lines", 1)[1].split("</table>", 1)[0]
    first = table.split("<tr data-pbi-line>", 2)[1]
    assert (
        "Not in NeuroDB: not in the Datamart read" in first and "The Datamart did not send this row." in first
    )
    assert "Download the Excel file and send it to the NeuroDB team." in first and "Check in PRP" in first
    assert '<th scope="col">Power BI</th>' in table and '<th scope="col">NeuroDB</th>' in table
    assert "data-pbi-pages-nav" in html  # 120 lines: 100 a page
    second = checker.get(url, {"page": 2}).content.decode()
    assert second.count("Only in Power BI") >= 20
    owed = checker.get(url, {"reason": "owed"}).content.decode()
    assert "Not in NeuroDB: not in the Datamart read" not in owed.split("data-pbi-lines", 1)[1]
    only = checker.get(url, {"kind": "only_neurodb"}).content.decode().split("data-pbi-lines", 1)[1]
    assert "Only in NeuroDB" in only and "Only in Power BI" not in only  # QPR2 and QPR3: not in the export
    # the Excel file: every difference, texts as texts, dates as dates, numbers as numbers
    response = checker.get(reverse("admin:datamart_powerbi_xlsx", args=[upload.pk]))
    assert response.status_code == 200 and response["Content-Disposition"].startswith(
        'attachment; filename="powerbi-comparison_'
    )
    import io

    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(response.content))
    assert book.sheetnames[0] == "Summary" and len(book.sheetnames) == 2
    sheet = book.worksheets[1]
    header = [c.value for c in sheet[1]]
    assert header[:5] == ["Kind", "Group", "Reason", "What to do", "Detail"]
    rows = [dict(zip(header, (c.value for c in r), strict=False)) for r in sheet.iter_rows(min_row=2)]
    formula = next(r for r in rows if str(r["PD #"]).startswith("=SUM"))
    assert formula["PD #"] == "=SUM(1,2)" and formula["Report #"] == "+1-2"
    assert isinstance(formula["Period End"], datetime.datetime)
    assert "<f>" not in zipfile_text(response.content)
    assert len(rows) == 122  # every line (118 owed, 2 not in the read, 2 only in NeuroDB), one row each: wide


def zipfile_text(data: bytes) -> str:
    import io
    import zipfile

    with zipfile.ZipFile(io.BytesIO(data)) as package:
        return "".join(package.read(n).decode("utf-8", "replace") for n in package.namelist())


def test_the_excel_writes_a_different_line_one_row_per_column_and_stops_at_its_caps(
    checker, started, fake_world, monkeypatch
):
    import datetime
    import io

    import openpyxl

    from neurodb.datamart import powerbi_results, raw

    rows = [_report(report_status="Submitted", partner="SOMEONE ELSE", report="QPR1")]
    rows += [
        _report(report=f"HR{n}", report_status="Overdue", period_end=datetime.date(2026, 8, 31))
        for n in range(1, 30)
    ]
    drop = tuple(n for n in REPORT_DROP if n != "partner")
    upload = pf.compared(checker, ("list.xlsx", pf.export("report_list", rows, drop=drop)))
    monkeypatch.setattr(powerbi_results, "XLSX_SHEET_ROWS", 10)
    freed = []
    real = raw.take_slot

    def slot():
        found = real()
        free = found.free
        found.free = lambda: freed.append(1) or free()
        return found

    monkeypatch.setattr(raw, "take_slot", slot)
    response = checker.get(reverse("admin:datamart_powerbi_xlsx", args=[upload.pk]))
    assert freed == [1]  # the slot is freed once the file is written
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).worksheets[1]
    header = [c.value for c in sheet[1]]
    rows = [dict(zip(header, (c.value for c in r), strict=False)) for r in sheet.iter_rows(min_row=2)]
    different = [r for r in rows if r["Kind"] == "Different"]
    assert {r["Column"] for r in different} == {"Report Status", "Partner"}  # long form: one row per column
    assert len(rows) == 10 and rows[-1]["Kind"].startswith("Differences counted but not listed here: ")
    monkeypatch.setattr(raw, "take_slot", lambda: None)
    busy = checker.get(reverse("admin:datamart_powerbi_xlsx", args=[upload.pk]))
    assert busy.status_code == 503 and busy.content.decode() == words.XLSX_BUSY


def test_the_upload_page_says_what_to_do_now_and_one_verdict_per_power_bi_page(checker, started, fake_world):
    import datetime

    rows = [
        _report(),
        _report(report="QPR2", report_status="Submitted", period_end=datetime.date(2026, 6, 30)),
        _report(report="QPR3", report_status="Overdue", period_end=datetime.date(2026, 9, 30)),
        _report(pd=pf.WORLD_ENDED),
    ]
    upload = pf.compared(
        checker,
        ("list.xlsx", pf.export("report_list", rows, drop=REPORT_DROP)),
        (
            "cards.xlsx",
            pf.export("partner_cards", [pf.row("partner_cards", accepted=1, submitted=1, due=0, overdue=1)]),
        ),
    )
    html = checker.get(reverse("admin:datamart_powerbi_upload", args=[upload.pk])).content.decode()
    todo = _todo(html)
    assert (
        todo[0].startswith("Check 1 PD in eTools: NeuroDB has it as not running.") and "?group=3" in todo[0]
    )
    assert 1 <= len(todo) <= 4
    verdicts = dict(
        re.findall(
            r'<th scope="row" rowspan="\d+">(.*?)</th>\s*<td class="nd-pbi__long" rowspan="\d+" '
            r'data-pbi-verdict="(\w+)"',
            html,
        )
    )
    assert verdicts["PD/SPD Report"] == "check" and verdicts["Partner Performance"] == "same"
    assert verdicts["Program Performance"] == "not_uploaded"
    assert words.XLSX_BUTTON in html and words.AGAIN_BUTTON in html
    history = checker.get(INDEX).content.decode().split("data-pbi-history", 1)[1]
    assert "<td>0</td>" in history  # no line needs a look
    del datetime


def test_what_to_do_now_never_says_agree_while_an_export_is_empty_or_of_another_country(
    checker, started, fake_world, settings
):
    """Every other export the same, one exported empty while NeuroDB's page has rows for its filter (or of
    another country): What to do now says so, on the page and in the Excel file, never that they agree."""
    import io

    import openpyxl

    from neurodb.datamart import powerbi_results

    settings.ETOOLS_DATAMART_COUNTRY = "Fakeland"
    here = ("country_name is Fakeland", "status is active")
    cards = pf.export(
        "partner_cards", [pf.row("partner_cards", accepted=1, submitted=1, due=0, overdue=1)], filters=here
    )
    elsewhere = ("country_name is Elsewhere", "status is active")
    for other, line in (
        (pf.export("report_list", [], filters=here), words.TODO_EMPTY_ONE),
        (
            pf.export("report_list", filters=elsewhere),
            words.TODO_OTHER_COUNTRY_ONE.format(country="Fakeland"),
        ),
    ):
        upload = pf.compared(checker, ("cards.xlsx", cards), ("list.xlsx", other))
        exports = list(pc.exports_of(upload))
        lines = [found["text"] for found in powerbi_results.todo(upload, exports, [])]
        assert line in lines and words.TODO_NOTHING not in lines
        page = checker.get(reverse("admin:datamart_powerbi_upload", args=[upload.pk])).content.decode()
        assert any(t.startswith(line.replace("'", "&#x27;")) for t in _todo(page))
        workbook = checker.get(reverse("admin:datamart_powerbi_xlsx", args=[upload.pk]))
        sheet = openpyxl.load_workbook(io.BytesIO(workbook.content)).worksheets[0]
        cells = [r[0] for r in sheet.iter_rows(values_only=True)]
        assert line in cells and words.TODO_NOTHING not in cells
    assert "Export it again with Country name Fakeland." in page


def test_the_excel_s_what_to_do_now_says_what_the_page_says(checker, started):
    """The files not read (with why), the visuals not uploaded and the filters to clear: in the Excel file's
    Summary as on the upload page."""
    import html as html_module
    import io

    import openpyxl

    filters = ("country_name is Lebanon", "name is Zq Unknownperson")
    upload = _done(checker, ("sector.xlsx", pf.export("by_sector", filters=filters)), ("notes.txt", b"fake"))
    page = checker.get(reverse("admin:datamart_powerbi_upload", args=[upload.pk])).content.decode()
    todo = [html_module.unescape(re.sub(r" <a .*", "", t)) for t in _todo(page)]
    response = checker.get(reverse("admin:datamart_powerbi_xlsx", args=[upload.pk]))
    cells = [r[0] for r in openpyxl.load_workbook(io.BytesIO(response.content)).worksheets[0].iter_rows(values_only=True)]  # fmt: skip
    start = cells.index("What to do now") + 1
    assert cells[start : start + len(todo)] == todo
    assert any("visuals are not uploaded" in t for t in todo) and any("Filters pane" in t for t in todo)
    assert any("file was not read" in t for t in todo)
    assert "  " + words.NOT_EXCEL.format(name="notes.txt") in cells  # why, as "see why below" says


def test_a_trimmed_upload_links_no_line_and_offers_no_excel(checker, started):
    """Its lines are removed: no link to them, no Excel file (404), its export page says the details were
    removed (not "No differences"), and no sentence asks to send the Excel file."""
    uploads = []
    for n in range(4):
        uploads.append(_done(checker, (f"{n}.xlsx", pf.export("by_sector"))))  # needs a look: empty NeuroDB
    first = uploads[0]
    first.refresh_from_db()
    assert first.trimmed and (first.summary or {}).get("look")
    html = checker.get(reverse("admin:datamart_powerbi_upload", args=[first.pk])).content.decode()
    assert "?group=" not in html and "data-pbi-xlsx" not in html
    assert words.TRIMMED_ACTION in html and "Download the Excel file" not in html
    assert checker.get(reverse("admin:datamart_powerbi_xlsx", args=[first.pk])).status_code == 404
    export = PowerBIExport.objects.get(upload=first)
    page = checker.get(reverse("admin:datamart_powerbi_export", args=[first.pk, export.pk])).content.decode()
    assert words.DETAILS_REMOVED in page and words.NO_LINES not in page and "data-pbi-line-filter" not in page


def test_how_to_export_and_the_country_follow_neurodb_s_history_and_settings(checker, settings, monkeypatch):
    import datetime

    from neurodb.datamart import powerbi_results, reached, services

    monkeypatch.setattr(
        reached,
        "history_from",
        lambda today=None: reached.HistoryFrom(datetime.date(2024, 1, 1), "read_from"),
    )
    settings.ETOOLS_DATAMART_COUNTRY = "Fakeland"
    html = checker.get(INDEX).content.decode()
    year = words.HOW_YEAR.format(year=services.reporting_today().year, day="1 Jan 2024")
    assert year.replace("'", "&#x27;") in html
    assert "Lebanon" not in html
    export = PowerBIExport(state="not_compared", verdict="other_country")
    assert powerbi_results.verdict_words(export)[1] == "Export it again with Country name Fakeland."
    export = PowerBIExport(state="not_compared", verdict="")
    assert powerbi_results.verdict_words(export)[1] == "Export it again with Country name Fakeland."


def test_a_counted_visual_s_differences_are_called_differences_not_rows(checker, started):
    """Two sector rows against an empty NeuroDB: two rows compared, eight differences (four columns each)."""
    rows = [
        pf.row("by_sector", section="Section Alpha", accepted=1, submitted=1, due=1, overdue=1),
        pf.row("by_sector", section="Section Beta", accepted=1, submitted=1, due=1, overdue=1),
    ]
    upload = _done(checker, ("sector.xlsx", pf.export("by_sector", rows)))
    export = PowerBIExport.objects.get(upload=upload)
    assert export.counts["rows"] == 2 and export.counts["groups"] == {"4": 8}
    todo = _todo(checker.get(reverse("admin:datamart_powerbi_upload", args=[upload.pk])).content.decode())
    assert todo[0].startswith("Needs a look: 8 differences.") and "rows" not in todo[0]


def test_an_export_s_differences_stay_readable_at_phone_width():
    """Why and What to do keep a column wide enough to read on a phone, the Column heading takes no more than
    its words, a file name without spaces wraps, and the line filter's choices stay in their card."""
    css = (ROOT / "web/static/css/admin.css").read_text(encoding="utf-8")
    assert '.nd-pbi__lines th[scope="row"] { min-width: 0; }' in css
    assert ".nd-pbi__lines td.nd-pbi__long { min-width: 14rem; max-width: 28rem; }" in css
    assert ".nd-pbi__note { min-width: 0; overflow-wrap: anywhere; }" in css
    assert (
        "[data-pbi-line-filter] label, [data-pbi-line-filter] select { max-width: 100%; min-width: 0; }"
        in css
    )
    page = (ROOT / "datamart/templates/admin/datamart/powerbi/export.html").read_text(encoding="utf-8")
    assert "nd-pbi__lines" in page and page.count('class="nd-pbi__long"') == 4 and "nd-pbi__note" in page


def test_the_upload_page_reads_at_phone_width_and_its_lists_show_their_markers():
    """The table of Power BI's pages keeps its verdict and result columns wide enough to read on a phone (it
    scrolls sideways in its card, never the page), and the lists show their bullets and numbers (unfold's
    stylesheet has no list-disc or list-decimal)."""
    css = (ROOT / "web/static/css/admin.css").read_text(encoding="utf-8")
    assert '.nd-pbi__pages th[scope="row"] { min-width: 7rem; }' in css
    assert ".nd-pbi__pages td.nd-pbi__long { min-width: 12rem; max-width: 24rem; }" in css
    assert ".nd-pbi .list-disc { list-style-type: disc; }" in css
    assert ".nd-pbi .list-decimal { list-style-type: decimal; }" in css
    folder = ROOT / "datamart/templates/admin/datamart/powerbi"
    page = (folder / "upload.html").read_text(encoding="utf-8")
    assert "nd-pbi__pages" in page and page.count('class="nd-pbi__long"') == 2
    for name in ("index.html", "upload.html", "export.html"):
        assert 'class="nd-found nd-raw nd-pbi ' in (folder / name).read_text(encoding="utf-8"), name


def test_choosing_the_files_throws_no_script_error(checker):
    """unfold's admin script, on any file input's change, writes the file's name into the text input three
    levels up (its own file widget's): the form keeps the browser's own file input (which says "13 files") and
    gives the script a hidden text input there, so choosing the files throws no error in the console."""
    from html.parser import HTMLParser

    class Tree(HTMLParser):
        def __init__(self):
            super().__init__()
            self.stack, self.found = [], None

        def handle_starttag(self, tag, attrs):
            attrs = dict(attrs)
            if tag == "input" and attrs.get("type") == "file":
                self.found = list(self.stack)
            if tag not in ("input", "br", "img", "meta", "link"):
                self.stack.append((tag, attrs))

        def handle_endtag(self, tag):
            if self.stack and tag != "input":
                self.stack.pop()

    html = checker.get(INDEX).content.decode()
    tree = Tree()
    tree.feed(html)
    third = tree.found[-3]
    assert "data-pbi-file" in third[1], tree.found[-3:]
    block = html.split("data-pbi-file", 1)[1].split("</form>", 1)[0]
    assert (
        '<input type="text" hidden' in block and "disabled" in block.split('<input type="text" hidden', 1)[1]
    )


def test_the_upload_page_says_how_far_the_comparison_is(checker, started):
    pf.post(checker, ("cards.xlsx", pf.export("partner_cards")), ("sector.xlsx", pf.export("by_sector")))
    upload = PowerBIUpload.objects.get()
    PowerBIUpload.objects.filter(pk=upload.pk).update(
        status=PowerBIUpload.Status.COMPARING, summary={"compared": 1}
    )
    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(background, "lock_is_held", lambda lock: True)  # the job holds its lock: not stale
        html = checker.get(reverse("admin:datamart_powerbi_upload", args=[upload.pk])).content.decode()
    assert "Comparing… 1 of 2 exports" in html and '<meta http-equiv="refresh" content="5">' in html
    assert words.XLSX_BUTTON not in html
