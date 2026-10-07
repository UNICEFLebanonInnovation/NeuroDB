"""The document review's data and pipeline (``neurodb.knowledge.review``): findings read with a scripted
model, located without AI, key statements and action points, what a new analysis keeps, the budget and
the switches, the job's plumbing, the settings and topics in the admin, and the Ask NeuroDB tool."""

from __future__ import annotations

import datetime
import io
import json
import re
import zipfile
from types import SimpleNamespace

import pytest
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone

from neurodb.assistant import agent, tools, usage
from neurodb.assistant.models import AIUsage
from neurodb.core import admin_jobs, jobs
from neurodb.core.models import ScheduledJob, SyncRun
from neurodb.geo.models import DistrictLocation, GovernorateLocation
from neurodb.integrations import background
from neurodb.knowledge import review, review_data, review_locate, review_prompts
from neurodb.knowledge.models import (
    Document,
    DocumentActionPoint,
    DocumentFinding,
    DocumentReviewSettings,
    DocumentStatement,
    ReviewBatch,
    Topic,
    TopicProgramme,
    Verdict,
)

AI_ON = {"AI_ASSISTANT_ENABLED": True, "OPENAI_API_KEY": "x"}
PAGES = [
    "Annual report 2024. In Akkar, 1,200 children remained out of school in 2024 because of transport "
    "costs. The Ministry of Education will rehabilitate 20 schools by Q3 2025.",
    "Water supply in Zahle district fell by 30 percent during the summer of 2024. Solar pumping is "
    "recommended for the water establishments.",
    "Funding gaps reached 40 percent in 2024 according to the sector partners.",
]
OUT_OF_SCHOOL = {
    "category": "challenge",
    "tag": "Education › Access to learning › Out-of-school children",
    "text": "1,200 children were out of school in Akkar in 2024.",
    "quote": "In Akkar, 1,200 children remained out of school in 2024",
    "place": "Akkar",
    "date": "2024",
    "kind": "reported",
}
SCHOOLS = {
    "category": "action_point",
    "tag": "School infrastructure",
    "text": "The Ministry of Education will rehabilitate 20 schools by Q3 2025.",
    "quote": "The Ministry of Education will rehabilitate 20 schools by Q3 2025",
    "place": "",
    "date": "Q3 2025",
    "kind": "reported",
}
WATER = {
    "category": "observation",
    "tag": "Space travel",  # not a tag of the list: Other
    "text": "Water supply in Zahle fell by a third in the summer.",
    "quote": "Water supply in Zahle district fell by 30 percent",
    "place": "Zahle",
    "date": "summer 2024",
    "kind": "interpreted",
}
FUNDING = {
    "category": "recommendation",
    "tag": "Funding gaps",
    "text": "Donors should close the funding gap.",
    "quote": "Donors must close the gap before winter arrives",  # not in the text
    "place": "Mars",
    "date": "",
    "kind": "reported",
}
BY_PAGE = {"[Page 1]": [OUT_OF_SCHOOL, SCHOOLS], "[Page 2]": [WATER], "[Page 3]": [FUNDING]}


def _number(content: str, word: str) -> int:
    """The number the payload gives the finding whose line holds ``word``."""
    for line in content.splitlines():
        found = re.match(r"\[(\d+)\] ", line)
        if found and word in line:
            return int(found[1])
    return 999


def statements(content: str) -> dict:
    return {
        "statements": [
            {
                "text": "Out-of-school children in Akkar need urgent action.",
                "urgency": 140,  # held to 0-100
                "category": "challenge",
                "place": "Akkar",
                "date": "2024",
                "cites": [_number(content, "out of school"), 999],
            }
        ]
    }


def actions(content: str) -> dict:
    return {
        "action_points": [
            {
                "action": "Rehabilitate 20 schools.",
                "owner": "Ministry of Education",
                "deadline": "Q3 2025",
                "priority": "high",
                "cites": [_number(content, "rehabilitate")],
            },
            {
                "action": "Install solar pumping in the water establishments.",
                "owner": "",
                "deadline": "2025",
                "priority": "medium",
                "cites": [_number(content, "Water supply")],
            },
        ]
    }


class FakeReview:
    """``agent.client()`` for the document review: answers each stage from what was sent. ``broken``:
    a findings part whose text satisfies it gets a reply that is not JSON; ``errors``: raised in turn."""

    def __init__(self, broken=None, errors=(), during=None):
        self.broken = broken
        self.during = during  # called at the summary's call, as a person acts while the AI reads
        self.errors = list(errors)
        self.requests: list[dict] = []
        self.options: list[dict] = []

    def client(self):
        def with_options(**options):
            self.options.append(options)
            return api

        api = SimpleNamespace(with_options=with_options, responses=SimpleNamespace(create=self.create))
        return api

    def create(self, **params):
        self.requests.append(params)
        if self.errors:
            raise self.errors.pop(0)
        stage = params["text"]["format"]["name"]
        content = params["input"][0]["content"]
        if stage == "doc_review_findings":
            if self.broken and self.broken(content):
                return self._answer("{'findings': [", status="incomplete")
            out = {"findings": [f for marker, found in BY_PAGE.items() if marker in content for f in found]}
        elif stage == "doc_review_summary":
            if self.during:
                self.during()
            out = statements(content)
        else:
            out = actions(content)
        return self._answer(json.dumps(out))

    @staticmethod
    def _answer(text: str, status: str = "completed"):
        return SimpleNamespace(
            output_text=text,
            status=status,
            output=[],
            usage=SimpleNamespace(
                input_tokens=1000, input_tokens_details=SimpleNamespace(cached_tokens=0), output_tokens=200
            ),
        )

    def stages(self) -> list[str]:
        return [r["text"]["format"]["name"].removeprefix("doc_review_") for r in self.requests]


@pytest.fixture
def fake(monkeypatch):
    def install(**kw):
        found = FakeReview(**kw)
        monkeypatch.setattr(agent, "client", found.client)
        return found

    return install


@pytest.fixture
def ai_on():
    with override_settings(**AI_ON):
        yield


@pytest.fixture
def places(db):
    GovernorateLocation.objects.create(code="LB1", name="Akkar", ai_id=1)
    GovernorateLocation.objects.create(code="LB7", name="Bekaa", ai_id=7)
    DistrictLocation.objects.create(code="LB71", gov_code="LB7", name="Zahle")


@pytest.fixture
def setting(db):
    found = DocumentReviewSettings.load()
    found.enabled = True
    found.chunk_size = 2000
    found.save()
    return found


@pytest.fixture
def batch(db):
    return ReviewBatch.objects.create(name="Annual reports")


@pytest.fixture
def report(db, batch, places):
    document = Document.objects.create(
        title="Annual report 2024", text="\f".join(PAGES), status=Document.Status.READY, pages=3
    )
    review.put_in_batch([document], batch)
    return document


def _finding(document: Document, words: str) -> DocumentFinding:
    return document.findings.get(text__icontains=words)


# ------------------------------------------------------------------------------------------ the data
def test_the_topics_are_seeded_with_other_and_ten_programmes(db):
    names = set(TopicProgramme.objects.values_list("name", flat=True))
    assert {
        "Education", "Child Protection", "Health & Nutrition", "WASH", "Social Protection & Inclusion",
        "Adolescents & Youth", "Gender", "Emergency/Humanitarian response", "Partnerships & Funding",
        "Monitoring & Data", "Other",
    } == names  # fmt: skip
    other = Topic.other()
    assert other.is_other and other.path == "Other › Other › Other"
    assert Topic.objects.filter(subtopic__programme__name="Education").count() >= 6
    Topic.objects.filter(pk=other.pk).delete()
    assert Topic.other().is_other  # made again


def test_the_settings_start_switched_off_with_the_shipped_prompts(db):
    found = DocumentReviewSettings.load()
    assert not found.enabled and found.token_cap == 1_000_000 and found.statements_limit == 20
    for stage in DocumentReviewSettings.PROMPT_FIELDS:
        assert found.is_default(stage) and review_prompts.REQUIRED_SENTENCES[stage] in found.prompt(stage)
    found.tagging_prompt = "List the findings."  # the JSON sentence removed
    with pytest.raises(ValidationError) as refused:
        found.full_clean()
    assert "tagging_prompt" in refused.value.message_dict
    found.tagging_prompt = f"List the findings carefully.\n{review_prompts.REQUIRED_SENTENCES['tagging']}"
    found.full_clean()
    found.daily_token_cap = 5000
    assert found.token_cap == 5000


# ------------------------------------------------------------------------------------------ the pipeline
def test_a_document_is_read_into_findings_statements_and_action_points(report, setting, fake, ai_on):
    client = fake()
    run = review.run(review.PENDING, triggered_by="test")
    assert run.status == SyncRun.Status.SUCCEEDED and run.rows_written == 1, run.details
    assert client.stages() == ["findings", "summary", "enrichment"]
    assert client.options[0]["max_retries"] == 1
    sent = client.requests[0]
    assert sent["store"] is False and sent["text"]["format"]["strict"] is True
    assert (
        review_prompts.FIXED in sent["instructions"]
        and "Education › Access to learning" in sent["instructions"]
    )
    assert "[Page 1]" in sent["input"][0]["content"] and "[Page 3]" in sent["input"][0]["content"]
    report.refresh_from_db()
    assert report.review_status == Document.ReviewStatus.DONE and report.reviewed_text_hash
    notes = report.review_stage_notes
    assert [notes[s]["state"] for s in review.STAGES] == ["yes"] * 5 and notes["findings"][
        "read_share"
    ] == 100

    school = _finding(report, "out of school")
    assert school.topic.path == "Education › Access to learning › Out-of-school children"
    assert (school.page_label, school.page_from, school.exact_page, school.quote_found) == (
        "p. 1",
        1,
        True,
        True,
    )
    assert (school.place_match, school.governorate_name, school.finding_date) == (
        "governorate", "Akkar", datetime.date(2024, 1, 1),
    )  # fmt: skip
    assert school.evidence == 100  # found 45, reported 25, dated 10, placed 10, tagged 10
    water = _finding(report, "Water supply")
    assert water.topic.is_other and water.tag_text == "Space travel"  # an unknown tag: Other
    assert (water.place_match, water.district_name, water.governorate_name) == ("district", "Zahle", "Bekaa")
    assert water.page_label == "p. 2" and water.evidence == 45 + 10 + 10  # interpreted, Other
    funding = _finding(report, "funding gap")
    assert not funding.quote_found and not funding.exact_page and funding.page_label == "pp. 1–3"
    assert funding.place_match == "unmatched" and funding.place_text == "Mars" and funding.evidence == 25 + 10

    statement = DocumentStatement.objects.get(document=report)
    assert statement.urgency == 100 and list(statement.cites.all()) == [school]
    assert statement.topic == school.topic

    schools, solar = DocumentActionPoint.objects.filter(document=report).order_by("position")
    assert (schools.owner_text, schools.deadline_date, schools.priority, schools.derived) == (
        "Ministry of Education", datetime.date(2025, 9, 30), "high", False,
    )  # fmt: skip
    assert (solar.owner_text, solar.deadline_date, solar.derived) == (
        "Unassigned",
        datetime.date(2025, 12, 31),
        True,
    )
    derived = report.findings.get(derived=True)
    assert (
        derived.category == "action_point"
        and derived.quote == WATER["quote"]
        and derived in solar.cites.all()
    )
    assert AIUsage.objects.get(feature=usage.DOC_REVIEW).calls == 3
    assert run.details["calls"] == 3 and run.details["tokens"] == 3600


def test_a_broken_part_is_read_again_as_two_halves(report, setting, fake, ai_on):
    client = fake(broken=lambda text: "[Page 1]" in text and "[Page 3]" in text)  # the whole part only
    review.run(review.FULL, triggered_by="test")
    report.refresh_from_db()
    assert client.stages()[:3] == ["findings", "findings", "findings"]
    assert report.review_status == Document.ReviewStatus.DONE
    assert report.findings.filter(derived=False).count() == 4


def test_what_still_fails_is_left_out_and_the_document_is_partly_analysed(report, setting, fake, ai_on):
    fake(broken=lambda text: "[Page 3]" in text)  # the whole part and its second half
    run = review.run(review.FULL, triggered_by="test")
    report.refresh_from_db()
    assert report.review_status == Document.ReviewStatus.PARTLY and run.rows_written == 1
    note = report.review_stage_notes["findings"]
    assert note["state"] == "partly" and 0 < note["read_share"] < 100
    assert note["error"] == f"Only {note['read_share']}% read: some parts got no usable answer."
    assert not report.findings.filter(text__icontains="Water").exists()
    assert report.findings.filter(text__icontains="out of school").exists()


def test_nothing_read_fails_the_document_and_keeps_what_an_earlier_analysis_found(
    report, setting, fake, ai_on
):
    fake()
    review.run(review.FULL, triggered_by="test")
    fake(broken=lambda text: True)
    run = review.run(review.FULL, triggered_by="test")
    report.refresh_from_db()
    assert report.review_status == Document.ReviewStatus.FAILED and run.rows_failed == 1
    assert run.status == SyncRun.Status.PARTIAL
    assert report.findings.count() == 5  # the earlier analysis stays


def test_a_new_analysis_keeps_verdicts_manual_findings_and_action_point_statuses(
    report, setting, fake, ai_on, admin_user
):
    fake()
    review.run(review.FULL, triggered_by="test")
    now = timezone.now()
    school = _finding(report, "out of school")
    DocumentFinding.objects.filter(pk=school.pk).update(
        verdict=Verdict.ACCEPTED, reviewed_by=admin_user, reviewed_at=now
    )
    DocumentFinding.objects.filter(text__icontains="Water").update(
        verdict=Verdict.REJECTED, reviewed_by=admin_user
    )
    manual = DocumentFinding.objects.create(
        document=report, category="observation", topic=Topic.other(), text="Teachers were paid late in 2024.",
        quote="1,200 children remained out of school", manual=True, page_from=1, page_to=1,
    )  # fmt: skip
    DocumentActionPoint.objects.filter(action__icontains="schools").update(
        status="done", status_by=admin_user
    )
    DocumentStatement.objects.update(verdict=Verdict.ACCEPTED)

    client = fake()
    review.run(review.FULL, triggered_by="test")
    report.refresh_from_db()
    again = _finding(report, "out of school")
    assert again.pk != school.pk and again.verdict == Verdict.ACCEPTED and again.reviewed_by == admin_user
    assert _finding(report, "Water supply").verdict == Verdict.REJECTED
    assert DocumentFinding.objects.filter(pk=manual.pk, manual=True).exists()
    manual.refresh_from_db()
    assert manual.quote_found and manual.page_label == "p. 1"  # located again
    assert DocumentActionPoint.objects.get(action__icontains="schools").status == "done"
    assert DocumentStatement.objects.get(document=report).verdict == Verdict.ACCEPTED
    # a rejected finding is not read by the summary or the enrichment; a person's finding is
    summary = next(r for r in client.requests if r["text"]["format"]["name"] == "doc_review_summary")
    assert "Water supply" not in summary["input"][0]["content"]
    assert "Teachers were paid late" in summary["input"][0]["content"]


def test_a_verdict_given_while_the_ai_reads_is_kept(report, setting, fake, ai_on, admin_user):
    fake()
    review.run(review.FULL, triggered_by="test")

    def accept():
        DocumentFinding.objects.filter(text__icontains="Water").update(verdict=Verdict.ACCEPTED)
        DocumentActionPoint.objects.filter(action__icontains="solar").update(status="done")

    fake(during=accept)
    review.run(review.FULL, triggered_by="test")
    assert _finding(report, "Water supply").verdict == Verdict.ACCEPTED
    assert DocumentActionPoint.objects.get(action__icontains="solar").status == "done"


def test_the_enrichment_alone_draws_the_action_points_again_and_keeps_statuses(report, setting, fake, ai_on):
    fake()
    review.run(review.FULL, triggered_by="test")
    DocumentActionPoint.objects.filter(action__icontains="solar").update(status="dropped")
    findings_before = set(report.findings.filter(derived=False).values_list("pk", flat=True))
    client = fake()
    review.run(review.ENRICH, triggered_by="test")
    assert client.stages() == ["enrichment"]
    assert set(report.findings.filter(derived=False).values_list("pk", flat=True)) == findings_before
    assert DocumentActionPoint.objects.filter(document=report).count() == 2
    assert DocumentActionPoint.objects.get(action__icontains="solar").status == "dropped"
    assert report.findings.filter(derived=True).count() == 1


def test_locating_again_needs_no_ai(report, setting, fake, ai_on):
    fake()
    review.run(review.FULL, triggered_by="test")
    DocumentFinding.objects.update(evidence=0, page_label="", place_match="")
    DocumentReviewSettings.objects.update(enabled=False)  # no AI needed
    client = fake()
    run = review.run(review.LOCATE, triggered_by="test")
    assert client.requests == [] and run.details["findings_located"] == 5
    assert _finding(report, "out of school").evidence == 100
    assert _finding(report, "Water supply").place_match == "district"


def test_a_document_still_being_read_waits_and_one_without_text_fails(report, setting, fake, ai_on, batch):
    Document.objects.filter(pk=report.pk).update(status=Document.Status.INDEXING)
    empty = Document.objects.create(
        title="Scan", text="", status=Document.Status.FAILED, error="No text layer"
    )
    review.put_in_batch([empty], batch)
    client = fake()
    review.run(review.PENDING, triggered_by="test")
    report.refresh_from_db()
    empty.refresh_from_db()
    assert client.requests == [] and report.review_status == Document.ReviewStatus.PENDING
    assert "still reading" in report.review_stage_notes["waiting"]
    assert empty.review_status == Document.ReviewStatus.FAILED
    assert "No text layer" in empty.review_stage_notes["text"]["error"]


def test_only_documents_in_a_batch_and_not_references_are_analysed(report, setting, fake, ai_on, batch):
    loose = Document.objects.create(title="Loose", text=PAGES[0], status=Document.Status.READY)
    reference = Document.objects.create(title="Guidance", text=PAGES[0], status=Document.Status.READY)
    review.put_in_batch([reference], batch)
    review.set_reference(reference, True)
    archived = ReviewBatch.objects.create(name="Old", archived=True)
    old = Document.objects.create(title="Old", text=PAGES[0], status=Document.Status.READY)
    review.put_in_batch([old], archived)
    fake()
    run = review.run(review.FULL, triggered_by="test")
    assert run.rows_in == 1
    assert not loose.findings.exists() and not reference.findings.exists() and not old.findings.exists()
    review.take_out(report)
    report.refresh_from_db()
    assert report.review_status == Document.ReviewStatus.NOT_IN_REVIEW and report.review_batch is None


# ------------------------------------------------------------------------------------------ limits
def test_switched_off_nothing_is_analysed(report, fake, ai_on):
    client = fake()
    run = review.run(review.PENDING, triggered_by="test")  # off until an administrator turns it on
    assert client.requests == [] and run.details["skipped_reason"] == review.OFF
    report.refresh_from_db()
    assert report.review_status == Document.ReviewStatus.PENDING


def _used_today(tokens: int) -> None:
    used = SimpleNamespace(
        input_tokens=tokens, input_tokens_details=SimpleNamespace(cached_tokens=0), output_tokens=0
    )
    usage.record(usage.DOC_REVIEW, "test-model", used)


def test_the_daily_cap_stops_the_run_cleanly(report, setting, fake, ai_on):
    setting.daily_token_cap = 60_000
    setting.save()
    _used_today(59_000)  # the document fits a day, but not what is left of today
    client = fake()
    run = review.run(review.PENDING, triggered_by="test")
    report.refresh_from_db()
    assert client.requests == [] and run.status == SyncRun.Status.SUCCEEDED
    assert run.details["stopped"] == review.BUDGET and report.review_status == Document.ReviewStatus.PENDING
    assert report.review_stage_notes["waiting"] == review.BUDGET


def test_the_cap_reached_half_way_leaves_the_document_as_it_was(report, setting, fake, ai_on, monkeypatch):
    estimate = review.estimate

    def costly_summary(stage, payload, setting_):
        return 10**9 if stage == review.SUMMARY else estimate(stage, payload, setting_)

    monkeypatch.setattr(review, "estimate", costly_summary)
    client = fake()
    run = review.run(review.PENDING, triggered_by="test")
    report.refresh_from_db()
    assert client.stages() == ["findings"] and run.details["stopped"] == review.BUDGET
    assert report.review_status == Document.ReviewStatus.PENDING and not report.findings.exists()


def test_the_credit_pause_and_the_shared_cap_stop_it(report, setting, fake, ai_on, monkeypatch):
    client = fake()
    monkeypatch.setattr("neurodb.fmm.ai.budget.paused_until", lambda *a: timezone.now())
    run = review.run(review.PENDING, triggered_by="test")
    assert client.requests == [] and run.details["stopped"] == review.PAUSED
    monkeypatch.setattr("neurodb.fmm.ai.budget.paused_until", lambda *a: None)
    _used_today(79_000)  # the document fits 80% of the shared cap, but not what is left of it today
    with override_settings(AI_DAILY_TOKEN_SOFT_CAP=100_000):
        run = review.run(review.PENDING, triggered_by="test")
    report.refresh_from_db()
    assert client.requests == [] and run.details["stopped"] == review.SHARED_BUDGET
    assert report.review_status == Document.ReviewStatus.PENDING


def test_a_document_too_long_for_a_whole_day_fails_without_a_call(report, setting, fake, ai_on):
    setting.daily_token_cap = 5000  # less than one analysis of the report needs
    setting.save()
    client = fake()
    run = review.run(review.PENDING, triggered_by="test")
    report.refresh_from_db()
    assert client.requests == [] and run.rows_failed == 1
    assert report.review_status == Document.ReviewStatus.FAILED
    assert "Daily token cap" in report.review_stage_notes["findings"]["error"]
    with override_settings(AI_DAILY_TOKEN_SOFT_CAP=6000):  # 80% of the shared cap counts too
        setting.daily_token_cap = 0
        setting.save()
        review.run(review.PENDING, triggered_by="test")
    report.refresh_from_db()
    assert client.requests == [] and report.review_status == Document.ReviewStatus.FAILED
    error = report.review_stage_notes["findings"]["error"]
    assert "AI_DAILY_TOKEN_SOFT_CAP" in error and "4,800" in error and "Daily token cap" not in error
    setting.daily_token_cap = 0
    setting.save()
    review.run(review.PENDING, triggered_by="test")  # a failed document is tried again: it fits now
    report.refresh_from_db()
    assert report.review_status == Document.ReviewStatus.DONE and client.requests


def test_a_full_run_stopped_by_the_budget_leaves_the_rest_waiting(
    report, setting, fake, ai_on, batch, monkeypatch
):
    second = Document.objects.create(
        title="Donor report 2024", text="\f".join(PAGES), status=Document.Status.READY, pages=3
    )
    review.put_in_batch([second], batch)
    fake()
    review.run(review.FULL, triggered_by="test")
    assert set(Document.objects.values_list("review_status", flat=True)) == {Document.ReviewStatus.DONE}

    client = fake()
    blocked = review.blocked
    monkeypatch.setattr(  # the budget lasts for one document's three calls
        review,
        "blocked",
        lambda tokens, s: review.BUDGET if len(client.requests) >= 3 else blocked(tokens, s),
    )
    run = review.run(review.FULL, triggered_by="test")
    assert run.details["stopped"] == review.BUDGET and run.details["left_waiting"] == 1
    report.refresh_from_db()
    second.refresh_from_db()
    assert report.review_status == Document.ReviewStatus.DONE
    assert second.review_status == Document.ReviewStatus.PENDING and second.findings.count() == 5  # kept
    monkeypatch.setattr(review, "blocked", blocked)
    run = review.run(review.PENDING, triggered_by="test")  # the next night goes on with it
    second.refresh_from_db()
    assert run.rows_written == 1 and second.review_status == Document.ReviewStatus.DONE


def test_a_document_taken_out_while_the_ai_reads_keeps_nothing(report, setting, fake, ai_on):
    fake(during=lambda: review.take_out(Document.objects.get(pk=report.pk)))
    run = review.run(review.FULL, triggered_by="test")
    report.refresh_from_db()
    assert run.rows_written == 0 and run.details["waiting"] == 1
    assert report.review_status == Document.ReviewStatus.NOT_IN_REVIEW and report.review_batch is None
    assert not report.findings.exists() and not report.statements.exists()


def test_a_finding_a_person_deletes_while_the_ai_reads_is_not_cited(report, setting, fake, ai_on):
    DocumentFinding.objects.create(
        document=report,
        category="observation",
        topic=Topic.other(),
        manual=True,
        text="Children out of school were counted twice.",
        quote="",
    )  # the statement cites it: the first finding with "out of school"
    fake(during=lambda: DocumentFinding.objects.filter(manual=True).delete())
    run = review.run(review.FULL, triggered_by="test")
    report.refresh_from_db()
    assert run.rows_written == 1 and report.review_status == Document.ReviewStatus.DONE
    statement = DocumentStatement.objects.get(document=report)
    assert list(statement.cites.all()) == [] and not report.findings.filter(manual=True).exists()


def test_a_new_text_read_while_the_ai_reads_is_analysed_again(report, setting, fake, ai_on):
    def read_again():
        Document.objects.filter(pk=report.pk).update(text="\f".join([*PAGES, "A new page."]))
        assert not review.text_read(Document.objects.get(pk=report.pk))  # it is being analysed

    fake(during=read_again)
    run = review.run(review.FULL, triggered_by="test")
    report.refresh_from_db()
    assert report.review_status == Document.ReviewStatus.PENDING and run.details["waiting"] == 1
    assert report.review_stage_notes["waiting"] == review.TEXT_CHANGED
    assert report.findings.count() == 5  # kept meanwhile
    fake()
    run = review.run(review.PENDING, triggered_by="test")
    report.refresh_from_db()
    assert run.rows_written == 1 and report.review_status == Document.ReviewStatus.DONE
    assert report.reviewed_text_hash == review.text_hash(report)


def test_a_persons_edit_of_a_finding_is_kept_with_its_verdict(report, setting, fake, ai_on, admin_user):
    fake()
    review.run(review.FULL, triggered_by="test")
    water = _finding(report, "Water supply")
    DocumentFinding.objects.filter(pk=water.pk).update(
        text="Water supply in Zahle fell by 30% in summer 2024.",
        topic=Topic.objects.get(name="Water supply"),
        kind=DocumentFinding.Kind.REPORTED,
        verdict=Verdict.ACCEPTED,
        edited_by=admin_user,
        edited_at=timezone.now(),
    )
    fake()
    review.run(review.FULL, triggered_by="test")
    again = report.findings.get(text__icontains="by 30% in summer")
    assert again.pk != water.pk and again.verdict == Verdict.ACCEPTED and again.edited_by == admin_user
    # the person's tag and "reported" (the review page's edit form sets both) stay: tagged and reported now
    assert again.topic.name == "Water supply" and again.kind == DocumentFinding.Kind.REPORTED
    assert again.evidence == 45 + 25 + 10 + 10 + 10
    assert not report.findings.filter(text=WATER["text"]).exists()


def test_a_document_with_no_progress_for_30_minutes_is_marked_failed(report, setting, fake, ai_on):
    Document.objects.filter(pk=report.pk).update(
        review_status=Document.ReviewStatus.RUNNING,
        review_progress_at=timezone.now() - datetime.timedelta(minutes=31),
    )
    other = Document.objects.create(title="Busy", text="x", review_status=Document.ReviewStatus.RUNNING)
    Document.objects.filter(pk=other.pk).update(review_progress_at=timezone.now())
    assert review.left_behind() == 1
    report.refresh_from_db()
    assert (
        report.review_status == Document.ReviewStatus.FAILED
        and report.review_stage_notes["stopped"] == review.STALE
    )


def test_a_document_read_again_with_a_new_text_waits_to_be_analysed(report, setting, fake, ai_on):
    fake()
    review.run(review.FULL, triggered_by="test")
    report.refresh_from_db()
    assert not review.text_read(report)  # the same text
    report.text += "\fA new page."
    assert review.text_read(report)
    report.refresh_from_db()
    assert report.review_status == Document.ReviewStatus.PENDING


# ------------------------------------------------------------------------------------------ the job
def test_the_command_and_its_buttons_and_schedule(report, setting, fake, ai_on, started):
    fake()
    call_command("review_documents", "--document", str(report.pk), "--triggered-by", "test")
    run = SyncRun.objects.get(job=SyncRun.Job.DOC_REVIEW)
    assert run.target == f"full #{report.pk}" and run.triggered_by == "test"
    loose = Document.objects.create(title="Loose", text="x")
    with pytest.raises(CommandError):
        call_command("review_documents", "--document", str(loose.pk))
    waiting = Document.objects.create(title="Not analysed yet", text="x")
    review.put_in_batch([waiting], report.review_batch)
    with pytest.raises(CommandError):  # no analysis to draw the action points again from
        call_command("review_documents", "--enrich", "--document", str(waiting.pk))
    assert jobs.COMMANDS["doc_review"].args == ("review_documents", "--pending")
    buttons = {j.name: j.command for j in admin_jobs.BACKGROUND_JOBS if j.job == SyncRun.Job.DOC_REVIEW}
    assert buttons == {
        "run_doc_review": ("review_documents", "--pending"),
        "run_doc_review_full": ("review_documents", "--full"),
        "run_doc_review_locate": ("review_documents", "--locate"),
    }
    scheduled = ScheduledJob.objects.get(key="doc-review")
    assert (scheduled.command, scheduled.schedule, scheduled.enabled) == ("doc_review", "40 4 * * *", True)
    assert background.LOCK_IDS[SyncRun.Job.DOC_REVIEW] == review.LOCK_ID
    assert usage.DOC_REVIEW in usage.FEATURES
    review.start(report, triggered_by="editor")
    assert started[-1] == ("review_documents", "--document", str(report.pk), "--triggered-by", "editor")


def test_the_admin_settings_are_for_administrators_and_restore_the_shipped_prompt(
    admin_user, client, editor, report
):
    found = DocumentReviewSettings.load()
    found.summary_prompt = f"Short statements.\n{review_prompts.REQUIRED_SENTENCES['summary']}"
    found.save()
    url = reverse("admin:knowledge_documentreviewsettings_change", args=[found.pk])
    client.force_login(admin_user)
    page = client.get(url)
    assert page.status_code == 200 and "Restore the shipped key statements prompt" in page.content.decode()
    data = {
        "enabled": "on", "tagging_prompt": "No sentence here", "summary_prompt": found.summary_prompt,
        "enrichment_prompt": found.enrichment_prompt, "statements_per_document": 0,
        "max_findings_per_chunk": 25, "chunk_size": 12000, "daily_token_cap": 0,
    }  # fmt: skip
    refused = client.post(url, data)
    assert refused.status_code == 200 and "Keep the sentence" in refused.content.decode()
    data.update(tagging_prompt=found.tagging_prompt, restore_summary="on")
    assert client.post(url, data).status_code == 302
    found.refresh_from_db()
    assert found.enabled and found.is_default("summary") and found.updated_by == admin_user
    for name in ("topicprogramme", "topic", "reviewbatch", "documentfinding", "documentactionpoint"):
        assert client.get(reverse(f"admin:knowledge_{name}_changelist")).status_code == 200
    client.force_login(editor)
    assert client.get(url).status_code in (302, 403)


def test_putting_a_document_in_a_batch_in_the_admin_makes_it_wait(admin_client, batch, started):
    document = Document.objects.create(title="Evaluation", text="x", status=Document.Status.READY)
    url = reverse("admin:knowledge_document_change", args=[document.pk])
    page = admin_client.get(url)
    assert page.status_code == 200
    form = {
        "title": document.title,
        "source": "",
        "summary": "",
        "key_points": "[]",
        "review_batch": batch.pk,
    }
    form.update(
        {
            "links-TOTAL_FORMS": 0,
            "links-INITIAL_FORMS": 0,
            "links-MIN_NUM_FORMS": 0,
            "links-MAX_NUM_FORMS": 1000,
        }
    )
    response = admin_client.post(url, form)
    assert response.status_code == 302, response.content.decode()[:3000]
    document.refresh_from_db()
    assert document.review_batch == batch and document.review_status == Document.ReviewStatus.PENDING


# ------------------------------------------------------------------------------------------ Ask NeuroDB
def test_ask_neurodb_searches_the_findings_without_rejected_ones(report, setting, fake, ai_on):
    fake()
    review.run(review.FULL, triggered_by="test")
    assert "search_document_findings" in [d["name"] for d in tools.definitions()]
    assert "search_document_findings" in agent.SYSTEM_PROMPT
    out = tools.run("search_document_findings", {"query": "Akkar children"})
    first = out["findings"][0]
    assert first["finding"] == OUT_OF_SCHOOL["text"] and first["page"] == "p. 1"
    assert first["url"] == f"/knowledge/review/?tab=findings&view=document&document={report.pk}"
    assert first["batch"] == "Annual reports"
    assert first["evidence"] == 100 and "instructions" in out["note"]
    DocumentFinding.objects.filter(text__icontains="Akkar").update(verdict=Verdict.REJECTED)
    out = tools.run("search_document_findings", {"query": "Akkar children"})
    assert all("Akkar" not in f["finding"] for f in out["findings"])
    assert tools.run("search_document_findings", {"query": "water", "batch": "Annual"})["findings"]
    assert not tools.run("search_document_findings", {"query": "water", "batch": "Evaluations"})["findings"]
    assert tools.run("search_document_findings", {"query": "water", "batch": str(report.review_batch_id)})[
        "findings"
    ]
    assert not tools.run("search_document_findings", {"query": "water", "batch": "9" * 30})["findings"]
    review.take_out(report)
    assert not tools.run("search_document_findings", {"query": "water"})["findings"]


# ------------------------------------------------------------------------------------------ without AI
@pytest.mark.parametrize(
    ("text", "end", "expected"),
    [
        ("Q3 2025", True, datetime.date(2025, 9, 30)),
        ("Q3 2025", False, datetime.date(2025, 7, 1)),
        ("end of 2025", True, datetime.date(2025, 12, 31)),
        ("March 2025", True, datetime.date(2025, 3, 31)),
        ("H1 2024", True, datetime.date(2024, 6, 30)),
        ("2024-2025", True, datetime.date(2025, 12, 31)),
        ("30/06/2025", False, datetime.date(2025, 6, 30)),
        ("2025-06-30", False, datetime.date(2025, 6, 30)),
        ("first quarter of 2026", True, datetime.date(2026, 3, 31)),
        ("soon", True, None),
        ("", False, None),
    ],
)
def test_dates_are_read_as_periods(text, end, expected):
    assert review_locate.period(text, end=end) == expected


def test_places_are_matched_to_the_gazetteer(places):
    gazetteer = review_locate.Gazetteer()
    assert gazetteer.match("Akkar Governorate").match == "governorate"
    assert gazetteer.match("Zahle, Bekaa").district_name == "Zahle"
    assert gazetteer.match("nationwide").match == "general"
    assert gazetteer.match("Atlantis").match == "unmatched"
    assert gazetteer.match("").match == ""


def test_parts_carry_page_markers_and_quotes_are_found_across_lines():
    document = SimpleNamespace(text="\f".join(PAGES), file=None)
    parts = review_locate.parts(document, 120)
    assert parts[0].text.startswith("[Page 1]") and parts[-1].last == 3
    index = review_locate.TextIndex(document)
    assert index.find("water supply in ZAHLE\ndistrict fell by 30 percent") == 2
    assert index.find("Funding gaps … according to the sector partners") == 3
    assert index.find("not in this document at all, surely") is None
    assert index.find("Funding gaps reached 4") is None  # whole words: "4" is not "40"
    assert index.find("Funding gaps reached 40") == 3
    halves = review_locate.halves(review_locate.Part(1, "page", [(1, PAGES[0]), (2, PAGES[1])]))
    assert [h.first for h in halves] == [1, 2]


def test_the_admin_action_analyses_the_chosen_documents(admin_client, report, batch, started):
    loose = Document.objects.create(title="Loose", text="x", status=Document.Status.READY)
    url = reverse("admin:knowledge_document_changelist")
    response = admin_client.post(
        url, {"action": "analyse_in_review", "_selected_action": [report.pk, loose.pk]}, follow=True
    )
    page = response.content.decode()
    assert "Analysing 1 document(s)" in page and "1 document(s) left out" in page
    assert started == [("review_documents", "--document", str(report.pk), "--triggered-by", "admin")]
    second = Document.objects.create(title="Second", text="y", status=Document.Status.READY)
    review.put_in_batch([second], batch)
    admin_client.post(url, {"action": "analyse_in_review", "_selected_action": [report.pk, second.pk]})
    assert started[-1] == ("review_documents", "--pending", "--triggered-by", "admin")


# ================================================================================ the review page (D3.3)
REVIEW_URL = "/knowledge/review/"
TABS = ("documents", "findings", "dashboard", "synthesis", "actions", "report")


@pytest.fixture
def analysed(report, setting, fake, ai_on):
    """The annual report analysed by the scripted model (4 findings and a derived one, a statement, two
    action points)."""
    fake()
    review.run(review.FULL, triggered_by="test")
    report.refresh_from_db()
    return report


def _tiles(client, **params) -> dict:
    response = client.get(REVIEW_URL, {"tab": "dashboard", **params})
    assert response.status_code == 200
    return response.context["tiles"]


def _topics(n: int) -> list[Topic]:
    return list(
        Topic.objects.exclude(name=Topic.OTHER).select_related("subtopic__programme").order_by("pk")[:n]
    )


def _doc(batch: ReviewBatch, title: str, year: int | None = None) -> Document:
    document = Document.objects.create(
        title=title, text="x", status=Document.Status.READY, year=year, review_batch=batch,
        review_status=Document.ReviewStatus.DONE,
    )  # fmt: skip
    return document


def _add(document: Document, topic: Topic, text: str, year: int | None = None, **extra) -> DocumentFinding:
    return DocumentFinding.objects.create(
        document=document, topic=topic, text=text, category=extra.pop("category", "challenge"),
        finding_date=datetime.date(year, 1, 1) if year else None, page_label="p. 2", page_from=2, **extra,
    )  # fmt: skip


def test_every_tab_renders_for_a_viewer_read_only(analysed, client_viewer, editor, client):
    for tab in TABS:
        full = client_viewer.get(REVIEW_URL, {"tab": tab})
        assert full.status_code == 200, tab
        page = full.content.decode()
        assert "Document review" in page and 'aria-current="page"' in page
        assert "manage.py" not in page
        partial = client_viewer.get(REVIEW_URL, {"tab": tab}, HTTP_HX_REQUEST="true")
        assert partial.status_code == 200 and "<html" not in partial.content.decode()
    for view in ("all", "document", "statements", "index"):
        assert client_viewer.get(REVIEW_URL, {"tab": "findings", "view": view}).status_code == 200
    for view in ("themes", "time", "coverage", "repeated"):
        assert (
            client_viewer.get(REVIEW_URL, {"tab": "synthesis", "view": view, "min": "2"}).status_code == 200
        )
    findings = client_viewer.get(REVIEW_URL, {"tab": "findings"}).content.decode()
    assert "Akkar" in findings and "/verdict/" not in findings and "/edit/" not in findings
    documents = client_viewer.get(REVIEW_URL, {"tab": "documents"}).content.decode()
    assert "Annual report 2024" in documents and "/analyse/" not in documents
    client.force_login(editor)
    by_document = client.get(REVIEW_URL, {"tab": "findings", "view": "document", "document": analysed.pk})
    page = by_document.content.decode()
    assert "Accept all" in page and "/verdict/" in page and "Add a finding" in page
    assert "/analyse/" in client.get(REVIEW_URL, {"tab": "documents"}).content.decode()
    sidebar = client.get(reverse("knowledge:index")).content.decode()
    assert REVIEW_URL in sidebar and "Document review" in sidebar


def test_the_documents_tab_shows_stages_counts_and_partly_read(analysed, client_viewer):
    notes = analysed.review_stage_notes
    notes["findings"].update(
        state="partly", read_share=80, error="Only 80% read: some parts got no usable answer."
    )
    Document.objects.filter(pk=analysed.pk).update(review_stage_notes=notes, review_status="partly")
    response = client_viewer.get(REVIEW_URL, {"tab": "documents", "batch": analysed.review_batch_id})
    row = response.context["rows"][0]
    assert row.n_findings == 5 and row.n_statements == 1 and row.n_actions == 2 and row.read_share == 80
    assert [c["key"] for c in row.stage_chips] == [
        "succeeded",
        "partial",
        "succeeded",
        "succeeded",
        "succeeded",
    ]
    page = response.content.decode()
    assert "Only 80% read" in page and "Findings: Only 80% read" in page


def test_a_document_left_running_shows_as_failed_on_the_page(analysed, client_viewer):
    Document.objects.filter(pk=analysed.pk).update(
        review_status="running", review_progress_at=timezone.now() - datetime.timedelta(hours=1)
    )
    page = client_viewer.get(REVIEW_URL, {"tab": "documents"}).content.decode()
    analysed.refresh_from_db()
    assert analysed.review_status == "failed" and "No progress for 30 minutes" in page


def _post_urls(document: Document, finding: DocumentFinding, statement, point, batch, topic) -> list:
    return [
        ("knowledge:review_batch_new", [], {"name": "Evaluations"}),
        ("knowledge:review_batch_edit", [batch.pk], {"action": "rename", "name": "Renamed"}),
        ("knowledge:review_batch_add", [batch.pk], {"documents": []}),
        ("knowledge:review_analyse", [document.pk], {}),
        ("knowledge:review_reference", [document.pk], {"reference": "1"}),
        ("knowledge:review_bulk_verdict", [document.pk], {"verdict": "accepted"}),
        ("knowledge:review_finding_new", [document.pk], {"text": "A finding.", "topic": topic.pk,
                                                         "category": "challenge"}),
        ("knowledge:review_finding_verdict", [finding.pk], {"verdict": "accepted"}),
        ("knowledge:review_finding_edit", [finding.pk], {"text": "Edited.", "topic": topic.pk,
                                                          "category": "observation"}),
        ("knowledge:review_statement_verdict", [statement.pk], {"verdict": "rejected"}),
        ("knowledge:review_action_status", [point.pk], {"status": "done"}),
        ("knowledge:review_finding_delete", [finding.pk], {}),
        ("knowledge:review_remove", [document.pk], {}),
    ]  # fmt: skip


def test_every_change_needs_an_administrator_or_a_section_editor(analysed, client, viewer, editor, started):
    finding = analysed.findings.filter(derived=False).first()
    statement = analysed.statements.first()
    point = analysed.review_action_points.first()
    topic = _topics(1)[0]
    urls = _post_urls(analysed, finding, statement, point, analysed.review_batch, topic)
    client.force_login(viewer)
    for name, args, data in urls:
        assert client.post(reverse(name, args=args), data).status_code == 403, name
        if name in ("knowledge:review_finding_new", "knowledge:review_finding_edit"):
            assert client.get(reverse(name, args=args)).status_code == 403, name
    assert not ReviewBatch.objects.filter(name="Evaluations").exists()
    assert DocumentFinding.objects.filter(pk=finding.pk, verdict="unreviewed").exists()
    # every reader may turn Verified only on for themself
    assert client.post(reverse("knowledge:review_verified"), {"on": "1"}).status_code == 302
    client.force_login(editor)
    for name, args, data in urls:
        response = client.post(reverse(name, args=args), data)
        assert response.status_code in (200, 204, 302), (name, response.status_code)
    point.refresh_from_db()
    assert point.status == "done" and point.status_by == editor
    assert ReviewBatch.objects.filter(name="Evaluations").exists()
    analysed.refresh_from_db()
    assert analysed.review_batch is None and analysed.review_status == "not_in_review"


def test_a_reviewer_accepts_rejects_edits_adds_and_deletes(analysed, client, editor):
    client.force_login(editor)
    school = _finding(analysed, "out of school")
    water = _finding(analysed, "Water supply")
    url = reverse("knowledge:review_finding_verdict", args=[school.pk])
    row = client.post(url, {"verdict": "rejected"}, HTTP_HX_REQUEST="true")
    assert row.status_code == 200 and 'class="is-rejected"' in row.content.decode()
    school.refresh_from_db()
    assert school.verdict == "rejected" and school.reviewed_by == editor and school.reviewed_at
    client.post(url, {"verdict": "unreviewed"})
    school.refresh_from_db()
    assert school.verdict == "unreviewed" and school.reviewed_by is None
    # Accept all: the unreviewed only, a decided finding stays
    water.verdict = "rejected"
    water.save()
    client.post(reverse("knowledge:review_bulk_verdict", args=[analysed.pk]), {"verdict": "accepted"})
    water.refresh_from_db()
    school.refresh_from_db()
    assert water.verdict == "rejected" and school.verdict == "accepted"
    # an edit recomputes the page, place, date and evidence, and keeps the person's page
    akkar_topic = school.topic
    edit = reverse("knowledge:review_finding_edit", args=[water.pk])
    assert client.get(edit, HTTP_HX_REQUEST="true").status_code == 200
    response = client.post(
        edit,
        {"text": "Water supply fell in Akkar.", "quote": "In Akkar, 1,200 children", "topic": akkar_topic.pk,
         "category": "challenge", "kind": "reported", "place_text": "Akkar", "date_text": "2024", "page": ""},
        HTTP_HX_REQUEST="true",
    )  # fmt: skip
    assert response.status_code == 204 and response["HX-Redirect"]
    water.refresh_from_db()
    assert (water.text, water.page_label, water.place_match, water.evidence) == (
        "Water supply fell in Akkar.", "p. 1", "governorate", 100,
    )  # fmt: skip
    assert water.edited_by == editor and water.verdict == "rejected"  # the verdict stays
    refused = client.post(edit, {"text": "", "topic": "x", "category": "nope"})
    assert refused.status_code == 200 and "Write the finding" in refused.content.decode()
    # a person's finding: accepted, manual, located, with the page they gave
    new = reverse("knowledge:review_finding_new", args=[analysed.pk])
    client.post(
        new,
        {"text": "Teachers were not paid.", "quote": "not in the text", "topic": akkar_topic.pk,
         "category": "challenge", "kind": "reported", "place_text": "Zahle", "date_text": "", "page": "3"},
    )  # fmt: skip
    added = analysed.findings.get(manual=True)
    assert added.verdict == "accepted" and added.reviewed_by == editor
    assert (added.page_label, added.place_match, added.evidence) == ("p. 3", "district", 25 + 10 + 10)
    response = client.post(
        reverse("knowledge:review_finding_delete", args=[added.pk]), HTTP_HX_REQUEST="true"
    )
    assert response.status_code == 200 and not DocumentFinding.objects.filter(pk=added.pk).exists()
    statement = analysed.statements.get()
    row = client.post(
        reverse("knowledge:review_statement_verdict", args=[statement.pk]),
        {"verdict": "accepted", "compact": "1"},
        HTTP_HX_REQUEST="true",
    )
    assert row.status_code == 200 and "Accepted" in row.content.decode()


def test_verified_only_changes_the_counts_and_rejected_ones_are_always_out(analysed, client, editor):
    client.force_login(editor)
    first = client.get(REVIEW_URL)
    assert "Verified only" not in first.content.decode()  # no verdict yet: no switch
    before = _tiles(client)
    assert before["findings"] == 5 and before["statements"] == 1 and before["open_actions"] == 2
    school = _finding(analysed, "out of school")
    rehabilitate = analysed.findings.get(text__icontains="rehabilitate")
    DocumentFinding.objects.filter(pk=school.pk).update(verdict="rejected")
    DocumentFinding.objects.filter(pk=rehabilitate.pk).update(verdict="accepted")
    assert _tiles(client)["findings"] == 4  # rejected: always out
    assert "Verified only: off" in client.get(REVIEW_URL).content.decode()
    client.post(reverse("knowledge:review_verified"), {"on": "1", "next": f"{REVIEW_URL}?tab=dashboard"})
    after = _tiles(client)
    assert after["findings"] == 1 and after["statements"] == 0
    assert after["open_actions"] == 1  # the action point citing the accepted finding only
    counted = client.get(REVIEW_URL, {"tab": "findings", "counted": "1"})
    assert counted.context["page_obj"].paginator.count == after["findings"]
    report = client.get(REVIEW_URL, {"tab": "report"}).content.decode()
    assert "Verified only is on" in report
    client.post(reverse("knowledge:review_verified"), {"on": "0"})
    assert _tiles(client)["findings"] == 4
    # the Findings tab still lists the rejected one (struck through) for its review
    listed = client.get(REVIEW_URL, {"tab": "findings"})
    assert listed.context["page_obj"].paginator.count == 5 and "is-rejected" in listed.content.decode()


def test_the_dashboard_figures_open_the_rows_they_count(analysed, client_viewer):
    response = client_viewer.get(REVIEW_URL, {"tab": "dashboard"})
    figures = response.context["figures"]
    quality = {q["key"]: q for q in figures["quality"]}
    other = analysed.findings.filter(topic=Topic.other())  # Water (an unknown tag) and what derives from it
    assert quality["untagged"]["gap"] == other.count() >= 1 and quality["untagged"]["total"] == 5
    gap = client_viewer.get(REVIEW_URL, {"tab": "findings", "counted": "1", "gap": "untagged"})
    assert WATER["text"] in [f.text for f in gap.context["findings"]]
    for tile in figures["quality"][:5]:
        listed = client_viewer.get(REVIEW_URL, {"tab": "findings", "counted": "1", "gap": tile["key"]})
        assert listed.context["page_obj"].paginator.count == tile["gap"], tile["key"]
    for label, n, key in figures["charts"]["category"]:
        listed = client_viewer.get(REVIEW_URL, {"tab": "findings", "counted": "1", "category": key})
        assert listed.context["page_obj"].paginator.count == n, label
    for row in figures["charts"]["years"] + figures["charts"]["evidence"]:
        key = "year" if row in figures["charts"]["years"] else "band"
        listed = client_viewer.get(REVIEW_URL, {"tab": "findings", "counted": "1", key: row["drill"]})
        assert listed.context["page_obj"].paginator.count == row["value"], row
    for name, n, drill in figures["charts"]["places"]:
        listed = client_viewer.get(REVIEW_URL, {"tab": "findings", "counted": "1", "place": drill})
        assert listed.context["page_obj"].paginator.count == n, name
    assert {name for name, _n, _d in figures["charts"]["places"]} >= {"Akkar", "Zahle"}
    urgent = client_viewer.get(
        REVIEW_URL, {"tab": "findings", "view": "statements", "counted": "1", "urgent": "1"}
    )
    assert urgent.context["page_obj"].paginator.count == figures["tiles"]["urgent"] == 1
    assert figures["per_batch"][0]["findings"] == figures["tiles"]["findings"]
    assert "review-chart-data" in response.content.decode()


def test_synthesis_ranks_themes_by_distinct_documents(db, client_viewer):
    annual, donor = ReviewBatch.objects.create(name="Annual"), ReviewBatch.objects.create(name="Donor")
    a, b, c = _doc(annual, "Report A", 2022), _doc(annual, "Report B", 2023), _doc(donor, "Report C", 2024)
    loud, shared, spread = _topics(3)
    for n in range(5):  # one document saying it five times: one document's view
        _add(a, loud, f"Point number {n} about the loud topic.", 2022)
    _add(a, shared, "Teachers lack training in rural schools.", 2022)
    _add(b, shared, "Teachers lack training in rural schools of the north.", 2024)
    for document, year in ((a, 2022), (b, 2023), (c, 2024)):
        _add(document, spread, f"Water trucking costs rose in {year}.", year)
    found = review_data.synthesis(None, False, min_documents=2)
    assert [t.topic for t in found.themes] == [spread, shared]
    assert [t.n_documents for t in found.themes] == [3, 2] and loud not in [t.topic for t in found.themes]
    assert {t.topic: t.trend for t in found.themes} == {spread: "persistent", shared: "recurring"}
    assert [t.topic for t in found.coverage()] == [shared]  # Annual only
    assert [t.topic for t in review_data.synthesis(None, False, min_documents=3).themes] == [spread]
    assert (
        review_data.synthesis(None, False, min_documents=2, q=shared.name.lower()).themes[0].topic == shared
    )
    assert not review_data.synthesis(None, False, min_documents=2, challenges=False, q="zzz").themes
    groups = review_data.repeated(None, False)
    # the water findings (3 documents) then the two teacher findings: words at least 60% alike
    assert [g.documents for g in groups] == [3, 2] and "Teachers" in groups[1].lead.text
    assert review_data.trend({2024}, 2024) == "emerging" and review_data.trend({2021}, 2024) == "no_longer"
    page = client_viewer.get(REVIEW_URL, {"tab": "synthesis", "min": "2"}).content.decode()
    assert spread.path in page and "Write a paragraph" not in page  # AI off: no button
    time = client_viewer.get(REVIEW_URL, {"tab": "synthesis", "view": "time"}).content.decode()
    assert "Persistent" in time and "2022, 2023, 2024" in time
    assert (
        "Annual" in client_viewer.get(REVIEW_URL, {"tab": "synthesis", "view": "coverage"}).content.decode()
    )


def test_actions_cards_filters_status_and_exports(analysed, client, editor):
    client.force_login(editor)
    schools = analysed.review_action_points.get(action__icontains="schools")
    response = client.get(REVIEW_URL, {"tab": "actions", "period": "all"})
    figures = response.context["figures"]
    assert (figures["open"], figures["high"], figures["derived"], figures["overdue"]) == (2, 1, 1, 2)
    assert figures["by_owner"][0][1] == 1
    row = client.post(
        reverse("knowledge:review_action_status", args=[schools.pk]),
        {"status": "done"},
        HTTP_HX_REQUEST="true",
    )
    assert row.status_code == 200 and "selected" in row.content.decode()
    done = client.get(REVIEW_URL, {"tab": "actions", "period": "all", "status": "done"})
    assert [p.pk for p in done.context["points"]] == [schools.pk]
    overdue = client.get(REVIEW_URL, {"tab": "actions", "period": "all", "overdue": "1"})
    assert all(p.overdue for p in overdue.context["points"]) and len(overdue.context["points"]) == 1
    owner = client.get(REVIEW_URL, {"tab": "actions", "period": "all", "owner": "Unassigned"})
    assert [p.owner_text for p in owner.context["points"]] == ["Unassigned"]
    old = Document.objects.create(title="Old plan", text="x", year=2015, review_batch=analysed.review_batch,
                                  review_status="done")  # fmt: skip
    analysed.year = 2024
    analysed.save()
    DocumentActionPoint.objects.create(document=old, action="An old commitment.", action_key="old")
    current = client.get(REVIEW_URL, {"tab": "actions"})
    assert "An old commitment." not in current.content.decode()
    assert (
        "An old commitment." in client.get(REVIEW_URL, {"tab": "actions", "period": "all"}).content.decode()
    )
    csv_response = client.get(
        REVIEW_URL, {"tab": "actions", "period": "all", "status": "done", "export": "csv"}
    )
    body = b"".join(csv_response.streaming_content).decode("utf-8-sig")
    assert body.splitlines()[0].startswith("Batch,Document,Where,Action,Owner")
    assert (
        "Rehabilitate 20 schools." in body and "Ministry of Education" in body and len(body.splitlines()) == 2
    )
    xlsx = client.get(REVIEW_URL, {"tab": "actions", "period": "all", "export": "xlsx"})
    assert xlsx.status_code == 200 and xlsx["Content-Disposition"].endswith('.xlsx"')
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    assert sheet.max_row == 4 and sheet.cell(1, 4).value == "Action"


def test_the_findings_csv_holds_the_filter_and_guards_formulas(analysed, client_viewer):
    DocumentFinding.objects.filter(text__icontains="Water supply").update(text="=HYPERLINK(1)")
    response = client_viewer.get(REVIEW_URL, {"tab": "findings", "export": "csv", "category": "observation"})
    body = b"".join(response.streaming_content).decode("utf-8-sig")
    lines = body.splitlines()
    assert lines[0].startswith("Batch,Document,Page,Link") and len(lines) == 2
    assert "'=HYPERLINK(1)" in body
    index = client_viewer.get(REVIEW_URL, {"tab": "findings", "view": "index", "export": "csv"})
    assert "Annual report 2024" in b"".join(index.streaming_content).decode("utf-8-sig")
    statements_csv = client_viewer.get(REVIEW_URL, {"tab": "findings", "view": "statements", "export": "csv"})
    assert "(Annual report 2024, p. 1)" in b"".join(statements_csv.streaming_content).decode("utf-8-sig")


def test_the_desk_review_opens_and_cites_its_evidence(db, client_viewer):
    from neurodb.knowledge import text

    annual, donor = ReviewBatch.objects.create(name="Annual"), ReviewBatch.objects.create(name="Donor")
    a, b = _doc(annual, "Report A", 2023), _doc(donor, "Report B & <C>", 2024)
    topic = _topics(1)[0]
    _add(a, topic, "Teachers lack training in rural schools.", 2023, quote="Teachers lack training")
    _add(b, topic, "Teachers lack training in rural schools of the north.", 2024)
    rejected = _add(b, topic, "A rejected point that must not appear.", 2024, verdict="rejected")
    DocumentStatement.objects.create(document=a, text="Training is the main gap.", urgency=90)
    DocumentActionPoint.objects.create(document=a, action="Train 500 teachers.", action_key="t",
                                       owner_text="Ministry of Education", deadline_text="2025")  # fmt: skip
    response = client_viewer.get(reverse("knowledge:review_docx"), {"min": "2"})
    assert response.status_code == 200
    assert (
        response["Content-Disposition"] == f'attachment; filename="desk-review-{timezone.localdate()}.docx"'
    )
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        names = set(archive.namelist())
        styles = archive.read("word/styles.xml").decode()
    assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml", "word/styles.xml"} <= names
    assert all(
        f'w:styleId="{s}"' in styles for s in ("Title", "Heading1", "Heading2", "Heading3", "TableGrid")
    )
    words = text.docx_text(response.content)
    for heading in ("Desk review", "Scope", "Summary", "Recurring themes", "Recurring challenges",
                    "Findings repeated across documents", "Most urgent statements",
                    "Open action points by owner", "Coverage", "Method"):  # fmt: skip
        assert heading in words, heading
    assert topic.path in words and "Raised by 2 documents" in words
    assert "(Report A, p. 2)" in words and "(Report B & <C>, p. 2)" in words
    assert "Train 500 teachers." in words and "Ministry of Education (1)" in words
    assert "Training is the main gap." in words and rejected.text not in words
    assert "Evidence score (0–100)" in words


def test_write_a_paragraph_cites_real_findings_and_counts_against_a_quota(
    db, client_viewer, viewer, monkeypatch
):
    annual = ReviewBatch.objects.create(name="Annual")
    a, b = _doc(annual, "Report A", 2023), _doc(annual, "Report B", 2024)
    topic = _topics(1)[0]
    _add(a, topic, "Teachers lack training.", 2023)
    _add(b, topic, "Teachers lack training in the north.", 2024)
    url = reverse("knowledge:review_paragraph", args=[topic.pk])
    off = client_viewer.post(url, HTTP_HX_REQUEST="true")
    assert "AI is switched off" in off.content.decode()
    sent = []

    def create(**params):
        sent.append(params)
        answer = {"paragraph": "Two reports find teachers lack training [1][2], getting worse [99]."}
        return FakeReview._answer(json.dumps(answer))

    api = SimpleNamespace(responses=SimpleNamespace(create=create))
    monkeypatch.setattr(agent, "client", lambda: SimpleNamespace(with_options=lambda **kw: api))
    with override_settings(**AI_ON):
        page = client_viewer.get(REVIEW_URL, {"tab": "synthesis"}).content.decode()
        assert "Write a paragraph" in page and "0 of 50 today" in page
        answer = client_viewer.post(url, HTTP_HX_REQUEST="true").content.decode()
        assert "(Report A, p. 2)" in answer and "(Report B, p. 2)" in answer and "[99]" not in answer
        assert "1 of 50 today" in answer
        content = sent[0]["input"][0]["content"]
        assert topic.path in content and "[1]" in content and sent[0]["store"] is False
        assert AIUsage.objects.get(feature=usage.DOC_REVIEW).calls == 1
        from neurodb.knowledge.models import ReviewParagraph

        row = ReviewParagraph.objects.get()
        assert row.status == "done" and row.user == viewer and row.findings == 2
        ReviewParagraph.objects.bulk_create(
            [ReviewParagraph(user=viewer, topic=topic, status="done") for _ in range(49)]
        )
        refused = client_viewer.post(url, HTTP_HX_REQUEST="true").content.decode()
        assert "50 paragraphs; the count starts again" in refused and len(sent) == 1
        assert ReviewParagraph.objects.filter(status="limited").count() == 1


def test_batches_documents_and_the_analyse_button(report, client, editor, started, media):
    from django.core.files.uploadedfile import SimpleUploadedFile

    client.force_login(editor)
    client.post(reverse("knowledge:review_batch_new"), {"name": "  Evaluations  ", "description": "Mid-term"})
    evaluations = ReviewBatch.objects.get(name="Evaluations")
    assert evaluations.created_by == editor
    loose = Document.objects.create(title="Loose", text="x", status=Document.Status.READY)
    client.post(
        reverse("knowledge:review_batch_add", args=[evaluations.pk]), {"documents": [loose.pk, report.pk]}
    )
    loose.refresh_from_db()
    report.refresh_from_db()
    assert loose.review_batch == evaluations and loose.review_status == "pending"
    assert report.review_batch != evaluations  # already in a batch: left where it is
    analyse = reverse("knowledge:review_analyse", args=[loose.pk])
    response = client.post(analyse, follow=True)
    assert "switched off" in response.content.decode() and not started  # the review is off
    found = DocumentReviewSettings.load()
    found.enabled = True
    found.save()
    with override_settings(**AI_ON):
        client.post(analyse)
    assert started[-1] == ("review_documents", "--document", str(loose.pk), "--triggered-by", "page")
    client.post(reverse("knowledge:review_reference", args=[loose.pk]), {"reference": "1"})
    loose.refresh_from_db()
    assert loose.review_status == "reference"
    with override_settings(**AI_ON):
        refused = client.post(analyse, follow=True).content.decode()
    assert "reference document is not analysed" in refused
    client.post(reverse("knowledge:review_batch_edit", args=[evaluations.pk]), {"action": "archive"})
    evaluations.refresh_from_db()
    assert evaluations.archived
    # an archived batch takes no new document (it would never be analysed)
    other = Document.objects.create(title="Other", text="x", status=Document.Status.READY)
    client.post(reverse("knowledge:review_batch_add", args=[evaluations.pk]), {"documents": [other.pk]})
    other.refresh_from_db()
    assert other.review_batch is None
    page = client.get(REVIEW_URL, {"tab": "documents", "batch": evaluations.pk}).content.decode()
    assert "Upload documents into this batch" not in page and "Bring the batch back" in page
    client.post(reverse("knowledge:review_batch_edit", args=[evaluations.pk]), {"action": "restore"})
    # uploaded from the batch: the document goes in it
    upload = client.post(
        f"{reverse('knowledge:add')}?batch={evaluations.pk}",
        {"title": "Evaluation", "batch": evaluations.pk, "files": SimpleUploadedFile("eval.txt", b"text")},
    )
    assert upload.status_code == 302 and upload["Location"].startswith(REVIEW_URL)
    assert Document.objects.get(title="Evaluation").review_batch == evaluations


def test_a_pdf_opens_in_the_browser_at_its_page(db, client_viewer, media):
    from django.core.files.uploadedfile import SimpleUploadedFile

    pdf_doc = Document.objects.create(title="Report", file=SimpleUploadedFile("r.pdf", b"%PDF-1.4"))
    response = client_viewer.get(reverse("knowledge:file", args=[pdf_doc.pk]))
    assert response["Content-Disposition"].startswith("inline")
    text_doc = Document.objects.create(title="Notes", file=SimpleUploadedFile("n.txt", b"x"))
    assert client_viewer.get(reverse("knowledge:file", args=[text_doc.pk]))["Content-Disposition"].startswith(
        "attachment"
    )


def test_donors_never_reach_the_review(analysed):
    from django.test import Client

    from tests.donors.test_donor_access import make_donor

    donor = Client()
    donor.force_login(make_donor().user)
    assert donor.get(REVIEW_URL)["Location"] == reverse("donors:page")
    assert donor.get(REVIEW_URL, {"tab": "findings"}, HTTP_HX_REQUEST="true").status_code == 403
    assert donor.get(reverse("knowledge:review_docx")).status_code == 302
    finding = analysed.findings.first()
    response = donor.post(
        reverse("knowledge:review_finding_verdict", args=[finding.pk]), {"verdict": "rejected"}
    )
    assert response.status_code in (302, 403)
    finding.refresh_from_db()
    assert finding.verdict == "unreviewed"


def test_every_dashboard_figure_and_the_owner_chart_open_the_rows_they_count(analysed, client_viewer, batch):
    from urllib.parse import quote

    # in the batch but not counted as analysed: one waiting, one reference
    Document.objects.create(title="Waiting", text="x", review_batch=batch, review_status="pending")
    Document.objects.create(title="A reference", text="x", review_batch=batch, review_status="reference")
    response = client_viewer.get(REVIEW_URL, {"tab": "dashboard"})
    figures, tiles, links = (
        response.context["figures"],
        response.context["tiles"],
        response.context["tile_links"],
    )
    documents = client_viewer.get(links["documents"])
    assert [d.title for d in documents.context["rows"]] == [analysed.title] and tiles["documents"] == 1
    assert "as counted on the Dashboard" in documents.content.decode()
    index_csv = client_viewer.get(links["documents"] + "&export=csv")
    assert len(b"".join(index_csv.streaming_content).decode("utf-8-sig").splitlines()) == 2
    assert client_viewer.get(links["findings"]).context["page_obj"].paginator.count == tiles["findings"]
    assert client_viewer.get(links["statements"]).context["page_obj"].paginator.count == tiles["statements"]
    assert (
        client_viewer.get(links["open_actions"]).context["page_obj"].paginator.count == tiles["open_actions"]
    )
    for key, param in (("programme", "programme"), ("tags", "topic")):
        assert figures["charts"][key], key
        for label, n, drill in figures["charts"][key]:
            listed = client_viewer.get(REVIEW_URL, {"tab": "findings", "counted": "1", param: drill})
            assert listed.context["page_obj"].paginator.count == n, label
    # the owner chart counts the period and batch (as the cards): a bar opens those rows, whatever the
    # table's own filter
    actions = client_viewer.get(REVIEW_URL, {"tab": "actions", "period": "all", "priority": "high"})
    bars = actions.context["figures"]["by_owner"]
    assert len(bars) == 2
    for owner, n, drill in bars:
        listed = client_viewer.get(actions.context["owner_href"].replace("{drill}", quote(drill)))
        assert listed.context["page_obj"].paginator.count == n, owner


def test_accept_all_decides_only_the_findings_listed(analysed, client, editor):
    client.force_login(editor)
    page = client.get(
        REVIEW_URL, {"tab": "findings", "view": "document", "document": analysed.pk, "category": "challenge"}
    ).content.decode()
    form = re.search(r'name="verdict" value="accepted">(.*?)<button', page).group(1)
    listed = {int(pk) for pk in re.findall(r'name="finding" value="(\d+)"', form)}
    assert listed == set(analysed.findings.filter(category="challenge").values_list("pk", flat=True))
    client.post(
        reverse("knowledge:review_bulk_verdict", args=[analysed.pk]),
        {"verdict": "accepted", "finding": sorted(listed)},
    )
    assert set(analysed.findings.filter(verdict="accepted").values_list("pk", flat=True)) == listed
    assert (
        analysed.findings.filter(verdict="unreviewed").count() == analysed.findings.count() - len(listed) > 0
    )


def test_verified_only_the_desk_review_cites_accepted_findings_only(analysed, client, editor):
    from neurodb.knowledge import text

    client.force_login(editor)
    statement = analysed.statements.get()
    DocumentStatement.objects.filter(pk=statement.pk).update(verdict="accepted")
    client.post(reverse("knowledge:review_verified"), {"on": "1"})
    words = text.docx_text(client.get(reverse("knowledge:review_docx")).content)
    # the statement is used; the finding it cites was not accepted, so the report does not point to it
    assert statement.text in words and "(Annual report 2024)" in words
    assert "(Annual report 2024, p. 1)" not in words
    DocumentFinding.objects.filter(pk=_finding(analysed, "out of school").pk).update(verdict="accepted")
    words = text.docx_text(client.get(reverse("knowledge:review_docx")).content)
    assert f"{statement.text} (Annual report 2024, p. 1)" in words


def test_the_review_pages_never_load_the_documents_full_text(analysed, client, editor):
    from django.db import connection
    from django.test.utils import CaptureQueriesContext

    client.force_login(editor)
    DocumentFinding.objects.filter(text__icontains="out of school").update(verdict="accepted")
    asked = [{"tab": tab} for tab in TABS]
    asked += [{"tab": "findings", "view": v} for v in ("document", "statements", "index")]
    asked += [{"tab": "synthesis", "view": v, "min": "1"} for v in ("time", "coverage", "repeated")]
    asked += [
        {"tab": "findings", "export": "csv"},
        {"tab": "findings", "view": "statements", "export": "csv"},
    ]
    asked += [{"tab": "actions", "period": "all", "export": "csv"}]
    with CaptureQueriesContext(connection) as queries:
        for params in asked:
            response = client.get(REVIEW_URL, params)
            assert response.status_code == 200, params
            if response.streaming:
                b"".join(response.streaming_content)
        assert client.get(reverse("knowledge:review_docx")).status_code == 200
    heavy = [q["sql"] for q in queries.captured_queries if '"knowledge_document"."text"' in q["sql"]]
    assert not heavy, heavy[:1]
