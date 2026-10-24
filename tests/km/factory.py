"""Small generated KM workbooks with obviously fake names and texts (never the user's workbook), and a load
of one in the test's own process."""

from __future__ import annotations

import datetime
import hashlib
import io
from typing import Any

import openpyxl

from neurodb.km import columns as cols
from neurodb.km import loading
from neurodb.km.models import KMVersion

DOCUMENTS_SHEET = "Test Documents"
FINDINGS_SHEET = "Test Findings"
READ_ME_SHEET = "Test Read Me"
DICTIONARY_SHEET = "Test Dictionary"
FAKE_PATH = "C:\\Fake\\Folder\\fake_source.xlsx"
ENTITY = "Fakeorg Institute; Doe, J.; Roe, R. et al."
TOKEN = "zqkmtoken"
PREPARED = datetime.datetime(2026, 10, 1)


def record_id(n: int) -> str:
    return f"TST-{n:04d}"


def finding_id(doc: int, n: int) -> str:
    return f"TST-{doc:04d}-F{n:05d}"


def document_row(n: int, findings: int) -> dict[str, Any]:
    return {
        "Record ID": record_id(n),
        "Item Title": f"Fake short title {n}",
        "Document Title": f"Fake document title {n}",
        "Publishing Entity / Authors": ENTITY if n == 1 else "Fakeorg University",
        "Publication Year": "2021" if n != 3 else "Not reported",
        "Document / Study Type": "Fake study",
        "Text Used": "Full text",
        "Evidence Confidence Level": "Level B",
        "Number of Findings": findings,
        "Document Link": None,
        "Data Quality Status": "Ready",
    }


def finding_row(doc: int, n: int, number: int) -> dict[str, Any]:
    row: dict[str, Any] = {
        "Finding ID": finding_id(doc, n),
        "Item Title": f"Fake short title {doc}",
        "Record ID": record_id(doc),
        "Document Title": f"Fake document title {doc}",
        "Pillar": "P1 – Fake pillar one" if number % 2 else "P2 – Fake pillar two",
        "Analytical Domain": "1. Fake domain",
        "Research Question": "RQ1 - Fake question one",
        "Sub-question": "Fake sub-question",
        "Sub-sub-question": "Fake sub-sub-question",
        "Indicator / Evidence Requirement": "Fake indicator",
        "Indicator Component Answered": "Fake component",
        "Also Relevant To": None,
        "Measure": "Fake measure",
        "Key Finding": f"Fake key finding {number} about {TOKEN}; <5 years (February 2023).",
        "Finding Analysis": "Fake analysis, as Roe (2018) found (Fakeorg Institute, 2020; Doe, 2019).",
        "Age Category / Age Bracket": "Fake ages",
        "Limitations": "Fake limitations (Poe, 2017, p. 4).",
        "Evidence Origin": "Primary: fake origin text",
        "Evidence Confidence Level": "Level A",
        "Publishing Entity / Author": ENTITY if doc == 1 else "Fakeorg University",
        "Document / Study Type": "Fake study",
        "Publication Year": 2022,
        "Study / Data Collection Period": datetime.datetime(2023, 2, 1),
        "Population": "Fake population",
        "Methods and Sample": "Jane Doe et al. (2021) asked fake.example.org/x and +961 3 123 456; “…” – Fakename, 34",
        "Source Link": None,
        "Data Quality Status": "Ready",
    }
    for flag in cols.FLAGS:
        row[flag.header] = "Yes" if flag.field in ("early_adolescence", "female", "mhpss") else "No"
    return row


def workbook(
    documents: int = 3,
    findings: int = 4,
    *,
    header_row: int = 1,
    drop: tuple[str, ...] = (),
    edits: dict[tuple[str, int, str], Any] | None = None,
    read_me: bool = True,
    stated: tuple[int | None, int | None] | None = None,
    sheet_names: tuple[str, str] = (DOCUMENTS_SHEET, FINDINGS_SHEET),
    extra_headers: dict[str, tuple[str, ...]] | None = None,
) -> bytes:
    """An .xlsx with every column of the KM workbook and fake values: ``documents`` documents and ``findings``
    findings (spread over the documents in turn), a Read Me sheet (a prepared-on date, a fake path, the stated
    counts: ``stated`` overrides them) and a dictionary sheet. ``drop`` leaves those headers out;
    ``header_row`` puts title rows above the header; ``edits`` sets cells by (``"documents"`` or
    ``"findings"``, Excel row, header) once the rows are written (relative to the header in row 1)."""
    wb = openpyxl.Workbook()
    first = wb.active
    if read_me:
        first.title = READ_ME_SHEET
        first.append(["Fake workbook title"])
        first.append([cols.PREPARED_ON, PREPARED])
        first.append([cols.SOURCE_WORKBOOK, FAKE_PATH])
        first.append([])
        first.append(["Dataset", "Rows", "Note"])
        stated_findings, stated_documents = stated if stated is not None else (findings, documents)
        if stated_findings is not None:
            first.append([sheet_names[1], stated_findings, "fake note"])
        if stated_documents is not None:
            first.append([sheet_names[0], stated_documents, "fake note"])
    else:
        first.title = "Fake Cover"
    per_doc = {n: 0 for n in range(1, documents + 1)}
    finding_rows = []
    for number in range(1, findings + 1):
        doc = (number - 1) % documents + 1 if documents else 1
        per_doc[doc] = per_doc.get(doc, 0) + 1
        finding_rows.append(finding_row(doc, per_doc[doc], number))
    document_rows = [document_row(n, per_doc[n]) for n in range(1, documents + 1)]
    extra_headers = extra_headers or {}
    for title, kind, columns, data in (
        (sheet_names[0], "documents", cols.DOCUMENT_COLUMNS, document_rows),
        (sheet_names[1], "findings", cols.FINDING_COLUMNS, finding_rows),
    ):
        ws = wb.create_sheet(title)
        for _ in range(header_row - 1):
            ws.append(["Fake title row"])
        headers = [c.header for c in columns if c.header not in drop] + list(extra_headers.get(kind, ()))
        ws.append(headers)
        for row in data:
            ws.append([row.get(h, "fake extra") for h in headers])
        for (sheet, number, header), value in (edits or {}).items():
            if sheet == kind:
                ws.cell(row=number + header_row - 1, column=headers.index(header) + 1).value = value
    dictionary = wb.create_sheet(DICTIONARY_SHEET)
    dictionary.append(["Field", "Type", "Allowed values"])
    dictionary.append(["Fake field", "Text", "Fake values"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def version(
    data: bytes, *, fewer_ok: bool = False, user=None, name: str = "fake.xlsx", **fields
) -> KMVersion:
    fields.setdefault("status", KMVersion.Status.QUEUED)
    return KMVersion.objects.create(
        content=data,
        file_name=name,
        file_size=len(data),
        file_sha256=hashlib.sha256(data).hexdigest(),
        fewer_ok=fewer_ok,
        uploaded_by=user,
        **fields,
    )


def load(data: bytes, *, fewer_ok: bool = False, user=None) -> KMVersion:
    """Upload ``data`` as a new version and load it here (no background process)."""
    queued = version(data, fewer_ok=fewer_ok, user=user)
    loading.run(queued.pk, fewer_ok=fewer_ok)
    return KMVersion.objects.defer("content").get(pk=queued.pk)
