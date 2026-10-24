"""The KM workbook's upload check and how its cells are read (SPEC_km §2.4-2.6), on small generated workbooks
with fake values only."""

from __future__ import annotations

import datetime
import io
import zipfile

import openpyxl
import pytest

from neurodb.km import columns as cols
from neurodb.km import words, workbook
from neurodb.km.models import KMFinding
from tests.km import factory

FIND, DOCS = factory.FINDINGS_SHEET, factory.DOCUMENTS_SHEET


def _check(data: bytes, name: str = "fake.xlsx") -> workbook.Checked:
    return workbook.check_upload(data, name)


# ------------------------------------------------------------------------------------------- the sheets
@pytest.mark.parametrize(
    ("title", "kind"),
    [
        ("Test Documents", cols.DOCUMENTS),
        ("  fake   FINDINGS ", cols.FINDINGS),
        ("Findings", cols.FINDINGS),
        ("Test Read Me", cols.READ_ME),
        ("README", cols.READ_ME),
        ("Test Dictionary", ""),
        ("Findings of 2024", ""),  # matched by the last word only
    ],
)
def test_sheets_are_found_by_their_last_word(title, kind):
    assert cols.sheet_kind(title) == kind


def test_a_good_workbook_passes_the_upload_check_with_no_note():
    checked = _check(factory.workbook())
    assert checked.problems == [] and checked.notes == []
    assert checked.sheets[cols.FINDINGS].title == FIND and checked.read_me == factory.READ_ME_SHEET


# -------------------------------------------------------------------------------------------- refusals
def test_too_big_and_wrong_type_are_two_sentences(monkeypatch):
    monkeypatch.setattr(workbook, "MAX_UPLOAD_BYTES", 10)
    assert _check(b"x" * 11 * 1024 * 1024 + b"x", "fake.csv").problems == [
        "This file is 11.0 MB; the upload takes at most 25 MB.",
        "This is a CSV file: open it in Excel and save it as Excel Workbook (.xlsx).",
    ]
    assert _check(b"x" * 20).problems == ["This file is 0.0 MB; the upload takes at most 25 MB."]


def test_the_extension_is_compared_without_case_and_a_missing_one_is_named():
    assert _check(factory.workbook(), "FAKE.XLSX").problems == []
    assert _check(factory.workbook(), "fake").problems == [str(words.NO_EXTENSION)]
    assert _check(b"a,b\n1,2\n", "fake.csv").problems == [
        "This is a CSV file: open it in Excel and save it as Excel Workbook (.xlsx)."
    ]
    assert _check(b"a,b\n1,2\n", "C:\\fake\\folder\\fake.xls").problems == [
        "This is a XLS file: open it in Excel and save it as Excel Workbook (.xlsx)."
    ]


def test_a_protected_workbook_is_named_as_such():
    ole = workbook.OLE + b"\xa1\xb1\x1a\xe1" + b"\x00" * 512
    assert _check(ole).problems == [
        "This workbook is protected (a password or a sensitivity label that encrypts it). Save a copy without "
        "the protection, then upload that copy."
    ]


def _zip(**parts: bytes) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, body in parts.items():
            archive.writestr(name, body)
    return out.getvalue()


def test_a_file_that_is_not_an_excel_workbook_is_refused(monkeypatch):
    assert _check(b"not a zip at all").problems == [str(words.NOT_EXCEL)]
    assert _check(_zip(fake=b"fake")).problems == [str(words.NOT_EXCEL)]  # a zip, not a workbook
    good = factory.workbook()
    monkeypatch.setattr(workbook, "MAX_UNPACKED", 1000)  # a zip bomb: too large unpacked
    assert _check(good).problems == [str(words.NOT_EXCEL)]
    monkeypatch.setattr(workbook, "MAX_UNPACKED", 300 * 1024 * 1024)
    monkeypatch.setattr(workbook, "MAX_PARTS", 3)
    assert _check(good).problems == [str(words.NOT_EXCEL)]


def test_a_missing_sheet_lists_the_sheets_the_workbook_has():
    problems = _check(factory.workbook(sheet_names=(DOCS, "Test Results"))).problems
    assert problems == [
        "The workbook has no sheet whose name ends with “Findings”.",
        f"Its sheets are: {factory.READ_ME_SHEET}, {DOCS}, Test Results, {factory.DICTIONARY_SHEET}.",
    ]


def test_two_findings_sheets_are_refused():
    wb = openpyxl.load_workbook(io.BytesIO(factory.workbook()))
    wb.create_sheet("Old Findings")
    out = io.BytesIO()
    wb.save(out)
    problems = _check(out.getvalue()).problems
    assert problems[0] == "Two sheets end with “Findings”: keep one." and problems[1].startswith(
        "Its sheets are: "
    )


def test_missing_required_columns_are_listed_with_the_columns_not_recognised():
    data = factory.workbook(
        drop=("Finding ID", "Key Finding", "Record ID"), extra_headers={"findings": ("Fake Finding Code",)}
    )
    problems = _check(data).problems
    assert problems == [
        f"The sheet “{DOCS}” has no column “Record ID”.",  # the documents sheet has one too
        f"The sheet “{FIND}” has no column “Finding ID”, “Record ID”, “Key Finding”.",
        "Columns in that sheet that NeuroDB does not recognise: “Fake Finding Code”.",
    ]


def test_both_data_sheets_are_reported_together():
    problems = _check(factory.workbook(drop=("Record ID",))).problems
    assert problems == [
        f"The sheet “{DOCS}” has no column “Record ID”.",
        f"The sheet “{FIND}” has no column “Record ID”.",
    ]


def test_a_header_under_a_title_row_is_named_with_its_row():
    problems = _check(factory.workbook(header_row=3)).problems
    assert problems == [
        f"The column names must be in the first row of the sheet “{DOCS}”; they are in row 3.",
        f"The column names must be in the first row of the sheet “{FIND}”; they are in row 3.",
    ]


def test_a_repeated_column_is_refused():
    data = factory.workbook(extra_headers={"findings": ("key  finding",)})
    assert _check(data).problems == [f"The column “Key Finding” is twice in the sheet “{FIND}”."]


def test_a_sheet_with_no_row_under_its_header_is_refused():
    assert _check(factory.workbook(findings=0)).problems == [
        f"The sheet “{FIND}” has no rows under its header."
    ]


def test_an_optional_column_missing_loads_with_its_note():
    data = factory.workbook(drop=("MHPSS", "Measure"), extra_headers={"findings": ("Fake Extra Column",)})
    checked = _check(data)
    assert checked.problems == []
    assert checked.notes == [
        f"The sheet “{FIND}” has no column “Measure”: it is shown empty.",
        f"The sheet “{FIND}” has no column “MHPSS”: it is shown empty.",
        f"Columns of the sheet “{FIND}” that are not read: “Fake Extra Column”.",
    ]
    assert _check(factory.workbook(read_me=False)).notes == [
        "The prepared-on date is not known: the workbook has no Read Me sheet."
    ]


@pytest.mark.django_db
def test_an_optional_column_missing_loads_empty():
    version = factory.load(factory.workbook(drop=("MHPSS", "Measure")))
    assert version.status == "live"
    assert f"The sheet “{FIND}” has no column “MHPSS”: it is shown empty." in version.warnings
    assert set(KMFinding.objects.values_list("mhpss", "measure")) == {("", "")}


# ------------------------------------------------------------------------------------------- the headers
@pytest.mark.parametrize(
    ("written", "header"),
    [
        ("finding id", "Finding ID"),
        ("  Key   Finding ", "Key Finding"),
        ("Sub–question", "Sub-question"),  # an en dash
        ("Indicator/Evidence Requirement", "Indicator / Evidence Requirement"),
        ("Publishing Entity / Authors", "Publishing Entity / Author"),
        ("PUBLISHING ENTITY/AUTHOR", "Publishing Entity / Authors"),
    ],
)
def test_headers_match_after_folding(written, header):
    assert cols.fold(written) == cols.fold(header)


# ----------------------------------------------------------------------------------------- the cells
def _read(kind: str, raw, header: str = "Fake Column", size: int = 0) -> cols.Cells:
    cells = cols.Cells()
    cols.read(cols.Column("field", header, kind, size=size), raw, cells)
    return cols.finish(cells)


@pytest.mark.parametrize(
    ("raw", "code"),
    [
        ("Yes", "Y"),
        ("yes ", "Y"),
        ("Y", "Y"),
        (True, "Y"),
        (1, "Y"),
        ("No", "N"),
        ("FALSE", "N"),
        (0, "N"),
        ("Not reported", "R"),
        ("not applicable", "A"),
        ("N/A", "A"),
        (None, ""),
    ],
)
def test_flags_read_their_usual_spellings(raw, code):
    cells = _read("flag", raw)
    assert cells.values["field"] == code and cells.checks == {}


def test_an_unexpected_flag_is_a_value_to_check():
    cells = _read("flag", "Partly", "MHPSS")
    assert cells.values["field"] == "" and cells.checks == {"MHPSS": "Partly"}


@pytest.mark.parametrize(
    ("raw", "letter"), [("Level A", "A"), ("Level b", "B"), ("level  c", "C"), ("d", "D"), ("", "")]
)
def test_confidence_levels(raw, letter):
    cells = _read("confidence", raw)
    assert (
        cells.values["confidence"] == letter and cells.values["confidence_text"] == raw and not cells.checks
    )


def test_a_spilled_value_in_a_short_column_is_shown_as_given_and_marked():
    spill = ("Fake text spilled over from another column. " * 3)[:100]
    cells = _read("confidence", spill, "Evidence Confidence Level")
    assert cells.values["confidence"] == "" and cells.checks == {"Evidence Confidence Level": spill}
    cells = _read("type", spill, "Document / Study Type")
    assert cells.values["field"] == spill and cells.checks == {"Document / Study Type": spill}
    assert _read("type", "Fake study").checks == {}
    long = "y" * 250
    cells = _read("type", long, "Document / Study Type", size=200)
    assert cells.values["field"] == long[:200] and cells.checks == {"Document / Study Type": long[:300]}


@pytest.mark.parametrize(
    ("raw", "year", "text", "checked"),
    [
        (2022, 2022, "2022", False),
        (2022.0, 2022, "2022", False),
        ("2021", 2021, "2021", False),
        ("Not reported", None, "Not reported", False),
        ("", None, "", False),
        ("2020-2021", None, "2020-2021", True),
        (1850, None, "1850", True),
    ],
)
def test_years(raw, year, text, checked):
    cells = _read("year", raw, "Publication Year")
    assert cells.values["year"] == year and cells.values["year_text"] == text
    assert bool(cells.checks) is checked


def test_texts_numbers_and_dates_are_written_as_given():
    assert cols.as_text(datetime.datetime(2023, 2, 1)) == "2023-02-01"
    assert cols.as_text(datetime.date(2023, 2, 1)) == "2023-02-01"
    assert cols.as_text(2023.0) == "2023" and cols.as_text(12.5) == "12.5" and cols.as_text(7) == "7"
    assert cols.as_text("  line one\r\nline two\x00  ") == "line one\nline two"
    cells = _read("text", "x" * (cols.TEXT_MAX + 5))
    assert len(cells.values["field"]) == cols.TEXT_MAX and cells.checks


def test_codes_are_worked_out_beside_the_labels():
    assert _read("pillar", "P2 – Fake pillar two").values == {
        "pillar_code": "P2",
        "pillar": "P2 – Fake pillar two",
    }
    assert _read("pillar", "Fake pillar").values["pillar_code"] == ""
    assert _read("pillar", "Fake pillar").checks
    assert _read("domain", "3. Fake domain").values["domain_code"] == "D3"
    assert _read("rq", "RQ 4 - Fake question").values["rq_code"] == "RQ4"
    assert _read("rq", "Fake question").checks
    assert _read("origin", "Primary: fake origin text").values["origin_kind"] == "P"
    assert _read("origin", "Secondary: fake").values["origin_kind"] == "O"
    assert _read("origin", "").values["origin_kind"] == ""
    assert _read("quality", "Ready").values["needs_review"] is False
    assert _read("quality", "Review required: fake reason").values["needs_review"] is True
    assert _read("quality", "").values["needs_review"] is True
    assert _read("count", 12).values["field"] == 12
    assert _read("count", "twelve").checks and _read("count", "twelve").values["field"] is None


def test_the_question_key_is_the_pillar_and_the_wording():
    key = cols.question_key("P1", "RQ1 - Fake question one")
    assert key == cols.question_key("P1", "  rq1 -  FAKE question one ")  # the same wording, written apart
    assert key != cols.question_key("P2", "RQ1 - Fake question one")  # the same code in another pillar
    assert key != cols.question_key("P1", "RQ1 - Fake question two") and len(key) == 10
    assert cols.question_key("P1", "") == ""


@pytest.mark.django_db
def test_the_question_key_survives_a_new_version_with_the_same_wording():
    factory.load(factory.workbook())
    first = set(KMFinding.objects.values_list("finding_id", "question_key"))
    edits = {("findings", 2, "Research Question"): "rq1 -  fake QUESTION one"}
    factory.load(factory.workbook(edits=edits))
    assert set(KMFinding.objects.values_list("finding_id", "question_key")) == first


@pytest.mark.django_db
def test_the_read_me_gives_the_date_the_counts_and_the_source_name_only():
    version = factory.load(factory.workbook())
    assert version.prepared_on == datetime.date(2026, 10, 1)
    assert version.source_name == "fake_source.xlsx"  # never the folder
    assert (version.stated_findings, version.stated_documents) == (4, 3)
    finding = KMFinding.objects.get(finding_id=factory.finding_id(1, 1))
    assert finding.period == "2023-02-01" and finding.year == 2022 and finding.document_type == "Fake study"


# ---------------------------------------------------------------------------------------- the template
def test_the_excel_template_holds_every_header_and_no_other_text():
    wb = openpyxl.load_workbook(io.BytesIO(workbook.template()))
    by_kind = {cols.sheet_kind(title): wb[title] for title in wb.sheetnames}
    assert set(by_kind) == {cols.DOCUMENTS, cols.FINDINGS, cols.READ_ME}
    for kind in (cols.DOCUMENTS, cols.FINDINGS):
        ws = by_kind[kind]
        assert [c.value for c in ws[1]] == [column.header for column in cols.COLUMNS[kind]]
        assert ws.max_row == 1
        assert [bool(c.font.bold) for c in ws[1]] == [column.required for column in cols.COLUMNS[kind]]
    labels = [row[0] for row in by_kind[cols.READ_ME].iter_rows(values_only=True)]
    assert labels == [cols.PREPARED_ON, cols.SOURCE_WORKBOOK, *workbook.TEMPLATE_SHEETS.values()]
    assert all(len(row) < 2 or row[1] is None for row in by_kind[cols.READ_ME].iter_rows(values_only=True))
    # and it passes the upload check itself, but for its empty sheets
    problems = _check(workbook.template()).problems
    assert problems == [
        "The sheet “KM Documents” has no rows under its header.",
        "The sheet “KM Findings” has no rows under its header.",
    ]


def test_every_flag_has_its_model_field_and_a_plain_label():
    fields = {f.name for f in KMFinding._meta.get_fields()}
    assert len(cols.FLAGS) == 27 and {flag.field for flag in cols.FLAGS} <= fields
    assert {flag.group for flag in cols.FLAGS} == set(cols.FLAG_GROUPS)
    for flag in cols.FLAGS:
        assert "SEM" not in flag.label and flag.label != "MHPSS"
